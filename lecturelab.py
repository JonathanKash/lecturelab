"""LectureLab v2: turns lecture recordings into study material.

One module on purpose. Sections, top to bottom:
  config / logging / lock / tools  ->  database  ->  courses  ->  jobs  ->  ingest
  ->  inbox watcher  ->  v1 import  ->  web API
"""
from __future__ import annotations

import hashlib
import json
import logging
import logging.handlers
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import uuid
from contextlib import asynccontextmanager, contextmanager
from datetime import date as Date, datetime, timezone
from pathlib import Path
from typing import Literal

from dotenv import load_dotenv
# FastAPI names are module-level on purpose: with `from __future__ import annotations` every
# annotation is a string that FastAPI resolves against module globals, not create_app() locals.
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, PlainTextResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

# ---------------------------------------------------------------- config

ROOT = Path(__file__).resolve().parent
load_dotenv(ROOT / ".env")

DATA = Path(os.environ.get("LECTURELAB_DATA", ROOT / "data")).resolve()
DB_PATH = DATA / "lecturelab.db"
PORT = int(os.environ.get("LECTURELAB_PORT", "8765"))
LOG_PATH = Path(os.environ.get("LECTURELAB_LOG", ROOT / "server.log"))
VERSION = "2.0.0"
HOME = Path.home()
LOCAL_INBOX = Path(os.environ.get("LECTURELAB_INBOX", HOME / "Documents" / "LectureLab Inbox"))
MEDIA_EXTS = {".mp4", ".mov", ".mkv", ".webm", ".m4a", ".mp3", ".wav"}
NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0)

log = logging.getLogger("lecturelab")


class AppendCloseHandler(logging.FileHandler):
    """Opens server.log only for the moment of each write.

    Several processes (supervisor, server, sync) log to the same file, but only the server
    rotates it. On Windows a rename fails while another process holds the file open, so
    the non-rotating writers must not keep a handle.
    """

    def __init__(self, path):
        super().__init__(path, encoding="utf-8", delay=True)

    def emit(self, record):
        super().emit(record)
        self.close()


def setup_logging(console: bool = False, rotate: bool = False) -> None:
    """server.log for every process. Only the server passes rotate=True (5 MB x 3)."""
    if getattr(setup_logging, "_done", False):
        return
    fmt = logging.Formatter("%(asctime)s %(levelname)s [%(process)d] %(message)s")
    if rotate:
        fh = logging.handlers.RotatingFileHandler(LOG_PATH, maxBytes=5_000_000, backupCount=3, encoding="utf-8")
    else:
        fh = AppendCloseHandler(LOG_PATH)
    fh.setFormatter(fmt)
    root = logging.getLogger()
    root.setLevel(logging.INFO)
    root.addHandler(fh)
    if console:
        sh = logging.StreamHandler()
        sh.setFormatter(fmt)
        root.addHandler(sh)
    setup_logging._done = True


class FileLock:
    """OS-level exclusive lock (msvcrt on Windows, fcntl elsewhere).

    The OS drops the lock when the process dies, so a crash never leaves a stale lock.
    A threading.Lock covers threads inside the same process.
    """

    _thread_locks: dict[str, threading.Lock] = {}

    def __init__(self, path: Path):
        self.path = Path(path)
        self._fh = None
        self._tl = self._thread_locks.setdefault(str(self.path), threading.Lock())

    def acquire(self) -> bool:
        if not self._tl.acquire(blocking=False):
            return False
        fh = None
        try:
            self.path.parent.mkdir(parents=True, exist_ok=True)
            fh = open(self.path, "a+")
            if os.name == "nt":
                import msvcrt
                fh.seek(0)
                msvcrt.locking(fh.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(fh.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError:
            if fh:
                fh.close()
            self._tl.release()
            return False
        self._fh = fh
        return True

    def release(self) -> None:
        if not self._fh:
            return
        try:
            if os.name == "nt":
                import msvcrt
                self._fh.seek(0)
                msvcrt.locking(self._fh.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                import fcntl
                fcntl.flock(self._fh.fileno(), fcntl.LOCK_UN)
        finally:
            self._fh.close()
            self._fh = None
            self._tl.release()

    def __enter__(self):
        if not self.acquire():
            raise RuntimeError(f"lock busy: {self.path}")
        return self

    def __exit__(self, *exc):
        self.release()


def pid_alive(pid: int) -> bool:
    """True if a process with this pid is running. (os.kill(pid, 0) would KILL it on Windows.)"""
    if os.name != "nt":
        try:
            os.kill(pid, 0)
            return True
        except OSError:
            return False
    import ctypes
    k32 = ctypes.windll.kernel32
    h = k32.OpenProcess(0x1000, False, pid)  # PROCESS_QUERY_LIMITED_INFORMATION
    if not h:
        return False
    code = ctypes.c_ulong()
    ok = k32.GetExitCodeProcess(h, ctypes.byref(code))
    k32.CloseHandle(h)
    return bool(ok) and code.value == 259  # STILL_ACTIVE


# ---------------------------------------------------------------- external tools

def find_tool(name: str) -> str:
    """ffmpeg/ffprobe/soffice. winget adds them to PATH only for new shells, and the logon
    task may have an old PATH, so also look where winget and LibreOffice install them."""
    hit = shutil.which(name)
    if hit:
        return hit
    globs = {
        "ffmpeg": [HOME / "AppData/Local/Microsoft/WinGet/Packages", "Gyan.FFmpeg*/*/bin/ffmpeg.exe"],
        "ffprobe": [HOME / "AppData/Local/Microsoft/WinGet/Packages", "Gyan.FFmpeg*/*/bin/ffprobe.exe"],
        "soffice": [Path("C:/Program Files/LibreOffice/program"), "soffice.exe"],
    }
    if name in globs:
        base, pattern = globs[name]
        found = sorted(base.glob(pattern)) if base.exists() else []
        if found:
            return str(found[-1])
    raise FileNotFoundError(f"{name} not found; install it (see README)")


def run(cmd: list, timeout: float = 2 * 3600, low: bool = False, **kw) -> subprocess.CompletedProcess:
    """subprocess.run with no console window, no stdin, captured output and a deadline (so one
    corrupt file can't hang the worker forever); raises with stderr on failure. low=True runs it
    at below-normal priority, for background work that shouldn't slow the laptop down."""
    kw.setdefault("capture_output", True)
    flags = NO_WINDOW | (getattr(subprocess, "BELOW_NORMAL_PRIORITY_CLASS", 0) if low else 0)
    r = subprocess.run([str(c) for c in cmd], creationflags=flags, stdin=subprocess.DEVNULL,
                       timeout=timeout, **kw)
    if r.returncode != 0:
        dec = lambda b: b.decode("utf-8", "replace") if isinstance(b, bytes) else (b or "")  # noqa: E731
        # Some tools (the colab CLI) print their errors on stdout, so include text stdout too. Binary
        # stdout (ffmpeg piping PCM) is data, not a message: leave it out. The colab CLI's error
        # URLs carry short-lived runtime tokens: never let those reach server.log or the UI.
        out = r.stdout[-4000:] if isinstance(r.stdout, str) else ""
        err = (out.strip() + "\n" + dec(r.stderr).strip()).strip()  # stderr last: it survives the cut below
        err = re.sub(r"(?i)\b((?:token|authuser|key|code|state)=)[^&\s'\"]+", r"\1<redacted>", err)
        err = re.sub(r"(?i)(\"(?:\w*token|key|secret)\"\s*:\s*\")[^\"]+", r"\1<redacted>", err)
        # wsl.exe -d <distro> -e <tool> ...: name the tool, not "wsl.exe" or "-d"
        tool = cmd[4] if Path(str(cmd[0])).name == "wsl.exe" and len(cmd) > 4 else cmd[0]
        raise RuntimeError(f"{Path(str(tool)).name} failed ({r.returncode}): {err[-1500:]}")
    return r


def ffprobe(path: Path) -> dict:
    out = run([find_tool("ffprobe"), "-v", "error", "-print_format", "json", "-show_format", "-show_streams", path],
              timeout=120).stdout
    info = json.loads(out)
    streams = info.get("streams", [])
    video = [s for s in streams if s.get("codec_type") == "video" and not s.get("disposition", {}).get("attached_pic")]
    audio = [s for s in streams if s.get("codec_type") == "audio"]
    tags = {k.lower(): v for k, v in (info.get("format", {}).get("tags") or {}).items()}
    dur = info.get("format", {}).get("duration")
    return {"duration": float(dur) if dur not in (None, "N/A") else None, "has_video": bool(video),
            "has_audio": bool(audio), "creation_time": tags.get("creation_time"), "title": tags.get("title")}


# ---------------------------------------------------------------- database

# Append-only: migration N runs once when PRAGMA user_version < N. Never edit a shipped one.
MIGRATIONS = [
    # IF NOT EXISTS only here: databases from before versioning already have these tables.
    """
    CREATE TABLE IF NOT EXISTS courses (
        id          INTEGER PRIMARY KEY,
        code        TEXT NOT NULL UNIQUE,      -- folder name in inboxes, e.g. COP4020
        name        TEXT NOT NULL,
        vocabulary  TEXT NOT NULL DEFAULT '',  -- one term per line; fed to Whisper and the note prompts
        canvas_id   INTEGER,
        auto_pull   INTEGER NOT NULL DEFAULT 0,
        created_at  TEXT NOT NULL DEFAULT (datetime('now'))
    );
    CREATE TABLE IF NOT EXISTS meta (key TEXT PRIMARY KEY, value TEXT);
    """,
    """
    CREATE TABLE lectures (
        id            INTEGER PRIMARY KEY AUTOINCREMENT,  -- never reuse an id: media/<id>/ may linger
        course_id     INTEGER REFERENCES courses(id) ON DELETE SET NULL,
        title         TEXT NOT NULL,
        date          TEXT,                    -- YYYY-MM-DD, from source metadata when available
        duration_sec  REAL,
        audio_path    TEXT,                    -- relative to DATA: media/<id>/audio.flac (16 kHz mono)
        video_path    TEXT,                    -- relative to DATA: media/<id>/original.<ext>, when there is video
        status        TEXT NOT NULL,           -- see STATUSES; 'ingesting' until its files are in place
        error         TEXT,
        sha_source    TEXT,                    -- sha256 of the file as received
        sha_pcm       TEXT,                    -- sha256 of the normalized PCM: same audio, different container
        fingerprint   BLOB,                    -- per-second loudness envelope (float16): same lecture, re-encoded
        dup_of        INTEGER REFERENCES lectures(id) ON DELETE SET NULL,
        transcript_source TEXT,                -- colab-large-v3 | local-<model> | v1
        created_at    TEXT NOT NULL DEFAULT (datetime('now')),
        updated_at    TEXT NOT NULL DEFAULT (datetime('now'))
    );
    CREATE INDEX lectures_sha_source ON lectures(sha_source);
    CREATE INDEX lectures_sha_pcm ON lectures(sha_pcm);
    CREATE TABLE sources (
        id          INTEGER PRIMARY KEY,
        lecture_id  INTEGER NOT NULL REFERENCES lectures(id) ON DELETE CASCADE,
        kind        TEXT NOT NULL,             -- file | inbox | drive | zoom | canvas | url | v1
        url         TEXT,
        added_at    TEXT NOT NULL DEFAULT (datetime('now')),
        UNIQUE(lecture_id, url)
    );
    CREATE TABLE segments (
        id          INTEGER PRIMARY KEY,
        lecture_id  INTEGER NOT NULL REFERENCES lectures(id) ON DELETE CASCADE,
        idx         INTEGER NOT NULL,
        start       REAL NOT NULL,
        end         REAL NOT NULL,
        text        TEXT NOT NULL,
        words       TEXT,                      -- JSON [[start, end, word], ...]
        slide_no    INTEGER,
        UNIQUE(lecture_id, idx)
    );
    CREATE TABLE jobs (
        id          INTEGER PRIMARY KEY,
        kind        TEXT NOT NULL,             -- ingest | ...
        lecture_id  INTEGER REFERENCES lectures(id) ON DELETE CASCADE,
        status      TEXT NOT NULL DEFAULT 'queued',  -- queued | running | done | error
        payload     TEXT NOT NULL DEFAULT '{}',
        result      TEXT,
        progress    TEXT,
        error       TEXT,
        owner_pid   INTEGER,
        attempts    INTEGER NOT NULL DEFAULT 0,
        retry_at    TEXT,
        created_at  TEXT NOT NULL DEFAULT (datetime('now')),
        updated_at  TEXT NOT NULL DEFAULT (datetime('now'))
    );
    CREATE INDEX jobs_status ON jobs(status, kind);
    CREATE TABLE seen_files (
        path        TEXT PRIMARY KEY,
        size        INTEGER NOT NULL,
        mtime       REAL NOT NULL,
        job_id      INTEGER,
        tries       INTEGER NOT NULL DEFAULT 1,  -- failed ingests are retried up to 3 times
        seen_at     TEXT NOT NULL DEFAULT (datetime('now'))
    );
    """,
    """
    CREATE TABLE canvas_items (
        key         TEXT NOT NULL,             -- e.g. file:123, kaltura:1_abcd, panopto:<guid>
        course_id   INTEGER NOT NULL REFERENCES courses(id) ON DELETE CASCADE,
        kind        TEXT NOT NULL,             -- file | kaltura | panopto | mediasite | youtube | zoom | media_object
        title       TEXT,
        url         TEXT NOT NULL,             -- what we hand to the downloader
        page_url    TEXT,                      -- where it was found, for the user
        created_at  TEXT,
        status      TEXT NOT NULL DEFAULT 'new',  -- new | queued | pulled | ignored | error
        lecture_id  INTEGER REFERENCES lectures(id) ON DELETE SET NULL,
        error       TEXT,
        found_at    TEXT NOT NULL DEFAULT (datetime('now')),
        PRIMARY KEY (course_id, key)           -- the same embed can appear in two courses
    );
    """,
    """
    CREATE TABLE transcribe_batches (
        id            INTEGER PRIMARY KEY AUTOINCREMENT,
        mode          TEXT NOT NULL,           -- colab | local
        gpu           TEXT,                    -- T4 / L4 on Colab, the model name locally
        lecture_ids   TEXT NOT NULL,           -- JSON list
        done_ids      TEXT NOT NULL DEFAULT '[]',
        audio_sec     REAL,                    -- total audio in the batch
        units_before  REAL,
        units_after   REAL,
        wall_sec      REAL,
        ok            INTEGER,
        error         TEXT,
        started_at    TEXT NOT NULL DEFAULT (datetime('now')),
        ended_at      TEXT
    );
    """,
    """
    -- billing rate while the session ran (0 on the free tier)
    ALTER TABLE transcribe_batches ADD COLUMN units_per_hour REAL;
    """,
    """
    -- the Canvas page an embed was found on; opened in the LectureLab browser before a
    -- Kaltura/Panopto download so the player's own login cookies get set
    ALTER TABLE canvas_items ADD COLUMN page_link TEXT;
    """,
    """
    -- Stage 3: slide decks, their pages, and which deck each lecture uses
    CREATE TABLE decks (
        id           INTEGER PRIMARY KEY AUTOINCREMENT,  -- never reuse: slide-images/<id>/ may linger
        course_id    INTEGER REFERENCES courses(id) ON DELETE CASCADE,
        key          TEXT NOT NULL,           -- canvas-file:<id> | drive:<id> | upload:<sha256> | ocr:<lecture id>
        source       TEXT NOT NULL,           -- canvas | drive | upload | ocr
        title        TEXT NOT NULL,
        url          TEXT,                    -- download handle: canvas-file:<id> or a Drive link
        found_at     TEXT,                    -- where on Canvas it was listed
        created_at   TEXT,                    -- the file's own date, used for matching
        path         TEXT,                    -- relative to DATA: Slides/<Course>/<file>
        pages        INTEGER,
        thumb_w      INTEGER,                 -- every page of a deck renders at this size (NULL: text only)
        thumb_h      INTEGER,
        extracted_at REAL,                    -- unix time; versions the image URLs
        status       TEXT NOT NULL DEFAULT 'new',  -- new (not downloaded) | ready | error
        error        TEXT,
        added_at     TEXT NOT NULL DEFAULT (datetime('now')),
        UNIQUE(course_id, key)
    );
    CREATE TABLE slides (
        deck_id  INTEGER NOT NULL REFERENCES decks(id) ON DELETE CASCADE,
        no       INTEGER NOT NULL,            -- 1-based page number (= segments.slide_no)
        title    TEXT,
        text     TEXT NOT NULL DEFAULT '',
        times    TEXT,                        -- video-frame decks: JSON [[start, end], ...] on screen
        PRIMARY KEY (deck_id, no)
    );
    ALTER TABLE lectures ADD COLUMN deck_id INTEGER REFERENCES decks(id) ON DELETE SET NULL;
    -- auto | manual (your pick; NULL deck = "no slides") | nothing (the video had no readable screens)
    ALTER TABLE lectures ADD COLUMN deck_how TEXT;
    ALTER TABLE lectures ADD COLUMN deck_score REAL;
    ALTER TABLE lectures ADD COLUMN slides_aligned_at TEXT;
    ALTER TABLE lectures ADD COLUMN slides_error TEXT;
    """,
    # 8: Whisper sometimes emits a lone "." for chatter it can't make out; such lines say nothing
    """
    DELETE FROM segments WHERE trim(text, ' .,;:?!-''"…') = '';
    """,
    """
    -- Stage 4: notes, flashcards, quiz, search, and a log of every claude call
    CREATE TABLE claude_calls (
        id           INTEGER PRIMARY KEY,
        at           TEXT NOT NULL DEFAULT (datetime('now')),
        lecture_id   INTEGER,
        kind         TEXT,                    -- outline | section | study | check
        model        TEXT,
        input_hash   TEXT,                    -- = the cache file's name
        cached       INTEGER NOT NULL,        -- 1: answered from the cache, no plan usage
        ms           INTEGER,
        in_tokens    INTEGER,
        cache_tokens INTEGER,
        out_tokens   INTEGER,
        error        TEXT
    );
    CREATE TABLE notes (
        lecture_id        INTEGER PRIMARY KEY REFERENCES lectures(id) ON DELETE CASCADE,
        outline           TEXT NOT NULL,      -- JSON: every section, off-topic ones too
        sections          TEXT NOT NULL,      -- JSON [{title, start, end, kind, slides, gist, markdown}]
        summary           TEXT NOT NULL,
        open_questions    TEXT NOT NULL DEFAULT '[]',
        issues            TEXT NOT NULL DEFAULT '[]',  -- what the check pass removed or flagged, and why
        removed           INTEGER NOT NULL DEFAULT 0,
        flagged           INTEGER NOT NULL DEFAULT 0,
        generator         TEXT NOT NULL,      -- claude | offline
        transcript_source TEXT,               -- what the notes were written from: a change means they're stale
        deck_id           INTEGER,
        slides_aligned_at TEXT,
        created_at        TEXT NOT NULL,
        exported_at       TEXT,
        export_paths      TEXT               -- JSON: the files written, so a renamed lecture leaves no old copy
    );
    CREATE TABLE cards (
        id            INTEGER PRIMARY KEY,
        lecture_id    INTEGER NOT NULL REFERENCES lectures(id) ON DELETE CASCADE,
        kind          TEXT NOT NULL,          -- cloze ({{c1::...}} in front) | qa
        front         TEXT NOT NULL,
        back          TEXT NOT NULL,
        t             REAL,                   -- where it's taught, seconds
        box           INTEGER NOT NULL DEFAULT 1,  -- Leitner box 1 to 5
        due           TEXT,                   -- YYYY-MM-DD; NULL = new, due now
        reviews       INTEGER NOT NULL DEFAULT 0,
        lapses        INTEGER NOT NULL DEFAULT 0,
        last_reviewed TEXT
    );
    CREATE INDEX cards_lecture ON cards(lecture_id);
    CREATE TABLE quiz (
        id          INTEGER PRIMARY KEY,
        lecture_id  INTEGER NOT NULL REFERENCES lectures(id) ON DELETE CASCADE,
        idx         INTEGER NOT NULL,
        kind        TEXT NOT NULL,            -- mcq | short
        question    TEXT NOT NULL,
        choices     TEXT NOT NULL DEFAULT '[]',
        answer      TEXT NOT NULL,
        explanation TEXT NOT NULL DEFAULT '',
        t           REAL
    );
    CREATE INDEX quiz_lecture ON quiz(lecture_id);
    CREATE TABLE quiz_attempts (
        id          INTEGER PRIMARY KEY,
        lecture_id  INTEGER NOT NULL REFERENCES lectures(id) ON DELETE CASCADE,
        answers     TEXT NOT NULL,            -- JSON {question id: {answer, correct}}
        score       INTEGER NOT NULL,
        total       INTEGER NOT NULL,
        at          TEXT NOT NULL DEFAULT (datetime('now'))
    );
    CREATE VIRTUAL TABLE search USING fts5(text, kind UNINDEXED, lecture_id UNINDEXED, t UNINDEXED,
                                           tokenize='porter unicode61');
    ALTER TABLE lectures ADD COLUMN notes_error TEXT;
    ALTER TABLE lectures ADD COLUMN transcribe_fails INTEGER NOT NULL DEFAULT 0;
    """,
]

# Lecture lifecycle. Stages add work for each step; the UI shows these as chips.
STATUSES = ["queued_transcribe", "transcribing", "transcribed", "aligned", "notes_queued",
            "notes_running", "notes_done", "exported", "duplicate", "error", "waiting_quota"]

# Starter vocab so Whisper and the note prompts know the jargon on day one. Edit in the UI.
SEED_COURSES = [
    ("COP4020", "Programming Language Concepts", """BNF
EBNF
context-free grammar
parse tree
abstract syntax tree
lexer
recursive descent
ambiguity
precedence
associativity
lambda calculus
beta reduction
alpha conversion
currying
closure
higher-order function
pattern matching
type inference
Hindley-Milner
polymorphism
static scoping
dynamic scoping
activation record
garbage collection
lazy evaluation
Haskell
OCaml
Racket
Prolog
unification"""),
    ("COP4533", "Algorithm Abstraction and Design", """asymptotic notation
Big-O
Big-Theta
Big-Omega
recurrence
Master theorem
divide and conquer
dynamic programming
memoization
greedy algorithm
Dijkstra
Bellman-Ford
Floyd-Warshall
Kruskal
Prim
minimum spanning tree
BFS
DFS
topological sort
Ford-Fulkerson
max flow
NP-complete
reduction
amortized analysis"""),
    ("CNT4007", "Computer Network Fundamentals", """TCP
UDP
IP
HTTP
DNS
DHCP
ARP
Ethernet
MAC address
subnet
CIDR
NAT
OSPF
BGP
congestion control
flow control
sliding window
Go-Back-N
selective repeat
checksum
socket
three-way handshake
Wireshark
throughput
bandwidth-delay product"""),
    ("CIS4930", "Special Topics", ""),
    ("PHI3681", "Ethics, Data and Technology", """utilitarianism
consequentialism
deontology
Kant
categorical imperative
virtue ethics
Aristotle
Mill
informed consent
privacy
surveillance
algorithmic bias
fairness
autonomy
GDPR"""),
    ("MML", "Math for Machine Learning", """eigenvalue
eigenvector
SVD
PCA
gradient
Jacobian
Hessian
chain rule
convex
Lagrange multiplier
gradient descent
determinant
rank
null space
orthogonal projection
inner product
norm
Bayes' theorem
expectation
variance
Gaussian
maximum likelihood
KL divergence"""),
]

_db_init_lock = threading.Lock()
_db_ready = False


def connect() -> sqlite3.Connection:
    con = sqlite3.connect(DB_PATH, timeout=30, check_same_thread=False)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys=ON")
    con.execute("PRAGMA journal_mode=WAL")
    return con


@contextmanager
def db():
    """Short-lived connection per unit of work; commits on success."""
    init_db()
    con = connect()
    try:
        yield con
        con.commit()
    finally:
        con.close()


def init_db() -> None:
    global _db_ready
    if _db_ready:
        return
    with _db_init_lock:
        if _db_ready:
            return
        DATA.mkdir(parents=True, exist_ok=True)
        # The server and a CLI command can start together; only one may migrate at a time,
        # and the version is read after taking the lock so the second sees the first's work.
        mlock = FileLock(DATA / "migrate.lock")
        for _ in range(600):
            if mlock.acquire():
                break
            time.sleep(0.05)
        else:
            raise RuntimeError("timed out waiting for another process to migrate the database")
        con = connect()
        try:
            version = con.execute("PRAGMA user_version").fetchone()[0]
            for n, sql in enumerate(MIGRATIONS, start=1):
                if n <= version:
                    continue
                con.executescript(f"BEGIN; {sql}; PRAGMA user_version={n}; COMMIT;")
                log.info("db migrated to version %s", n)
                if n == 1:  # seed once, so a course you delete later stays deleted
                    con.executemany("INSERT OR IGNORE INTO courses(code, name, vocabulary) VALUES (?,?,?)", SEED_COURSES)
                    con.commit()
        finally:
            con.close()
            mlock.release()
        _db_ready = True


def rows(cur) -> list[dict]:
    return [dict(r) for r in cur.fetchall()]


# ---------------------------------------------------------------- courses

def list_courses() -> list[dict]:
    with db() as con:
        return rows(con.execute("SELECT * FROM courses ORDER BY code"))


def update_course(course_id: int, name: str | None = None, vocabulary: str | None = None) -> dict:
    with db() as con:
        cur = con.execute("SELECT * FROM courses WHERE id=?", (course_id,)).fetchone()
        if not cur:
            raise KeyError(course_id)
        if name is not None:
            if not name.strip():
                raise ValueError("course name can't be empty")
            con.execute("UPDATE courses SET name=? WHERE id=?", (name.strip(), course_id))
        if vocabulary is not None:
            terms = [t.strip() for t in vocabulary.splitlines() if t.strip()]
            con.execute("UPDATE courses SET vocabulary=? WHERE id=?", ("\n".join(terms), course_id))
        return dict(con.execute("SELECT * FROM courses WHERE id=?", (course_id,)).fetchone())


def find_course(key: str | int | None) -> dict | None:
    """Course by id (int), or by code or name (str, case-insensitive). Inbox folder names go
    through here as strings, so a folder called '2024' is never mistaken for course id 2024."""
    if key in (None, ""):
        return None
    with db() as con:
        if isinstance(key, int):
            r = con.execute("SELECT * FROM courses WHERE id=?", (int(key),)).fetchone()
            if r:
                return dict(r)
        k = str(key).strip().lower()
        for c in rows(con.execute("SELECT * FROM courses")):
            if k in (c["code"].lower(), c["name"].lower()) or k.startswith(c["code"].lower() + " "):
                return c
    return None


def guess_course(*texts: str) -> dict | None:
    """A course code mentioned in a title/filename/URL, e.g. 'COP4020 Lecture 7.mp4'."""
    blob = " ".join(t for t in texts if t).upper().replace(" ", "")
    for c in list_courses():
        if c["code"].upper() in blob:
            return c
    return None


# ---------------------------------------------------------------- jobs

def enqueue(kind: str, payload: dict, lecture_id: int | None = None) -> int:
    with db() as con:
        return con.execute("INSERT INTO jobs(kind, payload, lecture_id) VALUES (?,?,?)",
                           (kind, json.dumps(payload), lecture_id)).lastrowid


def job_update(job_id: int, **fields) -> None:
    if not fields:
        return
    for k in ("result", "payload"):
        if k in fields and not isinstance(fields[k], str):
            fields[k] = json.dumps(fields[k])
    sets = ", ".join(f"{k}=?" for k in fields) + ", updated_at=datetime('now')"
    with db() as con:
        con.execute(f"UPDATE jobs SET {sets} WHERE id=?", (*fields.values(), job_id))


def get_job(job_id: int) -> dict | None:
    with db() as con:
        r = con.execute("SELECT * FROM jobs WHERE id=?", (job_id,)).fetchone()
    if not r:
        return None
    j = dict(r)
    for k in ("payload", "result"):
        j[k] = json.loads(j[k]) if j[k] else None
    return j


def claim_job(kinds: list[str]) -> dict | None:
    """Atomically take the oldest queued job of these kinds (safe across processes)."""
    marks = ",".join("?" * len(kinds))
    with db() as con:
        r = con.execute(
            f"""UPDATE jobs SET status='running', owner_pid=?, attempts=attempts+1, updated_at=datetime('now')
                WHERE id=(SELECT id FROM jobs WHERE status='queued' AND kind IN ({marks})
                          AND (retry_at IS NULL OR retry_at<=datetime('now')) ORDER BY id LIMIT 1)
                RETURNING id""", (os.getpid(), *kinds)).fetchone()
    return get_job(r["id"]) if r else None


def run_job_now(job_id: int, wait_retry: bool = True) -> dict:
    """Claim one specific queued job and run it here (CLI use). If the server's worker got to it
    first, wait for that run to finish instead. A job that put itself off until later (RetryLater)
    is waited out too, unless wait_retry=False: then it comes back still queued, with retry_at."""
    with db() as con:
        ok = con.execute("UPDATE jobs SET status='running', owner_pid=?, attempts=attempts+1 "
                         "WHERE id=? AND status='queued' RETURNING id", (os.getpid(), job_id)).fetchone()
    if ok:
        run_job(get_job(job_id))
    while (j := get_job(job_id))["status"] in ("queued", "running"):
        if not wait_retry and j["status"] == "queued" and j["retry_at"]:
            break
        time.sleep(1)
    return j


MAX_ATTEMPTS = 3


def requeue_orphaned_jobs() -> int:
    """Jobs left 'running' by a process that died go back to the queue, unless they've already
    been tried MAX_ATTEMPTS times (a job that kills its process would otherwise loop forever)."""
    with db() as con:
        stuck = rows(con.execute("SELECT id, owner_pid, attempts FROM jobs WHERE status='running'"))
        dead = [j for j in stuck if not j["owner_pid"] or not pid_alive(j["owner_pid"])]
        for j in dead:
            if j["attempts"] >= MAX_ATTEMPTS:
                con.execute("UPDATE jobs SET status='error', owner_pid=NULL, error=? WHERE id=?",
                            (f"gave up after {j['attempts']} attempts (the process running it died each time)", j["id"]))
            else:
                con.execute("UPDATE jobs SET status='queued', owner_pid=NULL WHERE id=?", (j["id"],))
    if dead:
        log.warning("recovered %s orphaned job(s): %s", len(dead), [j["id"] for j in dead])
    return len(dead)


JOB_HANDLERS: dict[str, callable] = {}  # kind -> fn(job) -> result dict


SECRET_KEYS = ("passcode",)  # scrubbed from a job's stored payload once it has run


def run_job(job: dict) -> None:
    scrubbed = {k: v for k, v in job["payload"].items() if k not in SECRET_KEYS}
    if scrubbed.get("url"):  # a Zoom link carries its passcode as ?pwd=
        scrubbed["url"] = clean_source_url(scrubbed["url"])
    try:
        result = JOB_HANDLERS[job["kind"]](job)
        job_update(job["id"], status="done", result=result or {}, error=None, progress=None, payload=scrubbed)
        log.info("job %s %s done: %s", job["id"], job["kind"], result)
    except RetryLater as e:  # the job re-queued itself with a retry time; leave that in place
        log.info("job %s %s deferred: %s", job["id"], job["kind"], e)
    except Exception as e:
        log.exception("job %s %s failed", job["id"], job["kind"])
        job_update(job["id"], status="error", error=str(e)[:4000], progress=None, payload=scrubbed)


class Worker(threading.Thread):
    """Background thread in the server: runs queued jobs of the given kinds, one at a time."""

    def __init__(self, kinds: list[str], name: str):
        super().__init__(daemon=True, name=name)
        self.kinds, self.wake = kinds, threading.Event()

    def run(self):
        while True:
            try:
                job = claim_job(self.kinds)
                if job:
                    run_job(job)
                    continue
            except Exception:
                log.exception("worker %s loop error", self.name)
            self.wake.wait(5)
            self.wake.clear()


# ---------------------------------------------------------------- ingest

class Duplicate(Exception):
    def __init__(self, lecture_id: int, how: str):
        super().__init__(f"duplicate of lecture {lecture_id} ({how})")
        self.lecture_id, self.how = lecture_id, how


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def normalize_audio(src: Path, dst: Path) -> None:
    """16 kHz mono FLAC: lossless, about half the size of WAV, plays in Chrome, decodes everywhere."""
    tmp = dst.with_suffix(".tmp.flac")
    run([find_tool("ffmpeg"), "-v", "error", "-y", "-i", src, "-vn", "-ac", "1", "-ar", "16000",
         "-sample_fmt", "s16", "-c:a", "flac", tmp])
    tmp.replace(dst)


def pcm_signature(flac: Path) -> tuple[str, "bytes", float]:
    """One pass over the decoded audio: sha256 of the PCM, per-second loudness envelope, duration.
    Raises if ffmpeg fails, so a truncated decode is never stored as a valid signature."""
    import numpy as np
    proc = subprocess.Popen([find_tool("ffmpeg"), "-nostdin", "-v", "error", "-i", str(flac), "-f", "s16le",
                             "-ac", "1", "-ar", "16000", "-"], stdin=subprocess.DEVNULL, stdout=subprocess.PIPE,
                            stderr=subprocess.PIPE, creationflags=NO_WINDOW)
    h, env, total = hashlib.sha256(), [], 0
    try:
        while True:
            chunk = proc.stdout.read(32000)  # one second of 16 kHz s16
            if not chunk:
                break
            h.update(chunk)
            total += len(chunk)
            x = np.frombuffer(chunk[: len(chunk) // 2 * 2], np.int16).astype(np.float32) / 32768
            env.append(float(np.sqrt(np.mean(x * x))) if x.size else 0.0)
        err = proc.stderr.read().decode("utf-8", "replace")
        if proc.wait(timeout=60) != 0:
            raise RuntimeError(f"ffmpeg could not decode {flac.name}: {err.strip()[-500:]}")
    finally:
        if proc.poll() is None:
            proc.kill()  # never leave ffmpeg holding the file open (Windows can't delete it then)
        proc.stdout.close()
        proc.stderr.close()
    loud = np.log10(np.array(env, np.float32) + 1e-4).astype(np.float16)
    return h.hexdigest(), loud.tobytes(), total / 32000


MAX_LAG = 30  # seconds: the same lecture from Zoom and Canvas is often trimmed differently


def envelope_match(a: bytes, b: bytes) -> float:
    """Best correlation of two per-second loudness envelopes over lags of up to +-MAX_LAG seconds."""
    import numpy as np
    x, y = np.frombuffer(a, np.float16).astype(np.float32), np.frombuffer(b, np.float16).astype(np.float32)
    best = 0.0
    for lag in range(-MAX_LAG, MAX_LAG + 1):
        xs, ys = (x[lag:], y) if lag >= 0 else (x, y[-lag:])
        n = min(len(xs), len(ys))
        if n < 60:
            continue
        xs, ys = xs[:n], ys[:n]
        if xs.std() < 1e-3 or ys.std() < 1e-3:
            continue  # silence matches everything; never call that a duplicate
        best = max(best, float(np.corrcoef(xs, ys)[0, 1]))
    return best


def find_duplicate(con, sha_source=None, sha_pcm=None, fingerprint=None, duration=None):
    """Exact matches consider every finished lecture, parked duplicates included (re-adding a
    parked file must not park it again). The fuzzy envelope match only considers real lectures."""
    done = "audio_path IS NOT NULL AND status != 'ingesting'"
    if sha_source:
        r = con.execute(f"SELECT id FROM lectures WHERE sha_source=? AND {done}", (sha_source,)).fetchone()
        if r:
            return r["id"], "identical file"
    if sha_pcm:
        r = con.execute(f"SELECT id FROM lectures WHERE sha_pcm=? AND {done}", (sha_pcm,)).fetchone()
        if r:
            return r["id"], "identical audio"
    if fingerprint and duration:
        tol = 2 * MAX_LAG + duration * 0.01
        for r in con.execute(f"SELECT id, fingerprint FROM lectures WHERE fingerprint IS NOT NULL AND {done}"
                             " AND status != 'duplicate' AND ABS(duration_sec - ?) <= ?", (duration, tol)):
            score = envelope_match(fingerprint, r["fingerprint"])
            if score >= 0.95:
                return r["id"], f"same recording, re-encoded (envelope r={score:.3f})"
    return None


def clean_source_url(url: str | None) -> str | None:
    """Drop secrets from a URL before storing or showing it (Zoom puts the passcode in ?pwd=)."""
    if not url or not url.startswith("http"):
        return url
    from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit
    u = urlsplit(url)
    q = [(k, v) for k, v in parse_qsl(u.query, keep_blank_values=True) if k.lower() not in ("pwd", "passcode", "token")]
    return urlunsplit(u._replace(query=urlencode(q)))


def add_source(con, lecture_id: int, kind: str, url: str | None) -> None:
    con.execute("INSERT OR IGNORE INTO sources(lecture_id, kind, url) VALUES (?,?,?)",
                (lecture_id, kind, clean_source_url(url)))


def parse_date(*candidates) -> str | None:
    """First usable date among yt-dlp upload_date (YYYYMMDD), ISO timestamps, epoch seconds and
    dates embedded in names such as Zoom's '2026-04-08 19.47.47 ...' folders.
    Full timestamps (ffprobe creation_time, Canvas created_at) are UTC: an evening class must
    land on its local date, not tomorrow's, so those are converted to local time first."""
    for c in candidates:
        if c in (None, ""):
            continue
        if isinstance(c, (int, float)):
            return datetime.fromtimestamp(c, timezone.utc).astimezone().date().isoformat()
        s = str(c).strip()
        if re.fullmatch(r"\d{4}-\d\d-\d\dT\d\d:\d\d(:\d\d(\.\d+)?)?(Z|[+-]\d\d:?\d\d)?", s):
            try:
                dt = datetime.fromisoformat(s.replace("Z", "+00:00"))
                if dt.tzinfo is None:
                    dt = dt.replace(tzinfo=timezone.utc)
                return dt.astimezone().date().isoformat()
            except ValueError:
                pass
        m = re.search(r"(20\d\d)[-_.]?(0[1-9]|1[0-2])[-_.]?(0[1-9]|[12]\d|3[01])", s)
        if m:
            try:
                return Date(int(m[1]), int(m[2]), int(m[3])).isoformat()
            except ValueError:
                continue
    return None


def _move_into_place(stage: Path, final: Path) -> None:
    """Rename the staging dir to media/<id>. Windows briefly locks fresh files (Defender, the
    indexer), so retry; and clear a leftover dir from a lecture whose delete couldn't finish."""
    if final.exists():
        log.warning("removing leftover %s before reuse", final)
        shutil.rmtree(final)
    for attempt in range(10):
        try:
            stage.rename(final)
            return
        except PermissionError:
            if attempt == 9:
                raise
            time.sleep(0.5 * (attempt + 1))


def ingest_file(path: Path, *, kind: str, url: str | None = None, course_id: int | None = None,
                title: str | None = None, date: str | None = None, meta: dict | None = None,
                keep_original: bool = True) -> dict:
    """Store one recording: dedupe, normalize to 16 kHz mono FLAC, keep the original video.

    Returns {"lecture_id", "duplicate": bool, "how", "parked_dup_of"}. Raises on unreadable media.
    The row is 'ingesting' until its files are in place and is removed if anything fails, so a
    failure never leaves a half-made lecture that blocks the retry.
    """
    path = Path(path)
    meta = meta or {}
    url = url or str(path)
    if date and not valid_date(date):
        raise ValueError(f"date must be YYYY-MM-DD, got {date!r}")
    if course_id is not None and not find_course(int(course_id)):
        raise ValueError(f"no course with id {course_id}")
    sha = sha256_file(path)
    with db() as con:
        hit = find_duplicate(con, sha_source=sha)
        if hit:
            add_source(con, hit[0], kind, url)
            log.info("ingest %s: %s of lecture %s", path.name, hit[1], hit[0])
            return {"lecture_id": hit[0], "duplicate": True, "how": hit[1]}

    probe = ffprobe(path)
    if not probe["has_audio"]:
        raise ValueError(f"{path.name} has no audio stream")
    stage = DATA / "media" / f"incoming-{uuid.uuid4().hex[:8]}"
    stage.mkdir(parents=True)
    lid = None
    try:
        normalize_audio(path, stage / "audio.flac")
        sha_pcm, fingerprint, duration = pcm_signature(stage / "audio.flac")
        title = (title or meta.get("title") or probe["title"] or path.stem).strip()
        date = date or parse_date(meta.get("upload_date"), meta.get("timestamp"), probe["creation_time"],
                                  path.name, path.parent.name) \
            or datetime.fromtimestamp(path.stat().st_mtime).date().isoformat()
        if course_id is None:
            c = guess_course(title, path.name, url)
            course_id = c["id"] if c else None
        with db() as con:
            # IMMEDIATE: the check and the insert are one step, so two processes ingesting the
            # same lecture at once can't both decide it's new.
            con.execute("BEGIN IMMEDIATE")
            hit = find_duplicate(con, sha_source=sha, sha_pcm=sha_pcm, fingerprint=fingerprint, duration=duration)
            if hit and hit[1] in ("identical file", "identical audio"):
                add_source(con, hit[0], kind, url)
                con.execute("UPDATE lectures SET sha_source=COALESCE(sha_source, ?) WHERE id=?", (sha, hit[0]))
                log.info("ingest %s: %s of lecture %s", path.name, hit[1], hit[0])
                return {"lecture_id": hit[0], "duplicate": True, "how": hit[1]}
            # A near match (re-encoded) still gets its own row, parked as 'duplicate', so a
            # false positive can be released with "Not a duplicate" without re-downloading.
            lid = con.execute(
                """INSERT INTO lectures(course_id, title, date, duration_sec, status, sha_source, sha_pcm,
                   fingerprint, dup_of) VALUES (?,?,?,?,'ingesting',?,?,?,?)""",
                (course_id, title, date, duration, sha, sha_pcm, fingerprint, hit[0] if hit else None)).lastrowid
            add_source(con, lid, kind, url)
        final = DATA / "media" / str(lid)
        _move_into_place(stage, final)
        video_rel = None
        if probe["has_video"] and keep_original:
            dst = final / f"original{path.suffix.lower()}"
            shutil.copy2(path, dst)
            video_rel = dst.relative_to(DATA).as_posix()
        with db() as con:
            con.execute("UPDATE lectures SET audio_path=?, video_path=?, status=?, updated_at=datetime('now') "
                        "WHERE id=?", ((final / "audio.flac").relative_to(DATA).as_posix(), video_rel,
                                       "duplicate" if hit else "queued_transcribe", lid))
        log.info("ingested lecture %s %r (%s, %.0fs, video=%s)%s", lid, title, kind, duration, bool(video_rel),
                 f" flagged duplicate of {hit[0]}: {hit[1]}" if hit else "")
        return {"lecture_id": lid, "duplicate": bool(hit), "how": hit[1] if hit else None,
                "parked_dup_of": hit[0] if hit else None}
    except BaseException:
        if lid is not None:  # roll back the half-made lecture so a retry starts clean
            with db() as con:
                con.execute("DELETE FROM lectures WHERE id=? AND status='ingesting'", (lid,))
            shutil.rmtree(DATA / "media" / str(lid), ignore_errors=True)
        raise
    finally:
        if stage.exists():
            shutil.rmtree(stage, ignore_errors=True)


def classify_url(url: str) -> str:
    from urllib.parse import urlsplit
    host = (urlsplit(url).hostname or "").lower()
    if host.endswith(("drive.google.com", "docs.google.com")):
        return "drive"
    if host.endswith(("zoom.us", "zoom.com")):
        return "zoom"
    if host.endswith("instructure.com"):
        return "canvas"
    return "url"


_DRIVE_FILE_ID = re.compile(r"(?:/file/(?:u/\d+/)?d/|[?&]id=)([\w-]{20,})")


def download_url(url: str, dest: Path, passcode: str | None = None, progress=None,
                 cookiefile: Path | None = None) -> list[tuple[Path, dict]]:
    """Fetch a URL into dest. Returns [(file, metadata)]. Errors carry the tool's own message."""
    dest.mkdir(parents=True, exist_ok=True)
    if classify_url(url) == "drive":
        import gdown
        if "/folders/" in url:
            files = gdown.download_folder(url, output=str(dest), quiet=True) or []
            got = [Path(f) for f in files if Path(f).suffix.lower() in MEDIA_EXTS]
        else:
            m = _DRIVE_FILE_ID.search(url)
            if not m:
                raise RuntimeError(f"Google Drive: can't find a file id in {url}")
            f = gdown.download(id=m[1], output=str(dest) + os.sep, quiet=True)
            got = [Path(f)] if f else []
        if not got:
            raise RuntimeError("Google Drive: nothing downloadable at that link. Is it shared as "
                               "'Anyone with the link'? Private files sync through the Drive inbox folder instead.")
        return [(f, {"title": f.stem}) for f in got]

    import yt_dlp
    opts = {
        "outtmpl": str(dest / "%(title).80B [%(id)s].%(ext)s"),
        # Up to 720p keeps slides legible for OCR without multi-GB downloads.
        "format": "bv*[height<=720]+ba/b[height<=720]/bv*+ba/b",
        "merge_output_format": "mp4",
        "ffmpeg_location": str(Path(find_tool("ffmpeg")).parent),
        # One link, one recording: a watch?v=X&list=Y link must not pull the whole playlist.
        "noplaylist": True, "playlistend": 1, "quiet": True, "no_warnings": True, "noprogress": True,
        "js_runtimes": {"node": {}},  # YouTube needs a JS runtime; Node is installed
        "retries": 5, "fragment_retries": 10,
    }
    if passcode:
        opts["videopassword"] = passcode
    if cookiefile:
        opts["cookiefile"] = str(cookiefile)
    if progress:
        last = [0.0]

        def hook(d):
            if d.get("status") != "downloading" or time.time() - last[0] < 2:
                return
            last[0] = time.time()
            total = d.get("total_bytes") or d.get("total_bytes_estimate")
            progress(f"downloading {d.get('downloaded_bytes', 0) / total:.0%}" if total else "downloading")
        opts["progress_hooks"] = [hook]
    with yt_dlp.YoutubeDL(opts) as ydl:
        try:
            info = ydl.extract_info(url, download=True)
        except yt_dlp.utils.DownloadError as e:
            msg = f"yt-dlp: {str(e).removeprefix('ERROR: ')}"
            if cookiefile and re.search(r"(?i)\b(401|403|login|sign in|private|permission)\b", msg):
                msg += ("\nThe player didn't accept the LectureLab browser's UF session. If Canvas shows"
                        " 'needs sign-in', click Connect Canvas; otherwise open the video once in that browser.")
            elif classify_url(url) == "zoom":
                msg += ("\nFor Zoom this usually means the link expired or was deleted, the passcode is missing or"
                        " wrong, or the host limited it to signed-in UF users. Open the link in a browser: if it"
                        " only plays after you sign in, download it there and drop the file in the inbox.")
            raise RuntimeError(msg) from None
    entries = info.get("entries") or [info]
    out = []
    for e in entries:
        if not e:
            continue
        files = [Path(r["filepath"]) for r in e.get("requested_downloads") or [] if r.get("filepath")]
        f = next((p for p in files if p.exists()), None)
        if f:
            out.append((f, {"title": e.get("title"), "upload_date": e.get("upload_date"),
                            "timestamp": e.get("timestamp"), "url": e.get("webpage_url") or url}))
    if not out:
        raise RuntimeError("yt-dlp finished but produced no media file")
    return out


_LOGIN_HOSTS = ("kaltura.com", "panopto.com", "instructure.com", "ufl.edu")


def needs_login_cookies(url: str) -> bool:
    """Links that may sit behind UF login: matched on the host, not anywhere in the URL."""
    from urllib.parse import urlsplit
    u = urlsplit(url)
    host = (u.hostname or "").lower()
    return any(host == h or host.endswith("." + h) for h in _LOGIN_HOSTS) or "/mediasite/" in u.path.lower()


def job_cookie_file(job_id: int) -> Path:
    """Per-job cookie file for yt-dlp; the job deletes it when it's done."""
    return DATA / "secrets" / f"cookies-job{job_id}.txt"


def _login_cookies_for(url: str, path: Path) -> Path | None:
    """A pasted Kaltura/Panopto/Mediasite link may need the UF session too, if Canvas is connected."""
    if not needs_login_cookies(url) or canvas_state().get("state") != "connected":
        return None
    try:
        with CanvasSession() as cs:
            return cs.export_cookies(path)
    except Exception as e:
        log.warning("could not export Canvas cookies: %s", e)
        return None


def job_ingest(job: dict) -> dict:
    """payload: {path | url, kind?, course_id?, title?, date?, passcode?, delete_after?}"""
    p = job["payload"]
    say = lambda msg: job_update(job["id"], progress=msg)  # noqa: E731
    results = []
    if p.get("url"):
        kind = p.get("kind") or classify_url(p["url"])
        tmp = DATA / "downloads" / f"job{job['id']}"
        cookies = job_cookie_file(job["id"])
        try:
            say("downloading")
            files = download_url(p["url"], tmp, passcode=p.get("passcode"), progress=say,
                                 cookiefile=_login_cookies_for(p["url"], cookies))
            for i, (f, meta) in enumerate(files, 1):
                say(f"processing {i}/{len(files)}: {f.name}")
                title = p.get("title") if len(files) == 1 else None
                results.append(ingest_file(f, kind=kind, url=meta.get("url") or p["url"], course_id=p.get("course_id"),
                                           title=title, date=p.get("date"), meta=meta))
        finally:
            cookies.unlink(missing_ok=True)
            shutil.rmtree(tmp, ignore_errors=True)
    else:
        say("processing")
        path = Path(p["path"])
        try:
            results.append(ingest_file(path, kind=p.get("kind", "file"), url=p.get("source_url"),
                                       course_id=p.get("course_id"), title=p.get("title"), date=p.get("date")))
        finally:
            if p.get("delete_after"):
                path.unlink(missing_ok=True)
    first = results[0] if results else {}
    if first.get("lecture_id"):
        with db() as con:
            con.execute("UPDATE jobs SET lecture_id=? WHERE id=?", (first["lecture_id"], job["id"]))
    return {"lectures": results}


JOB_HANDLERS["ingest"] = job_ingest


# ---------------------------------------------------------------- lectures

def list_lectures(course_id: int | None = None) -> list[dict]:
    q = """SELECT l.id, l.course_id, l.title, l.date, l.duration_sec, l.status, l.error, l.dup_of,
                  l.video_path IS NOT NULL AS has_video, l.transcript_source, c.code AS course_code
           FROM lectures l LEFT JOIN courses c ON c.id=l.course_id"""
    with db() as con:
        if course_id:
            return rows(con.execute(q + " WHERE l.course_id=? ORDER BY l.date DESC, l.id DESC", (course_id,)))
        return rows(con.execute(q + " ORDER BY l.date DESC, l.id DESC"))


def get_lecture(lecture_id: int) -> dict | None:
    with db() as con:
        r = con.execute("""SELECT l.*, c.code AS course_code, c.name AS course_name FROM lectures l
                           LEFT JOIN courses c ON c.id=l.course_id WHERE l.id=?""", (lecture_id,)).fetchone()
        if not r:
            return None
        lec = dict(r)
        lec.pop("fingerprint", None)
        lec["sources"] = rows(con.execute("SELECT kind, url, added_at FROM sources WHERE lecture_id=? ORDER BY id",
                                          (lecture_id,)))
        lec["segment_count"] = con.execute("SELECT COUNT(*) FROM segments WHERE lecture_id=?",
                                           (lecture_id,)).fetchone()[0]
    return lec


def valid_date(date: str) -> bool:
    """YYYY-MM-DD and a real day (no 2024-02-30)."""
    try:
        return bool(re.fullmatch(r"\d{4}-\d\d-\d\d", date)) and Date.fromisoformat(date) is not None
    except ValueError:
        return False


def update_lecture(lecture_id: int, **fields) -> dict:
    """Only the fields given change. course_id or date given as None clears it ("No course", no
    date); a title can't be cleared."""
    allowed = {k: v for k, v in fields.items() if k in ("title", "date", "course_id")
               and (v is not None or k != "title")}
    if "title" in allowed:
        allowed["title"] = " ".join(str(allowed["title"]).split())  # no stray spaces or line breaks
        if not allowed["title"]:
            raise ValueError("title can't be empty")
    if allowed.get("date") == "":
        allowed["date"] = None
    if allowed.get("date") and not valid_date(allowed["date"]):
        raise ValueError("date must be a real day as YYYY-MM-DD")
    with db() as con:
        if not con.execute("SELECT 1 FROM lectures WHERE id=?", (lecture_id,)).fetchone():
            raise KeyError(lecture_id)
        for k, v in allowed.items():
            con.execute(f"UPDATE lectures SET {k}=?, updated_at=datetime('now') WHERE id=?", (v, lecture_id))
    return get_lecture(lecture_id)


def release_duplicate(lecture_id: int) -> dict:
    """'Not a duplicate': a near-match flagged by the envelope check goes into the pipeline."""
    with db() as con:
        con.execute("UPDATE lectures SET status='queued_transcribe', dup_of=NULL, updated_at=datetime('now') "
                    "WHERE id=? AND status='duplicate'", (lecture_id,))
    return get_lecture(lecture_id)


def delete_lecture(lecture_id: int) -> None:
    with db() as con:
        exported = con.execute("SELECT export_paths FROM notes WHERE lecture_id=?", (lecture_id,)).fetchone()
        others = _others_exports(con, lecture_id)
        con.execute("DELETE FROM search WHERE lecture_id=?", (lecture_id,))
        con.execute("DELETE FROM lectures WHERE id=?", (lecture_id,))
    for p in json.loads((exported or [None])[0] or "[]"):
        if p.lower() in others:
            continue
        try:
            Path(p).unlink(missing_ok=True)
        except OSError as e:
            log.warning("couldn't remove the export %s: %s", p, e)
    with db() as con:
        ocr = con.execute("DELETE FROM decks WHERE key=? RETURNING id", (f"ocr:{lecture_id}",)).fetchone()
    shutil.rmtree(DATA / "media" / str(lecture_id), ignore_errors=True)
    if ocr:  # its slides were read off its own video
        shutil.rmtree(DATA / "slide-images" / str(ocr[0]), ignore_errors=True)
    log.info("deleted lecture %s", lecture_id)


# ---------------------------------------------------------------- inboxes

_drive_cache: tuple[float, Path | None] = (0.0, None)


def drive_inbox() -> Path | None:
    """My Drive/LectureLab/Inbox on whatever letter Drive for Desktop mounted.
    LECTURELAB_DRIVE_INBOX overrides it (a path, or "off"; tests use this).
    Cached for 5 minutes: probing dead drive letters can block for seconds."""
    global _drive_cache
    override = os.environ.get("LECTURELAB_DRIVE_INBOX")
    if override:
        return None if override == "off" else Path(override)
    if time.time() - _drive_cache[0] < 300:
        return _drive_cache[1]
    found = None
    for letter in "GHIJKLMNOPQRSTUVWXYZ":
        p = Path(f"{letter}:/My Drive")
        try:
            if p.is_dir():
                found = p / "LectureLab" / "Inbox"
                break
        except OSError:
            continue
    _drive_cache = (time.time(), found)
    return found


def inbox_roots() -> list[tuple[Path, str]]:
    roots = [(LOCAL_INBOX, "inbox")]
    d = drive_inbox()
    if d:
        roots.append((d, "drive"))
    return roots


_pending_sizes: dict[str, tuple[int, float]] = {}


def _still_being_written(f: Path) -> bool:
    """Windows refuses to open a file for writing while another process (a copy, Drive sync)
    still has it open for writing; size+mtime alone miss a copy that stalls between scans."""
    try:
        with open(f, "ab"):
            return False
    except OSError:
        return True


def scan_inboxes(require_stable: bool = True) -> list[int]:
    """Queue new media in <inbox>/<Course>/. Files are never moved (Drive would fight that).

    A file is queued once its size is unchanged between two scans and nothing is still writing
    it. A file whose ingest failed is retried on later scans, up to MAX_ATTEMPTS times.
    Returns the new job ids.
    """
    jobs = []
    with db() as con:
        seen = {r["path"]: dict(r) for r in con.execute(
            "SELECT s.*, j.status AS job_status FROM seen_files s LEFT JOIN jobs j ON j.id=s.job_id")}
    for root, kind in inbox_roots():
        try:
            root.mkdir(parents=True, exist_ok=True)
            for c in list_courses():  # a folder per course, so there's an obvious place to drop files
                (root / c["code"]).mkdir(exist_ok=True)
            files = [f for f in root.rglob("*") if f.is_file() and f.suffix.lower() in MEDIA_EXTS]
        except OSError as e:
            log.warning("inbox %s unreadable: %s", root, e)
            continue
        for f in files:
            try:
                st = f.stat()
            except OSError:
                continue
            key = str(f.resolve())
            prev = seen.get(key)
            unchanged = prev and prev["size"] == st.st_size and abs(prev["mtime"] - st.st_mtime) < 1
            if unchanged and not (prev["job_status"] == "error" and prev["tries"] < MAX_ATTEMPTS):
                continue
            if require_stable and not unchanged and _pending_sizes.get(key) != (st.st_size, st.st_mtime):
                _pending_sizes[key] = (st.st_size, st.st_mtime)
                continue
            if _still_being_written(f):
                continue
            _pending_sizes.pop(key, None)
            rel = f.relative_to(root).parts
            course = find_course(rel[0]) if len(rel) > 1 else None
            jid = enqueue("ingest", {"path": key, "kind": kind, "course_id": course["id"] if course else None})
            tries = prev["tries"] + 1 if unchanged else 1
            with db() as con:
                con.execute("INSERT OR REPLACE INTO seen_files(path, size, mtime, job_id, tries) VALUES (?,?,?,?,?)",
                            (key, st.st_size, st.st_mtime, jid, tries))
            log.info("inbox: queued %s (course %s, job %s, try %s)", f, course and course["code"], jid, tries)
            jobs.append(jid)
    return jobs


class InboxWatcher(threading.Thread):
    def __init__(self, interval: float = 30, on_new=None):
        super().__init__(daemon=True, name="inbox")
        self.interval, self.on_new = interval, on_new

    def run(self):
        while True:
            try:
                if scan_inboxes() and self.on_new:
                    self.on_new()
            except Exception:
                log.exception("inbox scan failed")
            time.sleep(self.interval)


# ---------------------------------------------------------------- Canvas

class CanvasError(RuntimeError):
    pass


class CanvasNeedsSignIn(CanvasError):
    pass


# UF students can't make API tokens, so LectureLab keeps its own browser profile: you sign in
# there once with GatorLink + Duo, and LectureLab reads the Canvas API with that session the way
# Canvas's own web pages do (GET only). No password is ever asked for or stored.
CANVAS_BASE = os.environ.get("CANVAS_BASE_URL", "https://ufl.instructure.com").rstrip("/")
PROFILE_DIR = DATA / "browser-profile"
# Canvas and UF login use session-only cookies, which Chromium drops when the browser closes.
# They're kept here between launches so one sign-in lasts until Canvas itself expires it.
SESSION_COOKIES = DATA / "secrets" / "canvas-session.json"
_canvas_local = threading.local()


def canvas_state() -> dict:
    return meta_get("canvas_state", {"state": "not_connected"})


def set_canvas_state(state: str, **info) -> None:
    prev = canvas_state().get("state")
    meta_set("canvas_state", {"state": state, "checked_at": datetime.now().isoformat(timespec="seconds"), **info})
    if state == "connected" and prev != "connected":  # deck downloads waiting for you to sign in: go now
        with db() as con:
            con.execute("UPDATE jobs SET retry_at=NULL WHERE kind='deck_pull' AND status='queued'")
    if state == "needs_signin" and prev != "needs_signin":
        log.warning("Canvas session expired: needs sign-in (inbox folders keep working)")
        notify("LectureLab: Canvas needs sign-in",
               "Your Canvas session expired. Open LectureLab and click Connect Canvas. Inbox folders still work.")


def notify(title: str, body: str) -> None:
    """Best-effort Windows toast. Off in tests (LECTURELAB_NOTIFY=off)."""
    if os.environ.get("LECTURELAB_NOTIFY") == "off" or os.name != "nt":
        return
    esc_ = lambda t: t.replace("&", "&amp;").replace("<", "&lt;").replace("'", "''")  # noqa: E731
    script = f"""
[Windows.UI.Notifications.ToastNotificationManager, Windows.UI.Notifications, ContentType = WindowsRuntime] | Out-Null
[Windows.Data.Xml.Dom.XmlDocument, Windows.Data.Xml.Dom.XmlDocument, ContentType = WindowsRuntime] | Out-Null
$x = New-Object Windows.Data.Xml.Dom.XmlDocument
$x.LoadXml('<toast><visual><binding template="ToastGeneric"><text>{esc_(title)}</text><text>{esc_(body)}</text></binding></visual></toast>')
$app = '{{1AC14E77-02E7-4E5D-B744-2EB1AE5198B7}}\\WindowsPowerShell\\v1.0\\powershell.exe'
[Windows.UI.Notifications.ToastNotificationManager]::CreateToastNotifier($app).Show([Windows.UI.Notifications.ToastNotification]::new($x))
"""
    try:
        subprocess.Popen(["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
                         creationflags=NO_WINDOW, stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                         stderr=subprocess.DEVNULL)
    except Exception as e:
        log.warning("notification failed: %s", e)


def _prepare_profile() -> None:
    """Turn off the profile's password manager so it never offers to save the UF password."""
    prefs = PROFILE_DIR / "Default" / "Preferences"
    prefs.parent.mkdir(parents=True, exist_ok=True)
    try:
        data = json.loads(prefs.read_text(encoding="utf-8")) if prefs.exists() else {}
    except ValueError:
        data = {}
    data["credentials_enable_service"] = False
    data.setdefault("profile", {})["password_manager_enabled"] = False
    prefs.write_text(json.dumps(data), encoding="utf-8")


class CanvasSession:
    """The LectureLab browser profile, open headless, used to read Canvas. One at a time
    (a profile can't be opened twice); nested `with` blocks in a thread reuse the open one."""

    def __init__(self, headless: bool = True):
        self.headless = headless
        self.lock = FileLock(DATA / "browser.lock")
        self.nested = False

    def __enter__(self):
        if getattr(_canvas_local, "session", None):
            self.nested = True
            return _canvas_local.session
        for _ in range(240):
            if self.lock.acquire():
                break
            time.sleep(0.5)
        else:
            raise CanvasError("the LectureLab browser is busy (a sign-in window may be open)")
        self.pw = self.ctx = None
        try:
            from playwright.sync_api import sync_playwright
            _prepare_profile()
            self.pw = sync_playwright().start()
            self.ctx = self.pw.chromium.launch_persistent_context(
                str(PROFILE_DIR), headless=self.headless, accept_downloads=False,
                args=["--disable-features=PasswordManagerOnboarding,AutofillServerCommunication"])
            self.restore_session_cookies()
        except Exception:
            self._shutdown(save=False)  # a half-open browser would keep the profile locked
            raise
        _canvas_local.session = self
        return self

    def __exit__(self, *exc):
        if self.nested:
            return
        _canvas_local.session = None
        self._shutdown(save=True)

    def _shutdown(self, save: bool) -> None:
        """Close the browser and Playwright and free the lock, each step even if an earlier one fails."""
        try:
            if save and self.ctx:
                try:
                    self.save_session_cookies()
                except Exception as e:
                    log.warning("could not save the Canvas session cookies: %s", e)
            for step in (self.ctx and self.ctx.close, self.pw and self.pw.stop):
                if step:
                    try:
                        step()
                    except Exception as e:
                        log.warning("closing the LectureLab browser: %s", e)
        finally:
            self.lock.release()

    def get(self, path: str, params: dict | None = None, paginate: bool = True, quiet_auth: bool = False):
        """GET /api/v1/<path> with the session, following Link: rel=next pagination.
        quiet_auth: a signed-out answer raises without marking Canvas 'needs sign-in' (used
        while you're signing in, so it doesn't notify you mid-login)."""
        from urllib.parse import urlencode
        if path.startswith("http"):
            if not path.startswith(CANVAS_BASE + "/api/v1/"):
                raise CanvasError(f"refusing a non-API Canvas URL: {path}")
            url = path
        else:
            url = f"{CANVAS_BASE}/api/v1/{path.lstrip('/')}?" + urlencode({"per_page": 100, **(params or {})}, doseq=True)
        out = []
        while url:
            for attempt in range(5):
                r = self.ctx.request.get(url, headers={"Accept": "application/json+canvas-string-ids, application/json"},
                                         max_redirects=0, timeout=30_000)
                # Canvas also answers 403 when throttling; that must not look like "no items".
                if r.status == 403 and "rate limit" in r.text().lower():
                    time.sleep(2 ** attempt)
                    continue
                break
            else:
                raise CanvasError("Canvas kept rate-limiting; try again in a few minutes.")
            if r.status in (403, 404):
                return [] if paginate else None  # tab hidden by the instructor, or no access: not an error
            # Signed out: the API answers 401, or bounces to the SSO login page. Anything else that
            # isn't JSON (a 5xx, a maintenance page) is an outage, not an expired session: it must
            # not flip the sticky "needs sign-in" state.
            html = "json" not in (r.headers.get("content-type") or "")
            if r.status == 401 or 300 <= r.status < 400 or (
                    r.ok and html and re.search(r"(?i)log\s*in|sign\s*in|/login|login\.ufl\.edu", r.text())):
                if not quiet_auth:
                    set_canvas_state("needs_signin", reason=f"HTTP {r.status} from {url.split('?')[0]}")
                raise CanvasNeedsSignIn("Canvas needs sign-in: click Connect Canvas in LectureLab")
            if not r.ok or html:
                raise CanvasError(f"Canvas API answered {r.status}{' (not JSON)' if html else ''} for "
                                  f"{url.split('?')[0]}; Canvas may be down, try again later")
            data = r.json()
            if not paginate or not isinstance(data, list):
                return data
            out.extend(data)
            url = _next_link(r.headers.get("link"))
        return out

    def whoami(self, quiet_auth: bool = False) -> dict:
        me = self.get("users/self", paginate=False, quiet_auth=quiet_auth) or {}
        set_canvas_state("connected", user=me.get("name"))
        return me

    def warm(self, url: str) -> None:
        """Open a Canvas page so embedded players (Kaltura, Panopto) run their LTI sign-in and
        set their own cookies in this profile."""
        page = self.ctx.new_page()
        try:
            page.goto(url, wait_until="networkidle", timeout=60_000)
            page.wait_for_timeout(3000)
        except Exception as e:
            log.warning("warming %s: %s", url, e)
        finally:
            page.close()

    def save_session_cookies(self, path: Path = SESSION_COOKIES) -> None:
        keep = [c for c in self.ctx.cookies() if c.get("expires", -1) < 0]
        path.parent.mkdir(parents=True, exist_ok=True)
        tmp = path.with_suffix(".tmp")
        tmp.write_text(json.dumps(keep), encoding="utf-8")
        os.replace(tmp, path)

    def restore_session_cookies(self, path: Path = SESSION_COOKIES) -> int:
        """Add back saved session cookies the profile lost on close (never overwrites newer ones)."""
        try:
            saved = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            return 0
        have = {(c["name"], c["domain"], c["path"]) for c in self.ctx.cookies()}
        add = [c for c in saved if (c["name"], c["domain"], c["path"]) not in have]
        if add:
            self.ctx.add_cookies(add)
        return len(add)

    def export_cookies(self, path: Path) -> Path:
        """Write the profile's cookies as a Netscape cookies.txt for yt-dlp --cookies. Callers use a
        per-job path (job_cookie_file) and delete it afterwards: no long-lived cookie file."""
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(netscape_cookies(self.ctx.cookies()), encoding="utf-8")
        return path

    def httpx_cookies(self):
        import httpx
        jar = httpx.Cookies()
        for c in self.ctx.cookies():
            jar.set(c["name"], c["value"], domain=c["domain"], path=c["path"])
        return jar


def _next_link(header: str | None) -> str | None:
    for part in (header or "").split(","):
        m = re.match(r'\s*<([^>]+)>;\s*rel="next"', part)
        if m:
            return m[1]
    return None


def netscape_cookies(cookies: list[dict]) -> str:
    lines = ["# Netscape HTTP Cookie File", "# Written by LectureLab from its browser profile. Do not share."]
    for c in cookies:
        domain = c["domain"]
        lines.append("\t".join([domain, "TRUE" if domain.startswith(".") else "FALSE", c.get("path") or "/",
                                "TRUE" if c.get("secure") else "FALSE",
                                str(int(c["expires"])) if c.get("expires", -1) and c.get("expires", -1) > 0 else "0",
                                c["name"], c["value"]]))
    return "\n".join(lines) + "\n"


def canvas_get(path: str, params: dict | None = None, paginate: bool = True):
    """GET a Canvas API path with the LectureLab profile's session (reuses an open one)."""
    if not getattr(_canvas_local, "session", None):
        canvas_ready()  # not connected or needs sign-in: say so without opening a browser
    with CanvasSession() as cs:
        return cs.get(path, params, paginate)


def canvas_login(timeout_min: float = 10) -> dict:
    """Open the LectureLab browser so you can sign in with GatorLink + Duo yourself. Returns once
    the Canvas API answers as you, or after timeout_min. Never fills in anything."""
    with CanvasSession(headless=False) as cs:
        page = cs.ctx.pages[0] if cs.ctx.pages else cs.ctx.new_page()
        page.goto(CANVAS_BASE + "/", wait_until="domcontentloaded")
        end = time.time() + timeout_min * 60
        while time.time() < end:
            try:
                me = cs.get("users/self", paginate=False, quiet_auth=True)
                if me and me.get("id"):
                    set_canvas_state("connected", user=me.get("name"))
                    log.info("Canvas connected as %s", me.get("name"))
                    page.wait_for_timeout(1500)
                    return me
            except CanvasNeedsSignIn:
                pass
            if not cs.ctx.pages:
                break  # you closed the window
            time.sleep(2)
        set_canvas_state("needs_signin" if canvas_state().get("state") != "not_connected" else "not_connected",
                         reason="sign-in window closed or timed out")
        raise CanvasError("Canvas sign-in didn't finish (window closed or timed out)")


def canvas_courses() -> list[dict]:
    """Current Canvas courses, each matched to a LectureLab course by course code."""
    mine = list_courses()
    out = []
    for c in canvas_get("courses", {"enrollment_state": "active", "include[]": "term"}):
        if not c.get("name"):
            continue
        blob = f"{c.get('course_code', '')} {c.get('name', '')}".upper().replace(" ", "")
        match = next((m for m in mine if m["code"].upper() in blob), None)
        out.append({"canvas_id": c["id"], "name": c["name"], "code": c.get("course_code"),
                    "term": (c.get("term") or {}).get("name"), "course_id": match and match["id"],
                    "course_code": match and match["code"]})
        if match and match["canvas_id"] != c["id"]:
            with db() as con:
                con.execute("UPDATE courses SET canvas_id=? WHERE id=?", (c["id"], match["id"]))
    return out


# Embedded players we know how to fetch. Each regex captures enough to build a stable key.
_EMBEDS = [
    ("kaltura", re.compile(r"(?:entry_?id[/=]|entryid%2F|/entryId/)(\d_[a-z0-9]{8})", re.I)),
    ("panopto", re.compile(r"https?://[\w.-]*panopto\.com/Panopto/Pages/(?:Viewer|Embed)\.aspx\?(?:[^\"'<> ]*&)?id=([0-9a-f-]{36})", re.I)),
    ("mediasite", re.compile(r"https?://[\w.-]+/Mediasite/Play/([0-9a-f]{32,34})", re.I)),
    ("youtube", re.compile(r"(?:youtube(?:-nocookie)?\.com/(?:embed/|watch\?v=)|youtu\.be/)([\w-]{11})", re.I)),
    ("zoom", re.compile(r"https?://[\w.-]*zoom\.us/rec/(?:share|play)/[\w.\-~%]+", re.I)),
    ("media_object", re.compile(r"media_objects(?:_iframe)?/(m-[\w]+)", re.I)),
    # Some courses keep recordings in a shared Google Drive folder linked from a module.
    ("drive_folder", re.compile(r"drive\.google\.com/drive/(?:u/\d+/)?folders/([\w-]{20,})", re.I)),
    ("drive", re.compile(r"drive\.google\.com/(?:file/(?:u/\d+/)?d/|open\?id=|uc\?(?:[^\"'<> ]*&)?id=)([\w-]{20,})", re.I)),
]
_KALTURA_PARTNER = re.compile(r"(\d{6,8})\.kaf\.kaltura\.com|/p/(\d{6,8})/|partner_?id[/=](\d{6,8})", re.I)


def scan_html_for_media(html: str, page_url: str) -> list[dict]:
    from html import unescape
    from urllib.parse import unquote
    text = unquote(unescape(html or ""))
    found = []
    for kind, rx in _EMBEDS:
        for m in rx.finditer(text):
            if kind == "kaltura":
                pm = _KALTURA_PARTNER.search(text)
                partner = next((g for g in pm.groups() if g), None) if pm else None
                # yt-dlp's Kaltura extractor takes kaltura:<partner>:<entry> directly. Without a
                # partner id there's no reliable URL, so mark it rather than guess a host.
                url = f"kaltura:{partner}:{m[1]}" if partner else f"kaltura-unresolved:{m[1]}"
                found.append({"key": f"kaltura:{m[1]}", "kind": kind, "url": url})
            elif kind == "panopto":
                found.append({"key": f"panopto:{m[1]}", "kind": kind, "url": m[0].replace("Embed.aspx", "Viewer.aspx")})
            elif kind == "youtube":
                found.append({"key": f"youtube:{m[1]}", "kind": kind, "url": f"https://www.youtube.com/watch?v={m[1]}"})
            elif kind == "media_object":
                found.append({"key": f"media_object:{m[1]}", "kind": kind, "url": f"media_object:{m[1]}"})
            elif kind == "drive_folder":
                found.append({"key": f"drive_folder:{m[1]}", "kind": kind,
                              "url": f"https://drive.google.com/drive/folders/{m[1]}"})
            elif kind == "drive":
                found.append({"key": f"drive:{m[1]}", "kind": kind, "url": f"https://drive.google.com/file/d/{m[1]}/view"})
            else:
                found.append({"key": f"{kind}:{m[0] if kind == 'zoom' else m[1]}", "kind": kind, "url": m[0]})
    for f in found:
        f["page_url"] = page_url
    return found


def drive_list_folder(folder_id: str) -> list[dict]:
    """The files in a public Drive folder (recursively), listed without downloading anything."""
    import gdown
    with tempfile.TemporaryDirectory() as tmp:
        files = gdown.download_folder(id=folder_id, output=tmp, skip_download=True, quiet=True) or []
    return [{"id": f.id, "path": f.path} for f in files]


def expand_drive_folders(items: dict[str, dict], decks: dict[str, dict] | None = None) -> None:
    """Replace each linked Drive folder with one item per recording in it, so recordings are pulled
    one at a time (a semester folder is many GB). A folder that can't be listed stays as it is.
    Slide decks in the folder go into `decks` when it's given."""
    for key, it in list(items.items()):
        if it["kind"] != "drive_folder":
            continue
        try:
            listing = drive_list_folder(key.split(":", 1)[1])
        except Exception as e:
            log.warning("couldn't list the Drive folder linked at %s: %s", it.get("page_url"), e)
            continue
        del items[key]
        for f in listing:
            if Path(f["path"]).suffix.lower() in MEDIA_EXTS:
                items.setdefault(f"drive:{f['id']}", {
                    "key": f"drive:{f['id']}", "kind": "drive", "title": Path(f["path"]).stem,
                    "url": f"https://drive.google.com/file/d/{f['id']}/view",
                    "page_url": f"{it.get('page_url')} > {f['path']}", "page_link": it.get("page_link")})
            elif decks is not None and Path(f["path"]).suffix.lower() in SLIDE_EXTS:
                decks.setdefault(f"drive:{f['id']}", {
                    "key": f"drive:{f['id']}", "source": "drive", "title": Path(f["path"]).stem,
                    "url": f"https://drive.google.com/file/d/{f['id']}/view",
                    "found_at": f"{it.get('page_url')} > {f['path']}"})


def canvas_scan_course(course: dict, progress=None) -> dict:
    """Find recordings and slide decks in a course's Files, Modules, Pages and Announcements.
    Records new ones; new decks are downloaded right away (they're small)."""
    canvas_ready()  # not connected or needs sign-in: no browser, no "expired" toast
    with CanvasSession():
        return _canvas_scan_course(course, progress)


def _canvas_scan_course(course: dict, progress=None) -> dict:
    from urllib.parse import quote
    cid = course["canvas_id"]
    say = progress or (lambda m: None)
    items: dict[str, dict] = {}
    decks: dict[str, dict] = {}

    def add(it):
        items.setdefault(it["key"], it)

    def add_file(f, where):
        mime = f.get("content-type") or f.get("mime_class") or ""
        name = f.get("display_name", "")
        if mime.startswith(("video/", "audio/")) or mime in ("video", "audio") or \
                Path(name).suffix.lower() in MEDIA_EXTS:
            add({"key": f"file:{f['id']}", "kind": "file", "title": Path(name).stem,
                 "url": f"canvas-file:{f['id']}", "page_url": where, "created_at": f.get("created_at")})
        elif (mime in SLIDE_MIMES or Path(name).suffix.lower() in SLIDE_EXTS) and \
                (f.get("size") or 0) <= MAX_DECK_MB << 20:
            decks.setdefault(f"canvas-file:{f['id']}", {
                "key": f"canvas-file:{f['id']}", "source": "canvas", "title": Path(name).stem,
                "url": f"canvas-file:{f['id']}", "found_at": where, "created_at": f.get("created_at")})

    say("files")
    for f in canvas_get(f"courses/{cid}/files", {"content_types[]": ["video", "audio"]}):
        add_file(f, "course files")
    for f in canvas_get(f"courses/{cid}/files", {"content_types[]": list(SLIDE_MIMES)}):
        add_file(f, "course files")
    say("modules")
    pages_to_read = set()
    for mod in canvas_get(f"courses/{cid}/modules", {"include[]": "items"}):
        for it in mod.get("items") or []:
            where = f"module: {mod.get('name')} > {it.get('title')}"
            if it.get("type") == "File" and it.get("content_id"):
                f = canvas_get(f"files/{it['content_id']}", paginate=False)
                if f:
                    add_file(f, where)
            elif it.get("type") in ("ExternalUrl", "ExternalTool"):
                for m in scan_html_for_media(it.get("external_url") or "", where):
                    m.setdefault("title", it.get("title"))
                    m["page_link"] = it.get("html_url")
                    add(m)
            elif it.get("type") == "Page" and it.get("page_url"):
                pages_to_read.add(it["page_url"])
    say("pages")
    for p in canvas_get(f"courses/{cid}/pages"):
        pages_to_read.add(p["url"])
    for slug in sorted(pages_to_read):
        p = canvas_get(f"courses/{cid}/pages/{quote(slug, safe='')}", paginate=False)  # a slug must not escape the path
        if p:
            for m in scan_html_for_media(p.get("body") or "", f"page: {p.get('title')}"):
                m.setdefault("title", p.get("title"))
                m.setdefault("created_at", p.get("created_at"))
                m["page_link"] = p.get("html_url") or f"{CANVAS_BASE}/courses/{cid}/pages/{slug}"
                add(m)
    say("announcements")
    for a in canvas_get(f"courses/{cid}/discussion_topics", {"only_announcements": "true"}):
        for m in scan_html_for_media(a.get("message") or "", f"announcement: {a.get('title')}"):
            m.setdefault("title", a.get("title"))
            m.setdefault("created_at", a.get("posted_at"))
            m["page_link"] = a.get("html_url")
            add(m)
    say("drive folders")
    expand_drive_folders(items, decks)
    new = 0
    with db() as con:
        for it in items.values():
            cur = con.execute("""INSERT OR IGNORE INTO canvas_items(key, course_id, kind, title, url, page_url,
                                 created_at, page_link) VALUES (?,?,?,?,?,?,?,?)""",
                              (it["key"], course["id"], it["kind"], it.get("title"), it["url"], it.get("page_url"),
                               it.get("created_at"), it.get("page_link")))
            new += cur.rowcount
    with db() as con:
        lecs = rows(con.execute("SELECT title, date FROM lectures WHERE course_id=?", (course["id"],)))
    skipped = [d["title"] for d in decks.values() if not looks_like_lecture_deck(d["title"], lecs)]
    decks = {k: d for k, d in decks.items() if looks_like_lecture_deck(d["title"], lecs)}
    decks_new = record_decks(course["id"], decks)
    slides_scanned(course["id"])
    log.info("canvas scan %s: %s items (%s new), %s decks (%s new); not lecture slides: %s", course["code"],
             len(items), new, len(decks), decks_new, skipped)
    return {"found": len(items), "new": new, "decks": len(decks), "decks_new": decks_new}


def canvas_file_source(file_id: int) -> tuple[str, str, dict]:
    """(download URL, file name, metadata) for a Canvas File. Needs the open session."""
    f = canvas_get(f"files/{file_id}", paginate=False)
    if not f:
        raise CanvasError(f"Canvas file {file_id} is not accessible")
    return f["url"], Path(f["display_name"]).name, {"title": Path(f["display_name"]).stem,
                                                    "upload_date": f.get("created_at")}


def canvas_media_source(mid: str) -> tuple[str, str, dict]:
    """Canvas-native media (Canvas Studio/media comments): the API lists direct source URLs."""
    mo = canvas_get(f"media_objects/{mid}", paginate=False) or {}
    sources = sorted(mo.get("media_sources") or [], key=lambda s: int(s.get("bitrate") or 0))
    if not sources:
        raise CanvasError(f"Canvas media {mid}: no downloadable sources")
    src = sources[-1]
    return src["url"], f"{mid}.{src.get('fileExt') or 'mp4'}", {"title": mo.get("title") or mid,
                                                              "upload_date": mo.get("created_at")}


def stream_download(url: str, out: Path, jar=None, max_bytes: int | None = None) -> Path:
    """Raises ValueError as soon as the file passes max_bytes."""
    import httpx
    out.parent.mkdir(parents=True, exist_ok=True)
    with httpx.stream("GET", url, follow_redirects=True, timeout=120, cookies=jar) as r:
        r.raise_for_status()
        size = 0
        with open(out, "wb") as fh:
            for chunk in r.iter_bytes(1 << 20):
                size += len(chunk)
                if max_bytes and size > max_bytes:
                    raise ValueError(f"download is over {max_bytes >> 20} MB")
                fh.write(chunk)
    return out


def canvas_source_url(it: dict) -> str:
    """A real link for the lecture's source, not the internal downloader handle."""
    base = CANVAS_BASE
    course = find_course(it["course_id"]) or {}
    if it["kind"] == "file":
        return f"{base}/courses/{course.get('canvas_id')}/files/{it['url'].split(':', 1)[1]}"
    if it["kind"] == "media_object":
        return f"{base}/media_objects/{it['url'].split(':', 1)[1]}"
    return it["url"]


def job_canvas_pull(job: dict) -> dict:
    """payload: {course_id, key}. Download one Canvas item and ingest it."""
    cid, key = job["payload"]["course_id"], job["payload"]["key"]
    with db() as con:
        it = dict(con.execute("SELECT * FROM canvas_items WHERE course_id=? AND key=?", (cid, key)).fetchone())
    tmp = DATA / "downloads" / f"job{job['id']}"
    cookie_file = job_cookie_file(job["id"])
    try:
        job_update(job["id"], progress=f"downloading {it['title'] or key}")
        if it["url"].startswith("kaltura-unresolved:"):
            raise CanvasError("This Kaltura embed doesn't say which Kaltura account it belongs to, so it can't be "
                              f"fetched automatically. Open it from {it['page_url']}, copy the video's share link "
                              "and add that under Add recording > Link.")
        # Only Canvas-hosted items need the browser, and only to look things up and copy its
        # cookies: it's closed again before the (possibly multi-GB) download, so Connect Canvas
        # and scans aren't blocked meanwhile.
        source = cookies = None
        if it["kind"] in ("file", "media_object", "kaltura", "panopto", "mediasite"):
            canvas_ready()
            with CanvasSession() as cs:
                ref = it["url"].split(":", 1)[1]
                if it["kind"] == "file":
                    source = (*canvas_file_source(int(ref)), cs.httpx_cookies())
                elif it["kind"] == "media_object":
                    source = (*canvas_media_source(ref), cs.httpx_cookies())
                else:
                    # Videos behind UF login: let the player sign in inside our profile, then hand
                    # yt-dlp that profile's cookies (Chrome's own cookies can't be read on Windows).
                    if it.get("page_link"):
                        job_update(job["id"], progress="opening the Canvas page to sign in to the player")
                        cs.warm(it["page_link"])
                    cookies = cs.export_cookies(cookie_file)
        if source:
            url, name, meta, jar = source
            # Canvas names can hold ? : " | that Windows refuses (or, for ':', turns into a hidden stream)
            files = [(stream_download(url, tmp / safe_filename(name, "recording"), jar), meta)]
        else:
            files = download_url(it["url"], tmp, cookiefile=cookies,
                                 progress=lambda m: job_update(job["id"], progress=m))
        results = []
        for f, meta in files:
            job_update(job["id"], progress=f"processing {f.name}")
            meta["upload_date"] = meta.get("upload_date") or it["created_at"]
            results.append(ingest_file(f, kind="canvas", url=canvas_source_url(it), course_id=it["course_id"],
                                       title=it["title"] or meta.get("title"), meta=meta))
        with db() as con:
            con.execute("UPDATE canvas_items SET status='pulled', lecture_id=?, error=NULL WHERE course_id=? AND key=?",
                        (results[0]["lecture_id"], cid, key))
        return {"lectures": results}
    except Exception as e:
        with db() as con:
            con.execute("UPDATE canvas_items SET status='error', error=? WHERE course_id=? AND key=?",
                        (str(e)[:2000], cid, key))
        raise
    finally:
        cookie_file.unlink(missing_ok=True)
        shutil.rmtree(tmp, ignore_errors=True)


def canvas_pull(course_id: int, keys: list[str]) -> list[int]:
    jobs = []
    for k in keys:
        with db() as con:
            n = con.execute("UPDATE canvas_items SET status='queued' WHERE course_id=? AND key=? "
                            "AND status IN ('new','error','ignored')", (course_id, k)).rowcount
            title = con.execute("SELECT title FROM canvas_items WHERE course_id=? AND key=?", (course_id, k)).fetchone()
        if n:
            # The title is only for the jobs list; the job reads the item itself.
            jobs.append(enqueue("canvas_pull", {"course_id": course_id, "key": k, "title": title and title[0]}))
    return jobs


def canvas_auto_pull() -> list[int]:
    """For courses with auto_pull on: scan, then pull everything new. Used by sync."""
    jobs = []
    for c in list_courses():
        if c["auto_pull"] and c["canvas_id"]:
            canvas_scan_course(c)
            with db() as con:
                keys = [r["key"] for r in con.execute("SELECT key FROM canvas_items WHERE course_id=? AND status='new'",
                                                      (c["id"],))]
            jobs += canvas_pull(c["id"], keys)
    return jobs


def job_canvas_scan(job: dict) -> dict:
    course = find_course(job["payload"]["course_id"])
    if not course["canvas_id"]:
        canvas_courses()  # links LectureLab courses to Canvas ids by course code
        course = find_course(course["id"])
    if not course["canvas_id"]:
        raise CanvasError(f"No current Canvas course matches {course['code']}")
    return canvas_scan_course(course, progress=lambda m: job_update(job["id"], progress=f"scanning {m}"))


def canvas_ready() -> None:
    """Raise a clear error unless the LectureLab browser has a Canvas session."""
    st = canvas_state().get("state")
    if st == "not_connected":
        raise CanvasError("Canvas isn't connected yet: click Connect Canvas and sign in with GatorLink.")
    if st == "needs_signin":
        raise CanvasNeedsSignIn("Canvas needs sign-in: click Connect Canvas. Inbox folders keep working meanwhile.")


JOB_HANDLERS["canvas_pull"] = job_canvas_pull
JOB_HANDLERS["canvas_scan"] = job_canvas_scan


# ---------------------------------------------------------------- transcription

WSL_DISTRO = os.environ.get("LECTURELAB_WSL_DISTRO", "Ubuntu")  # the default distro is docker-desktop
COLAB_SCRIPT = ROOT / "colab" / "transcribe.py"
COLAB_REMOTE = "/content/lecturelab"
# Runs colab/transcribe.py as a child process on the VM. A native crash (e.g. CTranslate2 not
# finding cuDNN) would otherwise kill the kernel silently; this way its stderr and exit code
# come back in the exec output. The pip cuDNN/cuBLAS libraries go on LD_LIBRARY_PATH first.
COLAB_LAUNCHER = f"""
import glob, os, subprocess, sys
env = dict(os.environ)
libs = sorted({{d for site in sys.path for pat in ("nvidia/cudnn/lib", "nvidia/cublas/lib")
               for d in glob.glob(os.path.join(site, pat))}})
env["LD_LIBRARY_PATH"] = ":".join(libs + [env.get("LD_LIBRARY_PATH", "")])
p = subprocess.Popen([sys.executable, "-u", "transcribe.py"], cwd={COLAB_REMOTE!r}, env=env,
                     stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True)
for line in p.stdout:
    print(line, end="", flush=True)
print(f"[worker exit {{p.wait()}}]", flush=True)
"""
COLAB_GPUS = ["T4", "L4"]  # T4 is cheapest; L4 when no T4 is free
MIN_UNITS = float(os.environ.get("LECTURELAB_MIN_UNITS", 20))
SESSION_PREFIX = "lecturelab-"
LOCAL_MODELS = ["large-v3-turbo", "distil-large-v3"]
# Tests run with LECTURELAB_COLAB=off: a test server must never touch the real Colab account
# (its orphan-session sweep would stop a real batch).
COLAB_ENABLED = os.environ.get("LECTURELAB_COLAB", "on") != "off"


class LowUnits(Exception):
    pass


def meta_get(key: str, default=None):
    with db() as con:
        r = con.execute("SELECT value FROM meta WHERE key=?", (key,)).fetchone()
    return json.loads(r["value"]) if r else default


def meta_set(key: str, value) -> None:
    with db() as con:
        con.execute("INSERT OR REPLACE INTO meta(key, value) VALUES (?,?)", (key, json.dumps(value)))


def initial_prompt(lecture: dict) -> str:
    """Whisper's initial_prompt: course name, title, the course vocabulary and the matched slide
    text, so jargon, names and function names are spelled right. Whisper only reads the last
    ~224 tokens of a prompt, so it's capped at ~800 characters, vocabulary first."""
    course = find_course(lecture["course_id"]) if lecture.get("course_id") else None
    parts = [f"{course['code']} {course['name']} lecture: {lecture['title']}." if course
             else f"Lecture: {lecture['title']}."]
    vocab = [t for t in (course or {}).get("vocabulary", "").splitlines() if t.strip()]
    if vocab:
        parts.append("Terms: " + ", ".join(vocab) + ".")
    slides = lecture_slide_terms(lecture["id"])
    if slides:
        parts.append("Slides: " + slides)
    text = " ".join(parts)
    return text[:800].rsplit(" ", 1)[0] if len(text) > 800 else text


def store_transcript(lecture_id: int, result: dict, source: str) -> int:
    """Save a transcript (Colab or local JSON shape) and mark the lecture transcribed."""
    if not get_lecture(lecture_id):  # deleted while it was being transcribed: nothing to store it on
        log.info("lecture %s was deleted mid-transcription; its transcript is dropped", lecture_id)
        return 0
    path = DATA / "transcripts" / f"{lecture_id}.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(result), encoding="utf-8")
    # Whisper sometimes emits a lone "." for chatter it can't make out: nothing to read there.
    segs = [s for s in result["segments"] if re.search(r"\w", s["text"])]
    with db() as con:
        con.execute("DELETE FROM segments WHERE lecture_id=?", (lecture_id,))
        con.executemany("INSERT INTO segments(lecture_id, idx, start, end, text, words) VALUES (?,?,?,?,?,?)",
                        [(lecture_id, i, s["start"], s["end"], s["text"], json.dumps(s.get("words") or []))
                         for i, s in enumerate(segs)])
        con.execute("UPDATE lectures SET status='transcribed', transcript_source=?, error=NULL, transcribe_fails=0, "
                    "slides_aligned_at=NULL, updated_at=datetime('now') WHERE id=?", (source, lecture_id))
    log.info("lecture %s transcribed by %s: %s segments", lecture_id, source, len(segs))
    index_lecture(lecture_id)
    enqueue_slides(lecture_id)  # (re)align: the old slide numbers went with the old segments
    return len(segs)


def load_audio_array(path: Path):
    """16 kHz float32 via ffmpeg (not PyAV: some PyAV releases break faster-whisper's decoder)."""
    import numpy as np
    raw = run([find_tool("ffmpeg"), "-nostdin", "-v", "error", "-i", path, "-f", "f32le", "-ac", "1",
               "-ar", "16000", "-"]).stdout
    return np.frombuffer(raw, np.float32)


_local_models: dict = {}


def transcribe_local(lecture: dict, progress=None, model_name: str | None = None) -> dict:
    """CPU fallback: faster-whisper int8 with the model the benchmark picked."""
    from faster_whisper import WhisperModel
    name = model_name or os.environ.get("LECTURELAB_LOCAL_MODEL") or meta_get("local_model", "large-v3-turbo")
    if name not in _local_models:
        _local_models.clear()  # one model in RAM at a time; they're ~1.5 GB each
        _local_models[name] = WhisperModel(name, device="cpu", compute_type="int8", cpu_threads=os.cpu_count())
    model = _local_models[name]
    audio = load_audio_array(DATA / lecture["audio_path"])
    t = time.time()
    segments, info = model.transcribe(audio, language="en", beam_size=5, vad_filter=True,
                                      vad_parameters={"min_silence_duration_ms": 500}, word_timestamps=True,
                                      initial_prompt=initial_prompt(lecture) or None)
    segs, last = [], 0.0
    for s in segments:
        segs.append({"start": round(s.start, 2), "end": round(s.end, 2), "text": s.text.strip(),
                     "words": [[round(w.start, 2), round(w.end, 2), w.word, round(w.probability, 3)]
                               for w in (s.words or [])]})
        if progress and time.time() - last > 3:
            last = time.time()
            progress(f"transcribing locally ({name}): {s.end / max(info.duration, 1):.0%}")
    return {"lecture_id": lecture["id"], "model": name, "device": "cpu-int8", "language": info.language,
            "duration": round(info.duration, 2), "elapsed_sec": round(time.time() - t, 1), "segments": segs}


# --- Colab through the official CLI, inside WSL

def wsl(*args, timeout: float = 600) -> subprocess.CompletedProcess:
    return run(["wsl.exe", "-d", WSL_DISTRO, "-e", *args], timeout=timeout, text=True, encoding="utf-8",
               errors="replace")


_colab_bin: str | None = None


def colab_bin() -> str:
    global _colab_bin
    if not _colab_bin:
        out = wsl("sh", "-c", 'command -v colab || echo "$HOME/.local/bin/colab"', timeout=60).stdout.strip()
        _colab_bin = out.splitlines()[-1]
    return _colab_bin


def colab(*args, timeout: float = 900) -> str:
    """Run one colab CLI command in WSL. stdin is closed, so if the CLI was never logged in it
    fails fast with its OAuth prompt in the error instead of hanging."""
    return wsl(colab_bin(), *args, timeout=timeout).stdout


UPLOAD_CHUNK_MB = int(os.environ.get("LECTURELAB_UPLOAD_CHUNK_MB", 8))


def colab_exec_code(session: str, code: str, workdir: Path, name: str, timeout: float = 600) -> str:
    """Run a small generated script on the VM (`colab exec -f`)."""
    f = workdir / name
    f.write_text(code, encoding="utf-8")
    return colab_exec(session, wslpath(f), timeout)


def colab_exec(session: str, remote_file: str, timeout: float) -> str:
    """`colab exec` gives up after 30 s by default and still exits 0, so always pass --timeout
    (a little under ours, so the CLI returns its own output before we kill it)."""
    return colab("exec", "-s", session, "-f", remote_file, "--timeout", str(int(max(timeout - 60, 30))),
                 timeout=timeout)


def colab_upload(session: str, local: Path, remote: str, workdir: Path) -> None:
    """Upload through the CLI. Its contents API rejects big files (an 86 MB FLAC fails with an SSL
    EOF), so anything over UPLOAD_CHUNK_MB goes up in parts and is joined and checksummed on the VM."""
    size, chunk = local.stat().st_size, UPLOAD_CHUNK_MB << 20
    if size <= chunk:
        colab("upload", "-s", session, wslpath(local), remote, timeout=900)
        return
    parts_dir = workdir / "parts"
    parts_dir.mkdir(parents=True, exist_ok=True)
    digest, parts = hashlib.sha256(), []
    try:
        with open(local, "rb") as fh:
            while block := fh.read(chunk):
                digest.update(block)
                part = parts_dir / f"{local.parent.name}-{len(parts):03d}.part"
                parts.append(part)
                part.write_bytes(block)
        for i, part in enumerate(parts):
            for attempt in (1, 2):  # one retry: the contents API sometimes drops a connection (SSL EOF)
                try:
                    colab("upload", "-s", session, wslpath(part), f"{remote}.{i:03d}.part", timeout=900)
                    break
                except RuntimeError as e:
                    if attempt == 2:
                        raise
                    log.warning("upload of part %s of %s failed, retrying: %s", i, local.name, str(e)[-200:])
    finally:
        for part in parts:
            part.unlink(missing_ok=True)
    out = colab_exec_code(session, f"""
import hashlib, os
remote, n, want = {remote!r}, {len(parts)}, {digest.hexdigest()!r}
h = hashlib.sha256()
with open(remote, "wb") as out:
    for i in range(n):
        p = f"{{remote}}.{{i:03d}}.part"
        with open(p, "rb") as f:
            b = f.read()
        h.update(b); out.write(b); os.remove(p)
print("JOINED_OK" if h.hexdigest() == want else "JOINED_BAD " + h.hexdigest())
""", workdir, f"join-{local.parent.name}.py")
    if "JOINED_OK" not in out:
        raise RuntimeError(f"upload of {local.name} was corrupted on the way to Colab: {out.strip()[-300:]}")


def wslpath(p: Path) -> str:
    return wsl("wslpath", "-a", str(Path(p).resolve()), timeout=60).stdout.strip()


def parse_colab_units(out: str) -> float | None:
    """`colab usage` prints "Current balance: 0.00 compute units". Anything else is None (unknown),
    never a guess: a wrong number here could spend paid units."""
    m = re.search(r"(?im)^\s*current balance:\s*([\d,]+(?:\.\d+)?)\s*compute units\s*$", out)
    return float(m[1].replace(",", "")) if m else None


def parse_colab_rate(out: str) -> float | None:
    m = re.search(r"(?im)^\s*usage rate:\s*([\d.]+)\s*/\s*hr\s*$", out)
    return float(m[1]) if m else None


def colab_units() -> float | None:
    """Remaining compute units from `colab usage`, or None if the output can't be read."""
    out = colab("usage", timeout=120)
    meta_set("colab_usage_raw", out[-2000:])
    return parse_colab_units(out)


def colab_rate() -> float | None:
    """Units per hour the running session is billed at ("Usage rate: 1.07/hr")."""
    return parse_colab_rate(colab("usage", timeout=120))


def colab_sessions() -> list[str]:
    out = colab("sessions", timeout=120)
    return sorted(set(re.findall(rf"\b({re.escape(SESSION_PREFIX)}\d+)\b", out)))


def stop_orphan_sessions(keep: str | None = None) -> list[str]:
    """An idle runtime keeps burning compute units. If a batch process died before its finally
    ran, its session is still up: stop every LectureLab session that isn't the current one."""
    stopped = []
    try:
        for s in colab_sessions():
            if s != keep:
                try:
                    colab("stop", "-s", s, timeout=180)
                    stopped.append(s)
                except Exception as e:
                    log.warning("could not stop orphan Colab session %s: %s", s, e)
    except Exception as e:
        log.warning("could not list Colab sessions: %s", e)
    if stopped:
        log.warning("stopped orphaned Colab session(s): %s", stopped)
    return stopped


def sweep_orphans_if_idle() -> list[str]:
    """Stop leftover sessions only when no batch is running anywhere (we can take the transcribe
    lock). Otherwise a server restart would kill the session another process is using."""
    lk = FileLock(DATA / "transcribe.lock")
    if not lk.acquire():
        return []
    try:
        return stop_orphan_sessions()
    finally:
        lk.release()


def colab_new_session(name: str) -> str:
    errors = []
    for gpu in COLAB_GPUS:
        try:
            colab("new", "-s", name, "--gpu", gpu, timeout=900)
            return gpu
        except Exception as e:
            errors.append(f"{gpu}: {e}")
            log.warning("colab new --gpu %s failed: %s", gpu, e)
    raise RuntimeError("no Colab GPU available: " + " | ".join(errors))


def transcribe_batch_colab(lectures: list[dict], batch_id: int, progress=None, force: bool = False) -> list[int]:
    """One Colab session for the whole batch: new -> upload -> exec -> download -> stop (always).
    Each finished transcript is stored as soon as it's downloaded, so an interrupted batch keeps
    its finished lectures and the rest stay queued. force: you asked for Colab explicitly, so an
    unreadable balance doesn't stop it (a known low balance still does)."""
    say = progress or (lambda m: None)
    session = f"{SESSION_PREFIX}{batch_id}"
    stop_orphan_sessions()
    units_before = colab_units()
    # 0 units means the free tier: there's nothing to burn, so a free T4 is fine. A small paid
    # balance is protected instead: below MIN_UNITS the batch goes to the local CPU.
    if units_before is None and not force:
        raise LowUnits("couldn't read the Colab compute-unit balance from `colab usage`; not risking paid units")
    if units_before is not None and 0 < units_before < MIN_UNITS:
        raise LowUnits(f"only {units_before:g} Colab compute units left (minimum {MIN_UNITS:g})")
    with db() as con:
        con.execute("UPDATE transcribe_batches SET units_before=? WHERE id=?", (units_before, batch_id))
    manifest = {"model": "large-v3", "items": [
        {"lecture_id": l["id"], "audio": f"{l['id']}.flac", "initial_prompt": initial_prompt(l)} for l in lectures]}
    mpath = DATA / "colab_batches" / str(batch_id) / "manifest.json"
    mpath.parent.mkdir(parents=True, exist_ok=True)
    mpath.write_text(json.dumps(manifest, indent=1), encoding="utf-8")
    done: list[int] = []
    gpu = None

    def fetch_results(budget: float = 900):
        # A transcript is a few MB at most. The budget bounds how long a hung WSL/CLI can keep
        # the (billed) session alive before the stop in the finally below.
        end = time.time() + budget
        for l in lectures:
            if l["id"] in done:
                continue
            left = end - time.time()
            if left < 10:
                log.warning("batch %s: out of time downloading results; the rest stay queued", batch_id)
                return
            local = mpath.parent / f"{l['id']}.json"
            try:
                colab("download", "-s", session, f"{COLAB_REMOTE}/out/{l['id']}.json", wslpath(local),
                      timeout=min(180, left))
            except Exception:
                continue  # not finished; it stays queued
            if local.exists():
                try:
                    store_transcript(l["id"], json.loads(local.read_text(encoding="utf-8")), f"colab-large-v3 ({gpu})")
                except Exception as e:  # one bad file mustn't cost the others (or mask the real error)
                    log.error("batch %s: couldn't store the transcript of lecture %s: %s", batch_id, l["id"], e)
                    continue
                done.append(l["id"])
                with db() as con:
                    con.execute("UPDATE transcribe_batches SET done_ids=? WHERE id=?", (json.dumps(done), batch_id))

    try:
        say("starting a Colab GPU")
        gpu = colab_new_session(session)
        with db() as con:
            con.execute("UPDATE transcribe_batches SET gpu=? WHERE id=?", (gpu, batch_id))
        try:
            rate = colab_rate()
            log.info("Colab %s session %s: billed at %s units/hr", gpu, session, rate)
            with db() as con:
                con.execute("UPDATE transcribe_batches SET units_per_hour=? WHERE id=?", (rate, batch_id))
        except Exception as e:
            log.warning("could not read the Colab usage rate: %s", e)
        say(f"uploading to Colab ({gpu})")
        # The contents API won't create folders (a 500 if the parent is missing), so make them first.
        colab_exec_code(session, f"import os; os.makedirs({COLAB_REMOTE + '/out'!r}, exist_ok=True)",
                        mpath.parent, "prep.py")
        colab("upload", "-s", session, wslpath(mpath), f"{COLAB_REMOTE}/manifest.json", timeout=600)
        colab("upload", "-s", session, wslpath(COLAB_SCRIPT), f"{COLAB_REMOTE}/transcribe.py", timeout=600)
        t_up = time.time()
        for i, l in enumerate(lectures, 1):
            say(f"uploading audio {i}/{len(lectures)} to Colab ({gpu})")
            colab_upload(session, DATA / l["audio_path"], f"{COLAB_REMOTE}/{l['id']}.flac", mpath.parent)
        log.info("batch %s: uploaded %s file(s) in %.0fs", batch_id, len(lectures), time.time() - t_up)
        say(f"transcribing {len(lectures)} lecture(s) on Colab {gpu}")
        out = colab_exec_code(session, COLAB_LAUNCHER, mpath.parent, "run.py", timeout=4 * 3600)
        (mpath.parent / "exec.log").write_text(out, encoding="utf-8")
        log.info("colab exec output (batch %s):\n%s", batch_id, out[-3000:])
        say("downloading transcripts")
        fetch_results()
        if len(done) < len(lectures):
            # exec can end early without an error (timeout, kernel crash): that's a failed batch.
            tail = out.strip().splitlines()[-3:] or ["(no output)"]
            raise RuntimeError(f"Colab returned {len(done)} of {len(lectures)} transcript(s); "
                               f"last output: {' | '.join(tail)}")
    finally:
        try:
            if gpu:  # a session exists: keep whatever finished before a failure
                fetch_results(budget=300)
        finally:
            for attempt in (1, 2):
                try:
                    colab("stop", "-s", session, timeout=300)
                    break
                except Exception as e:
                    log.error("colab stop -s %s failed (attempt %s): %s%s", session, attempt, e,
                              " (will retry before the next batch)" if attempt == 2 else "")
    return done


def finish_batch(batch_id: int, t0: float, ok: bool, error: str | None = None, units_after: float | None = None):
    with db() as con:
        con.execute("UPDATE transcribe_batches SET ok=?, error=?, units_after=?, wall_sec=?, ended_at=datetime('now') "
                    "WHERE id=?", (int(ok), error, units_after, round(time.time() - t0, 1), batch_id))
        b = dict(con.execute("SELECT * FROM transcribe_batches WHERE id=?", (batch_id,)).fetchone())
    used = (b["units_before"] - units_after) if (b["units_before"] is not None and units_after is not None) else None
    per75 = used / (b["audio_sec"] / 4500) if used is not None and b["audio_sec"] else None
    log.info("transcribe batch %s (%s %s): ok=%s, %s lecture(s), %.1f min audio, %.1f min wall, units %s -> %s "
             "(used %s, %s per 75-min lecture)%s", batch_id, b["mode"], b["gpu"], ok, len(json.loads(b["done_ids"])),
             (b["audio_sec"] or 0) / 60, (time.time() - t0) / 60, b["units_before"], units_after,
             f"{used:.2f}" if used is not None else "?", f"{per75:.2f}" if per75 is not None else "?",
             f" error: {error}" if error else "")


def stay_awake(on: bool) -> bool:
    """Keep Windows from sleeping while a batch runs: in standby the network drops and the Colab VM
    with it, and the batch hangs until the laptop wakes. This laptop has Modern Standby, which sleeps
    when the screen turns off, so the screen has to stay on too. Closing the lid still sleeps.
    Per thread: call on and off from the same one. Never raises: a failure here mustn't strand the
    transcription lock."""
    if os.name != "nt":
        return False
    try:
        import ctypes
        ES_CONTINUOUS, ES_SYSTEM_REQUIRED, ES_DISPLAY_REQUIRED = 0x80000000, 0x1, 0x2
        flags = ES_CONTINUOUS | (ES_SYSTEM_REQUIRED | ES_DISPLAY_REQUIRED if on else 0)
        ok = bool(ctypes.windll.kernel32.SetThreadExecutionState(ctypes.c_uint(flags)))
    except Exception:
        log.exception("SetThreadExecutionState failed")
        ok = False
    if on and not ok:
        log.warning("could not keep the laptop awake; if it sleeps, this batch's Colab VM is lost")
    return ok


def job_transcribe(job: dict) -> dict:
    """payload: {lecture_ids?: [...], mode: auto|colab|local}. With no ids, takes every queued lecture.
    auto = Colab, unless units are low or Colab failed twice in a row; then local CPU."""
    mode = job["payload"].get("mode", "auto")
    if mode == "colab" and not COLAB_ENABLED:  # off means off, even when Colab is asked for by name
        raise RuntimeError("Colab is turned off on this server (LECTURELAB_COLAB=off); transcribe locally instead.")
    say = lambda m: job_update(job["id"], progress=m)  # noqa: E731
    lock = FileLock(DATA / "transcribe.lock")
    if not lock.acquire():
        # attempts=0: waiting for another batch isn't a failed try, and a long wait would
        # otherwise run up attempts until a later crash gets the job given up on
        job_update(job["id"], status="queued", retry_at=_sqlnow(120), owner_pid=None, attempts=0)
        raise RetryLater("another transcription batch is running")
    stay_awake(True)
    try:
        with db() as con:
            con.execute("BEGIN IMMEDIATE")
            ids = job["payload"].get("lecture_ids")
            if ids is not None:
                # Lectures you picked: a failed one ('error') may be retried on purpose.
                lectures = rows(con.execute(
                    "SELECT * FROM lectures WHERE status IN ('queued_transcribe','error') AND audio_path IS NOT NULL "
                    f"AND id IN ({','.join('?' * len(ids)) or 'NULL'}) ORDER BY id", ids))
            else:
                # "Everything waiting": only the queue, never lectures that already failed.
                lectures = rows(con.execute("SELECT * FROM lectures WHERE status='queued_transcribe' "
                                            "AND audio_path IS NOT NULL ORDER BY id"))
            missing = [l for l in lectures if not (DATA / l["audio_path"]).exists()]
            for l in missing:  # a local problem: don't send it anywhere, and don't blame Colab for it
                con.execute("UPDATE lectures SET status='error', error=? WHERE id=?",
                            (f"audio file missing: {l['audio_path']}", l["id"]))
            lectures = [l for l in lectures if l not in missing]
            if not lectures:
                return {"transcribed": [], "note": "nothing waiting"}
            for l in lectures:
                con.execute("UPDATE lectures SET status='transcribing', error=NULL WHERE id=?", (l["id"],))
            audio_sec = sum(l["duration_sec"] or 0 for l in lectures)
        use_local = mode == "local"
        reason = "requested" if use_local else None
        if mode == "auto" and not COLAB_ENABLED:
            use_local, reason = True, "Colab disabled (LECTURELAB_COLAB=off)"
        if mode == "auto" and colab_backed_off():
            use_local, reason = True, f"Colab failed twice in a row; trying it again after {COLAB_BACKOFF_H:g} h"
        done, t0 = [], time.time()
        batch_id = None
        try:
            if not use_local:
                with db() as con:
                    batch_id = con.execute("INSERT INTO transcribe_batches(mode, lecture_ids, audio_sec) VALUES "
                                           "('colab',?,?)", (json.dumps([l["id"] for l in lectures]), audio_sec)).lastrowid
                try:
                    done = transcribe_batch_colab(lectures, batch_id, progress=say, force=mode == "colab")
                    meta_set("colab_failures", 0)
                    finish_batch(batch_id, t0, ok=True, units_after=_safe_units())
                except LowUnits as e:
                    finish_batch(batch_id, t0, ok=False, error=str(e))
                    if mode == "colab":
                        raise
                    use_local, reason = True, str(e)
                except Exception as e:
                    meta_set("colab_failures", meta_get("colab_failures", 0) + 1)
                    meta_set("colab_failed_at", time.time())
                    finish_batch(batch_id, t0, ok=False, error=str(e)[:2000], units_after=_safe_units())
                    raise
            if use_local:
                todo = [l for l in lectures if l["id"] not in done]
                model = os.environ.get("LECTURELAB_LOCAL_MODEL") or meta_get("local_model", "large-v3-turbo")
                with db() as con:
                    batch_id = con.execute("INSERT INTO transcribe_batches(mode, gpu, lecture_ids, audio_sec) VALUES "
                                           "('local',?,?,?)", (model, json.dumps([l["id"] for l in todo]),
                                                               sum(l["duration_sec"] or 0 for l in todo))).lastrowid
                log.info("transcribing %s lecture(s) locally (%s): %s", len(todo), model, reason)
                t0 = time.time()
                try:
                    for l in todo:
                        res = transcribe_local(l, progress=say)
                        store_transcript(l["id"], res, f"local-{res['model']}")
                        done.append(l["id"])
                except Exception as e:
                    finish_batch(batch_id, t0, ok=False, error=str(e)[:2000])  # never leave it looking busy
                    raise
                finish_batch(batch_id, t0, ok=True)
            return {"transcribed": done, "mode": "local" if use_local else "colab", "reason": reason}
        finally:
            # Anything not finished goes back to the queue for the next run; nothing is redone.
            with db() as con:
                for l in lectures:
                    if l["id"] not in done:
                        con.execute("UPDATE lectures SET status='queued_transcribe', transcribe_fails=transcribe_fails+1 "
                                    "WHERE id=? AND status='transcribing'", (l["id"],))
    finally:
        stay_awake(False)
        lock.release()


COLAB_BACKOFF_H = float(os.environ.get("LECTURELAB_COLAB_BACKOFF_H", "6"))


def colab_backed_off() -> bool:
    """After 2 Colab failures in a row, auto mode uses the CPU for a while, then tries Colab again
    (a success resets the count). Without the time limit it would never try Colab again."""
    if meta_get("colab_failures", 0) < 2:
        return False
    return time.time() - meta_get("colab_failed_at", 0) < COLAB_BACKOFF_H * 3600


def _safe_units():
    try:
        return colab_units()
    except Exception as e:
        log.warning("colab usage failed: %s", e)
        return None


def _sqlnow(plus_sec: int = 0) -> str:
    return datetime.fromtimestamp(time.time() + plus_sec, timezone.utc).strftime("%Y-%m-%d %H:%M:%S")


class RetryLater(Exception):
    """Raised by a job that re-queued itself (it set its own status and retry_at)."""


def reset_stuck_transcribing() -> int:
    """Lectures left 'transcribing' by a process that died go back to the queue, and that
    process's batch row is closed (an open row makes the UI think a batch is still running)."""
    lk = FileLock(DATA / "transcribe.lock")
    if not lk.acquire():
        return 0  # a batch is genuinely running
    try:
        with db() as con:
            n = con.execute("UPDATE lectures SET status='queued_transcribe', transcribe_fails=transcribe_fails+1 "
                            "WHERE status='transcribing'").rowcount
            con.execute("UPDATE transcribe_batches SET ok=0, error='interrupted (the process stopped mid-batch)', "
                        "ended_at=datetime('now') WHERE ended_at IS NULL")
    finally:
        lk.release()
    if n:
        log.warning("requeued %s lecture(s) left mid-transcription", n)
    return n


JOB_HANDLERS["transcribe"] = job_transcribe


# ---------------------------------------------------------------- slides
SLIDE_EXTS = {".pdf", ".pptx", ".ppt"}
SLIDE_MIMES = ("application/pdf", "application/vnd.openxmlformats-officedocument.presentationml.presentation",
               "application/vnd.ms-powerpoint")
SLIDES_DIR = DATA / "Slides"          # decks as downloaded, per course: Slides/<Course>/<file>
SLIDE_IMAGES = DATA / "slide-images"  # page renders per deck (not "slides": Windows paths ignore case)
THUMB_W, FULL_W = 320, 1280
MAX_DECK_MB = int(os.environ.get("LECTURELAB_MAX_DECK_MB", 200))
WINDOW_SEC = 30                       # transcript window compared against slide text
MATCH_MIN, MATCH_MARGIN = 0.45, 0.1   # auto-match only a clear winner
OCR_MAX_SCREENS = 150
MAX_DECK_PAGES = 300                  # more is a textbook or reader, not a lecture's slides
OCR_MAX_PAGES = 60                    # per deck: pages past this that need OCR keep what text they have
OCR_THREADS = 2                       # of the laptop's cores, so OCR in the background doesn't hog it
DECK_RETRY_SEC, DECK_MAX_TRIES = 600, 5  # a download that failed for a passing reason: wait 10, 20, ... min


class DeckError(Exception):
    """A deck that can't be used. The message says what happened and how to fix it."""


def user_error(e: Exception, what: str) -> str:
    """What to show for a failure: our own messages as they are. Anything else (ffmpeg's stderr,
    a timeout's command line, an OS error) can carry local paths, so it gets a generic line and
    the detail goes to server.log."""
    if isinstance(e, (DeckError, CanvasError, ClaudeError, QuotaExceeded)):
        return str(e)
    log.warning("%s failed: %r", what, e)
    return f"{what} failed unexpectedly ({type(e).__name__}). The details are in server.log."


def get_deck(deck_id: int | None) -> dict | None:
    if not deck_id:
        return None
    with db() as con:
        r = con.execute("SELECT * FROM decks WHERE id=?", (deck_id,)).fetchone()
    return dict(r) if r else None


def safe_filename(name: str, fallback: str = "slides") -> str:
    """A file name that's safe on Windows: no path parts, reserved characters or device names."""
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", Path(name or "").name).strip(" .") or fallback
    stem, suffix = Path(name).stem[:120], Path(name).suffix.lower()[:8]
    if re.fullmatch(r"(?i)(con|prn|aux|nul|com\d|lpt\d)", stem.split(".")[0]):  # CON.tar.pdf is CON too
        stem = "_" + stem
    return stem + suffix


# Lecture codes in titles and file names, each optionally a range ("M3TL2-3", "Lectures 4-5"):
# the last group is the end of the range. Two titles with the same kind of code must overlap.
_CODES = [re.compile(p, re.I) for p in (
    r"\bm\s*0*(\d{1,2})\s*[-_ .]?\s*(t?)l\s*0*(\d{1,2})(?:\s*-\s*0*(\d{1,2}))?\b",  # M2L1, M2TL1 (a different
                                                                                   # kind of lecture), M3TL2-3
    r"\b(?:lectures?|lect|lec|l)\s*[-_#.]?\s*0*(\d{1,3})(?:\s*-\s*0*(\d{1,3}))?\b",  # Lecture 7, Lec07, L07
    r"\b(?:week|wk)\s*[-_#.]?\s*0*(\d{1,2})(?:\s*-\s*0*(\d{1,2}))?\b",
    r"\b(?:chapter|ch)\s*[-_#.]?\s*0*(\d{1,2})(?:\s*-\s*0*(\d{1,2}))?\b")]
_TITLE_STOP = frozenset("a an and are as at be by for from in into is it of on or the to with lecture lec "
                        "slides slide notes part pdf pptx ppt".split())


def title_codes(text: str) -> dict[int, frozenset]:
    """{kind: set of codes}. "M3TL2-3" is {("3", "t", 2), ("3", "t", 3)}: one deck, two lectures."""
    out = {}
    for i, rx in enumerate(_CODES):
        m = rx.search(text or "")
        if not m:
            continue
        *head, first, last = m.groups()
        first = int(first)
        last = int(last) if last and 0 < int(last) - first <= 5 else first
        out[i] = frozenset((*(h.lower() for h in head), n) for n in range(first, last + 1))
    return out


def deck_capacity(title: str) -> int:
    """How many lectures a deck can serve: one per lecture code its name covers."""
    codes = title_codes(title)
    return max([len(v) for k, v in codes.items() if k in (0, 1)] or [1])


def _title_words(text: str) -> set[str]:
    for rx in _CODES:
        text = rx.sub(" ", text or "")
    return {w for w in re.findall(r"[a-z0-9]+", text.lower()) if w not in _TITLE_STOP and len(w) > 1}


def match_score(deck: dict, lec: dict) -> float:
    """How likely this deck goes with this lecture, about 0 to 1: lecture codes in the names
    (agreeing or contradicting), title word overlap, and how close the dates are."""
    a, b = title_codes(deck["title"]), title_codes(lec["title"])
    shared = a.keys() & b.keys()
    code = (1.0 if all(a[k] & b[k] for k in shared) else -1.0) if shared else 0.0
    wa, wb = _title_words(deck["title"]), _title_words(lec["title"])
    words = len(wa & wb) / len(wa | wb) if wa and wb else 0.0
    near = 0.0
    try:
        days = abs((Date.fromisoformat((deck.get("created_at") or "")[:10]) -
                    Date.fromisoformat((lec.get("date") or "")[:10])).days)
        near = max(0.0, 1 - days / 14)
    except ValueError:
        pass
    return round(0.5 * code + 0.4 * words + 0.15 * near + (0.1 if words == 1 and code >= 0 else 0), 3)


def _twin_key(title: str) -> str:
    """Decks with the same name (the PDF and the PPTX of one lecture) are one candidate."""
    return re.sub(r"\s+", " ", re.sub(r"(?i)\.(pdf|pptx?)$", "", (title or "").strip())).lower()


def match_decks(course_id: int) -> list[int]:
    """Auto-match a course's decks to its lectures: one deck per lecture, best pairs first, and
    only clear winners (a near-tie is left for you to pick). A deck serves one lecture unless its
    name covers several codes ("M3TL2-3"). Your own picks are never touched and use up a deck's
    room. A match already made survives a new near-tie (a second copy of the deck, a review
    deck); only a clearly better pairing replaces it. Returns lectures whose deck changed."""
    with db() as con:
        lecs = rows(con.execute("SELECT id, title, date, deck_id, deck_how FROM lectures WHERE course_id=? "
                                "AND status NOT IN ('duplicate','ingesting')", (course_id,)))
        decks = rows(con.execute("SELECT id, title, created_at, source, path FROM decks WHERE course_id=? "
                                 "AND status!='error'", (course_id,)))
    manual: dict[int, int] = {}
    for l in lecs:
        if l["deck_how"] == "manual" and l["deck_id"]:
            manual[l["deck_id"]] = manual.get(l["deck_id"], 0) + 1
    auto = [l for l in lecs if l["deck_how"] != "manual"]
    current = {l["id"]: l["deck_id"] for l in auto if l["deck_id"]}
    in_use = set(current.values()) | set(manual)
    twins: dict[str, list[dict]] = {}
    for d in decks:
        if d["source"] != "ocr":
            twins.setdefault(_twin_key(d["title"]), []).append(d)
    # of each set of twins: the one already in use, else the PDF, else the oldest
    real = [min(g, key=lambda d: (d["id"] not in in_use, not (d["path"] or "").lower().endswith(".pdf"), d["id"]))
            for g in twins.values()]
    room = {d["id"]: deck_capacity(d["title"]) - manual.get(d["id"], 0) for d in real}
    real = [d for d in real if room[d["id"]] > 0]
    scores = {(l["id"], d["id"]): match_score(d, l) for l in auto for d in real}
    pick: dict[int, tuple[int, float]] = {}
    for (lid, did), s in sorted(scores.items(), key=lambda kv: -kv[1]):
        if s < MATCH_MIN:
            break
        if lid in pick or not room[did]:
            continue
        # Ambiguous: another free deck fits this lecture about as well, or this deck (with room
        # for just one more) fits another unmatched lecture about as well.
        rivals = [v for (l2, d2), v in scores.items() if l2 == lid and d2 != did and room[d2]]
        if room[did] == 1:
            rivals += [v for (l2, d2), v in scores.items() if d2 == did and l2 != lid and l2 not in pick]
        if rivals and max(rivals) > s - MATCH_MARGIN and current.get(lid) != did:
            continue
        pick[lid] = (did, s)
        room[did] -= 1
    ocr = {d["id"] for d in decks if d["source"] == "ocr"}
    changed = []
    # Each write only lands if the lecture is as read above: a pick you made meanwhile wins.
    still = "AND deck_id IS ? AND COALESCE(deck_how,'')!='manual'"
    with db() as con:
        for l in auto:
            new = pick.get(l["id"])
            if new and l["deck_id"] != new[0]:
                cur = con.execute(f"UPDATE lectures SET deck_id=?, deck_how='auto', deck_score=?, slides_aligned_at=NULL, "
                                  f"slides_error=NULL WHERE id=? {still}", (new[0], new[1], l["id"], l["deck_id"]))
            elif new:
                con.execute(f"UPDATE lectures SET deck_score=? WHERE id=? {still}", (new[1], l["id"], l["deck_id"]))
                continue
            elif l["deck_id"] and l["deck_id"] not in ocr:  # its old match lost to a better lecture
                cur = con.execute(f"UPDATE lectures SET deck_id=NULL, deck_how=NULL, deck_score=NULL, "
                                  f"slides_aligned_at=NULL, slides_error=NULL WHERE id=? {still}",
                                  (l["id"], l["deck_id"]))
            else:
                continue
            if cur.rowcount:
                con.execute("UPDATE segments SET slide_no=NULL WHERE lecture_id=?", (l["id"],))
                changed.append(l["id"])
    if changed:
        log.info("slides: course %s re-matched lectures %s", course_id, changed)
    return changed


def enqueue_slides(lecture_id: int, ocr: bool = False) -> int | None:
    """Queue a slides job (find the deck, read it, align it) unless one is already waiting.
    ocr=True (you asked) reads slides off the video even before the course's Canvas scan.
    Returns the job id, the waiting one's if there is one."""
    with db() as con:
        if q := con.execute("SELECT id FROM jobs WHERE kind='slides' AND lecture_id=? AND status='queued'",
                            (lecture_id,)).fetchone():
            if ocr:
                con.execute("UPDATE jobs SET payload=json_set(payload, '$.ocr', 1) WHERE id=?", (q[0],))
            return q[0]
        r = con.execute("SELECT d.title, d.source FROM lectures l LEFT JOIN decks d ON d.id=l.deck_id "
                        "WHERE l.id=?", (lecture_id,)).fetchone()
    title = ("video frames" if r["source"] == "ocr" else r["title"]) if r and r["title"] else "slides"
    jid = enqueue("slides", {"lecture_id": lecture_id, "title": title, **({"ocr": 1} if ocr else {})},
                  lecture_id=lecture_id)
    wake_workers()
    return jid


# --- reading decks
def render_fixed(page, w: int, h: int):
    """Render a PDF page into exactly w x h (letterboxed on white if its shape differs), so every
    thumbnail of a deck is the same size and the UI can reserve the space before it loads."""
    import pymupdf as fitz
    zoom = min(w / page.rect.width, h / page.rect.height)
    pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
    if (pix.width, pix.height) == (w, h):
        return pix
    out = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, w, h), False)
    out.clear_with(255)
    pix.set_origin((w - pix.width) // 2, (h - pix.height) // 2)
    out.copy(pix, pix.irect)
    return out


def page_title(page, text: str) -> str:
    """The line set in the largest type, else the first line."""
    best, size = "", 0.0
    for b in page.get_text("dict").get("blocks", []):
        for line in b.get("lines", []):
            t = "".join(s["text"] for s in line["spans"]).strip()
            sz = max((s["size"] for s in line["spans"]), default=0)
            if t and sz > size + 0.5:
                best, size = t, sz
    best = best or next((l.strip() for l in text.splitlines() if l.strip()), "")
    return best[:120]


_ocr_engine = None
_ocr_lock = threading.Lock()


def _ocr(path: Path, read: bool = True):
    global _ocr_engine
    with _ocr_lock:
        if _ocr_engine is None:
            from rapidocr import RapidOCR
            logging.getLogger("RapidOCR").setLevel(logging.WARNING)
            _ocr_engine = RapidOCR(params={"EngineConfig.onnxruntime.intra_op_num_threads": OCR_THREADS,
                                           "EngineConfig.onnxruntime.inter_op_num_threads": 1})
        return _ocr_engine(str(path), use_cls=None if read else False, use_rec=read)


def ocr_image(path: Path) -> str:
    """Text in an image, line by line (RapidOCR: pip-installed, runs on the CPU, models bundled)."""
    res = _ocr(path)
    return "\n".join(t.strip() for t in (res.txts or ()) if t and t.strip())


def text_lines(path: Path, height: int, read: bool = True) -> list[tuple[str, float, float]]:
    """Each line of text in an image as (text, height, top), height and top as shares of the
    image's height. read=False only finds the lines (text ""), which is several times faster."""
    res = _ocr(path, read)
    boxes = [] if res.boxes is None else list(res.boxes)
    texts = list(res.txts or ()) if read else [""] * len(boxes)
    return [(t.strip(), float(b[:, 1].max() - b[:, 1].min()) / height, float(b[:, 1].min()) / height)
            for b, t in zip(boxes, texts)]


def looks_like_slide(lines: list[tuple[str, float, float]]) -> bool:
    """A slide is a few lines, usually under a big title. A desktop (a browser, an IDE) is a
    hundred small ones: tabs, menus, a sidebar, the taskbar. On real COP4020 and MIT recordings
    slides had 3 to 39 lines, browser and IntelliJ screens 46 to 137."""
    tallest = max((h for _, h, _ in lines), default=0.0)
    return bool(lines) and (len(lines) <= 30 or (len(lines) <= 45 and tallest >= 0.06))


def pptx_to_pdf(src: Path, out_dir: Path) -> Path:
    """LibreOffice headless, with its own throwaway profile so an open LibreOffice window (or a
    second conversion) can't block it. On Windows soffice.exe is a launcher for soffice.bin, so a
    timeout kills the whole tree; killing just the launcher would leave soffice.bin running."""
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as prof:
        p = subprocess.Popen([str(c) for c in (find_tool("soffice"), f"-env:UserInstallation={Path(prof).as_uri()}",
                                               "--headless", "--convert-to", "pdf", "--outdir", out_dir, src)],
                             creationflags=NO_WINDOW, stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                             stderr=subprocess.DEVNULL)
        try:
            code = p.wait(timeout=300)
        except subprocess.TimeoutExpired:
            if os.name == "nt":
                subprocess.run(["taskkill", "/T", "/F", "/PID", str(p.pid)], creationflags=NO_WINDOW,
                               capture_output=True)
            p.kill()
            p.wait()
            raise RuntimeError("LibreOffice took over 5 minutes to convert the deck")
        if code != 0:
            raise RuntimeError(f"LibreOffice failed ({code})")
    pdf = out_dir / f"{src.stem}.pdf"
    if not pdf.is_file():
        raise RuntimeError("LibreOffice finished without writing a PDF")
    return pdf


def pptx_text(src: Path) -> list[dict]:
    from pptx import Presentation
    pages = []
    for i, slide in enumerate(Presentation(str(src)).slides, 1):
        texts = [sh.text_frame.text.strip() for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        pages.append({"no": i, "text": "\n".join(texts), "title": (title or (texts[0] if texts else ""))[:120]})
    return pages


def _picture_share(page) -> float:
    """How much of the page pictures cover (a scanned slide is one page-sized picture)."""
    area = page.rect.get_area()
    if not area:
        return 0.0
    covered = sum((type(page.rect)(i["bbox"]) & page.rect).get_area() for i in page.get_image_info())
    return min(1.0, covered / area)


def extract_deck(deck_id: int, progress=None) -> dict:
    """Read a downloaded deck: each page's text and title, plus PNG renders 320 and 1280 px wide.
    A PPTX goes through LibreOffice to PDF; if that fails, python-pptx gives text without images.
    A page with almost no text (a scan) is read with OCR."""
    import pymupdf as fitz
    say = progress or (lambda m: None)
    deck = get_deck(deck_id)
    src = DATA / deck["path"]
    out = SLIDE_IMAGES / str(deck_id)
    shutil.rmtree(out, ignore_errors=True)  # best effort: an image open in a browser stays (and is overwritten)
    out.mkdir(parents=True, exist_ok=True)
    pdf, pages, size = src, [], (None, None)
    if src.suffix.lower() in (".pptx", ".ppt"):
        say("converting to PDF")
        try:
            pdf = pptx_to_pdf(src, out)
        except Exception as e:
            if src.suffix.lower() == ".ppt":
                log.warning("deck %s: LibreOffice failed on the .ppt: %s", deck_id, e)
                raise DeckError("This .ppt couldn't be converted. Save it as PDF or .pptx and attach that.") from e
            log.warning("deck %s: LibreOffice failed, using the PPTX text only: %s", deck_id, e)
            pdf = None
    if pdf:
        try:
            doc = fitz.open(pdf)
        except Exception as e:  # the library's message names our file path: log it, don't show it
            log.warning("deck %s: can't open %s: %s", deck_id, pdf, e)
            raise DeckError("This PDF couldn't be opened; it may be damaged or not really a PDF. "
                            "Export it again and attach that.") from e
        with doc:
            if doc.page_count == 0:
                raise DeckError("This PDF has no pages. Export the deck again and attach that.")
            if doc.page_count > MAX_DECK_PAGES:
                raise DeckError(f"This PDF has {doc.page_count} pages, more than a lecture's slides (it may be a "
                                "textbook or a reader). Attach the lecture's own deck instead.")
            r0 = doc[0].rect
            size = (THUMB_W, round(THUMB_W * r0.height / r0.width))
            ocr_left = OCR_MAX_PAGES
            for i, page in enumerate(doc, 1):
                say(f"extracting page {i} of {doc.page_count}")
                for w in (THUMB_W, FULL_W):
                    render_fixed(page, w, round(w * r0.height / r0.width)).save(out / f"{i}-{w}.png")
                text = page.get_text().strip()
                # A scan: almost no text, or a page-sized picture with only a footer's worth.
                if ocr_left and (len(text) < 20 or (len(text) < 200 and _picture_share(page) > 0.5)):
                    ocr_left -= 1
                    text = ocr_image(out / f"{i}-{FULL_W}.png") or text  # it reads the footer too
                pages.append({"no": i, "text": text, "title": page_title(page, text)})
    else:
        try:
            pages = pptx_text(src)
        except Exception as e:
            log.warning("deck %s: python-pptx can't read it: %s", deck_id, e)
            raise DeckError("This PowerPoint couldn't be read. Export it as PDF and attach that.") from e
        if not pages:
            raise DeckError("This PowerPoint has no slides. Export the deck again and attach that.")
    with db() as con:
        con.execute("DELETE FROM slides WHERE deck_id=?", (deck_id,))
        con.executemany("INSERT INTO slides(deck_id, no, title, text) VALUES (?,?,?,?)",
                        [(deck_id, p["no"], p["title"], p["text"]) for p in pages])
        con.execute("UPDATE decks SET pages=?, thumb_w=?, thumb_h=?, extracted_at=?, status='ready', error=NULL "
                    "WHERE id=?", (len(pages), *size, time.time(), deck_id))
    log.info("deck %s (%s): %s pages", deck_id, deck["title"], len(pages))
    return {"pages": len(pages)}


# --- slides read off the video, for lectures with no deck
SCREEN_W, SCREEN_H = 160, 90  # keyframes are compared at this size: big enough to see slide text change


def find_screens(frames, ts: list[float], end: float, change: float = 0.005, min_sec: float = 8.0) -> list[dict]:
    """Group keyframes (small grey images, one per entry of ts) into runs that stay put, drop
    runs shorter than min_sec (scrolling, transitions), and fold runs that look the same into one
    screen, so going back to a slide is a revisit. Two frames differ when more than `change` of
    their pixels moved by over 12 grey levels, ignoring pixels that change most of the time (a
    webcam inset, a clock). Mean difference doesn't work: two slides from one template differ
    in a few lines of text, under 1% of the picture. Screens come back in first-seen order, each
    with the keyframe index to grab (the run's last frame: a slide that builds up is complete)."""
    import numpy as np
    frames = np.asarray(frames)  # kept as given (uint8 from ffmpeg): pairs are compared one at a time
    n = len(frames)
    if n == 0:
        return []
    moved = lambda i: np.abs(frames[i].astype(np.int16) - frames[i - 1]) > 12  # noqa: E731
    counts = np.zeros(frames.shape[1:], np.int32)
    for i in range(1, n):
        counts += moved(i)
    live = ~(counts > 0.3 * (n - 1))
    differ = lambda a, b: float(((np.abs(a - b.astype(np.int16)) > 12) & live).mean()) > change  # noqa: E731
    runs, start = [], 0
    for i in range(1, n + 1):
        if i == n or float((moved(i) & live).mean()) > change:
            runs.append((start, i - 1))
            start = i
    screens: list[dict] = []
    for a, b in runs:
        t0, t1 = ts[a], (ts[b + 1] if b + 1 < len(ts) else end)
        if t1 - t0 < min_sec:
            continue
        for s in screens:
            if not differ(s["img"], frames[b]):
                s["times"].append([t0, t1])
                break
        else:
            screens.append({"img": frames[b].astype(np.int16), "grab": b, "grab_t": ts[b], "times": [[t0, t1]]})
    return screens


def _merge_times(times: list[list[float]]) -> list[list[float]]:
    out: list[list[float]] = []
    for s, e in sorted(times):
        if out and s - out[-1][1] < 1.0:
            out[-1][1] = max(out[-1][1], e)
        else:
            out.append([s, e])
    return [[round(s, 2), round(e, 2)] for s, e in out]


def video_screens(lec: dict) -> tuple[list[dict], int]:
    """The lecture video's distinct screens (find_screens) and how many keyframes were read.
    Only keyframes are decoded, small and grey, at most one per 2 s (a recording with a keyframe
    every frame would otherwise be gigabytes): about 35 s for 50 minutes of 1440p."""
    import numpy as np
    r = run([find_tool("ffmpeg"), "-hide_banner", "-nostdin", "-v", "info", "-skip_frame", "nokey",
             "-i", DATA / lec["video_path"], "-an", "-vf",
             f"select='isnan(prev_selected_t)+gte(t-prev_selected_t,2)',scale={SCREEN_W}:{SCREEN_H},"
             "format=gray,showinfo", "-fps_mode", "passthrough", "-f", "rawvideo", "-"], timeout=3600, low=True)
    ts: list[float] = []
    for x in re.findall(rb"pts_time:(\S+)", r.stderr):  # one per frame, even an odd "nan" or negative one
        try:
            t = float(x)
        except ValueError:
            t = float("nan")
        ts.append(t if t == t else (ts[-1] if ts else 0.0))
    frames = np.frombuffer(r.stdout, np.uint8).reshape(-1, SCREEN_H, SCREEN_W)
    if len(ts) != len(frames):
        log.warning("lecture %s: %s keyframes but %s timestamps; using the shorter", lec["id"], len(frames), len(ts))
    n = min(len(ts), len(frames))
    return find_screens(frames[:n], ts[:n], lec.get("duration_sec") or (ts[n - 1] if n else 0)), n


def ocr_video_deck(lec: dict, progress=None) -> int | None:
    """A deck made from the video for a lecture without slides: one page per distinct screen,
    read with OCR. Only keyframes are decoded (small and grey) to find screens, so a 50-minute
    1440p recording takes about half a minute; just the chosen frames are decoded at full size.
    Returns the deck id, or None when no screen has readable text (a talking head, a dark room)."""
    import pymupdf as fitz
    say = progress or (lambda m: None)
    video, ff = DATA / lec["video_path"], find_tool("ffmpeg")
    say("finding the distinct screens in the video")
    screens, n = video_screens(lec)
    if len(screens) > OCR_MAX_SCREENS:  # keep the ones on screen longest
        screens = sorted(screens, key=lambda s: -sum(e - b for b, e in s["times"]))[:OCR_MAX_SCREENS]
        screens.sort(key=lambda s: s["times"][0][0])
    if not screens:
        return None
    key = f"ocr:{lec['id']}"
    with db() as con:  # look up first: UNIQUE(course_id, key) doesn't hold when course_id is NULL
        r = con.execute("SELECT id FROM decks WHERE key=?", (key,)).fetchone()
        did = r[0] if r else con.execute("INSERT INTO decks(course_id, key, source, title) VALUES (?,?,?,?)",
                                         (lec["course_id"], key, "ocr", "Video frames")).lastrowid
    out = SLIDE_IMAGES / str(did)
    shutil.rmtree(out, ignore_errors=True)
    with tempfile.TemporaryDirectory() as tmp:
        kept: list[dict] = []
        for i, s in enumerate(screens, 1):
            say(f"reading screen {i} of {len(screens)}")
            # The time is a keyframe's, so seeking before -i lands on exactly that frame, fast.
            img = Path(tmp) / f"{i}.png"
            run([ff, "-hide_banner", "-nostdin", "-v", "error", "-ss", f"{s['grab_t']:.3f}", "-i", video, "-an",
                 "-frames:v", "1", "-vf", f"scale={FULL_W}:-2", img], timeout=300, low=True)
            if not img.is_file():
                continue
            height = fitz.Pixmap(str(img)).height
            if not looks_like_slide(text_lines(img, height, read=False)):  # a browser, an IDE: not a slide
                continue
            lines = [ln for ln in text_lines(img, height) if ln[0]]
            text = "\n".join(t for t, _, _ in lines)
            words = set(re.findall(r"[a-z0-9]{3,}", text.lower()))
            if len(words) < 3:
                continue
            for k in kept:  # the same slide with a cursor or webcam change: one page
                if len(words & k["words"]) / len(words | k["words"]) >= 0.85:
                    k["times"] += s["times"]
                    break
            else:
                title = max(lines, key=lambda ln: (round(ln[1], 2), -ln[2]))[0]  # the biggest line, topmost
                kept.append({"img": img, "text": text, "title": title, "words": words, "times": s["times"]})
        if len(kept) < 2:
            with db() as con:
                con.execute("DELETE FROM decks WHERE id=?", (did,))
            return None
        out.mkdir(parents=True, exist_ok=True)
        size = None
        for no, k in enumerate(kept, 1):
            full = fitz.Pixmap(str(k["img"]))
            size = size or (THUMB_W, round(THUMB_W * full.height / full.width))
            full.save(out / f"{no}-{FULL_W}.png")
            fitz.Pixmap(full, *size).save(out / f"{no}-{THUMB_W}.png")
    with db() as con:
        con.execute("DELETE FROM slides WHERE deck_id=?", (did,))
        con.executemany("INSERT INTO slides(deck_id, no, title, text, times) VALUES (?,?,?,?,?)",
                        [(did, no, k["title"][:120], k["text"], json.dumps(_merge_times(k["times"])))
                         for no, k in enumerate(kept, 1)])
        con.execute("UPDATE decks SET pages=?, thumb_w=?, thumb_h=?, extracted_at=?, status='ready', error=NULL "
                    "WHERE id=?", (len(kept), *size, time.time(), did))
    log.info("lecture %s: %s screen(s) read from the video (%s keyframes)", lec["id"], len(kept), n)
    return did


# --- alignment
_ALIGN_STOP = frozenset((
    "the and for are but not you all any can her was one our out has have had him his how its may new now "
    "see two way who did get let put say she too use that this with from they will what when your which "
    "their there then them these those were been being into just like also more some than very only over "
    "such here where would could should about after before because between each other does doing going "
    "okay yeah right really thing things kind sort want know think gonna mean look well good so").split())


def _tokens(text: str) -> list[str]:
    out = []
    for w in re.findall(r"[a-z][a-z0-9_]+", (text or "").lower()):
        if len(w) < 3 or w in _ALIGN_STOP:
            continue
        if len(w) > 4 and w.endswith("s") and not w.endswith("ss"):
            w = w[:-1]
        out.append(w)
    return out


def tfidf_sim(win_texts: list[str], slide_texts: list[str]):
    """Cosine similarity (windows x slides) over the slides' own vocabulary, idf from the slides:
    words on every slide (the course name in a footer) count for little."""
    import numpy as np
    docs = [_tokens(t) for t in slide_texts]
    vocab: dict[str, int] = {}
    for d in docs:
        for w in d:
            vocab.setdefault(w, len(vocab))
    if not vocab:
        return np.zeros((len(win_texts), len(slide_texts)))
    df = np.zeros(len(vocab))
    for d in docs:
        for w in set(d):
            df[vocab[w]] += 1
    idf = np.log((1 + len(docs)) / (1 + df)) + 1

    def vec(tokens):
        v = np.zeros(len(vocab))
        for w in tokens:
            if w in vocab:
                v[vocab[w]] += 1
        v *= idf
        n = np.linalg.norm(v)
        return v / n if n else v

    return np.stack([vec(_tokens(t)) for t in win_texts]) @ np.stack([vec(d) for d in docs]).T


# Costs, in units of a typical window's best match. Picked on three real COP4020 lectures by
# scoring text-only alignment against the slides their videos showed: 77% of segments exact.
ALIGN_SKIP, ALIGN_SKIP_MAX, ALIGN_BACK = 0.6, 1.5, 0.8


def align_windows(sim, visual=None, off=None) -> list[int | None] | None:
    """One slide per window (0-based), by dynamic programming: slide numbers mostly go forward, a
    skip costs more the further it goes, and stepping back one slide is allowed at a price
    (lecturers go back). Windows with no evidence keep the current slide. `visual` (windows x
    slides, 0 to 1) adds how much the screen looks like each slide; `off` windows showed no slide
    at all (a coding demo) and get None. None when nothing matches anything."""
    import numpy as np
    W, S = sim.shape
    if W == 0 or S == 0:
        return [] if W == 0 else None
    peak = sim.max(axis=1)
    gain = sim / peak[peak > 0].mean() if (peak > 0).any() else np.zeros_like(sim)
    if visual is not None:
        gain = gain + 2 * visual  # a screen that looks like the slide is the strongest evidence there is
    keep = [w for w in range(W) if not (off and off[w])]
    if not keep or not gain[keep].any():
        return None
    gain = gain[keep]
    k = np.arange(S)
    step = k[None, :] - k[:, None]  # step[i, j]: going from slide i to slide j
    trans = np.full((S, S), -np.inf)
    trans[step == 0] = 0.0
    trans[step == 1] = -0.05
    big = step >= 2
    trans[big] = -np.minimum(ALIGN_SKIP + 0.1 * (step[big] - 2), ALIGN_SKIP_MAX)
    trans[step == -1] = -ALIGN_BACK
    score = gain[0] - np.minimum(0.05 * k, 0.5)  # a lecture usually starts near the first slide
    back = np.zeros((len(keep), S), dtype=int)
    for w in range(1, len(keep)):
        cand = score[:, None] + trans
        back[w] = cand.argmax(axis=0)
        score = cand.max(axis=0) + gain[w]
    path = [int(score.argmax())]
    for w in range(len(keep) - 1, 0, -1):
        path.append(int(back[w][path[-1]]))
    out: list[int | None] = [None] * W
    for w, s in zip(keep, reversed(path)):
        out[w] = s
    return out


VISUAL_NEAR, VISUAL_FAR = 0.15, 0.30  # share of pixels that differ: "this slide" / "not a slide"


def slide_greys(deck: dict):
    """Each page's render, grey and the size keyframes are compared at."""
    import numpy as np
    import pymupdf as fitz
    out = []
    for no in range(1, deck["pages"] + 1):
        p = fitz.Pixmap(fitz.Pixmap(fitz.csGRAY, fitz.Pixmap(str(SLIDE_IMAGES / str(deck["id"]) / f"{no}-{THUMB_W}.png"))),
                        SCREEN_W, SCREEN_H)
        out.append(np.frombuffer(p.samples, np.uint8).reshape(p.height, p.stride)[:, :SCREEN_W])
    return np.stack(out).astype(np.int16)


def visual_evidence(lec: dict, deck: dict, segs: list[dict], wins: list[list[int]]):
    """Compare what the video shows with the deck's pages. Returns (windows x slides closeness,
    off-slide flags), or None when the video doesn't show the slides full screen (a camera on a
    projector, audio only), in which case the text decides alone."""
    import numpy as np
    if not lec.get("video_path") or not deck.get("thumb_w"):
        return None
    try:
        screens, _ = video_screens(lec)
        greys = slide_greys(deck)
    except Exception as e:
        log.warning("lecture %s: can't compare the video with the slides: %s", lec["id"], e)
        return None
    if not screens:
        return None
    dist = np.array([[float((np.abs(s["img"] - g) > 24).mean()) for g in greys] for s in screens])
    spans = [(t0, t1, i) for i, s in enumerate(screens) for t0, t1 in s["times"]]
    shown = sum(t1 - t0 for t0, t1, i in spans if dist[i].min() < VISUAL_NEAR)
    if shown < 0.1 * sum(t1 - t0 for t0, t1, _ in spans):
        return None
    vis, off = np.zeros((len(wins), len(greys))), [False] * len(wins)
    for w, idx in enumerate(wins):
        a, b = segs[idx[0]]["start"], segs[idx[-1]]["end"]
        cover = {}
        for t0, t1, i in spans:
            if min(b, t1) > max(a, t0):
                cover[i] = cover.get(i, 0) + min(b, t1) - max(a, t0)
        if not cover or max(cover.values()) < 0.5 * (b - a):
            continue  # mostly transitions: no visual evidence
        d = dist[max(cover, key=cover.get)]
        if d.min() > VISUAL_FAR:
            off[w] = True
        else:
            vis[w] = np.clip(1 - d / VISUAL_NEAR, 0, 1)
    log.info("lecture %s: the video shows the slides %.0f%% of the time, %s of %s windows off-slide",
             lec["id"], 100 * shown / max(1, sum(t1 - t0 for t0, t1, _ in spans)), sum(off), len(wins))
    return vis, off


def make_windows(segs: list[dict], sec: float = WINDOW_SEC) -> list[list[int]]:
    """Consecutive segments grouped into windows of about `sec` seconds (indexes into segs)."""
    out: list[list[int]] = []
    for i, s in enumerate(segs):
        if out and s["end"] - segs[out[-1][0]]["start"] <= sec:
            out[-1].append(i)
        else:
            out.append([i])
    return out


def align_lecture(lecture_id: int) -> int:
    """Give every segment the slide on screen while it was said: from the video when it shows the
    slides, otherwise from the words. Segments during something that isn't a slide (a coding
    demo) get none. Returns how many segments got a slide."""
    lec = get_lecture(lecture_id)
    deck = get_deck(lec["deck_id"]) if lec else None
    if not deck:
        return 0  # the deck was changed or removed meanwhile: that change queued its own job
    with db() as con:
        segs = rows(con.execute("SELECT idx, start, end, text FROM segments WHERE lecture_id=? ORDER BY idx",
                                (lecture_id,)))
        slides = rows(con.execute("SELECT no, text, times FROM slides WHERE deck_id=? ORDER BY no", (deck["id"],)))
    if not slides:
        raise DeckError("The deck has no pages. Attach the deck again.")
    if deck["source"] == "ocr":  # the screens carry their own times
        spans = sorted((s0, s1, sl["no"]) for sl in slides for s0, s1 in json.loads(sl["times"] or "[]"))
        nos = []
        for sg in segs:
            mid = (sg["start"] + sg["end"]) / 2
            inside = [no for s0, s1, no in spans if s0 <= mid < s1]
            before = [no for s0, _, no in spans if s0 <= mid]
            nos.append(inside[-1] if inside else before[-1] if before else spans[0][2])
    else:
        wins = make_windows(segs)
        sim = tfidf_sim([" ".join(segs[i]["text"] for i in w) for w in wins], [sl["text"] for sl in slides])
        vis = visual_evidence(lec, deck, segs, wins)
        path = align_windows(sim, *(vis or (None, None)))
        if path is None:
            raise DeckError("Nothing on these slides matches what was said. Check that it's the right deck "
                            "(pick another under Slides), or attach a PDF with real text instead of images.")
        nos = [None] * len(segs)
        for w, s in zip(wins, path):
            for i in w:
                nos[i] = None if s is None else slides[s]["no"]
    with db() as con:  # only if the deck is still the one aligned to (a pick meanwhile wins)
        if not con.execute("UPDATE lectures SET slides_aligned_at=datetime('now'), slides_error=NULL "
                           "WHERE id=? AND deck_id=?", (lecture_id, deck["id"])).rowcount:
            return 0
        con.executemany("UPDATE segments SET slide_no=? WHERE lecture_id=? AND idx=?",
                        [(no, lecture_id, sg["idx"]) for sg, no in zip(segs, nos)])
    log.info("lecture %s aligned to deck %s: %s segments, %s slides used", lecture_id, deck["id"], len(segs),
             len(set(nos) - {None}))
    return sum(no is not None for no in nos)


def read_uploaded_deck(deck: dict, say, lecture_id: int) -> None:
    """Read a deck you attached (the request only checked it), then line up every other lecture
    that uses it. A file that can't be read marks the deck as an error, so the lecture shows what
    to do; anything else (OCR timing out, a locked database) fails only this job, and "Look for
    slides again" reads it again."""
    msg = None
    try:
        extract_deck(deck["id"], say)
    except DeckError as e:
        msg = str(e)
    if msg:
        with db() as con:
            con.execute("DELETE FROM slides WHERE deck_id=?", (deck["id"],))
            con.execute("UPDATE decks SET status='error', error=? WHERE id=?", (msg, deck["id"]))
        raise DeckError(msg)
    with db() as con:
        others = [r[0] for r in con.execute("SELECT id FROM lectures WHERE deck_id=? AND id!=?",
                                            (deck["id"], lecture_id))]
    for lid in others:  # this job's own lecture is lined up next, right here
        enqueue_slides(lid)


def reread_upload(deck_id: int, pages: int | None = None) -> None:
    """An attached deck goes back to being read: its error clears on every lecture using it, and
    each one is queued to line it up once the read is done."""
    with db() as con:
        con.execute("UPDATE decks SET status='new', error=NULL, pages=COALESCE(?, pages) WHERE id=?", (pages, deck_id))
        lids = [r[0] for r in con.execute("SELECT id FROM lectures WHERE deck_id=?", (deck_id,))]
        con.execute(f"UPDATE lectures SET slides_error=NULL WHERE id IN ({','.join('?' * len(lids))})", lids)
    for lid in lids:
        enqueue_slides(lid)


def job_slides(job: dict) -> dict:
    """payload: {lecture_id}. Find the lecture's deck (or read slides off its video), then align."""
    lid = job["payload"]["lecture_id"]
    say = lambda m: job_update(job["id"], progress=m)  # noqa: E731
    lec = get_lecture(lid)
    if not lec:
        return {"skipped": "lecture deleted"}
    with db() as con:  # a new attempt: the last one's error no longer applies
        con.execute("UPDATE lectures SET slides_error=NULL WHERE id=?", (lid,))
    try:
        if lec["course_id"] and lec["deck_how"] != "manual":
            for other in match_decks(lec["course_id"]):
                if other != lid:
                    enqueue_slides(other)
            lec = get_lecture(lid)
        transcribed = lec["segment_count"] > 0
        if not lec["deck_id"]:
            if lec["deck_how"] == "manual":
                return {"slides": "none (your choice)"}
            if not lec["video_path"] or not transcribed:
                return {"waiting": "no slides, and no transcript yet" if lec["video_path"] else "no slides"}
            with db() as con:
                r = con.execute("SELECT id FROM decks WHERE key=? AND status='ready'", (f"ocr:{lid}",)).fetchone()
            if not r and not job["payload"].get("ocr") and not slides_scan_done(lec["course_id"]):
                return {"waiting": "the course's Canvas scan for slides"}  # the scan queues it again
            did = r[0] if r else ocr_video_deck(lec, say)
            with db() as con:  # OCR takes minutes: a deck you picked or attached meanwhile wins
                con.execute("UPDATE lectures SET deck_id=?, deck_how=? WHERE id=? AND deck_id IS NULL "
                            "AND COALESCE(deck_how,'')!='manual'", (did, "auto" if did else "nothing", lid))
            lec = get_lecture(lid)
            if not lec["deck_id"]:
                return {"slides": "none found in the video" if lec["deck_how"] == "nothing" else "none (your choice)"}
        deck = get_deck(lec["deck_id"])
        if not deck:
            return {"skipped": "the deck was removed"}
        if deck["status"] == "new" and deck["source"] == "upload":
            read_uploaded_deck(deck, say, lid)
            deck = get_deck(deck["id"])
        if deck["status"] == "new":
            return {"waiting": "downloading the slides"}  # the download queues this lecture again
        if deck["status"] == "error":
            raise DeckError(deck["error"] or "The deck couldn't be read. Attach it again.")
        if not transcribed:
            return {"waiting": "no transcript yet"}  # storing the transcript queues it again
        say("aligning")
        return {"deck": deck["title"], "pages": deck["pages"], "segments": align_lecture(lid)}
    except Exception as e:
        with db() as con:
            con.execute("UPDATE lectures SET slides_error=? WHERE id=?", (user_error(e, "Lining up the slides"), lid))
        raise


def _download_deck(deck: dict, tmp: Path) -> Path:
    if deck["source"] == "canvas":
        canvas_ready()
        with CanvasSession() as cs:
            url, name, _ = canvas_file_source(int(deck["url"].split(":", 1)[1]))
            jar = cs.httpx_cookies()
        try:
            return stream_download(url, tmp / safe_filename(name), jar, max_bytes=MAX_DECK_MB << 20)
        except ValueError as e:
            raise DeckError(f"The deck is over {MAX_DECK_MB} MB. Attach a smaller PDF export instead.") from e
    if deck["source"] == "drive":
        import gdown
        tmp.mkdir(parents=True, exist_ok=True)
        got = gdown.download(id=deck["key"].split(":", 1)[1], output=str(tmp) + os.sep, quiet=True)
        if not got:
            raise RuntimeError("Google Drive didn't return the file")  # quota or a hiccup: retried
        return Path(got)
    raise DeckError(f"A {deck['source']} deck has nothing to download.")


def retry_deck(deck_id: int) -> bool:
    """Download a deck that failed again (you asked to retry), or read an attached one again if
    its file is still here. False if there's nothing to fetch."""
    deck = get_deck(deck_id)
    if deck and deck["status"] == "error" and deck["source"] == "upload" and deck["path"] \
            and (DATA / deck["path"]).is_file():
        reread_upload(deck_id)
        wake_workers()
        return True
    if not deck or deck["status"] != "error" or deck["source"] not in ("canvas", "drive"):
        return False
    with db() as con:
        con.execute("UPDATE decks SET status='new', error=NULL WHERE id=?", (deck_id,))
    enqueue("deck_pull", {"deck_id": deck_id, "title": deck["title"]})
    wake_workers()
    return True


def job_deck_pull(job: dict) -> dict:
    """payload: {deck_id}. Download a deck listed on Canvas or Drive, read it, then re-match its
    course and queue alignment for every lecture whose deck changed or was waiting on this one.
    A download that fails for a passing reason (Canvas signed out, the browser busy, the network,
    Drive's quota) leaves the deck waiting and tries again later; only a bad file is an error."""
    did = job["payload"]["deck_id"]
    deck = get_deck(did)
    if not deck:
        return {"skipped": "deck deleted"}
    say = lambda m: job_update(job["id"], progress=m)  # noqa: E731
    tmp = DATA / "downloads" / f"job{job['id']}"
    try:
        say(f"downloading {deck['title']}")
        try:
            f = _download_deck(deck, tmp)
        except DeckError:
            raise
        except Exception as e:
            signed_out = deck["source"] == "canvas" and canvas_state().get("state") != "connected"
            tries = job["payload"].get("tries", 0) + (0 if signed_out else 1)
            if tries < DECK_MAX_TRIES:  # signed out doesn't count: it waits until you sign in
                wait = 1800 if signed_out else DECK_RETRY_SEC * tries
                # attempts counts crashed runs (requeue_orphaned_jobs); these ran fine, so start over
                job_update(job["id"], status="queued", retry_at=_sqlnow(wait), owner_pid=None, attempts=0,
                           payload={**job["payload"], "tries": tries},
                           progress="waiting for Canvas sign-in" if signed_out else "download failed; trying again later")
                raise RetryLater(f"deck {did}: {e!r}") from e
            log.warning("deck %s: download failed %s times, last: %r", did, tries, e)
            raise DeckError(f"Downloading this deck failed {DECK_MAX_TRIES} times. Look for slides again to "
                            "retry, or download it yourself and attach it.") from e
        if f.suffix.lower() not in SLIDE_EXTS:
            raise DeckError(f"The download is a {f.suffix or 'file without an extension'}, not a PDF or PowerPoint.")
        if f.stat().st_size > MAX_DECK_MB << 20:
            raise DeckError(f"The deck is over {MAX_DECK_MB} MB. Attach a smaller PDF export instead.")
        course = find_course(deck["course_id"]) or {}
        folder = SLIDES_DIR / safe_filename(course.get("code") or "Unsorted", "Unsorted")
        folder.mkdir(parents=True, exist_ok=True)
        dest = folder / safe_filename(f.name)
        if dest.exists() and sha256_file(dest) != sha256_file(f):
            dest = dest.with_name(f"{dest.stem} ({did}){dest.suffix}")
        shutil.move(str(f), dest)
        with db() as con:
            con.execute("UPDATE decks SET path=? WHERE id=?", (dest.relative_to(DATA).as_posix(), did))
        res = extract_deck(did, say)
    except RetryLater:
        raise
    except Exception as e:
        with db() as con:
            con.execute("UPDATE decks SET status='error', error=? WHERE id=?", (user_error(e, "Getting the deck"), did))
        raise
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    lids = set(match_decks(deck["course_id"])) if deck["course_id"] else set()
    with db() as con:
        lids |= {r[0] for r in con.execute("SELECT id FROM lectures WHERE deck_id=?", (did,))}
    for lid in sorted(lids):
        enqueue_slides(lid)
    return {**res, "lectures": sorted(lids)}


JOB_HANDLERS["slides"] = job_slides
JOB_HANDLERS["deck_pull"] = job_deck_pull


def looks_like_lecture_deck(title: str, lectures: list[dict]) -> bool:
    """Canvas holds the syllabus, homework and readings too. Fetch a PDF/PPTX only if its name has
    a lecture code (M2L1, Lecture 7, Week 3), says it's slides, or is close to a recording's title.
    One skipped now is looked at again on the next scan, when its lecture may have been recorded."""
    if title_codes(title).keys() & {0, 1, 2}:
        return True
    if re.search(r"(?i)\b(slides?|deck)\b", title or ""):
        return True
    return any(match_score({"title": title}, l) >= MATCH_MIN for l in lectures)


def slides_scan_done(course_id: int | None) -> bool:
    """Has the course been searched for decks? Until then a lecture without one doesn't get slides
    read off its video (minutes of CPU), unless you ask: its deck may just not be found yet. A
    course that isn't on Canvas has nothing to search."""
    course = find_course(course_id) if course_id else None
    return not (course and course.get("canvas_id")) or bool(meta_get(f"slides_scanned:{course_id}"))


def slides_scanned(course_id: int) -> None:
    """After a scan: lectures that were waiting for it get looked at (their deck downloads,
    queued just before, run first)."""
    meta_set(f"slides_scanned:{course_id}", datetime.now().isoformat(timespec="seconds"))
    with db() as con:
        waiting = [r[0] for r in con.execute(
            "SELECT id FROM lectures l WHERE course_id=? AND deck_id IS NULL AND deck_how IS NULL AND slides_error "
            "IS NULL AND video_path IS NOT NULL AND EXISTS (SELECT 1 FROM segments s WHERE s.lecture_id=l.id)",
            (course_id,))]
    for lid in waiting:
        enqueue_slides(lid)


def record_decks(course_id: int, found: dict[str, dict]) -> int:
    """Store decks a Canvas scan listed and queue a download for each new one."""
    new = []
    with db() as con:
        for d in found.values():
            cur = con.execute("INSERT OR IGNORE INTO decks(course_id, key, source, title, url, found_at, created_at) "
                              "VALUES (?,?,?,?,?,?,?)", (course_id, d["key"], d["source"], d["title"], d["url"],
                                                          d.get("found_at"), d.get("created_at")))
            if cur.rowcount:
                new.append((cur.lastrowid, d["title"]))
    for did, title in new:
        enqueue("deck_pull", {"deck_id": did, "title": title})
    if new:
        wake_workers()
    return len(new)


# --- what the UI reads
def lecture_slides(lecture_id: int) -> dict:
    """The lecture's slide state, its deck's pages with the time ranges each was on screen, and
    the course's other decks to pick from."""
    lec = get_lecture(lecture_id)
    if not lec:
        raise KeyError(lecture_id)
    deck = get_deck(lec["deck_id"])
    with db() as con:
        job = con.execute("""SELECT status, progress FROM jobs WHERE status IN ('queued','running') AND
                             ((kind='slides' AND lecture_id=?) OR (kind='deck_pull' AND json_extract(payload,'$.deck_id')=?))
                             ORDER BY status='running' DESC, id LIMIT 1""", (lecture_id, lec["deck_id"] or -1)).fetchone()
        slides = rows(con.execute("SELECT no, title, text FROM slides WHERE deck_id=? ORDER BY no",
                                  (lec["deck_id"] or -1,)))
        segs = rows(con.execute("SELECT start, end, slide_no FROM segments WHERE lecture_id=? AND slide_no IS NOT NULL "
                                "ORDER BY idx", (lecture_id,)))
        cands = rows(con.execute("""SELECT d.id, d.title, d.created_at, d.pages, d.source, d.status,
                                    (SELECT l.id FROM lectures l WHERE l.deck_id=d.id AND l.id!=? ORDER BY l.id LIMIT 1) AS lid
                                    FROM decks d WHERE (d.course_id IS ? AND d.source!='ocr') OR d.key=?""",
                                 (lecture_id, lec["course_id"], f"ocr:{lecture_id}")))
        lids = [c["lid"] for c in cands if c["lid"]]
        titles = {r[0]: r[1] for r in con.execute(
            f"SELECT id, title FROM lectures WHERE id IN ({','.join('?' * len(lids))})", lids)} if lids else {}
    transcribed = lec["segment_count"] > 0
    state = reason = progress = error = None
    if lec["slides_error"] or (deck and deck["status"] == "error"):
        state, error = "error", lec["slides_error"] or deck["error"]
    elif job and job["status"] == "running":
        state, progress = "aligning", job["progress"]
    elif deck and lec["slides_aligned_at"]:
        state = "aligned"
    elif deck:
        state = "waiting"
        reason = (("reading the slides" if deck["source"] == "upload" else "downloading the slides")
                  if deck["status"] == "new" else "no transcript yet" if not transcribed
                  else "queued")
    elif lec["deck_how"] == "manual":
        state, reason = "none", "You chose no slides for this lecture."
    elif lec["deck_how"] == "nothing":
        state, reason = "none", "No deck matched, and the video has no readable slides."
    elif lec["video_path"] and not job:
        state = "waiting"
        reason = "no transcript yet" if not transcribed else "not checked yet"
    elif job:
        state, reason = "waiting", "queued"
    else:
        state, reason = "none", "No deck matched this lecture yet. Attach a PDF or PowerPoint, or pick one below."
    ranges: dict[int, list[list[float]]] = {}
    for sg in segs:
        rs = ranges.setdefault(sg["slide_no"], [])
        if rs and sg["start"] - rs[-1][1] < 1.0:
            rs[-1][1] = max(rs[-1][1], sg["end"])
        else:
            rs.append([sg["start"], sg["end"]])
    v = int(deck["extracted_at"] or 0) if deck else 0
    img = (lambda no, w: f"/media/slides/{deck['id']}/{no}-{w}.png?v={v}") if deck and deck["thumb_w"] \
        else (lambda no, w: None)
    return {
        "state": state, "reason": reason, "progress": progress, "error": error,
        "deck": deck and {"id": deck["id"], "title": deck["title"], "source": deck["source"],
                          "how": lec["deck_how"], "score": lec["deck_score"], "pages": deck["pages"],
                          "thumb_w": deck["thumb_w"], "thumb_h": deck["thumb_h"]},
        "slides": [{"no": s["no"], "title": s["title"] or "", "text": s["text"],
                    "first": ranges[s["no"]][0][0] if s["no"] in ranges else None,
                    "ranges": [[round(a, 2), round(b, 2)] for a, b in ranges.get(s["no"], [])],
                    "thumb": img(s["no"], THUMB_W), "image": img(s["no"], FULL_W)} for s in slides],
        "candidates": sorted(({"id": c["id"], "title": c["title"], "pages": c["pages"], "source": c["source"],
                               "status": c["status"],
                               "score": None if c["source"] == "ocr" else match_score(c, lec),
                               "matched_lecture": {"id": c["lid"], "title": titles.get(c["lid"])} if c["lid"] else None}
                              for c in cands), key=lambda c: -(c["score"] if c["score"] is not None else -1)),
    }


def set_lecture_deck(lecture_id: int, deck_id: int | None) -> dict:
    """Your pick: this deck (or none) for this lecture. It sticks through later re-matching."""
    lec = get_lecture(lecture_id)
    if not lec:
        raise KeyError(lecture_id)
    if deck_id is not None:
        deck = get_deck(deck_id)
        if not deck or (deck["course_id"] != lec["course_id"] and deck["key"] != f"ocr:{lecture_id}"
                        and deck["source"] != "upload"):
            raise ValueError("That deck isn't in this lecture's course.")
    with db() as con:
        con.execute("UPDATE lectures SET deck_id=?, deck_how='manual', deck_score=NULL, slides_aligned_at=NULL, "
                    "slides_error=NULL WHERE id=?", (deck_id, lecture_id))
        con.execute("UPDATE segments SET slide_no=NULL WHERE lecture_id=?", (lecture_id,))
    if deck_id:
        enqueue_slides(lecture_id)
    if lec["course_id"]:  # a deck you placed here stops being auto-matched elsewhere
        for other in match_decks(lec["course_id"]):
            enqueue_slides(other)
    return lecture_slides(lecture_id)


def check_deck_file(path: Path) -> int | None:
    """A quick look at an uploaded deck (open it, count its pages), so a bad file is turned away
    while you wait; reading and rendering the pages is the slides job's. Returns the page count
    (None for a .ppt, which only LibreOffice can open)."""
    import zipfile
    ext = path.suffix.lower()
    if ext == ".pdf":
        import pymupdf as fitz
        try:
            doc = fitz.open(path)
        except Exception as e:  # the library's message names our file path: log it, don't show it
            log.warning("upload: can't open %s: %s", path.name, str(e))  # text only: a kept record can't hold the file
            raise DeckError("This PDF couldn't be opened; it may be damaged or not really a PDF. "
                            "Export it again and attach that.") from e
        with doc:
            if doc.needs_pass:
                raise DeckError("This PDF is password-protected. Export a copy without a password and attach that.")
            n = doc.page_count
    elif ext == ".pptx":
        try:
            with zipfile.ZipFile(path) as z:
                names = z.namelist()
        except (zipfile.BadZipFile, OSError):
            names = []
        if "ppt/presentation.xml" not in names:
            raise DeckError("This PowerPoint couldn't be read. Export it as PDF and attach that.")
        n = sum(bool(re.fullmatch(r"ppt/slides/slide\d+\.xml", x)) for x in names)
    else:  # .ppt: an OLE compound file
        with open(path, "rb") as f:
            if f.read(8) != b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
                raise DeckError("This .ppt couldn't be read. Save it as PDF or .pptx and attach that.")
        return None
    if n == 0:
        raise DeckError("This deck has no pages. Export it again and attach that.")
    if n > MAX_DECK_PAGES:
        raise DeckError(f"This deck has {n} pages, more than a lecture's slides (it may be a textbook or a "
                        "reader). Attach the lecture's own deck instead.")
    return n


def attach_deck_upload(lecture_id: int, filename: str, fileobj) -> dict:
    """An uploaded PDF/PPTX for this lecture: checked right away (so a bad file is reported to you,
    not to the job queue), stored under Slides/<Course>/ and set as your pick. The slides job
    reads its pages before lining them up."""
    lec = get_lecture(lecture_id)
    if not lec:
        raise KeyError(lecture_id)
    name = safe_filename(filename or "slides.pdf")
    ext = Path(name).suffix.lower()
    if ext not in SLIDE_EXTS:
        raise HTTPException(400, f"Slides need to be a PDF or PowerPoint file, and this is "
                                 f"{ext or 'a file without an extension'}. Export it as PDF and attach that.")
    tmp = DATA / "uploads" / f"{uuid.uuid4().hex}{ext}"
    tmp.parent.mkdir(parents=True, exist_ok=True)
    try:
        size, h = 0, hashlib.sha256()
        with open(tmp, "wb") as out:
            while chunk := fileobj.read(1 << 20):
                size += len(chunk)
                if size > MAX_DECK_MB << 20:
                    raise HTTPException(413, f"This file is over {MAX_DECK_MB} MB, too big for a slide deck. "
                                             "Export a smaller PDF (without embedded video) and attach that.")
                h.update(chunk)
                out.write(chunk)
        msg = None
        try:
            pages = check_deck_file(tmp)
        except DeckError as e:
            msg = str(e)
        except Exception as e:
            log.warning("upload for lecture %s: can't check %s: %s", lecture_id, name, str(e))
            msg = "This file couldn't be read. Export it as PDF and attach that."
        if msg:  # out here, the failure's traceback (which holds the file open) is gone, so it can be removed
            raise HTTPException(422, msg)
        key = f"upload:{h.hexdigest()}"
        with db() as con:
            r = con.execute("SELECT id, status, path FROM decks WHERE course_id IS ? AND key=?",
                            (lec["course_id"], key)).fetchone()
        stored = bool(r and r["path"] and (DATA / r["path"]).is_file())
        if stored and r["status"] in ("ready", "new"):  # read already, or being read
            did = r["id"]
        elif stored:  # the same bytes failed before (or were interrupted): read them again
            did = r["id"]
            reread_upload(did, pages)
        else:
            folder = SLIDES_DIR / safe_filename(lec.get("course_code") or "Unsorted", "Unsorted")
            folder.mkdir(parents=True, exist_ok=True)
            dest = folder / name
            if dest.exists():
                dest = dest.with_name(f"{dest.stem} ({h.hexdigest()[:8]}){dest.suffix}")
            shutil.move(str(tmp), dest)
            rel = dest.relative_to(DATA).as_posix()
            try:
                with db() as con:
                    if r:  # its file is gone: this copy takes its place
                        did = r["id"]
                        con.execute("UPDATE decks SET path=?, pages=?, status='new', error=NULL WHERE id=?",
                                    (rel, pages, did))
                    else:
                        did = con.execute("INSERT INTO decks(course_id, key, source, title, path, pages, created_at) "
                                          "VALUES (?,?,?,?,?,?,date('now'))",
                                          (lec["course_id"], key, "upload", Path(name).stem, rel, pages)).lastrowid
            except sqlite3.IntegrityError:  # the same file attached twice at once: the other request stored it
                dest.unlink(missing_ok=True)
                with db() as con:
                    did = con.execute("SELECT id FROM decks WHERE course_id IS ? AND key=?",
                                      (lec["course_id"], key)).fetchone()[0]
            except BaseException:
                dest.unlink(missing_ok=True)
                raise
    finally:
        tmp.unlink(missing_ok=True)
    return set_lecture_deck(lecture_id, did)


def lecture_slide_terms(lecture_id: int) -> str:
    """For Whisper's prompt: slide titles and the deck's jargon (identifiers, acronyms, CamelCase,
    words it repeats), so they get spelled the way the slides spell them."""
    with db() as con:
        r = con.execute("SELECT deck_id FROM lectures WHERE id=?", (lecture_id,)).fetchone()
        slides = rows(con.execute("SELECT title, text FROM slides WHERE deck_id=? ORDER BY no", (r[0],))) \
            if r and r[0] else []
    if not slides:
        return ""
    counts: dict[str, int] = {}
    for s in slides:
        for w in re.findall(r"[A-Za-z][A-Za-z0-9_+#]{2,}", s["text"]):
            counts[w] = counts.get(w, 0) + 1
    jargon = [w for w, c in sorted(counts.items(), key=lambda kv: -kv[1])
              if w.lower() not in _ALIGN_STOP and w.lower() not in _TITLE_STOP and
              (re.search(r"[A-Z].*[A-Z]|[0-9_+#]|[a-z][A-Z]", w) or (c >= 2 and len(w) >= 6))][:40]
    titles = list(dict.fromkeys(s["title"] for s in slides if s["title"]))[:12]
    return "; ".join(titles) + (". " if titles else "") + ", ".join(jargon)


def normalize_words(text: str) -> list[str]:
    return re.findall(r"[a-z0-9']+", text.lower().replace("-", " "))


def wer(reference: str, hypothesis: str) -> float:
    """Word error rate: word-level edit distance / reference length."""
    r, h = normalize_words(reference), normalize_words(hypothesis)
    if not r:
        return 0.0 if not h else 1.0
    prev = list(range(len(h) + 1))
    for i, rw in enumerate(r, 1):
        cur = [i] + [0] * len(h)
        for j, hw in enumerate(h, 1):
            cur[j] = min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (rw != hw))
        prev = cur
    return prev[-1] / len(r)


def benchmark_local(lecture_id: int, start: float = 120, length: float = 120) -> dict:
    """Pick the CPU fallback model: transcribe the same 2-minute slice with each candidate and
    score it against that lecture's Colab large-v3 transcript (the best transcript we have).
    The most accurate model wins unless it's more than 2x slower than the fastest."""
    lec = get_lecture(lecture_id)
    if not (lec.get("transcript_source") or "").startswith("colab"):
        raise ValueError(f"lecture {lecture_id} needs a Colab large-v3 transcript to score against")
    with db() as con:
        segs = rows(con.execute("SELECT start, end, text, words FROM segments WHERE lecture_id=? ORDER BY idx",
                                (lecture_id,)))
    ref_words = [w for s in segs for w in json.loads(s["words"] or "[]") if start <= w[0] < start + length]
    reference = " ".join(w[2] for w in ref_words)
    raw = run([find_tool("ffmpeg"), "-nostdin", "-v", "error", "-ss", start, "-t", length, "-i", DATA / lec["audio_path"],
               "-f", "f32le", "-ac", "1", "-ar", "16000", "-"]).stdout
    import numpy as np
    from faster_whisper import WhisperModel
    audio = np.frombuffer(raw, np.float32)
    results = {}
    for name in LOCAL_MODELS:
        model = WhisperModel(name, device="cpu", compute_type="int8", cpu_threads=os.cpu_count())
        t = time.time()
        sgs, _ = model.transcribe(audio, language="en", beam_size=5, vad_filter=True,
                                  initial_prompt=initial_prompt(lec) or None)
        text = " ".join(x.text.strip() for x in sgs)
        took = time.time() - t
        results[name] = {"seconds": round(took, 1), "rtf": round(took / length, 3), "wer": round(wer(reference, text), 4),
                         "minutes_per_75min": round(took / length * 75, 1)}
        log.info("benchmark %s: %s", name, results[name])
        del model
    fastest = min(r["seconds"] for r in results.values())
    eligible = {n: r for n, r in results.items() if r["seconds"] <= 2 * fastest}
    pick = min(eligible, key=lambda n: (eligible[n]["wer"], eligible[n]["seconds"]))
    report = {"lecture_id": lecture_id, "slice": [start, start + length], "reference_words": len(ref_words),
              "results": results, "picked": pick, "cpu": os.cpu_count(), "at": datetime.now().isoformat(timespec="seconds")}
    meta_set("benchmark", report)
    meta_set("local_model", pick)
    return report


# ---------------------------------------------------------------- notes

# Notes come from Claude Code in headless mode on the Max plan: `claude -p`, never the API.
CLAUDE_WORK = DATA / "claude-work"    # an empty folder to run claude in: no project CLAUDE.md gets picked up
CLAUDE_CACHE = DATA / "claude-cache"  # <sha256 of the inputs>.json: asking the same thing again costs nothing
CLAUDE_TIMEOUT = float(os.environ.get("LECTURELAB_CLAUDE_TIMEOUT", 1800))
NOTES_MODEL, CHECK_MODEL = "sonnet", "opus"  # opus only for the check pass
# One claude process at a time by default: each is a few hundred MB and this laptop runs short.
_claude_slot = threading.BoundedSemaphore(max(1, int(os.environ.get("LECTURELAB_CLAUDE_PARALLEL", 1))))
API_KEY_VARS = ("ANTHROPIC_API_KEY", "ANTHROPIC_AUTH_TOKEN")  # either would bill the API, not the plan


class ClaudeError(Exception):
    """claude -p failed; the message is fit to show."""


class QuotaExceeded(Exception):
    """The plan's usage limit is reached: notes wait until it resets."""

    def __init__(self, msg: str, reset_at: float):
        super().__init__(msg)
        self.reset_at = reset_at


class ClaudeUnavailable(QuotaExceeded):
    """Claude Code is signed out or missing. Nothing is the lecture's fault: every call waits an hour
    and tries again, instead of marking each lecture's notes as failed."""


AUTH_RE = re.compile(r"(?i)invalid api key|/login|not (logged|signed) in|log ?in again|oauth token|authenticat")


def claude_exe() -> str:
    """claude.exe itself: a .cmd shim would mangle the multi-line system prompt and schema."""
    p = Path.home() / ".local" / "bin" / "claude.exe"
    found = str(p) if p.exists() else shutil.which("claude.exe")
    if not found:
        raise ClaudeError("Claude Code isn't installed. Install it and sign in with your Max plan.")
    return found


def _paid_route(name: str) -> bool:
    """Variables that would make claude bill something other than the Max plan: API keys, a custom
    base URL or headers, and the Bedrock, Vertex and Foundry switches."""
    k = name.upper()
    return k in API_KEY_VARS or k.startswith(("ANTHROPIC_", "CLAUDE_CODE_USE_")) or k == "AWS_BEARER_TOKEN_BEDROCK"


def claude_env() -> dict:
    """os.environ without any of those, so claude bills the Max plan. Refuses if one is still visible."""
    env = {k: v for k, v in os.environ.items() if not _paid_route(k)}
    if any(_paid_route(k) for k in env):
        raise RuntimeError("refusing to call claude: an Anthropic API key would be visible to it")
    return env


QUOTA_RE = re.compile(r"(?i)\b(usage|session|weekly|5-hour|opus|sonnet)\s+limit|limit reached|hit your (\w+ )?limit")


def quota_reset_at(text: str, now: float | None = None) -> float:
    """When the limit resets, from claude's message ("...limit reached|1793170800" or "resets 3pm");
    an hour from now when it doesn't say."""
    now = time.time() if now is None else now
    if m := re.search(r"\|(\d{10})\b", text):
        return float(m[1])
    if m := re.search(r"(?i)resets?\s+(?:at\s+)?(\d{1,2})(?::(\d\d))?\s*([ap])\.?m", text):
        hour = int(m[1]) % 12 + (12 if m[3].lower() == "p" else 0)
        t = datetime.fromtimestamp(now).replace(hour=hour, minute=int(m[2] or 0), second=0, microsecond=0)
        ts = t.timestamp()
        return (ts if ts > now else ts + 86400) + 120  # a little past the reset
    return now + 3600


def _claude_run(cmd: list[str], prompt: str, timeout: float) -> subprocess.CompletedProcess:
    CLAUDE_WORK.mkdir(parents=True, exist_ok=True)
    with _claude_slot:
        return subprocess.run(cmd, input=prompt, capture_output=True, text=True, encoding="utf-8", errors="replace",
                              env=claude_env(), cwd=CLAUDE_WORK, timeout=timeout, creationflags=NO_WINDOW)


def _log_call(lecture_id, kind, model, key, cached, ms=None, usage=None, error=None) -> None:
    u = usage or {}
    with db() as con:
        con.execute("INSERT INTO claude_calls(lecture_id, kind, model, input_hash, cached, ms, in_tokens, "
                    "cache_tokens, out_tokens, error) VALUES (?,?,?,?,?,?,?,?,?,?)",
                    (lecture_id, kind, model, key, int(cached), ms, u.get("input_tokens"),
                     (u.get("cache_read_input_tokens") or 0) + (u.get("cache_creation_input_tokens") or 0),
                     u.get("output_tokens"), error))
    if not cached:
        log.info("claude %s %s lecture %s: %s ms, %s out tokens%s", model, kind, lecture_id, ms,
                 u.get("output_tokens"), f", error: {error}" if error else "")


def claude_call(prompt: str, *, system: str, model: str = NOTES_MODEL, schema: dict | None = None,
                lecture_id: int | None = None, kind: str = "", timeout: float = CLAUDE_TIMEOUT):
    """The only way LectureLab calls Claude. Runs `claude -p` on the Max plan with no API key in its
    environment, cut off from your own Claude Code setup (no hooks, CLAUDE.md, MCP servers or tools:
    they'd add tens of thousands of tokens to every call). Never --bare: it forces API-key auth.
    The prompt goes in on stdin (a transcript is far past Windows' 32 KB command line).
    Returns the text, or the parsed object when a JSON schema is given. Cached by input hash.
    Raises QuotaExceeded when the plan's limit is hit, and keeps refusing until it resets."""
    key = hashlib.sha256(json.dumps([model, system, schema, prompt], sort_keys=True).encode()).hexdigest()
    path = CLAUDE_CACHE / f"{key}.json"
    if path.exists():
        _log_call(lecture_id, kind, model, key, True)
        return json.loads(path.read_text(encoding="utf-8"))["output"]
    if (until := meta_get("claude_quota_until", 0)) > time.time():
        raise QuotaExceeded(f"Claude's usage limit is reached; trying again after "
                            f"{datetime.fromtimestamp(until):%H:%M}.", until)
    if (down := meta_get("claude_down", {})).get("until", 0) > time.time():
        raise ClaudeUnavailable(down["msg"], down["until"])
    if os.environ.get("LECTURELAB_CLAUDE") == "off":  # test servers: never spend the plan
        raise ClaudeUnavailable("Claude is turned off on this server (LECTURELAB_CLAUDE=off).", time.time() + 3600)
    try:
        exe = claude_exe()
    except ClaudeError as e:
        raise _claude_down(str(e)) from None
    cmd = [exe, "-p", "--model", model, "--output-format", "json", "--no-session-persistence",
           "--setting-sources=", "--strict-mcp-config", "--tools=", "--system-prompt", system]
    if schema:
        cmd += ["--json-schema", json.dumps(schema)]
    t0 = time.time()
    try:
        r = _claude_run(cmd, prompt, timeout)
    except subprocess.TimeoutExpired as e:
        _log_call(lecture_id, kind, model, key, False, int((time.time() - t0) * 1000), error="timeout")
        raise ClaudeError(f"Claude didn't answer within {timeout / 60:.0f} minutes.") from e
    ms = int((time.time() - t0) * 1000)
    try:
        d = json.loads(r.stdout)
    except ValueError:
        d = {"is_error": True, "result": ((r.stdout or "")[-400:] + "\n" + (r.stderr or "")[-400:]).strip()}
    msg = str(d.get("result") or "")
    if d.get("is_error") or r.returncode:
        if QUOTA_RE.search(msg) or d.get("api_error_status") == 429:
            reset = quota_reset_at(msg) if QUOTA_RE.search(msg) else time.time() + 900
            meta_set("claude_quota_until", reset)
            _log_call(lecture_id, kind, model, key, False, ms, d.get("usage"), "usage limit")
            raise QuotaExceeded(f"Claude's usage limit is reached; trying again after "
                                f"{datetime.fromtimestamp(reset):%H:%M}.", reset)
        _log_call(lecture_id, kind, model, key, False, ms, d.get("usage"), msg[:300] or f"exit {r.returncode}")
        if AUTH_RE.search(msg):
            raise _claude_down("Claude Code is signed out. Run claude in a terminal and sign in with your "
                               "Max plan; notes carry on within the hour.")
        raise ClaudeError(f"Claude Code failed: {msg[:300] or f'exit code {r.returncode}'}")
    out = msg
    if schema:
        out = d.get("structured_output")
        if out is None:
            try:
                out = json.loads(msg)
            except ValueError:
                raise ClaudeError("Claude's answer wasn't the JSON that was asked for.") from None
    if not out:  # never cache an empty answer: it would come back empty forever
        _log_call(lecture_id, kind, model, key, False, ms, d.get("usage"), "empty answer")
        raise ClaudeError("Claude's answer was empty.")
    CLAUDE_CACHE.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps({"model": model, "kind": kind, "output": out}), encoding="utf-8")
    tmp.replace(path)
    _log_call(lecture_id, kind, model, key, False, ms, d.get("usage"))
    return out


def _claude_down(msg: str) -> ClaudeUnavailable:
    until = time.time() + 3600
    meta_set("claude_down", {"until": until, "msg": msg})
    log.warning("claude unavailable until %s: %s", datetime.fromtimestamp(until).strftime("%H:%M"), msg)
    return ClaudeUnavailable(msg, until)


def fmt_t(sec: float) -> str:
    s = int(sec)
    return f"{s // 3600}:{s % 3600 // 60:02d}:{s % 60:02d}" if s >= 3600 else f"{s // 60}:{s % 60:02d}"


def jump(sec: float) -> str:
    return f"[▶ {fmt_t(sec)}](#t={int(sec)})"


def transcript_text(segs: list[dict], titles: dict[int, str] | None = None, max_sec: float = 25,
                    max_chars: int = 320, clock=None) -> str:
    """Segments as "[m:ss] text" lines of up to ~25 s (Whisper's segments can be single words),
    with a "--- Slide N: title ---" line wherever the lecture moves to another slide."""
    out, cur, slide = [], None, None
    for s in segs:
        if s.get("slide_no") and s["slide_no"] != slide:
            slide = s["slide_no"]
            if cur:
                out.append(cur)
                cur = None
            title = (titles or {}).get(slide) or ""
            out.append(f"--- Slide {slide}{': ' + title if title else ''} ---")
        if cur and (s["start"] - cur[0] > max_sec or len(cur[1]) > max_chars):
            out.append(cur)
            cur = None
        cur = [cur[0], cur[1] + " " + s["text"].strip()] if cur else [s["start"], s["text"].strip()]
    if cur:
        out.append(cur)
    return "\n".join(x if isinstance(x, str) else f"[{(clock or fmt_t)(x[0])}] {x[1]}" for x in out)


def notes_inputs(lecture_id: int) -> dict:
    lec = get_lecture(lecture_id)
    if not lec:
        raise KeyError(lecture_id)
    course = find_course(lec["course_id"]) if lec["course_id"] else None
    with db() as con:
        segs = rows(con.execute("SELECT start, end, text, slide_no FROM segments WHERE lecture_id=? ORDER BY idx",
                                (lecture_id,)))
        slides = rows(con.execute("SELECT no, title, text FROM slides WHERE deck_id=? ORDER BY no",
                                  (lec["deck_id"] or -1,)))
    return {"lec": lec, "course": course, "segs": segs, "slides": {s["no"]: s for s in slides},
            "duration": max([lec["duration_sec"] or 0] + [s["end"] for s in segs])}


def _header(inp: dict) -> str:
    lec, c = inp["lec"], inp["course"]
    vocab = ", ".join(v.strip() for v in (c or {}).get("vocabulary", "").splitlines() if v.strip())
    return (f"Course: {c['code']} {c['name']}\n" if c else "") + (f"Course vocabulary: {vocab}\n" if vocab else "") + \
        f"Lecture: {lec['title']}" + (f" ({lec['date']})" if lec["date"] else "") + \
        f", {fmt_t(inp['duration'])} long\n"


def _slide_block(inp: dict, nos) -> str:
    out = []
    for no in sorted(set(nos)):
        s = inp["slides"].get(no)
        if s:
            body = re.sub(r"\n{3,}", "\n\n", (s["text"] or "").strip())[:2500]
            out.append(f"[Slide {no}] {s['title'] or ''}\n{body}".strip())
    return "\n\n".join(out)


OUTLINE_SYSTEM = """You outline university lecture recordings so study notes can be written one section at a time.

Split the lecture into sections, one per coherent topic, in order, covering the whole recording without gaps or overlaps. Aim for sections of roughly 5 to 12 minutes; a short recording may be one or two sections. Use the times in the transcript ([m:ss], seconds from the start) for start and end, in seconds. List the slide numbers each section uses ("--- Slide N ---" lines mark when the lecturer moved to slide N).

kind: "content" for teaching; "logistics" for announcements, deadlines, exam dates and course admin; "off-topic" for chatter before the lecture starts, breaks, or students talking among themselves. Keep logistics and off-topic as their own sections so they don't mix into the teaching.

title: a specific topic title a student would scan for ("Left recursion and why LL parsers can't handle it", not "Part 2"). gist: one sentence on what the section teaches.

The transcript is automatic speech recognition, so technical terms may be misheard; use the slides and the course vocabulary to read them correctly."""

OUTLINE_SCHEMA = {
    "type": "object", "required": ["sections"],
    "properties": {"sections": {"type": "array", "minItems": 1, "items": {
        "type": "object", "required": ["title", "start", "end", "kind", "slides", "gist"],
        "properties": {"title": {"type": "string"}, "start": {"type": "number"}, "end": {"type": "number"},
                       "kind": {"type": "string", "enum": ["content", "logistics", "off-topic"]},
                       "slides": {"type": "array", "items": {"type": "integer"}}, "gist": {"type": "string"}}}}}}


def clean_outline(sections: list[dict], duration: float, min_sec: float = 20) -> list[dict]:
    """Sorted, contiguous from 0 to the end, with slivers merged into their neighbour."""
    secs = sorted((dict(s) for s in sections if s.get("title")), key=lambda s: float(s["start"]))
    out: list[dict] = []
    for s in secs:
        s["start"] = min(max(0.0, float(s["start"])), duration)
        if out and s["start"] - out[-1]["start"] < min_sec:
            out[-1]["slides"] = sorted(set(out[-1]["slides"]) | set(s.get("slides") or []))
            continue
        s["slides"] = sorted({int(n) for n in s.get("slides") or [] if int(n) > 0})
        out.append(s)
    if not out:
        return [{"title": "Lecture", "start": 0.0, "end": duration, "kind": "content", "slides": [], "gist": ""}]
    out[0]["start"] = 0.0
    for a, b in zip(out, out[1:]):
        a["end"] = b["start"]
    out[-1]["end"] = duration
    if len(out) > 1 and out[-1]["end"] - out[-1]["start"] < min_sec:
        out[-2]["end"] = out[-1]["end"]
        out[-2]["slides"] = sorted(set(out[-2]["slides"]) | set(out[-1]["slides"]))
        out.pop()
    return out


SECTION_SYSTEM = """You write one section of study notes for a university lecture, from its transcript and slides. The student reading them missed the lecture and won't rewatch it, so the notes must carry everything the lecturer taught in this part: every concept, definition, example, caveat, and anything the lecturer stressed.

Rules:
- Use only what the lecturer said and what the slides show. Don't add facts, examples or explanations from your own knowledge. If you can't tell what was meant, say so in an [!unclear] callout instead of guessing.
- The transcript is automatic speech recognition. Silently fix misheard technical terms using the slides and the course vocabulary (for example "back us now are form" is "Backus-Naur Form").
- Write Markdown and start straight with the content: no heading for the section (one is added for you), no preamble, no closing summary.
- Cover only this section's part of the lecture. Other sections of the outline are written separately, so don't repeat what they cover.
- The slides given may include ones the lecturer skipped or hasn't reached yet. Write up a slide only as far as the transcript of this section discusses it.
- Use ### subheadings when the section covers more than one idea.
- Start each subtopic, or each paragraph where the lecture moves on, with a jump link like [▶ 12:34](#t=754): 12:34 is a time from the transcript lines and 754 is that same time in whole seconds. Only use times that appear in the transcript you're given.
- Cite slides inline, like (Slide 7) or (Slides 7 to 9), where a point comes from a slide.
- Definitions: **Term**: what it means, in the lecturer's sense.
- Worked examples: reproduce them step by step as the lecturer did, including the result.
- Code goes in fenced blocks with a language (```java, ```python, ```haskell, ...). Grammars go in ```ebnf or ```bnf blocks. Keep code as shown on the slides or dictated; don't write code the lecture didn't have.
- Math in KaTeX: $...$ inline, $$...$$ on its own line.
- Callouts are blockquotes that start with a tag:
  > [!exam] when the lecturer says something will be on an exam, quiz or homework, or must be known.
  > [!emphasis] when the lecturer stresses or repeats a point.
  > [!unclear] when the explanation or the audio is unclear, or the lecturer contradicts themselves; say what is unclear.
- Plain, direct sentences. Bullets for lists of facts, prose for reasoning. No filler."""

LOGISTICS_SYSTEM = """You write the course-logistics part of a lecture's study notes: announcements, deadlines, exam dates and format, grading, office hours. Use only what the transcript says. Markdown bullets, each starting with a jump link like [▶ 12:34](#t=754) (a time from the transcript and the same time in whole seconds). Put anything about what will be on an exam in a > [!exam] callout. No heading, no preamble. If nothing in it matters to a student, reply with exactly: (nothing to note)"""


def _section_prompt(inp: dict, outline: list[dict], i: int) -> str:
    sec = outline[i]
    segs = [s for s in inp["segs"] if sec["start"] - 1 <= s["start"] < sec["end"]]
    nos = set(sec["slides"]) | {s["slide_no"] for s in segs if s["slide_no"]}
    titles = {n: (s["title"] or "") for n, s in inp["slides"].items()}
    lines = "\n".join(f"{'►' if j == i else ' '} {j + 1}. {s['title']} ({fmt_t(s['start'])} to {fmt_t(s['end'])})"
                      for j, s in enumerate(outline))
    slides = _slide_block(inp, nos)
    return (_header(inp) + f"\nOutline of the whole lecture (write the section marked ►):\n{lines}\n"
            + (f"\nSlides for this section:\n{slides}\n" if slides else "\nThis lecture has no slides.\n")
            + "\nTranscript of this section (times are m:ss from the start of the recording):\n"
            + transcript_text(segs, titles))


def no_heading(md: str) -> str:
    return re.sub(r"\A\s*#{1,2} [^\n]*\n+", "", md.strip())


def tidy_section(md: str, start: float, end: float, margin: float = 90) -> str:
    """Drop a heading the model added anyway, and jump links that point outside this section.
    A callout written after a bullet or a jump link ("- [▶ 1:02](#t=62) > [!unclear] ...") isn't
    a callout in Markdown, so it's moved to the start of its own line."""
    md = no_heading(md)
    md = re.sub(r"(?m)^[ \t]*(?:[-*] +)?(\[▶ [0-9:]+\]\(#t=\d+\)):?[ \t]*> (\[![a-z]+\])", r"> \2 \1", md)

    def fix(m):
        t = int(m[2])
        return m[0] if start - margin <= t <= end + margin else f"({m[1]})"
    return re.sub(r"\[▶ ([0-9:]+)\]\(#t=(\d+)\)", fix, md).strip()


STUDY_SYSTEM = """You turn a lecture's study notes into a summary, flashcards and a practice quiz.

Use only what the notes say. Every card and question gets t: the time in seconds where its material is taught, taken from the nearest jump link before it in the notes ([▶ 12:34](#t=754) means t=754).

summary: Markdown. One or two sentences on what the lecture covers, then 3 to 7 bullets with the key takeaways.

open_questions: things the lecture left unclear or unanswered (from the [!unclear] callouts and gaps you notice), as short questions. Empty if there are none.

cards: one per definition, rule, distinction, step or result worth remembering. kind "cloze": front is a sentence from the material with the key term hidden as {{c1::term}} (a second hidden part in the same card is {{c2::...}}), back is a short note of context. A hidden part must never contain "::" (Anki reads it as a hint): for a name like std::vector or a Haskell type signature, write a qa card instead. kind "qa": a direct question on the front, a short complete answer on the back. Cards should be atomic and test understanding, not trivia. No cards about course logistics. Typically 10 to 30 for a full lecture, fewer for a short one.

quiz: 10 to 15 questions for a full lecture (4 to 6 if the recording is under 20 minutes), mixing kind "mcq" (exactly 4 choices, answer is the exact text of the right choice, distractors are plausible mistakes a student would make) and kind "short" (answer is a model answer of one to three sentences). explanation says why the answer is right, and for mcq why the most tempting wrong choice is wrong. Favour questions that apply the material (trace this grammar, which derivation, what does this code print) over recall.

Keep code and grammar formatting (```), and KaTeX math ($...$), as in the notes."""

STUDY_SCHEMA = {
    "type": "object", "required": ["summary", "open_questions", "cards", "quiz"],
    "properties": {
        "summary": {"type": "string"},
        "open_questions": {"type": "array", "items": {"type": "string"}},
        "cards": {"type": "array", "items": {
            "type": "object", "required": ["kind", "front", "back", "t"],
            "properties": {"kind": {"type": "string", "enum": ["cloze", "qa"]}, "front": {"type": "string"},
                           "back": {"type": "string"}, "t": {"type": "number"}}}},
        "quiz": {"type": "array", "items": {
            "type": "object", "required": ["kind", "question", "choices", "answer", "explanation", "t"],
            "properties": {"kind": {"type": "string", "enum": ["mcq", "short"]}, "question": {"type": "string"},
                           "choices": {"type": "array", "items": {"type": "string"}}, "answer": {"type": "string"},
                           "explanation": {"type": "string"}, "t": {"type": "number"}}}}}}

CHECK_SYSTEM = """You fact-check a lecture's study notes, flashcards and quiz against the lecture itself: its full transcript and its slides. The student will study from these instead of rewatching, so anything the lecture didn't say is a problem.

For each claim, example, definition, card and quiz answer: is it supported by the transcript or the slides? The transcript is automatic speech recognition, so a correctly fixed mishearing of a technical term counts as supported.
- Not supported, or contradicted: remove it. If removing it leaves a section that no longer reads well, repair the wording around the gap, without adding new material.
- Probably right but you can't confirm it (the audio is garbled there, or it depends on a slide you can't see): keep it, and put a callout right after it: > [!check] what to verify.
- Supported: leave it exactly as it is. Don't rewrite for style, don't reorder, don't add anything, and keep every jump link [▶ m:ss](#t=N), slide citation, code block and math as they are.

Return:
- edits: every change to a section, as an exact find-and-replace: the section's number, find (text copied character for character from that section, long enough to occur only once in it, usually the whole sentence or bullet) and replace (the new text, or "" to remove it). To flag something, replace it with itself followed by the callout. Never return a whole section.
- summary: the new summary if you changed it, otherwise null.
- drop_cards and drop_quiz: the numbers of cards and quiz questions that are wrong or unsupported. A quiz question whose stated answer is wrong is dropped, not fixed.
- issues: one entry per change: where ("section 3", "card 12", "quiz 4", "summary"), verdict ("removed" or "flagged"), the claim, and the reason."""

CHECK_SCHEMA = {
    "type": "object", "required": ["edits", "summary", "drop_cards", "drop_quiz", "issues"],
    "properties": {
        "edits": {"type": "array", "items": {"type": "object", "required": ["number", "find", "replace"],
                                              "properties": {"number": {"type": "integer"}, "find": {"type": "string"},
                                                             "replace": {"type": "string"}}}},
        "summary": {"type": ["string", "null"]},
        "drop_cards": {"type": "array", "items": {"type": "integer"}},
        "drop_quiz": {"type": "array", "items": {"type": "integer"}},
        "issues": {"type": "array", "items": {"type": "object", "required": ["where", "verdict", "claim", "reason"],
                                               "properties": {"where": {"type": "string"},
                                                              "verdict": {"type": "string", "enum": ["removed", "flagged"]},
                                                              "claim": {"type": "string"},
                                                              "reason": {"type": "string"}}}}}}


def _numbered(items: list[dict], fmt) -> str:
    return "\n".join(f"{n}. {fmt(x)}" for n, x in enumerate(items, 1))


def _card_line(c: dict) -> str:
    return f"[{c['kind']}] {c['front']} || {c['back']} (t={int(c['t'])})"


def _quiz_line(q: dict) -> str:
    ch = "".join(f"\n   - {x}" for x in q.get("choices") or [])
    return f"[{q['kind']}] {q['question']}{ch}\n   Answer: {q['answer']}\n   Why: {q['explanation']}"


def notes_markdown(sections: list[dict]) -> str:
    return "\n\n".join(f"## {s['title']}\n\n{jump(s['start'])}"
                       + (f" (Slides {s['slides'][0]} to {s['slides'][-1]})" if len(s["slides"]) > 1
                          else f" (Slide {s['slides'][0]})" if s["slides"] else "")
                       + f"\n\n{s['markdown']}" for s in sections if s.get("markdown"))


def generate_notes(lecture_id: int, say=lambda m: None) -> dict:
    """Outline, then each section from its transcript slice and slides, then summary, cards and
    quiz, then an Opus check against the transcript. Every call is cached, so a re-run after a
    failure (or the usage limit) picks up where it stopped."""
    inp = notes_inputs(lecture_id)
    if not inp["segs"]:
        raise ClaudeError("There's no transcript to write notes from yet.")
    titles = {n: (s["title"] or "") for n, s in inp["slides"].items()}
    full = transcript_text(inp["segs"], titles)
    say("outlining")
    slides = _slide_block(inp, inp["slides"])
    out = claude_call(_header(inp) + (f"\nSlides:\n{slides}\n" if slides else "\nThis lecture has no slides.\n")
                      + f"\nTranscript:\n{full}", system=OUTLINE_SYSTEM, schema=OUTLINE_SCHEMA,
                      lecture_id=lecture_id, kind="outline")
    outline = clean_outline(out["sections"], inp["duration"])
    todo = [i for i, s in enumerate(outline) if s["kind"] != "off-topic"]
    for n, i in enumerate(todo, 1):
        say(f"writing section {n} of {len(todo)}")
        system = LOGISTICS_SYSTEM if outline[i]["kind"] == "logistics" else SECTION_SYSTEM
        md = claude_call(_section_prompt(inp, outline, i), system=system, lecture_id=lecture_id, kind="section")
        md = tidy_section(md, outline[i]["start"], outline[i]["end"])
        outline[i]["markdown"] = "" if md.strip("() \n").lower() == "nothing to note" else md
    sections = [s for s in outline if s.get("markdown")]
    if not sections:
        raise ClaudeError("Nothing in this recording was teaching, so there are no notes to write.")
    say("writing the summary, flashcards and quiz")
    body = notes_markdown(sections)
    study = claude_call(_header(inp) + f"\nNotes:\n{body}", system=STUDY_SYSTEM, schema=STUDY_SCHEMA,
                        lecture_id=lecture_id, kind="study")
    cards = [c for c in study["cards"] if c["front"].strip() and c["back"].strip()]
    quiz = [q for q in study["quiz"] if q["question"].strip() and
            (q["kind"] != "mcq" or (len(q["choices"]) >= 2 and q["answer"] in q["choices"]))]
    summary = no_heading(study["summary"])
    say("checking the notes against the transcript")
    numbered = "\n\n".join(f"=== Section {n}: {s['title']} ===\n{s['markdown']}" for n, s in enumerate(sections, 1))
    check = claude_call(
        _header(inp) + (f"\nSlides:\n{slides}\n" if slides else "\nThis lecture has no slides.\n")
        + f"\nFull transcript:\n{full}\n\n### The notes to check\n\nSummary:\n{summary}\n\n{numbered}"
        + f"\n\nFlashcards (front || back):\n{_numbered(cards, _card_line)}"
        + f"\n\nQuiz:\n{_numbered(quiz, _quiz_line)}",
        system=CHECK_SYSTEM, model=CHECK_MODEL, schema=CHECK_SCHEMA, lecture_id=lecture_id, kind="check")
    issues = list(check["issues"])
    for n, s in apply_check_edits(sections, check["edits"]):
        issues.append({"where": f"section {n}", "verdict": "flagged", "claim": s[:300],
                       "reason": "The check wanted to change this but its edit didn't match the notes, so it "
                                 "wasn't applied. Verify it against the recording."})
    new = no_heading(check.get("summary") or "")
    if new and len(new) >= 0.4 * len(summary):  # a summary cut to a stub is a bad answer, not a correction
        summary = new
    cards = [c for n, c in enumerate(cards, 1) if n not in set(check["drop_cards"])]
    quiz = [q for n, q in enumerate(quiz, 1) if n not in set(check["drop_quiz"])]
    return {"outline": outline, "sections": sections, "summary": summary,
            "open_questions": [q.strip() for q in study["open_questions"] if q.strip()],
            "cards": cards, "quiz": quiz, "issues": issues,
            "removed": sum(i["verdict"] == "removed" for i in issues),
            "flagged": sum(i["verdict"] == "flagged" for i in issues)}


def apply_check_edits(sections: list[dict], edits: list[dict]) -> list[tuple[int, str]]:
    """Apply the check's find-and-replace edits in place. An edit whose text isn't in its section
    exactly once, or that would cut most of the section, is skipped (returned), so a confused
    check can never wipe out a section."""
    skipped, touched = [], set()
    for e in edits:
        n, find = e["number"], e["find"]
        if not (1 <= n <= len(sections)) or not find.strip():
            continue
        s = sections[n - 1]
        after = len(s["markdown"]) - len(find) + len(e["replace"])
        if s["markdown"].count(find) != 1 or after < 0.4 * len(s["markdown"]):
            skipped.append((n, find))
            continue
        s["markdown"] = s["markdown"].replace(find, e["replace"], 1)
        touched.add(n)
    for n in touched:
        s = sections[n - 1]
        s["markdown"] = tidy_section(re.sub(r"\n{3,}", "\n\n", s["markdown"]), s["start"], s["end"])
    return skipped


CLOZE_RE = re.compile(r"\{\{c\d+::(.*?)\}\}", re.S)  # Anki's is multi-line too


def fix_card(c: dict) -> dict:
    """Anki reads {{c1::answer::hint}}, so a cloze whose answer holds "::" (std::vector, a Haskell
    type) would show part of the answer as a hint. That card becomes a question and answer."""
    hidden = CLOZE_RE.findall(c["front"])
    if c["kind"] != "cloze" or not any("::" in h for h in hidden):
        return c
    back = "; ".join(f"`{h}`" if "::" in h and "`" not in h else h for h in hidden)
    return {**c, "kind": "qa", "front": "Fill in the blank: " + CLOZE_RE.sub("____", c["front"]),
            "back": back + ("\n\n" + c["back"].strip() if c["back"].strip() else "")}


def _card_key(front: str) -> str:
    return re.sub(r"\W+", " ", front).strip().lower()


def save_notes(lecture_id: int, notes: dict, generator: str) -> None:
    """Store the notes, keeping your review progress on any card whose front didn't change."""
    lec = get_lecture(lecture_id)
    dur = notes_inputs(lecture_id)["duration"]
    with db() as con:
        con.execute("BEGIN IMMEDIATE")
        old = {_card_key(r["front"]): dict(r) for r in con.execute(
            "SELECT front, box, due, reviews, lapses, last_reviewed FROM cards WHERE lecture_id=?", (lecture_id,))}
        con.execute("DELETE FROM cards WHERE lecture_id=?", (lecture_id,))
        con.execute("DELETE FROM quiz WHERE lecture_id=?", (lecture_id,))
        for c in map(fix_card, notes["cards"]):
            p = old.get(_card_key(c["front"]), {})
            con.execute("INSERT INTO cards(lecture_id, kind, front, back, t, box, due, reviews, lapses, last_reviewed) "
                        "VALUES (?,?,?,?,?,?,?,?,?,?)",
                        (lecture_id, c["kind"], c["front"].strip(), c["back"].strip(), min(max(0, c["t"]), dur),
                         p.get("box", 1), p.get("due"), p.get("reviews", 0), p.get("lapses", 0), p.get("last_reviewed")))
        for i, q in enumerate(notes["quiz"]):
            con.execute("INSERT INTO quiz(lecture_id, idx, kind, question, choices, answer, explanation, t) "
                        "VALUES (?,?,?,?,?,?,?,?)",
                        (lecture_id, i, q["kind"], q["question"].strip(), json.dumps(q.get("choices") or []),
                         q["answer"].strip(), q["explanation"].strip(), min(max(0, q["t"]), dur)))
        # an upsert, not a replace: export_paths survives, so a renamed lecture's old files get cleaned up
        con.execute("""INSERT INTO notes(lecture_id, outline, sections, summary, open_questions, issues,
                       removed, flagged, generator, transcript_source, deck_id, slides_aligned_at, created_at)
                       VALUES (?,?,?,?,?,?,?,?,?,?,?,?,datetime('now'))
                       ON CONFLICT(lecture_id) DO UPDATE SET outline=excluded.outline, sections=excluded.sections,
                       summary=excluded.summary, open_questions=excluded.open_questions, issues=excluded.issues,
                       removed=excluded.removed, flagged=excluded.flagged, generator=excluded.generator,
                       transcript_source=excluded.transcript_source, deck_id=excluded.deck_id,
                       slides_aligned_at=excluded.slides_aligned_at, created_at=excluded.created_at""",
                    (lecture_id, json.dumps(notes["outline"]), json.dumps(notes["sections"]), notes["summary"],
                     json.dumps(notes["open_questions"]), json.dumps(notes["issues"]), notes["removed"],
                     notes["flagged"], generator, lec["transcript_source"], lec["deck_id"], lec["slides_aligned_at"]))
        con.execute("UPDATE lectures SET notes_error=NULL, updated_at=datetime('now'), status=CASE WHEN status "
                    "IN ('queued_transcribe','transcribing','duplicate') THEN status ELSE 'notes_done' END WHERE id=?", (lecture_id,))
    try:  # the notes are safe in the database: search catches up later, and the next sync exports them
        index_lecture(lecture_id)
        export_lecture(lecture_id)
    except Exception:
        log.exception("indexing or exporting lecture %s failed", lecture_id)


def enqueue_notes(lecture_id: int, mode: str = "claude", force: bool = False) -> int:
    """Queue notes for a lecture unless they're already queued (returns that job then)."""
    with db() as con:
        gen = (con.execute("SELECT generator FROM notes WHERE lecture_id=?", (lecture_id,)).fetchone() or [None])[0]
        if mode == "offline" and gen == "claude" and not force:
            raise ValueError("This lecture already has notes from Claude. Offline notes would replace them "
                             "and their flashcards; ask again with force to do it anyway.")
        if q := con.execute("SELECT id FROM jobs WHERE kind='notes' AND lecture_id=? AND status='queued'",
                            (lecture_id,)).fetchone():
            return q[0]
        title = (con.execute("SELECT title FROM lectures WHERE id=?", (lecture_id,)).fetchone() or [""])[0]
        con.execute("UPDATE lectures SET status='notes_queued', notes_error=NULL WHERE id=? AND status NOT IN "
                    "('notes_running','queued_transcribe','transcribing','duplicate')", (lecture_id,))
    jid = enqueue("notes", {"lecture_id": lecture_id, "title": title, "mode": mode}, lecture_id=lecture_id)
    wake_workers()
    return jid


AUTO_NOTES = os.environ.get("LECTURELAB_AUTO_NOTES", "on") != "off"


def notes_after_slides(lecture_id: int) -> None:
    """Once a lecture's slides are settled (lined up, none, or failed), write its notes if it has
    none yet. A lecture whose deck is still downloading, or whose course hasn't been scanned for
    decks yet, waits: notes written now would miss the slides."""
    if not AUTO_NOTES:
        return
    with db() as con:
        busy = con.execute("SELECT 1 FROM jobs WHERE kind='notes' AND lecture_id=? AND status IN ('queued','running')",
                           (lecture_id,)).fetchone()
        have = con.execute("SELECT 1 FROM notes WHERE lecture_id=?", (lecture_id,)).fetchone()
        segs = con.execute("SELECT 1 FROM segments WHERE lecture_id=? LIMIT 1", (lecture_id,)).fetchone()
    if segs and not busy and not have:
        enqueue_notes(lecture_id)


# A notes run that stopped (failed, or Claude signed out) leaves the lecture as it was before it:
# a lecture that already has notes still has them, so it doesn't drop back to "transcribed".
_NOTES_STOPPED = ("(CASE WHEN EXISTS (SELECT 1 FROM notes n WHERE n.lecture_id=lectures.id AND n.exported_at IS NOT NULL) "
                  "THEN 'exported' WHEN EXISTS (SELECT 1 FROM notes n WHERE n.lecture_id=lectures.id) THEN 'notes_done' "
                  "ELSE 'transcribed' END)")


def job_notes(job: dict) -> dict:
    """payload: {lecture_id, mode: claude|offline}."""
    lid = job["payload"]["lecture_id"]
    say = lambda m: job_update(job["id"], progress=m)  # noqa: E731
    if not get_lecture(lid):
        return {"skipped": "lecture deleted"}
    mode = job["payload"].get("mode", "claude")
    with db() as con:
        con.execute("UPDATE lectures SET status='notes_running', notes_error=NULL WHERE id=? "
                    "AND status NOT IN ('queued_transcribe','transcribing','duplicate')", (lid,))
    try:
        notes = offline_notes(lid) if mode == "offline" else generate_notes(lid, say)
        save_notes(lid, notes, mode)
    except QuotaExceeded as e:  # the usage limit, or Claude Code signed out: try again later
        with db() as con:
            if isinstance(e, ClaudeUnavailable):
                con.execute(f"UPDATE lectures SET status={_NOTES_STOPPED} WHERE id=? AND status='notes_running'", (lid,))
            else:
                con.execute("UPDATE lectures SET status='waiting_quota' WHERE id=? AND status='notes_running'", (lid,))
        wait = max(60, int(e.reset_at - time.time()))
        job_update(job["id"], status="queued", retry_at=_sqlnow(wait), owner_pid=None, attempts=0, progress=str(e))
        raise RetryLater(str(e)) from e
    except Exception as e:
        with db() as con:
            con.execute(f"UPDATE lectures SET notes_error=?, status=CASE status WHEN 'notes_running' THEN "
                        f"{_NOTES_STOPPED} ELSE status END WHERE id=?", (user_error(e, "Writing the notes"), lid))
        raise
    return {"sections": len(notes["sections"]), "cards": len(notes["cards"]), "quiz": len(notes["quiz"]),
            "removed": notes["removed"], "flagged": notes["flagged"], "mode": mode}


def job_slides_then_notes(job: dict) -> dict:
    """The slides job, then notes once the slides are settled either way."""
    lid = job["payload"]["lecture_id"]
    try:
        res = job_slides(job)
    except Exception:
        notes_after_slides(lid)  # notes without slides beat no notes; regenerate once they're fixed
        raise
    if "segments" in res:  # newly lined up: slides are searchable at the time they were shown
        index_lecture(lid)
    wait = res.get("waiting")
    # "no slides": no deck and no video to read. Settled, unless the course's scan may still find one.
    if not wait or (wait == "no slides" and slides_scan_done((get_lecture(lid) or {}).get("course_id"))):
        notes_after_slides(lid)
    return res


JOB_HANDLERS["notes"] = job_notes
JOB_HANDLERS["slides"] = job_slides_then_notes


def stale_reason(n, lec: dict) -> str | None:
    if n["transcript_source"] != lec["transcript_source"]:
        return "The lecture was transcribed again after these notes were written."
    if n["deck_id"] != lec["deck_id"]:
        return "These notes were written with other slides, or before any slides were found."
    if n["slides_aligned_at"] != lec["slides_aligned_at"]:
        return "The slides were lined up again after these notes were written."
    return None


def lecture_notes(lecture_id: int) -> dict:
    lec = get_lecture(lecture_id)
    if not lec:
        raise KeyError(lecture_id)
    with db() as con:
        n = con.execute("SELECT * FROM notes WHERE lecture_id=?", (lecture_id,)).fetchone()
        job = con.execute("SELECT status, progress, retry_at FROM jobs WHERE kind='notes' AND lecture_id=? AND "
                          "status IN ('queued','running') ORDER BY status='running' DESC, id LIMIT 1",
                          (lecture_id,)).fetchone()
        counts = con.execute("SELECT (SELECT COUNT(*) FROM cards WHERE lecture_id=?), "
                             "(SELECT COUNT(*) FROM quiz WHERE lecture_id=?)", (lecture_id, lecture_id)).fetchone()
    if job and job["status"] == "running":
        state = "writing"
    elif lec["status"] == "waiting_quota":
        state = "waiting_quota"
    elif job:
        state = "queued"
    elif lec["notes_error"]:
        state = "error"
    elif n:
        state = "done"
    else:
        state = "none"
    out = {"state": state, "progress": job and job["progress"], "error": lec["notes_error"],
           "retry_at": job and job["retry_at"]}
    if n:
        out.update({
            "summary": n["summary"], "sections": json.loads(n["sections"]),
            "open_questions": json.loads(n["open_questions"]), "issues": json.loads(n["issues"]),
            "removed": n["removed"], "flagged": n["flagged"], "generator": n["generator"],
            "created_at": n["created_at"], "cards": counts[0], "quiz": counts[1],
            # written from an older transcript or other slides: worth regenerating
            "stale": stale_reason(n, lec) is not None, "stale_reason": stale_reason(n, lec)})
    return out


# ---------------------------------------------------------------- study: review, quiz, Anki, search

LEITNER_DAYS = [1, 2, 4, 8, 16]  # box n comes back after LEITNER_DAYS[n-1] days


def _today(offset: int = 0) -> str:
    return Date.fromordinal(Date.today().toordinal() + offset).isoformat()


def due_cards(course_id: int | None = None, lecture_id: int | None = None, limit: int = 100) -> dict:
    """Cards due today or new, oldest due first, plus counts per Leitner box. due is a local date
    (YYYY-MM-DD); a new card has none and isn't in a box until you first answer it."""
    where, args = ["1=1"], []
    if lecture_id:
        where.append("k.lecture_id=?")
        args.append(lecture_id)
    if course_id:
        where.append("l.course_id=?")
        args.append(course_id)
    w = " AND ".join(where)
    with db() as con:
        due = rows(con.execute(
            f"""SELECT k.id, k.lecture_id, k.kind, k.front, k.back, k.t, k.box, k.due, k.reviews, k.lapses,
                       l.title AS lecture_title, l.date AS lecture_date, c.code AS course_code
                FROM cards k JOIN lectures l ON l.id=k.lecture_id LEFT JOIN courses c ON c.id=l.course_id
                WHERE {w} AND (k.due IS NULL OR k.due<=?)
                ORDER BY k.due IS NULL, k.due, k.box, k.id LIMIT ?""", (*args, _today(), limit)))
        boxes = dict(con.execute(f"SELECT k.box, COUNT(*) FROM cards k JOIN lectures l ON l.id=k.lecture_id "
                                 f"WHERE {w} AND k.due IS NOT NULL GROUP BY k.box", args).fetchall())
        counts = con.execute(f"""SELECT COUNT(*), SUM(k.due IS NULL), SUM(k.due IS NOT NULL AND k.due<=?)
                                 FROM cards k JOIN lectures l ON l.id=k.lecture_id WHERE {w}""",
                             (_today(), *args)).fetchone()
    return {"cards": due, "total": counts[0], "new": counts[1] or 0, "due": counts[2] or 0,
            "boxes": [boxes.get(b, 0) for b in range(1, 6)]}


def review_card(card_id: int, correct: bool) -> dict:
    """Leitner: right moves the card up a box and pushes it out (1, 2, 4, 8, 16 days); wrong sends
    it back to box 1 for tomorrow. Returns the card as it was too, so the UI can undo."""
    with db() as con:
        r = con.execute("SELECT * FROM cards WHERE id=?", (card_id,)).fetchone()
        if not r:
            raise KeyError(card_id)
        before = dict(r)
        box = min(5, before["box"] + 1) if correct else 1
        due = _today(LEITNER_DAYS[box - 1])
        con.execute("UPDATE cards SET box=?, due=?, reviews=reviews+1, lapses=lapses+?, last_reviewed=datetime('now') "
                    "WHERE id=?", (box, due, 0 if correct else 1, card_id))
        after = dict(con.execute("SELECT * FROM cards WHERE id=?", (card_id,)).fetchone())
    return {"card": after, "before": {k: before[k] for k in ("box", "due", "reviews", "lapses", "last_reviewed")}}


def undo_review(card_id: int, before: dict) -> dict:
    if not {"box", "reviews", "lapses"} <= before.keys():
        raise ValueError("before needs box, reviews and lapses")
    if not 1 <= int(before["box"]) <= len(LEITNER_DAYS) or not re.fullmatch(r"(\d{4}-\d\d-\d\d)?",
                                                                              before.get("due") or ""):
        raise ValueError("before has a box outside 1 to 5 or a due date that isn't YYYY-MM-DD")
    with db() as con:
        n = con.execute("UPDATE cards SET box=?, due=?, reviews=?, lapses=?, last_reviewed=? WHERE id=?",
                        (int(before["box"]), before.get("due"), int(before["reviews"]), int(before["lapses"]),
                         before.get("last_reviewed"), card_id)).rowcount
        if not n:
            raise KeyError(card_id)
        return dict(con.execute("SELECT * FROM cards WHERE id=?", (card_id,)).fetchone())


def lecture_quiz(lecture_id: int) -> dict:
    with db() as con:
        qs = rows(con.execute("SELECT id, idx, kind, question, choices, t FROM quiz WHERE lecture_id=? ORDER BY idx",
                              (lecture_id,)))
        last = con.execute("SELECT score, total, at FROM quiz_attempts WHERE lecture_id=? ORDER BY id DESC LIMIT 1",
                           (lecture_id,)).fetchone()
    for q in qs:
        q["choices"] = json.loads(q["choices"])
    return {"questions": qs, "last_attempt": dict(last) if last else None}


def _norm_answer(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip().lower())


def grade_quiz(lecture_id: int, answers: list[dict], check: bool = False) -> dict:
    """answers: [{id, answer, correct?}]. Multiple choice is marked here; for a short answer you
    compare yours with the model answer and say whether you got it (correct). check: mark and show
    the model answers without saving an attempt (an unmarked short answer comes back correct None);
    otherwise the attempt is saved, and an unmarked short answer counts as wrong."""
    with db() as con:
        qs = {r["id"]: dict(r) for r in con.execute("SELECT * FROM quiz WHERE lecture_id=?", (lecture_id,))}
        given = {int(a["id"]): a for a in answers if int(a.get("id", -1)) in qs}
        results = []
        for qid, q in sorted(qs.items(), key=lambda kv: kv[1]["idx"]):
            a = given.get(qid, {})
            ok = (_norm_answer(a.get("answer")) == _norm_answer(q["answer"])) if q["kind"] == "mcq" \
                else None if check and a.get("correct") is None else bool(a.get("correct"))
            results.append({"id": qid, "kind": q["kind"], "your_answer": a.get("answer"), "correct": ok,
                            "answer": q["answer"], "explanation": q["explanation"], "t": q["t"]})
        score = sum(r["correct"] is True for r in results)
        unmarked = sum(r["correct"] is None for r in results)
        if check:
            return {"results": results, "score": score, "total": len(results), "unmarked": unmarked, "saved": False}
        con.execute("INSERT INTO quiz_attempts(lecture_id, answers, score, total) VALUES (?,?,?,?)",
                    (lecture_id, json.dumps({r["id"]: {"answer": r["your_answer"], "correct": r["correct"]}
                                             for r in results}), score, len(results)))
    return {"results": results, "score": score, "total": len(results), "unmarked": 0, "saved": True}


def lecture_link(lecture_id: int, t: float | None = None) -> str:
    """A link into LectureLab at a moment. ?t= goes before the hash: the hash is the UI's router."""
    return f"http://127.0.0.1:{PORT}/" + (f"?t={int(t)}" if t is not None else "") + f"#/lecture/{lecture_id}"


def _anki_html(s: str) -> str:
    import html
    return html.escape(s or "").replace("\t", "    ").replace("\n", "<br>")


def anki_tsv(course_id: int | None = None, lecture_id: int | None = None) -> str:
    """Anki's text import: Cloze and Basic notes in one file, one deck per lecture. The GUID column
    makes a re-import update the same notes instead of duplicating them."""
    where, args = ["1=1"], []
    if lecture_id:
        where.append("k.lecture_id=?")
        args.append(lecture_id)
    if course_id:
        where.append("l.course_id=?")
        args.append(course_id)
    with db() as con:
        cards = rows(con.execute(
            f"""SELECT k.*, l.title, l.date, c.code FROM cards k JOIN lectures l ON l.id=k.lecture_id
                LEFT JOIN courses c ON c.id=l.course_id WHERE {' AND '.join(where)}
                ORDER BY l.date, l.id, k.id""", args))
    out = ["#separator:tab", "#html:true", "#guid column:1", "#notetype column:2", "#deck column:3",
           "#tags column:6"]
    for k in cards:
        deck = "::".join(x.replace("::", ":") for x in ["LectureLab", k["code"] or "Unsorted",
                                                         f"{k['date'] or ''} {k['title']}".strip()])
        guid = "ll" + hashlib.sha1(f"{k['lecture_id']}:{_card_key(k['front'])}".encode()).hexdigest()[:16]
        back = _anki_html(k["back"])
        if k["t"] is not None:
            back += f'<br><a href="{lecture_link(k["lecture_id"], k["t"])}">▶ {fmt_t(k["t"])}</a>'
        tags = " ".join(filter(None, [re.sub(r"\s+", "_", k["code"] or ""), f"lecture{k['lecture_id']}"]))
        deck = re.sub(r"[\t\r\n]+", " ", deck)  # Anki reads the deck column as plain text
        out.append("\t".join([guid, "Cloze" if k["kind"] == "cloze" else "Basic", deck,
                              _anki_html(k["front"]), back, tags]))
    return "\n".join(out) + "\n"


# Full-text search over transcripts, notes and slides. Rebuilt per lecture whenever one changes.

def index_lecture(lecture_id: int) -> int:
    with db() as con:
        segs = rows(con.execute("SELECT start, end, text, slide_no FROM segments WHERE lecture_id=? ORDER BY idx",
                                (lecture_id,)))
        lec = con.execute("SELECT deck_id FROM lectures WHERE id=?", (lecture_id,)).fetchone()
        slides = rows(con.execute("SELECT no, title, text FROM slides WHERE deck_id=?",
                                  ((lec and lec[0]) or -1,)))
        n = con.execute("SELECT sections, summary FROM notes WHERE lecture_id=?", (lecture_id,)).fetchone()
    docs: list[tuple[str, str, float]] = []
    chunk = None
    for s in segs:  # ~30 s windows: Whisper's segments can be one word, too short to match a phrase
        if chunk and s["start"] - chunk[1] > 30:
            docs.append(("transcript", chunk[0], chunk[1]))
            chunk = None
        chunk = (chunk[0] + " " + s["text"], chunk[1]) if chunk else (s["text"], s["start"])
    if chunk:
        docs.append(("transcript", chunk[0], chunk[1]))
    first = {}
    for s in segs:
        if s["slide_no"] and s["slide_no"] not in first:
            first[s["slide_no"]] = s["start"]
    for sl in slides:
        text = f"{sl['title'] or ''}\n{sl['text'] or ''}".strip()
        if text:
            docs.append(("slide", f"Slide {sl['no']}: {text}", first.get(sl["no"])))
    if n:
        docs.append(("notes", n["summary"], 0.0))
        for sec in json.loads(n["sections"]):
            t = sec["start"]
            for para in re.split(r"\n\s*\n", sec.get("markdown") or ""):
                if m := re.search(r"\(#t=(\d+)\)", para):
                    t = float(m[1])
                if para.strip():
                    docs.append(("notes", f"{sec['title']}: {para.strip()}", t))
    with db() as con:
        con.execute("DELETE FROM search WHERE lecture_id=?", (lecture_id,))
        con.executemany("INSERT INTO search(text, kind, lecture_id, t) VALUES (?,?,?,?)",
                        [(text, kind, lecture_id, t) for kind, text, t in docs])
    return len(docs)


def index_missing() -> int:
    """Index lectures that have a transcript but nothing in the search index (first start after
    Stage 4, or an index lost to a crash)."""
    with db() as con:
        ids = [r[0] for r in con.execute(
            "SELECT DISTINCT lecture_id FROM segments WHERE lecture_id NOT IN "
            "(SELECT DISTINCT CAST(lecture_id AS INTEGER) FROM search)")]
    for lid in ids:
        index_lecture(lid)
    return len(ids)


def search(q: str, course_id: int | None = None, limit: int = 50) -> list[dict]:
    """Every word must match (stemmed, so 'parsing' finds 'parser'); the last one as a prefix, so
    results come while you type. Snippets mark the hits with <mark>, everything else escaped."""
    import html
    words = re.findall(r"\w+", q or "")
    if not words:
        return []
    match = " ".join(f'"{w}"' for w in words[:-1]) + f' "{words[-1]}"*'
    sql = """SELECT s.lecture_id, s.kind, s.t, snippet(search, 0, char(2), char(3), '…', 14) AS snip,
                    l.title, l.date, c.code
             FROM search s JOIN lectures l ON l.id=CAST(s.lecture_id AS INTEGER)
             LEFT JOIN courses c ON c.id=l.course_id
             WHERE search MATCH ?""" + (" AND l.course_id=?" if course_id else "") + " ORDER BY bm25(search) LIMIT ?"
    with db() as con:
        found = rows(con.execute(sql, (match, *([course_id] if course_id else []), limit)))
    for f in found:
        f["lecture_id"] = int(f["lecture_id"])
        f["snippet_html"] = html.escape(f.pop("snip")).replace("\x02", "<mark>").replace("\x03", "</mark>")
    return found


# The offline generator: no Claude at all, only when you ask for it. Extractive, so it can't make
# anything up, but it can't explain anything either.

def _timed_sentences(segs: list[dict]) -> list[tuple[str, float]]:
    """Sentences of six words or more across segment boundaries, each with the start time of the
    segment it begins in."""
    import bisect
    text, offs = "", []
    for s in segs:
        offs.append(len(text))
        text += s["text"].strip() + " "
    out = []
    for m in re.finditer(r"[^.?!]+[.?!]?", text):
        sent = m[0].strip()
        if len(sent.split()) >= 6:
            out.append((sent, segs[bisect.bisect_right(offs, m.start() + len(m[0]) - len(m[0].lstrip())) - 1]["start"]))
    return out


def offline_notes(lecture_id: int) -> dict:
    """Sections per slide change (or every 5 minutes without slides), each with its most telling
    sentences by TF-IDF; cloze cards on the course vocabulary; multiple choice with other
    vocabulary terms as the wrong answers."""
    import math
    import random
    from collections import Counter
    inp = notes_inputs(lecture_id)
    if not inp["segs"]:
        raise ClaudeError("There's no transcript to write notes from yet.")
    titles = {n: (s["title"] or "") for n, s in inp["slides"].items()}
    groups: list[dict] = []
    has_slides = any(x["slide_no"] for x in inp["segs"])
    for s in inp["segs"]:
        key = s["slide_no"] if has_slides else int(s["start"] // 300)
        if not groups or groups[-1]["key"] != key and s["start"] - groups[-1]["start"] >= 60:
            groups.append({"key": key, "start": s["start"], "segs": []})
        groups[-1]["segs"].append(s)
    tok = lambda t: re.findall(r"[a-z]{3,}", t.lower())  # noqa: E731
    sents = [_timed_sentences(g["segs"]) for g in groups]
    df = Counter(w for g in sents for w in {w for s, _ in g for w in tok(s)})
    n_docs = max(1, len(groups))

    def score(s: str) -> float:
        ws = tok(s)
        return sum(math.log(n_docs / df[w]) for w in set(ws) if df[w]) / (len(ws) ** 0.5 or 1)
    vocab = [v.strip() for v in ((inp["course"] or {}).get("vocabulary") or "").splitlines() if v.strip()]
    sections, cards, quiz, all_best = [], [], [], []
    for i, g in enumerate(groups):
        end = groups[i + 1]["start"] if i + 1 < len(groups) else inp["duration"]
        best = sorted(sorted(sents[i], key=lambda c: -score(c[0]))[:4], key=lambda c: c[1])
        if not best:
            continue
        all_best += best
        on_slide = has_slides and g["key"] in titles
        title = (titles[g["key"]] if on_slide else "") or f"Part {len(sections) + 1} ({fmt_t(g['start'])})"
        slides = [g["key"]] if on_slide else []
        sections.append({"title": title, "start": g["start"], "end": end, "kind": "content", "slides": slides,
                         "gist": "", "markdown": "\n".join(f"- {jump(t)} {s}" for s, t in best)})
    for s, t in all_best:
        for term in vocab:
            if re.search(rf"(?i)\b{re.escape(term)}\b", s) and len(cards) < 40:
                cards.append({"kind": "cloze", "front": re.sub(rf"(?i)\b({re.escape(term)})\b", r"{{c1::\1}}", s, count=1),
                              "back": "", "t": t})
                others = [v for v in vocab if v.lower() != term.lower()]
                if len(others) >= 3 and len(quiz) < 10:
                    rnd = random.Random(f"{lecture_id}:{s}")  # same quiz every run
                    choices = rnd.sample(others, 3) + [term]
                    rnd.shuffle(choices)
                    quiz.append({"kind": "mcq", "question": re.sub(rf"(?i)\b{re.escape(term)}\b", "_____", s, count=1),
                                 "choices": choices, "answer": term, "explanation": "From the transcript.", "t": t})
                break
    top = sorted(sorted(all_best, key=lambda c: -score(c[0]))[:5], key=lambda c: c[1])
    summary = "Offline notes: the lecture's most distinctive sentences, without Claude.\n\n" + \
        "\n".join(f"- {jump(t)} {s}" for s, t in top)
    return {"outline": sections, "sections": sections, "summary": summary, "open_questions": [],
            "cards": cards, "quiz": quiz, "issues": [], "removed": 0, "flagged": 0}



# ---------------------------------------------------------------- export and sync

# Notes land where Recall already indexes: AI Notes\Lectures\<Course>\<date> - <Title>.md
NOTES_DIR = Path(os.environ.get("LECTURELAB_NOTES_DIR", Path.home() / "Documents" / "AI Notes" / "Lectures"))


def fmt_hms(sec: float) -> str:
    s = int(sec)
    return f"{s // 3600}:{s % 3600 // 60:02d}:{s % 60:02d}"


def _export_name(s: str) -> str:
    # not safe_filename: "Lecture 3.1 Parsing" has no extension to keep
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", s).strip(" .")[:120].strip(" .")


def export_paths(lec: dict, taken: set[str] = frozenset()) -> tuple[Path, Path]:
    """Where a lecture's notes and transcript go. Two lectures with the same date and title get the
    lecture id added rather than overwrite each other ("taken": other lectures' files, lowercase)."""
    folder = NOTES_DIR / (_export_name(lec["course_code"] or "") or "Unsorted")
    date = lec["date"] or (lec["created_at"] or "")[:10]
    base = _export_name(f"{date} - {lec['title']}") or f"Lecture {lec['id']}"
    if str(folder / f"{base}.md").lower() in taken:
        base += f" ({lec['id']})"
    return folder / f"{base}.md", folder / f"{base} - Transcript.md"


def _others_exports(con, lecture_id: int) -> set[str]:
    return {p.lower() for (ep,) in con.execute("SELECT export_paths FROM notes WHERE lecture_id<>? AND "
                                               "export_paths IS NOT NULL", (lecture_id,)) for p in json.loads(ep)}


def _frontmatter(fields: dict) -> str:
    # JSON strings are valid YAML, and survive titles with colons, quotes or #
    return "---\n" + "".join(f"{k}: {json.dumps(v, ensure_ascii=False) if isinstance(v, str) else v}\n"
                             for k, v in fields.items() if v not in (None, "")) + "---\n\n"


def _absolute_links(md: str, lecture_id: int) -> str:
    """[▶ 12:34](#t=754) only means something inside LectureLab: point it at the running app."""
    return re.sub(r"\]\(#t=(\d+)\)", lambda m: f"]({lecture_link(lecture_id, int(m[1]))})", md)


def export_lecture(lecture_id: int) -> list[Path]:
    """Write the notes and the transcript as Markdown. Re-exporting overwrites them, and a lecture
    renamed since the last export leaves no old copy behind (only files written here are removed)."""
    lec = get_lecture(lecture_id)
    if not lec:
        return []
    with db() as con:
        n = con.execute("SELECT * FROM notes WHERE lecture_id=?", (lecture_id,)).fetchone()
        segs = rows(con.execute("SELECT start, end, text, slide_no FROM segments WHERE lecture_id=? ORDER BY idx",
                                (lecture_id,)))
        slides = dict(con.execute("SELECT no, title FROM slides WHERE deck_id=?", (lec["deck_id"] or -1,)).fetchall())
        others = _others_exports(con, lecture_id)
    if not n:
        return []
    notes_path, tr_path = export_paths(lec, others)
    deck = get_deck(lec["deck_id"])
    source = next((s["url"] for s in lec["sources"] if (s["url"] or "").startswith("http")), None)
    dur = max([lec["duration_sec"] or 0] + [s["end"] for s in segs])
    meta = {"course": lec["course_code"], "course_name": lec["course_name"], "title": lec["title"],
            "date": lec["date"], "source_url": source, "duration": fmt_hms(dur),
            "slides_file": str(DATA / deck["path"]) if deck and deck.get("path") else (deck or {}).get("title"),
            "lecturelab": lecture_link(lecture_id)}
    sections = json.loads(n["sections"])
    body = [f"# {lec['title']}\n", f"[Open in LectureLab]({lecture_link(lecture_id)})\n"
            f"[Transcript](<{tr_path.name}>)\n", "## Summary\n", n["summary"].strip() + "\n",
            _absolute_links(notes_markdown(sections), lecture_id) + "\n"]
    if oq := json.loads(n["open_questions"]):
        body += ["## Open questions\n", "\n".join(f"- {q}" for q in oq) + "\n"]
    notes_md = _frontmatter({**meta, "generator": n["generator"],
                             "checked": f"{n['removed']} removed, {n['flagged']} flagged"}) + "\n".join(body)
    lines = transcript_text(segs, slides, clock=fmt_hms)
    tr_md = _frontmatter({**meta, "kind": "transcript", "transcribed_by": lec["transcript_source"]}) + \
        f"# {lec['title']}: transcript\n\n" + re.sub(r"(?m)^--- (Slide .*) ---$", r"\n**\1**\n", lines) + "\n"
    notes_path.parent.mkdir(parents=True, exist_ok=True)
    for path, text in ((notes_path, notes_md), (tr_path, tr_md)):
        tmp = path.with_name(path.name + ".tmp")
        try:
            tmp.write_text(text, encoding="utf-8", errors="replace")
            tmp.replace(path)
        finally:  # a file open in another app: no stray .tmp left in the indexed folder
            tmp.unlink(missing_ok=True)
    new = [str(notes_path), str(tr_path)]
    for old in json.loads(n["export_paths"] or "[]"):
        if old.lower() not in {p.lower() for p in new} | others:
            Path(old).unlink(missing_ok=True)
    with db() as con:
        con.execute("UPDATE notes SET exported_at=datetime('now'), export_paths=? WHERE lecture_id=?",
                    (json.dumps(new), lecture_id))
        con.execute("UPDATE lectures SET status='exported' WHERE id=? AND status='notes_done'", (lecture_id,))
    return [notes_path, tr_path]


def export_pending() -> int:
    """Export notes that are new or changed since their last export, and lectures renamed, re-dated
    or moved to another course since (their files would be under the old name)."""
    with db() as con:
        ids = [r[0] for r in con.execute("SELECT lecture_id FROM notes WHERE exported_at IS NULL "
                                         "OR exported_at < created_at")]
        for lid, ep in con.execute("SELECT lecture_id, export_paths FROM notes WHERE exported_at >= created_at"):
            lec = get_lecture(lid)
            if lec and str(export_paths(lec, _others_exports(con, lid))[0]).lower() != (json.loads(ep or '[""]') or
                                                                                       [""])[0].lower():
                ids.append(lid)
    done = 0
    for lid in ids:
        try:
            done += bool(export_lecture(lid))
        except OSError as e:  # AI Notes on a drive that isn't there, a file open elsewhere: next sync
            log.warning("export of lecture %s failed: %s", lid, e)
    return done


SYNC_NOTES = int(os.environ.get("LECTURELAB_SYNC_NOTES", 3))


def queue_missing_notes(limit: int = SYNC_NOTES) -> int:
    """Notes for transcribed lectures that have none, once their slides are settled, at most
    `limit` at a time so a semester's backlog doesn't spend the week's plan usage in one go. A
    lecture whose notes failed isn't retried here (it would spend usage every hour): regenerate it."""
    with db() as con:
        ids = [r[0] for r in con.execute(
            """SELECT l.id FROM lectures l WHERE l.status NOT IN ('duplicate','ingesting','queued_transcribe',
               'transcribing') AND l.notes_error IS NULL
               AND EXISTS (SELECT 1 FROM segments s WHERE s.lecture_id=l.id)
               AND NOT EXISTS (SELECT 1 FROM notes n WHERE n.lecture_id=l.id)
               AND NOT EXISTS (SELECT 1 FROM jobs j WHERE j.lecture_id=l.id AND j.kind IN ('notes','slides')
                               AND j.status IN ('queued','running')) ORDER BY l.id""")]
    with db() as con:
        pending = con.execute("SELECT COUNT(*) FROM jobs WHERE kind='notes' AND status IN ('queued','running')").fetchone()[0]
    n, limit = 0, limit - pending
    for lid in ids:
        if n >= limit:
            break
        if lecture_slides(lid)["state"] in ("aligned", "none", "error"):
            enqueue_notes(lid)
            n += 1
    return n


def server_running() -> bool:
    import httpx
    try:
        r = httpx.get(f"http://127.0.0.1:{PORT}/api/health", timeout=3)
        return r.status_code == 200 and r.json().get("pid") != os.getpid()
    except Exception:
        return False


SYNC_ORDER = ["ingest", "canvas_scan", "canvas_pull", "deck_pull", "transcribe", "slides", "notes"]


def drain_queue(say=lambda m: None, kinds: list[str] = SYNC_ORDER) -> int:
    """Run queued jobs here until none is ready (the server isn't up to run them). A job waiting
    out a retry time (the usage limit, Canvas signed out) is left for later."""
    ran = 0
    while ran < 500:  # every self-requeue sets a retry time, but don't trust that forever
        job = next((j for k in kinds if (j := claim_job([k]))), None)
        if not job:
            return ran
        say(f"job {job['id']}: {job['kind']} {job['payload'].get('title') or ''}".strip())
        run_job(job)
        ran += 1
    return ran


def run_sync(say=lambda m: None) -> dict:
    """The whole cycle, safe to run every hour: Canvas, inbox folders, one transcription batch,
    notes, export. Only new work is done, so a second run straight after does nothing. With the
    server up it queues the work for the server; without it, it runs the queue itself."""
    lock = FileLock(DATA / "sync.lock")
    if not lock.acquire():
        return {"skipped": "another sync is running"}
    t0, did = time.time(), {}
    try:
        server = server_running()
        if not server:  # a run killed mid-job (shutdown, the task's time limit) would block it for good
            did["recovered"] = recover_orphans()
        if canvas_state().get("state") == "connected":
            say("Canvas")
            errors = []
            try:  # recordings for courses set to auto-pull
                did["canvas_pulls"] = len(canvas_auto_pull())
            except Exception as e:  # Canvas down or signed out: everything else still runs
                errors.append(user_error(e, "The Canvas scan"))
            for c in list_courses():  # slide decks for every other linked course, each on its own
                # (signed out mid-way marks the state, so the rest are skipped)
                if c["canvas_id"] and not c["auto_pull"] and canvas_state().get("state") == "connected":
                    try:
                        canvas_scan_course(c)
                    except Exception as e:
                        errors.append(user_error(e, f"The Canvas scan of {c['code']}"))
            if errors:
                did["canvas"] = " ".join(errors)
        else:
            did["canvas"] = "needs sign-in"
        say("inbox folders")
        # this process lives for one scan, so it can't wait for a second sighting of a file;
        # a file still being copied is skipped either way
        did["inbox"] = len(scan_inboxes(require_stable=False))
        # bring this hour's recordings in first, so they make this hour's batch (with the server up,
        # only the quick ingest jobs: its Canvas browser may be busy)
        did["ran"] = drain_queue(say, ["ingest"] if server else ["ingest", "canvas_scan", "canvas_pull", "deck_pull"])
        with db() as con:
            waiting = [r[0] for r in con.execute("SELECT id FROM lectures WHERE status='queued_transcribe' "
                                                 "AND audio_path IS NOT NULL AND transcribe_fails < 3")]
            stuck = con.execute("SELECT COUNT(*) FROM lectures WHERE status='queued_transcribe' "
                                "AND transcribe_fails >= 3").fetchone()[0]
            busy = con.execute("SELECT 1 FROM jobs WHERE kind='transcribe' AND status IN ('queued','running')").fetchone()
        if waiting and not busy:
            enqueue("transcribe", {"mode": "auto", "lecture_ids": waiting})
            did["transcribe"] = len(waiting)
        did["transcribe_failing"] = stuck  # failed 3 times: only a manual Transcribe retries them
        did["notes"] = queue_missing_notes()
        if not server:  # with the server up, its workers pick up queued jobs within seconds
            did["ran"] += drain_queue(say)
            did["notes"] += queue_missing_notes()  # transcribed just now
            did["ran"] += drain_queue(say)
        did["exported"] = export_pending()
        did = {k: v for k, v in did.items() if v}
    finally:
        lock.release()
    log.info("sync done in %.0f s: %s", time.time() - t0, did or "nothing new")
    return did



# ---------------------------------------------------------------- v1 import

V1_COURSE_MAP = {"ethics": "PHI3681", "programming language concepts": "COP4020",
                 "math for machine learning": "MML"}


def import_v1(v1_db: Path) -> list[dict]:
    """One-time: bring v1 recordings and transcripts in. Safe to re-run (dedupe by content hash).

    v1 transcripts come from Whisper 'base'; they're kept as transcript_source='v1' so the
    lecture is usable right away, and the lecture stays queued for a large-v3 re-transcription.
    """
    v1_db = Path(v1_db)
    audio_dir = v1_db.parent / "audio"
    src = sqlite3.connect(f"file:{v1_db}?mode=ro", uri=True)
    src.row_factory = sqlite3.Row
    try:
        return _import_v1_rows(src, audio_dir)
    finally:
        src.close()


def _import_v1_rows(src, audio_dir: Path) -> list[dict]:
    out = []
    for lec in src.execute("""SELECT l.*, c.name AS course_name FROM lectures l
                              LEFT JOIN courses c ON c.id=l.course_id ORDER BY l.id"""):
        audio = audio_dir / lec["audio_file"] if lec["audio_file"] else None
        if not audio or not audio.exists():
            # v1 sometimes lost the link while the file stayed on disk: lecture_<id>_<hash>.webm
            matches = sorted(audio_dir.glob(f"lecture_{lec['id']}_*"))
            audio = matches[0] if matches else None
        if not audio:
            out.append({"v1_id": lec["id"], "title": lec["title"], "skipped": "no audio file"})
            continue
        course = find_course(V1_COURSE_MAP.get((lec["course_name"] or "").lower(), lec["course_name"]))
        res = ingest_file(audio, kind="v1", url=f"v1://lecture/{lec['id']}", course_id=course and course["id"],
                          title=lec["title"], date=parse_date(lec["created_at"]))  # v1 stored epoch seconds
        segs = json.loads(lec["transcript"] or "[]")
        with db() as con:
            has = con.execute("SELECT COUNT(*) FROM segments WHERE lecture_id=?", (res["lecture_id"],)).fetchone()[0]
            if segs and not has:
                con.executemany("INSERT INTO segments(lecture_id, idx, start, end, text) VALUES (?,?,?,?,?)",
                                [(res["lecture_id"], i, s["start"], s["end"], s["text"].strip())
                                 for i, s in enumerate(segs)])
                con.execute("UPDATE lectures SET transcript_source='v1' WHERE id=?", (res["lecture_id"],))
        out.append({"v1_id": lec["id"], "title": lec["title"], "lecture_id": res["lecture_id"],
                    "duplicate": res["duplicate"], "segments": len(segs), "audio": audio.name})
    return out


# ---------------------------------------------------------------- web API

# Request models live at module level for the same reason as the FastAPI imports above.
class CoursePatch(BaseModel):
    name: str | None = None
    vocabulary: str | None = None


class RecordingIn(BaseModel):
    url: str
    course_id: int | None = None
    title: str | None = None
    date: str | None = None
    passcode: str | None = None


class LecturePatch(BaseModel):
    title: str | None = None
    date: str | None = None
    course_id: int | None = None


class CanvasPull(BaseModel):
    course_id: int
    keys: list[str]


class AutoPull(BaseModel):
    auto_pull: bool


class DeckPick(BaseModel):
    deck_id: int | None = None


class TranscribeIn(BaseModel):
    lecture_ids: list[int] | None = None
    mode: str = "auto"  # auto | colab | local


class ReviewRequest(BaseModel):
    correct: bool


class UndoReviewRequest(BaseModel):
    before: dict  # what POST /api/cards/{id}/review returned as "before"


class QuizAnswer(BaseModel):
    id: int
    answer: str | None = None
    correct: bool | None = None  # short answers: you say whether yours matched the model answer


class QuizAttempt(BaseModel):
    answers: list[QuizAnswer]
    check: bool = False  # mark it and show the model answers, but don't save an attempt yet


class NotesRequest(BaseModel):
    mode: Literal["claude", "offline"] = "claude"  # offline: no Claude, only when you ask for it
    force: bool = False  # offline over Claude's notes


def check_ingest_fields(course_id: int | None, date: str | None) -> None:
    """Reject bad fields up front, not after a 20-minute download."""
    if course_id is not None and not find_course(course_id):
        raise HTTPException(400, f"no course with id {course_id}")
    if date and not valid_date(date):
        raise HTTPException(400, "date must be a real day as YYYY-MM-DD")


WORKERS: dict[str, Worker] = {}


def wake_workers() -> None:
    for w in WORKERS.values():
        if isinstance(w, Worker):
            w.wake.set()


ORPHAN_SWEEP_SEC = float(os.environ.get("LECTURELAB_ORPHAN_SWEEP_SEC", 60))


def recover_orphans() -> int:
    """Clean up after a process that died mid-job (a crash, or Windows killing it for memory):
    its lectures go back to the queue, its batch is closed as interrupted, its Colab session is
    stopped, and its job is queued again. The server does this every minute, not just at
    startup, so a dead run never shows as running. Returns how many jobs were re-queued."""
    if reset_stuck_transcribing() and COLAB_ENABLED:
        sweep_orphans_if_idle()  # before the job is re-queued: a new batch would hold the lock
    n = requeue_orphaned_jobs()
    if n:
        wake_workers()
    return n


def _sweep_orphans_forever() -> None:
    while True:
        time.sleep(ORPHAN_SWEEP_SEC)
        try:
            recover_orphans()
        except Exception:
            log.exception("orphan sweep failed")


def start_background(watch_inboxes: bool = True) -> None:
    requeue_orphaned_jobs()
    reset_stuck_transcribing()
    if COLAB_ENABLED:  # off the startup path: WSL can take a few seconds to wake
        threading.Thread(target=sweep_orphans_if_idle, daemon=True, name="colab-sweep").start()
    if ORPHAN_SWEEP_SEC > 0 and "orphans" not in WORKERS:  # 0 turns it off (the unit tests)
        WORKERS["orphans"] = threading.Thread(target=_sweep_orphans_forever, daemon=True, name="orphan-sweep")
        WORKERS["orphans"].start()
    if "search" not in WORKERS:  # a one-off: lectures transcribed before search existed
        WORKERS["search"] = threading.Thread(target=index_missing, daemon=True, name="search-index")
        WORKERS["search"].start()
    if "transcribe" not in WORKERS:
        WORKERS["transcribe"] = Worker(["transcribe"], "transcribe-worker")
        WORKERS["transcribe"].start()
    if "ingest" not in WORKERS:
        WORKERS["ingest"] = Worker(["ingest"], "ingest-worker")
        WORKERS["ingest"].start()
        WORKERS["canvas"] = Worker(["canvas_scan", "canvas_pull"], "canvas-worker")
        WORKERS["canvas"].start()
        # Its own worker: a semester of queued recording pulls shouldn't hold up the slides.
        WORKERS["slides"] = Worker(["deck_pull", "slides"], "slides-worker")
        WORKERS["slides"].start()
        WORKERS["notes"] = Worker(["notes"], "notes-worker")
        WORKERS["notes"].start()
    if watch_inboxes and "inbox" not in WORKERS:
        WORKERS["inbox"] = InboxWatcher(interval=float(os.environ.get("LECTURELAB_INBOX_INTERVAL", 30)),
                                        on_new=wake_workers)
        WORKERS["inbox"].start()


def create_app():
    @asynccontextmanager
    async def lifespan(_app):
        init_db()
        start_background()
        log.info("server up on 127.0.0.1:%s pid %s (data=%s)", PORT, os.getpid(), DATA)
        yield
        log.info("server shutting down")

    app = FastAPI(title="LectureLab", version=VERSION, lifespan=lifespan)
    allowed_hosts = {f"127.0.0.1:{PORT}", f"localhost:{PORT}"}

    @app.middleware("http")
    async def local_only(request: Request, call_next):
        # Bound to 127.0.0.1, but a web page could still reach us through DNS rebinding
        # (wrong Host header) or a cross-site form post (foreign Origin). Refuse both.
        if request.headers.get("host") not in allowed_hosts:
            return PlainTextResponse("bad host", status_code=403)
        origin = request.headers.get("origin")
        if request.method not in ("GET", "HEAD") and origin and origin.split("//", 1)[-1] not in allowed_hosts:
            return PlainTextResponse("bad origin", status_code=403)
        return await call_next(request)

    @app.get("/api/health")
    def health():
        return {"ok": True, "version": VERSION, "pid": os.getpid(), "time": time.time()}

    @app.get("/api/courses")
    def api_courses():
        return list_courses()

    @app.put("/api/courses/{course_id}")
    def api_update_course(course_id: int, body: CoursePatch):
        try:
            return update_course(course_id, body.name, body.vocabulary)
        except KeyError:
            raise HTTPException(404, "course not found")
        except ValueError as e:
            raise HTTPException(400, str(e))

    # --- recordings / lectures
    @app.post("/api/recordings")
    def api_add_url(body: RecordingIn):
        url = body.url.strip()
        if not re.match(r"https?://", url):
            raise HTTPException(400, "that doesn't look like a link (expected http:// or https://)")
        check_ingest_fields(body.course_id, body.date)
        jid = enqueue("ingest", {"url": url, "course_id": body.course_id, "title": body.title or None,
                                 "date": body.date or None, "passcode": body.passcode or None})
        wake_workers()
        return {"job_id": jid}

    @app.post("/api/recordings/upload")
    def api_upload(file: UploadFile = File(...), course_id: int | None = Form(None),
                   title: str | None = Form(None), date: str | None = Form(None)):
        # Plain def: FastAPI runs it in a thread, so the copy and DB calls don't block the event loop.
        name = Path(file.filename or "upload").name
        if Path(name).suffix.lower() not in MEDIA_EXTS:
            raise HTTPException(400, f"unsupported file type; use one of {', '.join(sorted(MEDIA_EXTS))}")
        check_ingest_fields(course_id, date)
        # The stored name is ours (uuid + extension); the user's filename only travels as metadata,
        # so names like CON.mp4 or ../../x.mp4 never reach the filesystem.
        dest = DATA / "uploads" / f"{uuid.uuid4().hex}{Path(name).suffix.lower()}"
        dest.parent.mkdir(parents=True, exist_ok=True)
        try:
            with open(dest, "wb") as out:
                shutil.copyfileobj(file.file, out, 1 << 20)
        except BaseException:
            dest.unlink(missing_ok=True)  # an aborted upload leaves nothing behind
            raise
        jid = enqueue("ingest", {"path": str(dest), "kind": "file", "source_url": f"upload:{name}",
                                 "course_id": course_id, "title": title or Path(name).stem, "date": date or None,
                                 "delete_after": True})
        wake_workers()
        return {"job_id": jid}

    @app.get("/api/jobs/{job_id}")
    def api_job(job_id: int):
        j = get_job(job_id)
        if not j:
            raise HTTPException(404, "job not found")
        j.pop("payload", None)
        return j

    @app.get("/api/jobs")
    def api_jobs(active: bool = True):
        with db() as con:
            q = "SELECT id, kind, status, progress, error, lecture_id, payload, updated_at FROM jobs"
            q += " WHERE status IN ('queued','running') ORDER BY id" if active else " ORDER BY id DESC LIMIT 50"
            out = rows(con.execute(q))
        for j in out:
            p = json.loads(j.pop("payload") or "{}")
            if j["kind"] == "transcribe":
                ids = p.get("lecture_ids")
                j["label"] = f"{'lecture ' + ', '.join(map(str, ids)) if ids else 'everything waiting'} ({p.get('mode', 'auto')})"
            else:
                j["label"] = (clean_source_url(p.get("url")) or p.get("title") or p.get("key") or Path(p.get("path", "")).name
                              or (find_course(p.get("course_id")) or {}).get("code", ""))
        return out

    @app.get("/api/lectures")
    def api_lectures(course_id: int | None = None):
        return list_lectures(course_id)

    @app.get("/api/lectures/{lecture_id}")
    def api_lecture(lecture_id: int):
        lec = get_lecture(lecture_id)
        if not lec:
            raise HTTPException(404, "lecture not found")
        return lec

    @app.patch("/api/lectures/{lecture_id}")
    def api_patch_lecture(lecture_id: int, body: LecturePatch):
        if body.course_id is not None and not find_course(body.course_id):
            raise HTTPException(400, f"no course with id {body.course_id}")
        try:
            return update_lecture(lecture_id, **body.model_dump(exclude_unset=True))
        except KeyError:
            raise HTTPException(404, "lecture not found")
        except ValueError as e:
            raise HTTPException(400, str(e))

    @app.post("/api/lectures/{lecture_id}/not-duplicate")
    def api_not_duplicate(lecture_id: int):
        return release_duplicate(lecture_id)

    @app.delete("/api/lectures/{lecture_id}")
    def api_delete_lecture(lecture_id: int):
        delete_lecture(lecture_id)
        return {"ok": True}

    @app.get("/api/lectures/{lecture_id}/media/{which}")
    def api_media(lecture_id: int, which: str):
        lec = get_lecture(lecture_id)
        rel = lec and lec.get({"audio": "audio_path", "video": "video_path"}.get(which, ""))
        if not rel or not (DATA / rel).is_file():
            raise HTTPException(404, "no such media")
        return FileResponse(DATA / rel)  # supports Range requests, so the player can seek

    # --- Canvas (through LectureLab's own browser profile; you sign in there yourself)
    @app.get("/api/canvas/status")
    def api_canvas_status():
        st = canvas_state()
        lk = FileLock(DATA / "browser.lock")
        busy = not lk.acquire()
        if not busy:
            lk.release()
        return {**st, "browser_busy": busy, "base": CANVAS_BASE}

    @app.post("/api/canvas/connect")
    def api_canvas_connect():
        lk = FileLock(DATA / "browser.lock")
        if not lk.acquire():
            raise HTTPException(409, "The LectureLab browser is already open; finish signing in there.")
        lk.release()
        # Its own process: the sign-in window must be a normal visible window, and Playwright's
        # browser shouldn't live inside the server process.
        subprocess.Popen([sys.executable, str(ROOT / "cli.py"), "canvas-login"], cwd=ROOT,
                         stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        log.info("opened the LectureLab browser for Canvas sign-in")
        return {"opened": True}

    @app.get("/api/canvas/courses")
    def api_canvas_courses():
        try:
            return canvas_courses()
        except CanvasError as e:
            raise HTTPException(400, str(e))

    @app.post("/api/canvas/scan/{course_id}")
    def api_canvas_scan(course_id: int):
        if not find_course(course_id):
            raise HTTPException(404, "course not found")
        try:
            canvas_ready()
        except CanvasError as e:
            raise HTTPException(400, str(e))
        jid = enqueue("canvas_scan", {"course_id": course_id})
        wake_workers()
        return {"job_id": jid}

    @app.get("/api/canvas/items")
    def api_canvas_items(course_id: int):
        with db() as con:
            return rows(con.execute("SELECT * FROM canvas_items WHERE course_id=? ORDER BY created_at DESC, key",
                                    (course_id,)))

    @app.post("/api/canvas/pull")
    def api_canvas_pull(body: CanvasPull):
        jobs = canvas_pull(body.course_id, body.keys)
        wake_workers()
        return {"job_ids": jobs}

    @app.post("/api/canvas/ignore")
    def api_canvas_ignore(body: CanvasPull):
        with db() as con:
            con.executemany("UPDATE canvas_items SET status='ignored' WHERE course_id=? AND key=?",
                            [(body.course_id, k) for k in body.keys])
        return {"ok": True}

    @app.put("/api/courses/{course_id}/auto-pull")
    def api_auto_pull(course_id: int, body: AutoPull):
        if not find_course(course_id):
            raise HTTPException(404, "course not found")
        with db() as con:
            con.execute("UPDATE courses SET auto_pull=? WHERE id=?", (int(body.auto_pull), course_id))
        return find_course(course_id)

    # --- transcription
    @app.post("/api/transcribe")
    def api_transcribe(body: TranscribeIn):
        if body.mode not in ("auto", "colab", "local"):
            raise HTTPException(400, "mode must be auto, colab or local")
        with db() as con:
            busy = con.execute("SELECT id FROM jobs WHERE kind='transcribe' AND status IN ('queued','running')").fetchone()
            # the same lectures asked for twice (a double click): one batch, not two Colab sessions
            same = body.lecture_ids and next((r["id"] for r in con.execute(
                "SELECT id, payload FROM jobs WHERE kind='transcribe' AND status IN ('queued','running')")
                if sorted(json.loads(r["payload"]).get("lecture_ids") or []) == sorted(body.lecture_ids)), None)
        if busy and not body.lecture_ids:
            return {"job_id": busy["id"], "note": "a batch is already queued"}
        if same:
            return {"job_id": same, "note": "a batch is already queued"}
        jid = enqueue("transcribe", {"lecture_ids": body.lecture_ids, "mode": body.mode})
        wake_workers()
        return {"job_id": jid}

    @app.get("/api/transcribe/status")
    def api_transcribe_status():
        with db() as con:
            batches = rows(con.execute("SELECT * FROM transcribe_batches ORDER BY id DESC LIMIT 10"))
            queued = con.execute("SELECT COUNT(*) FROM lectures WHERE status='queued_transcribe'").fetchone()[0]
        return {"colab_enabled": COLAB_ENABLED, "colab_failures": meta_get("colab_failures", 0),
                "local_model": os.environ.get("LECTURELAB_LOCAL_MODEL") or meta_get("local_model", "large-v3-turbo"),
                "benchmark": meta_get("benchmark"), "min_units": MIN_UNITS, "queued": queued, "batches": batches}

    @app.get("/api/lectures/{lecture_id}/segments")
    def api_segments(lecture_id: int):
        with db() as con:
            segs = rows(con.execute("SELECT idx, start, end, text, words, slide_no FROM segments "
                                    "WHERE lecture_id=? ORDER BY idx", (lecture_id,)))
        for sg in segs:
            sg["words"] = json.loads(sg["words"]) if sg["words"] else []
        return segs

    # --- slides
    @app.get("/api/lectures/{lecture_id}/slides")
    def api_slides(lecture_id: int):
        try:
            return lecture_slides(lecture_id)
        except KeyError:
            raise HTTPException(404, "lecture not found")

    @app.put("/api/lectures/{lecture_id}/slides")
    def api_set_slides(lecture_id: int, body: DeckPick):
        try:
            return set_lecture_deck(lecture_id, body.deck_id)
        except KeyError:
            raise HTTPException(404, "lecture not found")
        except ValueError as e:
            raise HTTPException(400, str(e))

    @app.post("/api/lectures/{lecture_id}/slides")
    def api_upload_slides(lecture_id: int, file: UploadFile = File(...)):
        try:
            return attach_deck_upload(lecture_id, file.filename, file.file)
        except KeyError:
            raise HTTPException(404, "lecture not found")
        except ValueError as e:
            raise HTTPException(400, str(e))

    @app.post("/api/lectures/{lecture_id}/slides/align")
    def api_align_slides(lecture_id: int):
        lec = get_lecture(lecture_id)
        if not lec:
            raise HTTPException(404, "lecture not found")
        with db() as con:
            con.execute("UPDATE lectures SET slides_error=NULL WHERE id=?", (lecture_id,))
            failed = [r[0] for r in con.execute("SELECT id FROM decks WHERE status='error' AND (id=? OR course_id IS ?)",
                                                (lec["deck_id"] or -1, lec["course_id"]))]
        for did in failed:  # decks that failed to download get fetched again (a bad file just fails again)
            retry_deck(did)
        enqueue_slides(lecture_id, ocr=True)  # you asked: read the video if no deck turns up
        return lecture_slides(lecture_id)

    @app.get("/api/lectures/{lecture_id}/notes")
    def api_lecture_notes(lecture_id: int):
        try:
            return lecture_notes(lecture_id)
        except KeyError:
            raise HTTPException(404, "lecture not found")

    @app.post("/api/lectures/{lecture_id}/notes")
    def api_write_notes(lecture_id: int, req: NotesRequest):
        lec = get_lecture(lecture_id)
        if not lec:
            raise HTTPException(404, "lecture not found")
        if not lec["segment_count"]:
            raise HTTPException(409, "This lecture has no transcript yet, so there's nothing to write notes from.")
        try:
            enqueue_notes(lecture_id, req.mode, req.force)
        except ValueError as e:
            raise HTTPException(409, str(e)) from None
        return lecture_notes(lecture_id)

    @app.get("/api/lectures/{lecture_id}/cards")
    def api_lecture_cards(lecture_id: int):
        if not get_lecture(lecture_id):
            raise HTTPException(404, "lecture not found")
        with db() as con:
            return rows(con.execute("SELECT * FROM cards WHERE lecture_id=? ORDER BY t, id", (lecture_id,)))

    @app.get("/api/review")
    def api_review(course_id: int | None = None, lecture_id: int | None = None, limit: int = 100):
        return due_cards(course_id, lecture_id, max(1, min(limit, 500)))

    @app.post("/api/cards/{card_id}/review")
    def api_review_card(card_id: int, req: ReviewRequest):
        try:
            return review_card(card_id, req.correct)
        except KeyError:
            raise HTTPException(404, "card not found")

    @app.post("/api/cards/{card_id}/undo")
    def api_undo_review(card_id: int, req: UndoReviewRequest):
        try:
            return undo_review(card_id, req.before)
        except KeyError:
            raise HTTPException(404, "card not found")
        except (TypeError, ValueError):
            raise HTTPException(422, "before must be the card state the review returned")

    @app.get("/api/lectures/{lecture_id}/quiz")
    def api_quiz(lecture_id: int):
        if not get_lecture(lecture_id):
            raise HTTPException(404, "lecture not found")
        return lecture_quiz(lecture_id)

    @app.post("/api/lectures/{lecture_id}/quiz")
    def api_grade_quiz(lecture_id: int, req: QuizAttempt):
        if not get_lecture(lecture_id):
            raise HTTPException(404, "lecture not found")
        return grade_quiz(lecture_id, [a.model_dump() for a in req.answers], req.check)

    @app.get("/api/anki.tsv")
    def api_anki(course_id: int | None = None, lecture_id: int | None = None):
        name = "LectureLab" + (f"-{(find_course(course_id) or {}).get('code', course_id)}" if course_id else "") + \
            (f"-lecture{lecture_id}" if lecture_id else "")
        name = re.sub(r"[^\w.-]+", "_", name)
        return PlainTextResponse(anki_tsv(course_id, lecture_id), media_type="text/tab-separated-values",
                                 headers={"Content-Disposition": f'attachment; filename="{name}.txt"'})

    @app.get("/api/search")
    def api_search(q: str = "", course_id: int | None = None, limit: int = 50):
        return search(q, course_id, max(1, min(limit, 200)))

    @app.get("/media/slides/{deck_id}/{name}")
    def api_slide_image(deck_id: int, name: str):
        path = SLIDE_IMAGES / str(deck_id) / name
        if not re.fullmatch(rf"\d+-(?:{THUMB_W}|{FULL_W})\.png", name) or not path.is_file():
            raise HTTPException(404, "no such slide image")
        return FileResponse(path, headers={"Cache-Control": "no-cache"})

    @app.get("/api/inboxes")
    def api_inboxes():
        return [{"path": str(p), "kind": k, "exists": p.exists()} for p, k in inbox_roots()]

    @app.get("/")
    def index():
        return FileResponse(ROOT / "static" / "index.html", headers={"Cache-Control": "no-cache"})

    app.mount("/static", StaticFiles(directory=ROOT / "static"), name="static")
    return app
