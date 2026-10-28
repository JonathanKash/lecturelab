import importlib
import json
import os
import subprocess
import sys
import time
from datetime import datetime

import pytest


@pytest.fixture()
def ll(tmp_path, monkeypatch):
    """Fresh lecturelab module pointed at a throwaway data dir."""
    monkeypatch.setenv("LECTURELAB_DATA", str(tmp_path / "data"))
    monkeypatch.setenv("LECTURELAB_LOG", str(tmp_path / "test.log"))
    monkeypatch.setenv("LECTURELAB_ORPHAN_SWEEP_SEC", "0")  # no minute-by-minute thread outliving the test
    monkeypatch.setenv("LECTURELAB_NOTES_DIR", str(tmp_path / "AI Notes" / "Lectures"))  # never the real one
    monkeypatch.delenv("LECTURELAB_CLAUDE", raising=False)  # conftest's off switch: these tests fake claude
    monkeypatch.delenv("LECTURELAB_AUTO_NOTES", raising=False)
    import lecturelab
    return importlib.reload(lecturelab)


def test_courses_seeded_once(ll):
    codes = [c["code"] for c in ll.list_courses()]
    assert codes == sorted(["COP4020", "COP4533", "CNT4007", "CIS4930", "PHI3681", "MML"])
    ll._db_ready = False
    ll.init_db()  # re-running init must not duplicate or reset anything
    assert len(ll.list_courses()) == 6


def test_vocab_update_normalizes_lines(ll):
    cop = next(c for c in ll.list_courses() if c["code"] == "COP4020")
    out = ll.update_course(cop["id"], vocabulary="  monad \n\n\nfunctor\r\n  ")
    assert out["vocabulary"] == "monad\nfunctor"


def test_deleted_seed_course_stays_deleted(ll):
    with ll.db() as con:
        con.execute("DELETE FROM courses WHERE code='CIS4930'")
    ll._db_ready = False
    ll.init_db()
    assert "CIS4930" not in [c["code"] for c in ll.list_courses()]


def test_empty_course_name_rejected(ll):
    cop = next(c for c in ll.list_courses() if c["code"] == "COP4020")
    with pytest.raises(ValueError):
        ll.update_course(cop["id"], name="   ")


def test_filelock_freed_when_holder_dies(ll, tmp_path):
    path = tmp_path / "y.lock"
    holder = subprocess.Popen([sys.executable, "-c",
                               f"import os, lecturelab as ll; assert ll.FileLock(r'{path}').acquire(); os._exit(0)"])
    assert holder.wait() == 0  # exited without release()
    assert ll.FileLock(path).acquire()


def test_task_script_runs_hidden_and_restarts():
    import cli
    s = cli.server_task_script()
    assert "pythonw.exe" in s and "supervise" in s
    assert "-RestartCount 999" in s and "-Priority 4" in s and "-AtLogOn" in s
    assert cli._q("C:/it's here") == "'C:/it''s here'"


def test_filelock_excludes_other_process(ll, tmp_path):
    path = tmp_path / "x.lock"
    lock = ll.FileLock(path)
    assert lock.acquire()
    assert not ll.FileLock(path).acquire()  # same process, second holder
    probe = ("import sys, lecturelab as ll; "
             f"sys.exit(0 if ll.FileLock(r'{path}').acquire() else 3)")
    assert subprocess.run([sys.executable, "-c", probe]).returncode == 3  # other process blocked
    lock.release()
    assert subprocess.run([sys.executable, "-c", probe]).returncode == 0  # freed after release


def test_parse_date_sources(ll):
    assert ll.parse_date(None, "20260821") == "2026-08-21"  # yt-dlp upload_date
    assert ll.parse_date("2026-04-08 19.47.47 Jonathan Kashi's Zoom Meeting") == "2026-04-08"
    assert ll.parse_date("2026-13-40T00:00", "Lecture_2026_09_03.mp4") == "2026-09-03"  # invalid first, valid second
    assert ll.parse_date("no date here") is None


def test_classify_url(ll):
    assert ll.classify_url("https://drive.google.com/file/d/abc/view") == "drive"
    assert ll.classify_url("https://ufl.zoom.us/rec/share/xyz") == "zoom"
    assert ll.classify_url("https://ufl.instructure.com/courses/1/files/2") == "canvas"
    assert ll.classify_url("https://www.youtube.com/watch?v=x") == "url"


def test_envelope_match_tolerates_offset_but_not_different_audio(ll):
    import numpy as np
    rng = np.random.default_rng(0)
    talk = rng.normal(-1.5, 0.4, 900).astype(np.float16)  # 15 min of "speech" loudness
    trimmed = talk[7:].copy()  # same lecture, trimmed 7 s, slightly different encode
    trimmed += rng.normal(0, 0.02, trimmed.size).astype(np.float16)
    other = rng.normal(-1.5, 0.4, 900).astype(np.float16)
    assert ll.envelope_match(talk.tobytes(), trimmed.tobytes()) > 0.95
    assert ll.envelope_match(talk.tobytes(), other.tobytes()) < 0.5
    silence = np.full(900, -4, np.float16)
    assert ll.envelope_match(silence.tobytes(), silence.tobytes()) == 0.0


V1_DB = __import__("pathlib").Path.home() / "Documents/Projects/lecturelab/lecturelab.db"


@pytest.mark.skipif(not V1_DB.exists(), reason="v1 database not on this machine")
def test_import_v1_brings_both_lectures_and_is_rerunnable(ll):
    first = ll.import_v1(V1_DB)
    assert [r.get("skipped") for r in first] == [None, None]
    by_title = {ll.get_lecture(r["lecture_id"])["title"]: ll.get_lecture(r["lecture_id"]) for r in first}
    assert by_title["Lecture 1"]["course_code"] == "COP4020"  # its audio link was NULL in v1; relinked
    assert by_title["Discussion 1"]["course_code"] == "PHI3681"
    for lec in by_title.values():
        assert lec["segment_count"] > 50 and lec["transcript_source"] == "v1"
        assert lec["status"] == "queued_transcribe"  # still gets a large-v3 transcript
        assert (ll.DATA / lec["audio_path"]).exists()
    again = ll.import_v1(V1_DB)
    assert all(r["duplicate"] for r in again)
    assert len(ll.list_lectures()) == 2


def test_scan_html_finds_real_embed_shapes(ll):
    html = """
    <p><iframe src="https://ufl.instructure.com/courses/5/external_tools/retrieve?display=borderless&amp;url=https%3A%2F%2F2375821.kaf.kaltura.com%2Fbrowseandembed%2Findex%2Fmedia%2Fentryid%2F1_ab12cd34%2FshowDescription%2Ffalse"></iframe></p>
    <iframe src="https://ufl.hosted.panopto.com/Panopto/Pages/Embed.aspx?id=0f8fad5b-d9cb-469f-a165-70867728950e&autoplay=false"></iframe>
    <a href="https://www.youtube.com/watch?v=dQw4w9WgXcQ">video</a>
    <a href="https://ufl.zoom.us/rec/share/AbC-123_xyz.Q9">recording</a>
    <iframe src="/media_objects_iframe/m-4uoGqVdEqXhpqu2ZMytHgy9XRV4xNkB"></iframe>"""
    found = {f["key"]: f for f in ll.scan_html_for_media(html, "page: Week 1")}
    assert found["kaltura:1_ab12cd34"]["url"] == "kaltura:2375821:1_ab12cd34"
    assert found["panopto:0f8fad5b-d9cb-469f-a165-70867728950e"]["url"] ==         "https://ufl.hosted.panopto.com/Panopto/Pages/Viewer.aspx?id=0f8fad5b-d9cb-469f-a165-70867728950e"
    assert "youtube:dQw4w9WgXcQ" in found
    assert any(k.startswith("zoom:https://ufl.zoom.us/rec/share/") for k in found)
    assert found["media_object:m-4uoGqVdEqXhpqu2ZMytHgy9XRV4xNkB"]["page_url"] == "page: Week 1"


class FakeResponse:
    def __init__(self, status, body, ctype="application/json", link=None):
        self.status, self.body = status, body
        self.headers = {"content-type": ctype, **({"link": link} if link else {})}
        self.ok = 200 <= status < 300

    def text(self):
        return self.body if isinstance(self.body, str) else __import__("json").dumps(self.body)

    def json(self):
        return self.body


def fake_session(ll, responses):
    """A CanvasSession whose browser context answers from a list, recording the URLs asked for."""
    cs = ll.CanvasSession()
    asked = []

    class Req:
        def get(self, url, **kw):
            asked.append(url)
            return responses.pop(0)
    cs.ctx = type("Ctx", (), {"request": Req()})()
    return cs, asked


def test_canvas_api_paginates_and_treats_404_as_empty(ll):
    nxt = '<https://ufl.instructure.com/api/v1/courses?page=2&per_page=100>; rel="next"'
    cs, asked = fake_session(ll, [FakeResponse(200, [{"id": 1}], link=nxt), FakeResponse(200, [{"id": 2}]),
                                  FakeResponse(404, {"errors": []})])
    assert cs.get("courses", {"include[]": ["term", "teachers"]}) == [{"id": 1}, {"id": 2}]
    assert "include%5B%5D=term&include%5B%5D=teachers" in asked[0] and asked[1].endswith("page=2&per_page=100")
    assert cs.get("courses/9/files") == []  # Files tab hidden: not an error
    with pytest.raises(ll.CanvasError, match="non-API"):
        cs.get("https://evil.example/api/v1/x")


def test_expired_session_marks_needs_signin_and_notifies_once(ll, monkeypatch):
    sent = []
    monkeypatch.setattr(ll, "notify", lambda t, b: sent.append(t))
    ll.set_canvas_state("connected", user="Kashi")
    for resp in (FakeResponse(401, {"status": "unauthenticated"}),
                 FakeResponse(302, "", ctype="text/html"),  # bounced to the GatorLink login page
                 FakeResponse(200, "<html>login</html>", ctype="text/html")):
        cs, _ = fake_session(ll, [resp])
        with pytest.raises(ll.CanvasNeedsSignIn):
            cs.get("users/self", paginate=False)
    assert ll.canvas_state()["state"] == "needs_signin"
    assert len(sent) == 1  # one notification per expiry, not one per request


def test_signin_polling_is_quiet(ll, monkeypatch):
    sent = []
    monkeypatch.setattr(ll, "notify", lambda t, b: sent.append(t))
    cs, _ = fake_session(ll, [FakeResponse(401, {})])
    with pytest.raises(ll.CanvasNeedsSignIn):
        cs.get("users/self", paginate=False, quiet_auth=True)
    assert ll.canvas_state()["state"] == "not_connected" and not sent


def test_needs_signin_never_opens_the_browser(ll, monkeypatch):
    ll.meta_set("canvas_state", {"state": "needs_signin"})
    monkeypatch.setattr(ll, "CanvasSession", lambda *a, **k: (_ for _ in ()).throw(AssertionError("opened browser")))
    with pytest.raises(ll.CanvasNeedsSignIn):
        ll.canvas_get("courses")
    with pytest.raises(ll.CanvasNeedsSignIn):
        ll.canvas_ready()


def test_netscape_cookie_export(ll):
    text = ll.netscape_cookies([
        {"name": "canvas_session", "value": "abc", "domain": "ufl.instructure.com", "path": "/", "secure": True,
         "expires": -1},
        {"name": "kms_ks", "value": "x=1", "domain": ".kaltura.com", "path": "/", "secure": False,
         "expires": 1893456000.5}])
    lines = text.splitlines()
    assert lines[0] == "# Netscape HTTP Cookie File"
    assert lines[2].split("\t") == ["ufl.instructure.com", "FALSE", "/", "TRUE", "0", "canvas_session", "abc"]
    assert lines[3].split("\t") == [".kaltura.com", "TRUE", "/", "FALSE", "1893456000", "kms_ks", "x=1"]


# ---- ingest robustness (review findings)

@pytest.fixture()
def noise(ll, tmp_path):
    """A 70 s audio file for plumbing tests (fingerprints need at least 60 s)."""
    out = tmp_path / "noise.wav"
    subprocess.run([ll.find_tool("ffmpeg"), "-v", "error", "-f", "lavfi", "-i",
                    "anoisesrc=d=70:c=pink:a=0.3,volume='if(lt(mod(t,7),3),1,0.2)':eval=frame",
                    "-ac", "1", "-ar", "16000", str(out)], check=True)
    return out


def test_failed_finalize_rolls_back_so_retry_works(ll, noise, monkeypatch):
    real = ll._move_into_place
    monkeypatch.setattr(ll, "_move_into_place", lambda *a: (_ for _ in ()).throw(PermissionError("locked")))
    with pytest.raises(PermissionError):
        ll.ingest_file(noise, kind="file")
    assert ll.list_lectures() == []  # no poisoned half-row left behind
    monkeypatch.setattr(ll, "_move_into_place", real)
    r = ll.ingest_file(noise, kind="file")
    assert not r["duplicate"] and ll.get_lecture(r["lecture_id"])["status"] == "queued_transcribe"


def test_ids_never_reused_and_leftover_dirs_cleared(ll, noise, tmp_path):
    first = ll.ingest_file(noise, kind="file")["lecture_id"]
    ll.delete_lecture(first)
    (ll.DATA / "media" / str(first + 1)).mkdir(parents=True)  # a leftover the next id would hit
    (ll.DATA / "media" / str(first + 1) / "junk").write_text("x")
    second = ll.ingest_file(noise, kind="file")["lecture_id"]
    assert second == first + 1
    assert not (ll.DATA / "media" / str(second) / "junk").exists()


def test_exact_reingest_of_parked_duplicate_does_not_park_again(ll, noise, tmp_path):
    a = ll.ingest_file(noise, kind="file")["lecture_id"]
    mp3 = tmp_path / "copy.mp3"
    subprocess.run([ll.find_tool("ffmpeg"), "-v", "error", "-i", str(noise), "-b:a", "64k", str(mp3)], check=True)
    parked = ll.ingest_file(mp3, kind="file")
    assert parked["parked_dup_of"] == a
    again = ll.ingest_file(mp3, kind="file")
    assert again["duplicate"] and again["lecture_id"] == parked["lecture_id"]
    assert len(ll.list_lectures()) == 2


def test_corrupt_audio_raises_instead_of_storing_bad_signature(ll, tmp_path):
    bad = tmp_path / "bad.flac"
    bad.write_bytes(b"fLaC" + b"\x00" * 5000)
    with pytest.raises(RuntimeError, match="could not decode"):
        ll.pcm_signature(bad)


def test_claim_job_hands_each_job_out_once(ll):
    import threading
    ids = [ll.enqueue("noop", {}) for _ in range(20)]
    got, lock = [], threading.Lock()

    def grab():
        while (j := ll.claim_job(["noop"])):
            with lock:
                got.append(j["id"])
    ts = [threading.Thread(target=grab) for _ in range(4)]
    [t.start() for t in ts]
    [t.join() for t in ts]
    assert sorted(got) == ids


def test_orphaned_jobs_requeued_then_given_up(ll):
    jid = ll.enqueue("noop", {})
    with ll.db() as con:
        con.execute("UPDATE jobs SET status='running', owner_pid=999999, attempts=1 WHERE id=?", (jid,))
    live = ll.enqueue("noop", {})
    with ll.db() as con:
        con.execute("UPDATE jobs SET status='running', owner_pid=? WHERE id=?", (__import__("os").getpid(), live))
    ll.requeue_orphaned_jobs()
    assert ll.get_job(jid)["status"] == "queued" and ll.get_job(live)["status"] == "running"
    with ll.db() as con:
        con.execute("UPDATE jobs SET status='running', owner_pid=999999, attempts=3 WHERE id=?", (jid,))
    ll.requeue_orphaned_jobs()
    assert ll.get_job(jid)["status"] == "error"


def test_inbox_waits_for_stable_file_and_retries_failures(ll, noise, tmp_path, monkeypatch):
    import shutil
    monkeypatch.setattr(ll, "LOCAL_INBOX", tmp_path / "inbox")
    monkeypatch.setenv("LECTURELAB_DRIVE_INBOX", "off")
    ll.scan_inboxes()  # creates course folders
    shutil.copy(noise, tmp_path / "inbox" / "COP4020" / "lec.wav")
    assert ll.scan_inboxes() == []  # first sighting: wait for a stable size
    [jid] = ll.scan_inboxes()
    assert ll.get_job(jid)["payload"]["course_id"] == ll.find_course("COP4020")["id"]
    assert ll.scan_inboxes() == []  # unchanged and its job hasn't failed: leave it
    ll.job_update(jid, status="error", error="boom")
    [retry] = ll.scan_inboxes()  # failed job: try again
    assert retry != jid


def test_canvas_same_embed_in_two_courses_kept_for_both(ll, monkeypatch):
    page = {"body": '<iframe src="https://www.youtube.com/embed/dQw4w9WgXcQ"></iframe>', "title": "Wk1"}
    def fake_get(path, params=None, paginate=True):
        if path.endswith("/pages"):
            return [{"url": "wk1"}]
        if "/pages/" in path:
            return page
        return [] if paginate else None
    monkeypatch.setattr(ll, "canvas_get", fake_get)
    monkeypatch.setattr(ll, "CanvasSession", lambda *a, **k: __import__("contextlib").nullcontext())  # no browser
    ll.set_canvas_state("connected", user="Kashi")
    a, b = ll.find_course("COP4020"), ll.find_course("COP4533")
    for c, cid in ((a, 11), (b, 22)):
        c["canvas_id"] = cid
        assert ll.canvas_scan_course(c)["new"] == 1
    with ll.db() as con:
        rows_ = con.execute("SELECT page_link FROM canvas_items WHERE key='youtube:dQw4w9WgXcQ'").fetchall()
    assert [r[0] for r in rows_] == ["https://ufl.instructure.com/courses/11/pages/wk1",
                                     "https://ufl.instructure.com/courses/22/pages/wk1"]  # opened before a pull


def test_source_urls_drop_zoom_passcode(ll):
    assert ll.clean_source_url("https://ufl.zoom.us/rec/share/abc?pwd=s3cret&startTime=1") == \
        "https://ufl.zoom.us/rec/share/abc?startTime=1"


def test_utc_timestamps_use_local_date(ll):
    # 2026-09-03 01:30 UTC is the evening of Sep 2 in Gainesville (UTC-4 in September);
    # whatever this machine's zone is, the stored date must be the local one.
    from datetime import datetime, timezone
    expected = datetime(2026, 9, 3, 1, 30, tzinfo=timezone.utc).astimezone().date().isoformat()
    assert ll.parse_date("2026-09-03T01:30:00Z") == expected


def test_api_upload_rejects_bad_names_and_media_404(ll):
    from fastapi.testclient import TestClient
    c = TestClient(ll.create_app(), base_url=f"http://127.0.0.1:{ll.PORT}")
    r = c.post("/api/recordings/upload", files={"file": ("notes.exe", b"x")})
    assert r.status_code == 400
    r = c.post("/api/recordings/upload", files={"file": (r"..\..\CON.mp4", b"x")}, data={"course_id": "999"})
    assert r.status_code == 400 and "no course" in r.text
    assert c.get("/api/lectures/1/media/..%2F..%2Fx").status_code == 404
    with ll.db() as con:
        con.execute("INSERT INTO lectures(title, status, audio_path) VALUES ('t','queued_transcribe','media/9/audio.flac')")
    assert c.get("/api/lectures/1/media/audio").status_code == 404  # row points at a missing file
    assert c.get("/api/health", headers={"host": "evil.example"}).status_code == 403


def test_orphan_sweep_skips_while_a_batch_runs(ll, monkeypatch):
    calls = []
    monkeypatch.setattr(ll, "stop_orphan_sessions", lambda keep=None: calls.append(1) or [])
    batch = ll.FileLock(ll.DATA / "transcribe.lock")
    assert batch.acquire()  # a batch is running (in this or another process)
    ll.sweep_orphans_if_idle()
    assert calls == []  # must not stop the session that batch is using
    batch.release()
    ll.sweep_orphans_if_idle()
    assert calls == [1]


def test_colab_exec_that_returns_early_fails_the_batch_and_stops(ll, monkeypatch, tmp_path):
    calls = []

    def fake_colab(*args, timeout=0):
        calls.append(args)
        if args[0] == "download":
            raise RuntimeError("404")  # nothing was written on the VM
        return ""

    monkeypatch.setattr(ll, "colab", fake_colab)
    monkeypatch.setattr(ll, "wslpath", lambda p: str(p))
    monkeypatch.setattr(ll, "stop_orphan_sessions", lambda keep=None: [])
    monkeypatch.setattr(ll, "colab_units", lambda: 0.0)
    monkeypatch.setattr(ll, "colab_rate", lambda: 1.07)
    monkeypatch.setattr(ll, "colab_new_session", lambda s: "T4")
    monkeypatch.setattr(ll, "colab_upload", lambda *a: None)
    audio = tmp_path / "a.flac"
    audio.write_bytes(b"x")
    with ll.db() as con:
        bid = con.execute("INSERT INTO transcribe_batches(mode, lecture_ids, audio_sec) VALUES ('colab','[1]',60)").lastrowid
    lec = {"id": 1, "audio_path": str(audio), "course_id": None, "title": "t"}
    monkeypatch.setattr(ll, "initial_prompt", lambda l: "")
    with pytest.raises(RuntimeError, match="0 of 1 transcript"):
        ll.transcribe_batch_colab([lec], bid)
    execs = [a for a in calls if a[0] == "exec"]
    assert execs and all("--timeout" in a for a in execs)  # the CLI default (30 s) cut batch 2 short
    assert calls[-1][:3] == ("stop", "-s", f"{ll.SESSION_PREFIX}{bid}")


def test_canvas_session_cookies_survive_a_browser_restart(ll):
    class Ctx:
        def __init__(self, jar):
            self.jar = jar

        def cookies(self):
            return list(self.jar)

        def add_cookies(self, cs):
            self.jar += cs

    sess = {"name": "canvas_session", "value": "v1", "domain": "ufl.instructure.com", "path": "/", "expires": -1}
    keep = {"name": "_ga", "value": "g", "domain": ".ufl.edu", "path": "/", "expires": 1900000000}
    cs = ll.CanvasSession()
    cs.ctx = Ctx([sess, keep])
    cs.save_session_cookies()
    assert [c["name"] for c in ll.json.loads(ll.SESSION_COOKIES.read_text())] == ["canvas_session"]
    cs.ctx = Ctx([keep])  # a relaunch: Chromium dropped the session cookie
    assert cs.restore_session_cookies() == 1
    assert sess in cs.ctx.jar
    fresher = dict(sess, value="v2")
    cs.ctx = Ctx([fresher, keep])  # the profile already has a newer one: leave it alone
    assert cs.restore_session_cookies() == 0
    assert cs.ctx.jar == [fresher, keep]


def test_drive_folder_in_a_module_becomes_one_item_per_recording(ll, monkeypatch):
    folder = "https://drive.google.com/drive/folders/1FakeCourseFolder0000000000000000?usp=sharing"
    modules = [{"name": "Resources", "items": [
        {"type": "ExternalUrl", "title": "Lecture Slides/Recordings (Drive)", "external_url": folder,
         "html_url": "https://ufl.instructure.com/courses/11/modules/items/5"},
        {"type": "ExternalUrl", "title": "One clip", "external_url": "https://drive.google.com/file/d/1AbCdEfGhIjKlMnOpQrStUvWxYz/view"}]}]

    def fake_get(path, params=None, paginate=True):
        return modules if path.endswith("/modules") else ([] if paginate else None)

    listed = []
    monkeypatch.setattr(ll, "canvas_get", fake_get)
    monkeypatch.setattr(ll, "CanvasSession", lambda *a, **k: __import__("contextlib").nullcontext())
    ll.set_canvas_state("connected", user="Kashi")
    monkeypatch.setattr(ll, "drive_list_folder", lambda fid: listed.append(fid) or [
        {"id": "1FakeRecording00000000000000000-", "path": "M0L1 Course Overview.mkv"},
        {"id": "1FakeSlides000000000000000000000", "path": "M0L1 Course Overview.pdf"}])
    c = ll.find_course("COP4020")
    c["canvas_id"] = 11
    assert ll.canvas_scan_course(c)["new"] == 2
    assert listed == ["1FakeCourseFolder0000000000000000"]
    with ll.db() as con:
        got = {r[0]: r[1:] for r in con.execute("SELECT key, kind, title, url, page_link FROM canvas_items")}
    assert got["drive:1FakeRecording00000000000000000-"] == (
        "drive", "M0L1 Course Overview", "https://drive.google.com/file/d/1FakeRecording00000000000000000-/view",
        "https://ufl.instructure.com/courses/11/modules/items/5")
    assert "drive:1AbCdEfGhIjKlMnOpQrStUvWxYz" in got  # a single shared file link
    assert not any(k.startswith("drive_folder:") for k in got)  # the folder itself isn't an item
    assert ll.classify_url(got["drive:1FakeRecording00000000000000000-"][2]) == "drive"


def test_unlistable_drive_folder_is_kept_as_one_item(ll, monkeypatch):
    items = {"drive_folder:X": {"key": "drive_folder:X", "kind": "drive_folder", "url": "u", "page_url": "p"}}
    monkeypatch.setattr(ll, "drive_list_folder", lambda fid: (_ for _ in ()).throw(RuntimeError("private")))
    ll.expand_drive_folders(items)
    assert list(items) == ["drive_folder:X"]


def test_drive_file_id_from_share_link_shapes(ll):
    for url in ["https://drive.google.com/file/d/1FakeRecording00000000000000000-/view?usp=sharing",
                "https://drive.google.com/file/u/0/d/1FakeRecording00000000000000000-/view",
                "https://drive.google.com/open?id=1FakeRecording00000000000000000-",
                "https://drive.google.com/uc?export=download&id=1FakeRecording00000000000000000-"]:
        assert ll._DRIVE_FILE_ID.search(url)[1] == "1FakeRecording00000000000000000-", url


# ---- Stage 2 review fixes

def _lecture(ll, status, audio=True):
    with ll.db() as con:
        lid = con.execute("INSERT INTO lectures(title, status, duration_sec) VALUES ('t', ?, 60)", (status,)).lastrowid
        if audio:
            f = ll.DATA / "media" / str(lid) / "audio.flac"
            f.parent.mkdir(parents=True, exist_ok=True)
            f.write_bytes(b"x")
            con.execute("UPDATE lectures SET audio_path=? WHERE id=?", (f"media/{lid}/audio.flac", lid))
    return lid


def _fake_local(ll, monkeypatch, fail=False):
    def fake(lec, progress=None, model_name=None):
        if fail:
            raise RuntimeError("cpu blew up")
        return {"model": "base", "segments": [{"start": 0, "end": 1, "text": "hi", "words": []}]}
    monkeypatch.setattr(ll, "transcribe_local", fake)


def test_transcribe_all_waiting_skips_failed_lectures(ll, monkeypatch):
    _fake_local(ll, monkeypatch)
    queued, failed = _lecture(ll, "queued_transcribe"), _lecture(ll, "error")
    out = ll.job_transcribe({"id": 1, "payload": {"mode": "local"}})
    assert out["transcribed"] == [queued]
    assert ll.get_lecture(failed)["status"] == "error"
    assert ll.job_transcribe({"id": 2, "payload": {"mode": "local", "lecture_ids": []}})["transcribed"] == []
    assert ll.job_transcribe({"id": 3, "payload": {"mode": "local", "lecture_ids": [failed]}})["transcribed"] == [failed]


def test_missing_audio_is_marked_and_never_sent_to_colab(ll, monkeypatch):
    monkeypatch.setattr(ll, "transcribe_batch_colab", lambda *a, **k: pytest.fail("nothing should go to Colab"))
    lid = _lecture(ll, "queued_transcribe")
    (ll.DATA / "media" / str(lid) / "audio.flac").unlink()
    assert ll.job_transcribe({"id": 1, "payload": {"mode": "colab"}})["transcribed"] == []
    assert "audio file missing" in ll.get_lecture(lid)["error"]
    assert ll.meta_get("colab_failures", 0) == 0  # a local problem doesn't count against Colab


def test_colab_backoff_expires(ll):
    ll.meta_set("colab_failures", 2)
    ll.meta_set("colab_failed_at", ll.time.time())
    assert ll.colab_backed_off()
    ll.meta_set("colab_failed_at", ll.time.time() - (ll.COLAB_BACKOFF_H * 3600 + 1))
    assert not ll.colab_backed_off()  # auto mode tries Colab again


def test_failed_local_batch_is_closed_and_stale_rows_closed_at_startup(ll, monkeypatch):
    _fake_local(ll, monkeypatch, fail=True)
    _lecture(ll, "queued_transcribe")
    with pytest.raises(RuntimeError, match="cpu blew up"):
        ll.job_transcribe({"id": 1, "payload": {"mode": "local"}})
    with ll.db() as con:
        assert con.execute("SELECT COUNT(*) FROM transcribe_batches WHERE ended_at IS NULL").fetchone()[0] == 0
        con.execute("INSERT INTO transcribe_batches(mode, lecture_ids) VALUES ('colab', '[]')")  # a killed process
    ll.reset_stuck_transcribing()
    with ll.db() as con:
        assert con.execute("SELECT COUNT(*) FROM transcribe_batches WHERE ended_at IS NULL").fetchone()[0] == 0


def test_run_errors_name_the_wsl_tool_and_redact_tokens(ll, monkeypatch):
    class R:
        returncode, stderr = 1, ""
        stdout = 'Error: https://x/proxy?token=abc123&authuser=0 {"token": "sekrit"}'
    monkeypatch.setattr(ll.subprocess, "run", lambda *a, **k: R())
    with pytest.raises(RuntimeError) as e:
        ll.run(["wsl.exe", "-d", "Ubuntu", "-e", "/home/u/.local/bin/colab", "usage"])
    msg = str(e.value)
    assert msg.startswith("colab failed (1)"), msg
    assert "abc123" not in msg and "sekrit" not in msg


def test_colab_usage_parsing_uses_the_real_format(ll):
    real = "Current balance: 0.00 compute units\nUsage rate: 1.07/hr\nActive assignments: 1\n"
    assert ll.parse_colab_units(real) == 0.0
    assert ll.parse_colab_rate(real) == 1.07
    assert ll.parse_colab_units("Current balance: 1,234.50 compute units\n") == 1234.5
    for junk in ["", "Usage rate: 1.76 units/hr", "Error: not logged in", "available: 3 sessions"]:
        assert ll.parse_colab_units(junk) is None, junk


def test_unreadable_balance_fails_closed_unless_colab_was_asked_for(ll, monkeypatch):
    monkeypatch.setattr(ll, "stop_orphan_sessions", lambda keep=None: [])
    monkeypatch.setattr(ll, "colab_units", lambda: None)
    with pytest.raises(ll.LowUnits):
        ll.transcribe_batch_colab([], 1)
    monkeypatch.setattr(ll, "colab_new_session", lambda s: (_ for _ in ()).throw(RuntimeError("got past the guard")))
    monkeypatch.setattr(ll, "colab", lambda *a, **k: "")
    with pytest.raises(RuntimeError, match="got past the guard"):
        ll.transcribe_batch_colab([], 1, force=True)


def test_colab_stop_is_retried_once(ll, monkeypatch):
    stops = []

    def fake_colab(*args, timeout=0):
        if args[0] == "stop":
            stops.append(args)
            if len(stops) == 1:
                raise RuntimeError("wsl hiccup")
        return ""
    monkeypatch.setattr(ll, "colab", fake_colab)
    monkeypatch.setattr(ll, "stop_orphan_sessions", lambda keep=None: [])
    monkeypatch.setattr(ll, "colab_units", lambda: 0.0)
    monkeypatch.setattr(ll, "colab_new_session", lambda s: (_ for _ in ()).throw(RuntimeError("no GPU")))
    with pytest.raises(RuntimeError, match="no GPU"):
        ll.transcribe_batch_colab([], 7)
    assert len(stops) == 2


def test_canvas_outage_is_not_an_expired_session(ll, monkeypatch):
    monkeypatch.setattr(ll, "notify", lambda t, b: pytest.fail("no sign-in toast for an outage"))
    ll.set_canvas_state("connected", user="Kashi")
    for resp in (FakeResponse(502, "<html>Bad gateway</html>", ctype="text/html"),
                 FakeResponse(503, {"errors": "maintenance"}),
                 FakeResponse(200, "<html>Canvas is down for maintenance</html>", ctype="text/html")):
        cs, _ = fake_session(ll, [resp])
        with pytest.raises(ll.CanvasError) as e:
            cs.get("courses")
        assert not isinstance(e.value, ll.CanvasNeedsSignIn)
    assert ll.canvas_state()["state"] == "connected"


def test_canvas_get_never_opens_a_browser_before_you_connect(ll, monkeypatch):
    monkeypatch.setattr(ll, "CanvasSession", lambda *a, **k: pytest.fail("opened the browser"))
    monkeypatch.setattr(ll, "notify", lambda t, b: pytest.fail("no toast"))
    with pytest.raises(ll.CanvasError, match="isn't connected"):
        ll.canvas_get("courses")
    assert ll.canvas_state()["state"] == "not_connected"


def test_canvas_session_cleans_up_when_launch_half_fails(ll, monkeypatch):
    closed = []

    class Ctx:
        def cookies(self):
            return []

        def close(self):
            closed.append("ctx")

    class PW:
        def __init__(self):
            self.chromium = self

        def launch_persistent_context(self, *a, **k):
            return Ctx()

        def start(self):
            return self

        def stop(self):
            closed.append("pw")

    import playwright.sync_api as psa
    monkeypatch.setattr(psa, "sync_playwright", lambda: PW())
    monkeypatch.setattr(ll, "_prepare_profile", lambda: None)
    monkeypatch.setattr(ll.CanvasSession, "restore_session_cookies",
                        lambda self: (_ for _ in ()).throw(ValueError("bad cookie")))
    with pytest.raises(ValueError):
        with ll.CanvasSession():
            pass
    assert closed == ["ctx", "pw"]
    lock = ll.FileLock(ll.DATA / "browser.lock")
    assert lock.acquire()  # the lock was freed too
    lock.release()


def test_login_cookies_only_for_login_hosts(ll):
    yes = ["https://ufl.hosted.panopto.com/Panopto/Pages/Viewer.aspx?id=1", "https://cdnapisec.kaltura.com/p/1/x",
           "https://ufl.instructure.com/courses/1", "https://mediasite.example.edu/Mediasite/Play/abc"]
    no = ["https://evil.example/?r=ufl.edu", "https://notufl.edu.example.com/x", "https://www.youtube.com/watch?v=x"]
    assert all(ll.needs_login_cookies(u) for u in yes)
    assert not any(ll.needs_login_cookies(u) for u in no)


def test_canvas_pull_of_a_public_item_needs_no_browser_and_leaves_no_cookies(ll, monkeypatch, noise):
    monkeypatch.setattr(ll, "CanvasSession", lambda *a, **k: pytest.fail("YouTube/Drive/Zoom need no browser"))
    c = ll.find_course("COP4020")
    with ll.db() as con:
        con.execute("INSERT INTO canvas_items(key, course_id, kind, title, url) VALUES "
                    "('drive:abc', ?, 'drive', 'M0L1', 'https://drive.google.com/file/d/abc/view')", (c["id"],))

    def fake_download(url, dest, **kw):
        dest.mkdir(parents=True, exist_ok=True)
        f = dest / "clip.wav"
        f.write_bytes(noise.read_bytes())
        return [(f, {"title": "clip"})]
    monkeypatch.setattr(ll, "download_url", fake_download)
    out = ll.job_canvas_pull({"id": 5, "payload": {"course_id": c["id"], "key": "drive:abc"}})
    assert out["lectures"][0]["lecture_id"]
    assert not ll.job_cookie_file(5).exists()


def test_scan_never_opens_a_browser_when_signed_out(ll, monkeypatch):
    monkeypatch.setattr(ll, "CanvasSession", lambda *a, **k: pytest.fail("opened the browser"))
    monkeypatch.setattr(ll, "notify", lambda t, b: None)
    c = ll.find_course("COP4020")
    c["canvas_id"] = 11
    for state, err in (("not_connected", ll.CanvasError), ("needs_signin", ll.CanvasNeedsSignIn)):
        ll.set_canvas_state(state)
        with pytest.raises(err):
            ll.canvas_scan_course(c)


def test_run_error_keeps_stderr_when_stdout_is_binary(ll, monkeypatch):
    class R:
        returncode, stdout, stderr = 1, b"\x00" * 10_000_000, b"Invalid data found when processing input"
    monkeypatch.setattr(ll.subprocess, "run", lambda *a, **k: R())
    with pytest.raises(RuntimeError) as e:
        ll.run(["ffmpeg", "-i", "x", "-f", "f32le", "-"])
    assert str(e.value) == "ffmpeg failed (1): Invalid data found when processing input"
    class T:
        returncode, stdout, stderr = 2, "", "exit code: 2, see https://x/?code=abc&state=xyz"
    monkeypatch.setattr(ll.subprocess, "run", lambda *a, **k: T())
    with pytest.raises(RuntimeError) as e:
        ll.run(["tool"])
    assert "exit code: 2" in str(e.value) and "abc" not in str(e.value) and "xyz" not in str(e.value)


def test_canvas_pull_jobs_are_labelled_with_the_recording_title(ll):
    from fastapi.testclient import TestClient
    c = ll.find_course("COP4020")
    with ll.db() as con:
        con.execute("INSERT INTO canvas_items(key, course_id, kind, title, url) VALUES "
                    "('drive:abc', ?, 'drive', 'M2L1 Lexer Introduction', 'u')", (c["id"],))
    ll.canvas_pull(c["id"], ["drive:abc"])
    client = TestClient(ll.create_app(), base_url=f"http://127.0.0.1:{ll.PORT}")
    jobs = client.get("/api/jobs").json()
    assert [j["label"] for j in jobs if j["kind"] == "canvas_pull"] == ["M2L1 Lexer Introduction"]


# ---------------------------------------------------------------- slides (Stage 3)
SLIDE_TEXT = [
    "Course overview: grading policy, office hours, exams and the syllabus schedule",
    "Lexical analysis: a lexer turns characters into tokens using regular expressions",
    "Finite automata: a DFA has states, an alphabet, transitions and accepting states",
    "Parsing: a recursive descent parser builds a syntax tree from the token stream",
    "Type checking: the checker walks the syntax tree and rejects ill typed programs",
    "Interpreters: evaluate the syntax tree directly with an environment of bindings",
    "Garbage collection: mark and sweep frees heap objects that are unreachable",
]


def _pdf(path, pages, sizes=None):
    import pymupdf as fitz
    doc = fitz.open()
    for i, text in enumerate(pages):
        w, h = (sizes or {}).get(i, (960, 540))
        p = doc.new_page(width=w, height=h)
        p.insert_text((40, 80), text.split(":")[0], fontsize=32)
        p.insert_textbox(fitz.Rect(40, 120, w - 40, h - 40), text, fontsize=16)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def _deck(ll, title, course_id, path=None, source="drive", status="new", created_at=None):
    with ll.db() as con:
        return con.execute("INSERT INTO decks(course_id, key, source, title, path, status, created_at) "
                           "VALUES (?,?,?,?,?,?,?)", (course_id, f"{source}:{title}", source, title, path, status,
                                                      created_at)).lastrowid


def _ready_deck(ll, title, course_id, pages=SLIDE_TEXT):
    rel = f"Slides/test/{title}.pdf"
    _pdf(ll.DATA / rel, pages)
    did = _deck(ll, title, course_id, path=rel)
    ll.extract_deck(did)
    return did


def _lec(ll, title, course_id, segs=(), video=None, date=None):
    with ll.db() as con:
        lid = con.execute("INSERT INTO lectures(title, course_id, date, status, duration_sec, video_path) "
                          "VALUES (?,?,?,'transcribed',600,?)", (title, course_id, date, video)).lastrowid
        con.executemany("INSERT INTO segments(lecture_id, idx, start, end, text) VALUES (?,?,?,?,?)",
                        [(lid, i, s, e, t) for i, (s, e, t) in enumerate(segs)])
    return lid


def _said(slide_order, seg_sec=15):
    """Two 15 s segments per 30 s window, each window talking about one slide."""
    segs = []
    for w, s in enumerate(slide_order):
        body = SLIDE_TEXT[s].split(":")[1] if s is not None else "okay so yeah let me just grab some water"
        for half in range(2):
            t = (2 * w + half) * seg_sec
            segs.append((t, t + seg_sec, f"so um{body} right"))
    return segs


def test_match_score_reads_codes_titles_and_dates(ll):
    s = lambda d, l, **kw: ll.match_score({"title": d, **kw}, {"title": l, "date": kw.get("date")})  # noqa: E731
    assert s("M2L1 Lexer Introduction", "M2L1 Lexer Introduction") >= 1.0
    assert s("L07 slides", "Lecture 7: Parsing") >= ll.MATCH_MIN          # same code, nothing else
    assert s("Regular Languages", "M2L2 Regular Languages") >= ll.MATCH_MIN  # same title, one side has a code
    assert s("Lecture 8 Parsing", "Lecture 7: Parsing") < 0                # codes contradict
    assert s("Syllabus", "M2L1 Lexer Introduction") < ll.MATCH_MIN
    assert s("M2TL1 Error Handling", "M2TL1 Error Handling") >= 1.0
    assert s("M2L1 Error Handling", "M2TL1 Error Handling") < 0  # TL is a different lecture from L
    assert s("M3TL2-3 Syntax Review", "M3TL3 Syntax Review (cont)") >= ll.MATCH_MIN  # a range covers it
    near =ll.match_score({"title": "Week notes", "created_at": "2026-09-01T10:00:00Z"},
                          {"title": "Notes", "date": "2026-09-02"})
    far = ll.match_score({"title": "Week notes", "created_at": "2026-06-01T10:00:00Z"},
                         {"title": "Notes", "date": "2026-09-02"})
    assert near > far


def test_match_decks_is_one_to_one_skips_ties_and_respects_your_pick(ll):
    cid = ll.find_course("COP4020")["id"]
    a, b = _lec(ll, "M0L1 Course Overview", cid), _lec(ll, "M2L1 Lexer Introduction", cid)
    c, d = _lec(ll, "Review session", cid), _lec(ll, "Review session", cid)
    d1, d2 = _deck(ll, "M0L1 Course Overview", cid), _deck(ll, "M2L1 Lexer Introduction", cid)
    d3 = _deck(ll, "Review session", cid)
    assert sorted(ll.match_decks(cid)) == [a, b]
    deck_of = lambda lid: ll.get_lecture(lid)["deck_id"]  # noqa: E731
    assert (deck_of(a), deck_of(b), deck_of(c), deck_of(d)) == (d1, d2, None, None)  # c/d tie: you pick
    assert ll.match_decks(cid) == []  # stable
    ll.set_lecture_deck(b, d1)  # you say deck 1 belongs to lecture b
    assert ll.get_lecture(b)["deck_how"] == "manual"
    assert deck_of(a) is None  # deck 1 is no longer auto-assigned to a
    ll.match_decks(cid)
    assert (deck_of(a), deck_of(b)) == (None, d1)
    ll.set_lecture_deck(c, d3)
    assert deck_of(d) is None and ll.lecture_slides(c)["deck"]["how"] == "manual"
    # one deck for two lectures when its name says so, but still only one per code
    e, f = _lec(ll, "M3TL2 Syntax Review", cid), _lec(ll, "M3TL3 Syntax Review (cont)", cid)
    g = _lec(ll, "M3TL4 Syntax Review again", cid)
    both = _deck(ll, "M3TL2-3 Syntax Review", cid)
    ll.match_decks(cid)
    assert (deck_of(e), deck_of(f), deck_of(g)) == (both, both, None)


def test_align_follows_skips_one_step_back_and_quiet_stretches(ll):
    truth = [0, 0, 1, 1, 2, 2, 3, 2, 3, 4, 4, 6, 6]  # skips slide 5, goes back from 3 to 2 once
    wins = [SLIDE_TEXT[s].split(":")[1] + " um so basically" for s in truth]
    wins.insert(5, "okay so yeah let me just grab some water")  # nothing slide-related said
    expect = truth[:5] + [truth[4]] + truth[5:]
    assert ll.align_windows(ll.tfidf_sim(wins, SLIDE_TEXT)) == expect


def test_align_gives_up_when_no_slide_text_matches(ll):
    import numpy as np
    assert ll.align_windows(np.zeros((4, 3))) is None
    cid = ll.find_course("COP4020")["id"]
    did = _ready_deck(ll, "Unrelated", cid, pages=["Medieval poetry: sonnets, ballads and heroic couplets"] * 2)
    lid = _lec(ll, "Lexers", cid, segs=_said([1, 2, 3]))
    with ll.db() as con:
        con.execute("UPDATE lectures SET deck_id=?, deck_how='manual' WHERE id=?", (did, lid))
    with pytest.raises(ll.DeckError, match="right deck"):
        ll.align_lecture(lid)


def test_pdf_pages_get_text_titles_and_same_size_renders(ll):
    import pymupdf as fitz
    cid = ll.find_course("COP4020")["id"]
    rel = "Slides/COP4020/mixed.pdf"
    _pdf(ll.DATA / rel, SLIDE_TEXT[:3], sizes={1: (540, 960)})  # page 2 is portrait
    did = _deck(ll, "mixed", cid, path=rel)
    assert ll.extract_deck(did) == {"pages": 3}
    deck = ll.get_deck(did)
    assert (deck["status"], deck["pages"], deck["thumb_w"], deck["thumb_h"]) == ("ready", 3, 320, 180)
    with ll.db() as con:
        titles = [r[0] for r in con.execute("SELECT title FROM slides WHERE deck_id=? ORDER BY no", (did,))]
    assert titles == ["Course overview", "Lexical analysis", "Finite automata"]
    for no in (1, 2, 3):
        for w, h in ((320, 180), (1280, 720)):
            pix = fitz.Pixmap(str(ll.SLIDE_IMAGES / str(did) / f"{no}-{w}.png"))
            assert (pix.width, pix.height) == (w, h), (no, w)


def _pptx(path, slides):
    from pptx import Presentation
    prs = Presentation()
    for title, body in slides:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        s.placeholders[1].text = body
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


def test_pptx_is_converted_with_libreoffice(ll):
    cid = ll.find_course("COP4020")["id"]
    _pptx(ll.DATA / "Slides/COP4020/deck.pptx", [("Lexers", "Tokens come from regular expressions"),
                                                 ("Parsers", "Recursive descent builds the tree")])
    did = _deck(ll, "deck", cid, path="Slides/COP4020/deck.pptx")
    assert ll.extract_deck(did) == {"pages": 2}
    assert ll.get_deck(did)["thumb_w"] == 320
    assert (ll.SLIDE_IMAGES / str(did) / "2-1280.png").is_file()


def test_pptx_falls_back_to_text_and_ppt_reports_how_to_fix(ll, monkeypatch):
    cid = ll.find_course("COP4020")["id"]
    monkeypatch.setattr(ll, "pptx_to_pdf", lambda src, out: (_ for _ in ()).throw(RuntimeError("soffice died")))
    _pptx(ll.DATA / "Slides/COP4020/deck.pptx", [("Lexers", "Tokens come from regular expressions")])
    did = _deck(ll, "deck", cid, path="Slides/COP4020/deck.pptx")
    ll.extract_deck(did)
    assert ll.get_deck(did)["thumb_w"] is None
    with ll.db() as con:
        assert con.execute("SELECT title, text FROM slides WHERE deck_id=?", (did,)).fetchone()[:] == (
            "Lexers", "Lexers\nTokens come from regular expressions")
    (ll.DATA / "Slides/COP4020/old.ppt").write_bytes(b"x")
    old = _deck(ll, "old", cid, path="Slides/COP4020/old.ppt")
    with pytest.raises(ll.DeckError, match="Save it as PDF"):
        ll.extract_deck(old)


def test_find_screens_drops_transients_and_folds_revisits(ll):
    import numpy as np
    a = np.full((90, 160), 250, np.uint8)
    b, c = a.copy(), a.copy()
    b[10:20] = 0                   # a different screen
    c[40:42, 20:120] = 30          # one new line of text: 1.4% of the picture
    frames = np.stack([a, a, a, b, c, c, c, a, a])
    frames[:, 70:90, 130:160] = np.random.default_rng(1).integers(0, 255, (9, 20, 30))  # a webcam inset
    ts = [0, 4, 8, 12, 16, 20, 24, 28, 32]
    screens = ll.find_screens(frames, ts, end=36)
    assert [s["times"] for s in screens] == [[[0, 12], [28, 36]], [[16, 28]]]  # b was on for 4 s
    assert [s["grab"] for s in screens] == [2, 6]  # the last frame of the first run
    assert ll.find_screens(frames[:1], ts[:1], end=10)[0]["times"] == [[0, 10]]


def test_slides_job_aligns_and_the_api_shows_revisits(ll):
    cid = ll.find_course("COP4020")["id"]
    did = _ready_deck(ll, "M2L1 Lexer Introduction", cid)
    lid = _lec(ll, "M2L1 Lexer Introduction", cid, segs=_said([0, 1, 1, 2, 1, 2, 3]))
    out = ll.job_slides({"id": 0, "payload": {"lecture_id": lid}})
    assert out["segments"] == 14
    with ll.db() as con:
        nos = [r[0] for r in con.execute("SELECT slide_no FROM segments WHERE lecture_id=? ORDER BY idx", (lid,))]
    assert nos == [n + 1 for n in (0, 0, 1, 1, 1, 1, 2, 2, 1, 1, 2, 2, 3, 3)]
    s = ll.lecture_slides(lid)
    assert (s["state"], s["deck"]["id"], s["deck"]["how"], s["deck"]["thumb_h"]) == ("aligned", did, "auto", 180)
    by_no = {x["no"]: x for x in s["slides"]}
    assert by_no[2]["ranges"] == [[30.0, 90.0], [120.0, 150.0]] and by_no[2]["first"] == 30.0
    assert by_no[7]["ranges"] == [] and by_no[7]["first"] is None  # never shown
    assert by_no[1]["thumb"] == f"/media/slides/{did}/1-320.png?v={int(ll.get_deck(did)['extracted_at'])}"
    assert s["candidates"][0]["id"] == did and s["candidates"][0]["pages"] == 7
    assert "Lexical analysis" in ll.lecture_slide_terms(lid)
    assert "Slides: " in ll.initial_prompt(ll.get_lecture(lid))


def test_slides_states_say_what_is_missing(ll):
    cid = ll.find_course("COP4020")["id"]
    state = lambda lid: (lambda s: (s["state"], s["reason"] or s["error"]))(ll.lecture_slides(lid))  # noqa: E731
    bare = _lec(ll, "Office hours", cid)
    assert state(bare)[0] == "none"
    video = _lec(ll, "Talk", cid, video="media/x/original.mp4")
    with ll.db() as con:
        con.execute("DELETE FROM segments WHERE lecture_id=?", (video,))
    assert state(video) == ("waiting", "no transcript yet")
    pending = _lec(ll, "M2L1 Lexer Introduction", cid, segs=_said([1]))
    _deck(ll, "M2L1 Lexer Introduction", cid)  # listed on Canvas, not downloaded yet
    ll.match_decks(cid)
    assert state(pending) == ("waiting", "downloading the slides")
    with ll.db() as con:
        con.execute("UPDATE lectures SET slides_error='The deck has no pages. Attach the deck again.' WHERE id=?",
                    (pending,))
    assert state(pending) == ("error", "The deck has no pages. Attach the deck again.")
    ll.set_lecture_deck(bare, None)
    assert state(bare) == ("none", "You chose no slides for this lecture.")


def test_api_slide_upload_errors_pick_and_images(ll, monkeypatch, tmp_path):
    from fastapi.testclient import TestClient
    c = TestClient(ll.create_app(), base_url=f"http://127.0.0.1:{ll.PORT}")
    cid = ll.find_course("COP4020")["id"]
    lid = _lec(ll, "Parsing day", cid, segs=_said([3, 3, 4]))
    url = f"/api/lectures/{lid}/slides"
    r = c.post(url, files={"file": ("notes.docx", b"x")})
    assert r.status_code == 400 and "PDF" in r.json()["detail"]
    r = c.post(url, files={"file": ("broken.pdf", b"this is not a pdf")})
    assert r.status_code == 422 and "attach" in r.json()["detail"].lower()
    good = _pdf(tmp_path / "Parsing.pdf", SLIDE_TEXT).read_bytes()
    monkeypatch.setattr(ll, "MAX_DECK_MB", 0)
    assert c.post(url, files={"file": ("big.pdf", good)}).status_code == 413
    monkeypatch.setattr(ll, "MAX_DECK_MB", 200)
    r = c.post(url, files={"file": ("..\\CON.pdf", good)})
    assert r.status_code == 200, r.text
    body = r.json()
    assert (body["deck"]["source"], body["deck"]["how"], body["deck"]["pages"]) == ("upload", "manual", 7)
    assert (body["state"], body["reason"]) == ("waiting", "reading the slides")  # the job reads it
    assert (ll.DATA / "Slides" / "COP4020" / "_CON.pdf").is_file()  # no device name, no path parts
    did = body["deck"]["id"]
    assert c.get(f"/media/slides/{did}/1-320.png").status_code == 404
    out = ll.job_slides({"id": 0, "payload": {"lecture_id": lid}})
    assert out["pages"] == 7 and c.get(url).json()["state"] == "aligned"
    img = c.get(f"/media/slides/{did}/1-320.png")
    assert img.status_code == 200 and img.headers["cache-control"] == "no-cache"
    assert c.get(f"/media/slides/{did}/..%5C..%5Clecturelab.db").status_code == 404
    _deck(ll, "Parsing day", cid, status="ready")  # a better-looking auto match appears
    ll.match_decks(cid)
    assert c.get(url).json()["deck"]["id"] == did  # your pick stays
    with ll.db() as con:
        jobs = con.execute("SELECT COUNT(*) FROM jobs WHERE kind='slides' AND lecture_id=? AND status='queued'",
                           (lid,)).fetchone()[0]
    assert jobs == 1  # queued once, not once per change
    r = c.put(url, json={"deck_id": None})
    assert r.json()["state"] == "none"
    assert c.put(url, json={"deck_id": 999}).status_code == 400


def test_slide_upload_turns_bad_files_away_at_once_and_reads_in_the_job(ll, monkeypatch, tmp_path):
    import io
    import zipfile
    import pymupdf as fitz
    from fastapi.testclient import TestClient
    c = TestClient(ll.create_app(), base_url=f"http://127.0.0.1:{ll.PORT}")
    cid = ll.find_course("COP4020")["id"]
    lid = _lec(ll, "Parsing day", cid, segs=_said([3, 3, 4]))
    url = f"/api/lectures/{lid}/slides"
    locked = tmp_path / "locked.pdf"
    doc = fitz.open(_pdf(tmp_path / "open.pdf", SLIDE_TEXT[:2]))
    doc.save(locked, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="u", owner_pw="o")
    doc.close()
    r = c.post(url, files={"file": ("locked.pdf", locked.read_bytes())})
    assert r.status_code == 422 and "password" in r.json()["detail"]
    for name, data in (("deck.pptx", b"PK not really a zip"), ("old.ppt", b"not an OLE file")):
        assert c.post(url, files={"file": (name, data)}).status_code == 422
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("ppt/presentation.xml", "<p/>")
    assert "no pages" in c.post(url, files={"file": ("empty.pptx", buf.getvalue())}).json()["detail"]
    monkeypatch.setattr(ll, "MAX_DECK_PAGES", 1)
    assert "textbook" in c.post(url, files={"file": ("book.pdf", locked.with_name("open.pdf").read_bytes())}).json()["detail"]
    monkeypatch.setattr(ll, "MAX_DECK_PAGES", 300)
    assert not list((ll.DATA / "Slides").rglob("*.*"))  # nothing turned away was kept

    # one file attached to two lectures: one deck, one stored copy
    good = _pdf(tmp_path / "Parsing.pdf", SLIDE_TEXT).read_bytes()
    other = _lec(ll, "Parsing day, section 2", cid, segs=_said([3, 3, 4]))
    real = ll.extract_deck

    def unreadable(did, say=None):
        raise ll.DeckError("This PDF couldn't be opened; export it again and attach that.")
    monkeypatch.setattr(ll, "extract_deck", unreadable)
    did = c.post(url, files={"file": ("Parsing.pdf", good)}).json()["deck"]["id"]
    assert c.post(f"/api/lectures/{other}/slides", files={"file": ("Parsing.pdf", good)}).json()["deck"]["id"] == did
    stored = lambda: list((ll.DATA / "Slides").rglob("*.pdf"))  # noqa: E731
    assert len(stored()) == 1

    # it can't be read: both lectures show why, and attaching it again reads it again for both
    with pytest.raises(ll.DeckError, match="export it again"):
        ll.job_slides({"id": 0, "payload": {"lecture_id": lid}})
    with pytest.raises(ll.DeckError):
        ll.job_slides({"id": 0, "payload": {"lecture_id": other}})
    assert [c.get(f"/api/lectures/{x}/slides").json()["state"] for x in (lid, other)] == ["error", "error"]
    assert ll.get_deck(did)["status"] == "error"
    monkeypatch.setattr(ll, "extract_deck", real)
    body = c.post(url, files={"file": ("Parsing.pdf", good)}).json()
    assert body["deck"]["id"] == did and (body["state"], body["reason"]) == ("waiting", "reading the slides")
    assert len(stored()) == 1 and c.get(f"/api/lectures/{other}/slides").json()["state"] == "waiting"
    ll.job_slides({"id": 0, "payload": {"lecture_id": lid}})
    ll.job_slides({"id": 0, "payload": {"lecture_id": other}})
    assert [c.get(f"/api/lectures/{x}/slides").json()["state"] for x in (lid, other)] == ["aligned", "aligned"]

    # a passing failure (OCR timing out) fails the job but not the deck: looking again reads it
    lid3 = _lec(ll, "Parsing day, again", cid, segs=_said([3, 3, 4]))
    good2 = _pdf(tmp_path / "Parsing2.pdf", SLIDE_TEXT[:5]).read_bytes()
    monkeypatch.setattr(ll, "extract_deck", lambda did, say=None: 1 / 0)
    did2 = c.post(f"/api/lectures/{lid3}/slides", files={"file": ("Parsing2.pdf", good2)}).json()["deck"]["id"]
    with pytest.raises(ZeroDivisionError):
        ll.job_slides({"id": 0, "payload": {"lecture_id": lid3}})
    assert c.get(f"/api/lectures/{lid3}/slides").json()["state"] == "error" and ll.get_deck(did2)["status"] == "new"
    monkeypatch.setattr(ll, "extract_deck", real)
    assert c.post(f"/api/lectures/{lid3}/slides/align").status_code == 200
    ll.job_slides({"id": 0, "payload": {"lecture_id": lid3}})
    assert c.get(f"/api/lectures/{lid3}/slides").json()["state"] == "aligned"


def test_test_servers_never_call_claude(ll, monkeypatch):
    monkeypatch.setenv("LECTURELAB_CLAUDE", "off")
    monkeypatch.setattr(ll, "claude_exe", lambda: pytest.fail("reached claude"))
    with pytest.raises(ll.ClaudeUnavailable, match="turned off"):
        ll.claude_call("p", system=ll.OUTLINE_SYSTEM, kind="outline")


def test_canvas_scan_lists_decks_separately_and_queues_downloads(ll, monkeypatch):
    folder = "https://drive.google.com/drive/folders/1FakeCourseFolder0000000000000000"
    files = {"video": [], "pdf": [{"id": 77, "display_name": "Lecture 7 Parsing.pdf", "content-type": "application/pdf",
                                   "size": 1 << 20, "created_at": "2026-09-01T12:00:00Z"},
                                  {"id": 78, "display_name": "Huge scan.pdf", "content-type": "application/pdf",
                                   "size": 900 << 20},
                                  {"id": 79, "display_name": "Syllabus Fall 2026.pdf", "content-type": "application/pdf",
                                   "size": 1 << 18},
                                  {"id": 80, "display_name": "Homework 3.pdf", "content-type": "application/pdf",
                                   "size": 1 << 18},
                                  {"id": 81, "display_name": "Type Checking.pdf", "content-type": "application/pdf",
                                   "size": 1 << 18}]}

    def fake_get(path, params=None, paginate=True):
        if path.endswith("/files"):
            return files["pdf"] if "application/pdf" in params["content_types[]"] else files["video"]
        if path.endswith("/modules"):
            return [{"name": "Resources", "items": [{"type": "ExternalUrl", "title": "Drive", "external_url": folder}]}]
        return [] if paginate else None

    monkeypatch.setattr(ll, "canvas_get", fake_get)
    monkeypatch.setattr(ll, "CanvasSession", lambda *a, **k: __import__("contextlib").nullcontext())
    monkeypatch.setattr(ll, "drive_list_folder", lambda fid: [
        {"id": "1FakeRecording00000000000000000-", "path": "M0L1 Course Overview.mkv"},
        {"id": "1FakeSlides000000000000000000000", "path": "M0L1 Course Overview.pdf"}])
    ll.set_canvas_state("connected", user="Kashi")
    c = ll.find_course("COP4020")
    c["canvas_id"] = 11
    _lec(ll, "Type Checking", c["id"])  # a recording with that title: its deck is fetched
    out = ll.canvas_scan_course(c)
    # the 900 MB scan is skipped, and so are the syllabus and the homework: not lecture slides
    assert (out["new"], out["decks"], out["decks_new"]) == (1, 3, 3)
    with ll.db() as con:
        decks = {r[0]: r[1:] for r in con.execute("SELECT key, source, title, status FROM decks")}
        items = [r[0] for r in con.execute("SELECT key FROM canvas_items")]
        pulls = con.execute("SELECT COUNT(*) FROM jobs WHERE kind='deck_pull'").fetchone()[0]
    assert decks == {"canvas-file:77": ("canvas", "Lecture 7 Parsing", "new"),
                     "canvas-file:81": ("canvas", "Type Checking", "new"),
                     "drive:1FakeSlides000000000000000000000": ("drive", "M0L1 Course Overview", "new")}
    assert items == ["drive:1FakeRecording00000000000000000-"]  # the Canvas tab only lists recordings
    assert pulls == 3
    assert ll.canvas_scan_course(c)["decks_new"] == 0  # a rescan doesn't queue them again


def test_deck_download_reads_it_and_queues_the_matching_lecture(ll, monkeypatch, tmp_path):
    import gdown
    cid = ll.find_course("COP4020")["id"]
    lid = _lec(ll, "M0L1 Course Overview", cid, segs=_said([0, 0, 1]))
    did = _deck(ll, "M0L1 Course Overview", cid)
    src = _pdf(tmp_path / "src.pdf", SLIDE_TEXT[:3])

    def fake_download(id, output, quiet):
        dest = __import__("pathlib").Path(output) / "M0L1 Course Overview.pdf"
        dest.write_bytes(src.read_bytes())
        return str(dest)

    monkeypatch.setattr(gdown, "download", fake_download)
    with ll.db() as con:
        con.execute("UPDATE decks SET key='drive:1FakeSlides000000000000000000000' WHERE id=?", (did,))
    out = ll.job_deck_pull({"id": 5, "payload": {"deck_id": did}})
    assert out == {"pages": 3, "lectures": [lid]}
    assert ll.get_deck(did)["path"] == "Slides/COP4020/M0L1 Course Overview.pdf"
    ll.job_slides({"id": 0, "payload": {"lecture_id": lid}})
    assert ll.lecture_slides(lid)["state"] == "aligned"


def test_slides_read_off_a_video_with_revisits(ll, tmp_path):
    import pymupdf as fitz
    pdf = fitz.open(_pdf(tmp_path / "s.pdf", SLIDE_TEXT[1:3]))
    imgs = []
    for i, page in enumerate(pdf):
        imgs.append(tmp_path / f"{i}.png")
        ll.render_fixed(page, 1280, 720).save(imgs[-1])
    video = ll.DATA / "media" / "v" / "original.mp4"
    video.parent.mkdir(parents=True)
    seq = [imgs[0], imgs[1], imgs[0]]  # slide A, slide B, back to A: 20 s each
    args = [ll.find_tool("ffmpeg"), "-v", "error"]
    for im in seq:
        args += ["-loop", "1", "-t", "20", "-framerate", "5", "-i", str(im)]
    args += ["-filter_complex", "".join(f"[{i}:v]" for i in range(len(seq))) + f"concat=n={len(seq)}:v=1[v]",
             "-map", "[v]", "-c:v", "libx264", "-g", "10", "-pix_fmt", "yuv420p", str(video)]
    subprocess.run(args, check=True)
    cid = ll.find_course("COP4020")["id"]
    lid = _lec(ll, "Board day", cid, video="media/v/original.mp4",
               segs=[(t, t + 10, "talking") for t in range(0, 60, 10)])
    out = ll.job_slides({"id": 0, "payload": {"lecture_id": lid}})
    s = ll.lecture_slides(lid)
    assert out["pages"] == 2 and s["deck"]["source"] == "ocr" and s["deck"]["thumb_h"] == 180
    assert "lexer" in s["slides"][0]["text"].lower() and "dfa" in s["slides"][1]["text"].lower()
    assert [len(x["ranges"]) for x in s["slides"]] == [2, 1]  # A came back
    with ll.db() as con:
        nos = [r[0] for r in con.execute("SELECT slide_no FROM segments WHERE lecture_id=? ORDER BY idx", (lid,))]
    assert nos == [1, 1, 2, 2, 1, 1]
    did = s["deck"]["id"]
    ll.delete_lecture(lid)
    assert ll.get_deck(did) is None and not (ll.SLIDE_IMAGES / str(did)).exists()


def _slide_video(ll, tmp_path, shots):
    """An h264 video of (png or None for a dark 'IDE' screen, seconds) shots, keyframe every 2 s."""
    args = [ll.find_tool("ffmpeg"), "-v", "error"]
    for img, sec in shots:
        args += (["-loop", "1", "-t", str(sec), "-framerate", "5", "-i", str(img)] if img else
                 ["-f", "lavfi", "-t", str(sec), "-i", "color=c=0x1e1f22:s=1280x720:r=5"])
    video = ll.DATA / "media" / "vv" / "original.mp4"
    video.parent.mkdir(parents=True, exist_ok=True)
    args += ["-filter_complex", "".join(f"[{i}:v]setsar=1[v{i}];" for i in range(len(shots))) +
             "".join(f"[v{i}]" for i in range(len(shots))) + f"concat=n={len(shots)}:v=1[v]",
             "-map", "[v]", "-c:v", "libx264", "-g", "10", "-pix_fmt", "yuv420p", str(video)]
    subprocess.run(args, check=True)
    return "media/vv/original.mp4"


def test_video_that_shows_the_slides_places_them_and_marks_demos(ll, tmp_path):
    import pymupdf as fitz
    cid = ll.find_course("COP4020")["id"]
    did = _ready_deck(ll, "M3L1 Parser Introduction", cid, pages=SLIDE_TEXT[1:4])
    pdf = fitz.open(ll.DATA / ll.get_deck(did)["path"])
    imgs = []
    for i, page in enumerate(pdf):
        imgs.append(tmp_path / f"p{i}.png")
        ll.render_fixed(page, 1280, 720).save(imgs[-1])
    video = _slide_video(ll, tmp_path, [(imgs[0], 20), (None, 40), (imgs[1], 20), (imgs[2], 20)])
    # nothing said matches any slide: only the video can place them
    lid = _lec(ll, "M3L1 Parser Introduction", cid, video=video,
               segs=[(t, t + 10, "okay so yeah") for t in range(0, 100, 10)])
    ll.job_slides({"id": 0, "payload": {"lecture_id": lid}})
    with ll.db() as con:
        nos = [r[0] for r in con.execute("SELECT slide_no FROM segments WHERE lecture_id=? ORDER BY idx", (lid,))]
    assert nos == [1, 1, 1, None, None, None, 2, 2, 2, 3]
    s = ll.lecture_slides(lid)
    assert s["state"] == "aligned" and [x["ranges"] for x in s["slides"]] == [[[0, 30]], [[60, 90]], [[90, 100]]]


def test_deck_download_that_fails_for_a_passing_reason_waits_and_retries(ll, monkeypatch):
    import gdown
    cid = ll.find_course("COP4020")["id"]
    did = _deck(ll, "M0L1 Course Overview", cid)
    with ll.db() as con:
        con.execute("UPDATE decks SET key='drive:1FakeSlides000000000000000000000' WHERE id=?", (did,))

    def boom(**kw):
        raise ConnectionError(rf"{ll.DATA}\downloads\job1 unreachable")

    monkeypatch.setattr(gdown, "download", boom)
    jid = ll.enqueue("deck_pull", {"deck_id": did, "title": "M0L1"})
    for n in range(ll.DECK_MAX_TRIES - 1):  # five tries in all
        with pytest.raises(ll.RetryLater):
            ll.job_deck_pull(ll.get_job(jid))
        j = ll.get_job(jid)
        assert (j["status"], j["payload"]["tries"], ll.get_deck(did)["status"]) == ("queued", n + 1, "new")
        assert j["retry_at"]
    with pytest.raises(ll.DeckError):
        ll.job_deck_pull(ll.get_job(jid))
    deck = ll.get_deck(did)
    assert deck["status"] == "error" and str(ll.DATA) not in deck["error"] and "attach it" in deck["error"]
    assert ll.retry_deck(did) and ll.get_deck(did)["status"] == "new"  # "look for slides" asks again
    # signed out of Canvas: it waits for you, however long, without using up its tries
    cdid = _deck(ll, "M2L1 Lexer Introduction", cid, source="canvas")
    with ll.db() as con:
        con.execute("UPDATE decks SET url='canvas_file:7' WHERE id=?", (cdid,))
    cj = ll.enqueue("deck_pull", {"deck_id": cdid, "title": "M2L1"})
    for _ in range(ll.DECK_MAX_TRIES + 2):
        with pytest.raises(ll.RetryLater):
            ll.job_deck_pull(ll.get_job(cj))
    assert ll.get_job(cj)["payload"]["tries"] == 0 and ll.get_deck(cdid)["status"] == "new"
    assert ll.get_job(cj)["progress"] == "waiting for Canvas sign-in" and ll.get_job(cj)["retry_at"]
    ll.set_canvas_state("connected", user="test")  # you sign in: it goes now, not in half an hour
    assert ll.get_job(cj)["retry_at"] is None


def test_cli_slides_does_not_hang_on_a_download_that_waits(ll, monkeypatch, capsys):
    import gdown
    sys.path.insert(0, str(ll.ROOT))
    import cli
    monkeypatch.setattr(cli, "ll", ll)
    monkeypatch.setattr(gdown, "download", lambda **kw: None)  # Drive returns nothing: a passing failure
    cid = ll.find_course("COP4020")["id"]
    did = _deck(ll, "M0L1 Course Overview", cid)
    jid = ll.enqueue("deck_pull", {"deck_id": did, "title": "M0L1"})
    assert cli.cmd_slides(type("A", (), {"ids": []})()) == 0
    assert "next try after" in capsys.readouterr().out
    assert ll.get_job(jid)["status"] == "queued"


def test_a_second_copy_of_a_deck_keeps_the_match_and_its_alignment(ll):
    cid = ll.find_course("COP4020")["id"]
    lid = _lec(ll, "M2L1 Lexer Introduction", cid, segs=_said([1, 1, 2]))
    pdf = _ready_deck(ll, "M2L1 Lexer Introduction", cid)
    ll.job_slides({"id": 0, "payload": {"lecture_id": lid}})
    assert ll.lecture_slides(lid)["state"] == "aligned"
    _deck(ll, "M2L1 Lexer Introduction", cid, source="canvas")  # the PPTX of it, posted on Canvas too
    assert ll.match_decks(cid) == []
    assert ll.get_lecture(lid)["deck_id"] == pdf and ll.lecture_slides(lid)["state"] == "aligned"
    # a new lecture whose deck was posted twice still gets one: the PDF
    l2 = _lec(ll, "M2L2 Regular Languages", cid)
    _deck(ll, "M2L2 Regular Languages", cid, source="canvas", path="Slides/x/M2L2 Regular Languages.pptx")
    b = _deck(ll, "M2L2 Regular Languages", cid, path="Slides/x/M2L2 Regular Languages.pdf")
    assert ll.match_decks(cid) == [l2] and ll.get_lecture(l2)["deck_id"] == b
    # a range deck you picked for one lecture still serves the other one in its range
    e, f = _lec(ll, "M3TL2 Syntax Review", cid), _lec(ll, "M3TL3 Syntax Review (cont)", cid)
    both = _deck(ll, "M3TL2-3 Syntax Review", cid)
    ll.set_lecture_deck(e, both)
    assert ll.get_lecture(f)["deck_id"] == both


def test_a_pick_made_while_the_video_is_read_wins(ll, monkeypatch):
    cid = ll.find_course("COP4020")["id"]
    lid = _lec(ll, "Guest talk", cid, segs=_said([1, 2]), video="media/x/original.mp4")
    did = _ready_deck(ll, "Some deck", cid)
    other = _deck(ll, "Video frames", cid, source="ocr", status="ready")

    def slow_ocr(lec, progress=None):
        ll.set_lecture_deck(lid, did)  # you attach a deck while the video is being read
        return other

    monkeypatch.setattr(ll, "ocr_video_deck", slow_ocr)
    ll.job_slides({"id": 0, "payload": {"lecture_id": lid}})
    lec = ll.get_lecture(lid)
    assert (lec["deck_id"], lec["deck_how"]) == (did, "manual")


def test_a_new_transcript_is_not_shown_as_aligned(ll):
    cid = ll.find_course("COP4020")["id"]
    lid = _lec(ll, "M2L1 Lexer Introduction", cid, segs=_said([1, 2]))
    _ready_deck(ll, "M2L1 Lexer Introduction", cid)
    ll.job_slides({"id": 0, "payload": {"lecture_id": lid}})
    assert ll.lecture_slides(lid)["state"] == "aligned"
    ll.store_transcript(lid, {"segments": [{"start": 0, "end": 10, "text": "a lexer makes tokens"}]}, "local")
    s = ll.lecture_slides(lid)
    assert s["state"] == "waiting" and all(x["first"] is None for x in s["slides"])


def test_big_and_scanned_pdfs_are_bounded(ll, monkeypatch, tmp_path):
    import pymupdf as fitz
    cid = ll.find_course("COP4020")["id"]
    rel = "Slides/test/book.pdf"
    _pdf(ll.DATA / rel, SLIDE_TEXT[:4])
    monkeypatch.setattr(ll, "MAX_DECK_PAGES", 3)
    with pytest.raises(ll.DeckError, match="4 pages"):
        ll.extract_deck(_deck(ll, "Textbook", cid, path=rel))
    # scans: two near-empty pages, and a page-sized picture with a footer; OCR stops at the cap
    pic = tmp_path / "pic.png"
    ll.render_fixed(fitz.open(_pdf(tmp_path / "p.pdf", SLIDE_TEXT[1:2]))[0], 960, 540).save(pic)
    doc = fitz.open()
    for _ in range(2):
        doc.new_page(width=960, height=540).insert_text((40, 80), "x")
    p = doc.new_page(width=960, height=540)
    p.insert_image(p.rect, filename=str(pic))
    p.insert_text((40, 530), "COP4020 Programming Language Concepts", fontsize=10)
    rel = "Slides/test/scan.pdf"
    doc.save(ll.DATA / rel)
    seen = []
    monkeypatch.setattr(ll, "ocr_image", lambda path: seen.append(path.name) or "read off the picture, lexer tokens")
    monkeypatch.setattr(ll, "OCR_MAX_PAGES", 2)
    did = _deck(ll, "Scanned", cid, path=rel)
    ll.extract_deck(did)
    assert seen == ["1-1280.png", "2-1280.png"]
    monkeypatch.setattr(ll, "OCR_MAX_PAGES", 5)
    seen.clear()
    ll.extract_deck(did)
    assert seen == ["1-1280.png", "2-1280.png", "3-1280.png"]  # the footer alone doesn't count as text
    with ll.db() as con:
        assert con.execute("SELECT text FROM slides WHERE deck_id=? AND no=3", (did,)).fetchone()[0].startswith("read")


def test_slide_errors_shown_to_you_carry_no_local_paths(ll, monkeypatch):
    cid = ll.find_course("COP4020")["id"]
    lid = _lec(ll, "M2L1 Lexer Introduction", cid, segs=_said([1, 2]))
    _ready_deck(ll, "M2L1 Lexer Introduction", cid)

    def broken(lecture_id):
        raise RuntimeError(rf"ffmpeg.exe failed (1): {ll.DATA}\media\5\original.mp4: Invalid data")

    monkeypatch.setattr(ll, "align_lecture", broken)
    with pytest.raises(RuntimeError):
        ll.job_slides({"id": 0, "payload": {"lecture_id": lid}})
    s = ll.lecture_slides(lid)
    assert s["state"] == "error" and str(ll.DATA) not in s["error"] and "server.log" in s["error"]


def test_video_slides_wait_for_the_canvas_scan_unless_you_ask(ll, monkeypatch):
    cid = ll.find_course("COP4020")["id"]
    with ll.db() as con:
        con.execute("UPDATE courses SET canvas_id=11 WHERE id=?", (cid,))
    lid = _lec(ll, "Guest talk", cid, segs=_said([1, 2]), video="media/x/original.mp4")
    read = []
    monkeypatch.setattr(ll, "ocr_video_deck", lambda lec, progress=None: read.append(lec["id"]))
    ll.job_slides({"id": 0, "payload": {"lecture_id": lid}})
    assert read == [] and ll.lecture_slides(lid)["reason"] == "not checked yet"  # its deck may turn up
    ll.slides_scanned(cid)  # the scan found nothing for it: now the video is read
    with ll.db() as con:
        jobs = [ll.get_job(r[0]) for r in con.execute("SELECT id FROM jobs WHERE kind='slides' AND lecture_id=?",
                                                       (lid,))]
    assert len(jobs) == 1
    ll.job_slides(jobs[0])
    assert read == [lid]
    # before any scan, "Look for slides" reads it right away
    other = _lec(ll, "Another talk", cid, segs=_said([1, 2]), video="media/y/original.mp4")
    with ll.db() as con:
        con.execute("DELETE FROM meta WHERE key=?", (f"slides_scanned:{cid}",))
    ll.job_slides(ll.get_job(ll.enqueue_slides(other, ocr=True)))
    assert read == [lid, other]


def test_screens_that_are_not_slides_are_left_out(ll, tmp_path):
    import pymupdf as fitz
    slide = tmp_path / "slide.png"
    ll.render_fixed(fitz.open(_pdf(tmp_path / "s.pdf", SLIDE_TEXT[1:2]))[0], 1280, 720).save(slide)
    doc = fitz.open()  # a desktop: a sidebar, menus and code, lots of small lines
    p = doc.new_page(width=960, height=540)
    for i in range(48):
        p.insert_text((20 + 480 * (i // 24), 20 + 21 * (i % 24)), f"def visit_{i}(self, node): return node.child_{i}",
                      fontsize=9)
    desk = tmp_path / "desk.png"
    ll.render_fixed(doc[0], 1280, 720).save(desk)
    lines = ll.text_lines(slide, 720)
    assert ll.looks_like_slide(ll.text_lines(slide, 720, read=False)) and lines
    assert max(lines, key=lambda ln: ln[1])[0].startswith("Lexical analysis")  # the title is the biggest line
    assert not ll.looks_like_slide(ll.text_lines(desk, 720, read=False))


def test_a_run_that_died_is_cleaned_up_while_the_server_runs(ll, monkeypatch):
    dead = subprocess.Popen([sys.executable, "-c", "pass"])
    dead.wait()  # a pid that existed and is gone, like a run killed for memory
    lid = _lecture(ll, "transcribing")
    with ll.db() as con:
        jid = con.execute("INSERT INTO jobs(kind, payload, status, owner_pid, attempts) VALUES "
                          "('transcribe', '{}', 'running', ?, 1)", (dead.pid,)).lastrowid
        bid = con.execute("INSERT INTO transcribe_batches(mode, lecture_ids) VALUES ('colab', ?)",
                          (f"[{lid}]",)).lastrowid
    swept = []
    monkeypatch.setattr(ll, "COLAB_ENABLED", True)
    monkeypatch.setattr(ll, "sweep_orphans_if_idle", lambda: swept.append(1) or [])
    assert ll.recover_orphans() == 1
    assert ll.get_job(jid)["status"] == "queued" and ll.get_lecture(lid)["status"] == "queued_transcribe"
    with ll.db() as con:
        assert tuple(con.execute("SELECT ended_at IS NOT NULL, ok FROM transcribe_batches WHERE id=?",
                                 (bid,)).fetchone()) == (1, 0)
    assert swept == [1]  # its Colab session is stopped too
    assert ll.recover_orphans() == 0 and swept == [1]  # nothing left: no WSL call every minute


def test_lines_whisper_left_as_just_punctuation_are_dropped(ll):
    lid = _lecture(ll, "transcribing")
    ll.store_transcript(lid, {"segments": [{"start": 0, "end": 2, "text": "Next year or ."},
                                           {"start": 2, "end": 3, "text": " . "},
                                           {"start": 3, "end": 4, "text": "?"},
                                           {"start": 4, "end": 6, "text": "Clase de hoy"}]}, "colab")
    with ll.db() as con:
        assert [r[0] for r in con.execute("SELECT text FROM segments WHERE lecture_id=? ORDER BY idx", (lid,))]             == ["Next year or .", "Clase de hoy"]


def test_a_transcribe_batch_keeps_the_laptop_awake_until_it_ends(ll, monkeypatch):
    calls = []
    monkeypatch.setattr(ll, "stay_awake", lambda on: calls.append(on))
    lid = _lecture(ll, "queued_transcribe")
    monkeypatch.setattr(ll, "transcribe_batch_colab", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("VM gone")))
    with pytest.raises(RuntimeError):
        ll.job_transcribe({"id": 0, "payload": {"lecture_ids": [lid], "mode": "colab"}})
    assert calls == [True, False]  # released even when the batch fails
    lock = ll.FileLock(ll.DATA / "transcribe.lock")
    assert lock.acquire()  # and so is the transcription lock
    lock.release()


@pytest.mark.skipif(os.name != "nt", reason="SetThreadExecutionState is Windows only")
def test_stay_awake_asks_windows_for_system_and_screen_then_lets_go(ll, monkeypatch):
    import ctypes
    sent = []
    monkeypatch.setattr(ctypes.windll.kernel32, "SetThreadExecutionState", lambda f: sent.append(f.value) or 1)
    assert ll.stay_awake(True) and ll.stay_awake(False)
    assert sent == [0x80000003, 0x80000000]
    monkeypatch.setattr(ctypes.windll.kernel32, "SetThreadExecutionState", lambda f: 1 / 0)
    assert ll.stay_awake(True) is False  # a failure is logged, never raised


# ---------------------------------------------------------------- Stage 4: notes

def _fake_claude(ll, monkeypatch, answers, calls=None, env_seen=None):
    """Stand in for `claude -p`: answers[kind] is what each pass returns (a dict for JSON passes,
    a string or a function of the prompt for sections)."""
    kinds = {ll.OUTLINE_SYSTEM: "outline", ll.SECTION_SYSTEM: "section", ll.LOGISTICS_SYSTEM: "logistics",
             ll.STUDY_SYSTEM: "study", ll.CHECK_SYSTEM: "check"}

    def run(cmd, prompt, timeout):
        kind = kinds[cmd[cmd.index("--system-prompt") + 1]]
        if calls is not None:
            calls.append((kind, cmd, prompt))
        if env_seen is not None:
            env_seen.append(ll.claude_env())
        a = answers[kind]
        a = a(prompt) if callable(a) else a
        if isinstance(a, Exception):
            raise a
        if isinstance(a, tuple):  # (is_error, message)
            out = {"type": "result", "is_error": True, "result": a[1]}
        elif isinstance(a, str):
            out = {"type": "result", "is_error": False, "result": a, "usage": {"output_tokens": 5}}
        else:
            out = {"type": "result", "is_error": False, "result": json.dumps(a), "structured_output": a}
        return subprocess.CompletedProcess(cmd, 0, json.dumps(out), "")
    monkeypatch.setattr(ll, "_claude_run", run)
    monkeypatch.setattr(ll, "claude_exe", lambda: "claude")


def _notes_lecture(ll):
    cid = ll.find_course("COP4020")["id"]
    segs = [(0, 20, "okay can everyone hear me in the back"), (20, 40, "so the midterm is next thursday"),
            (40, 70, "today back us now are form which is how we write grammars"),
            (70, 100, "a nonterminal goes on the left of the arrow"),
            (100, 130, "left recursion means the nonterminal calls itself first"),
            (130, 160, "and an LL parser loops forever on that")]
    return _lec(ll, "Grammars", cid, segs=segs)


OUTLINE = {"sections": [
    {"title": "Before class", "start": 0, "end": 20, "kind": "off-topic", "slides": [], "gist": "mic check"},
    {"title": "Midterm", "start": 20, "end": 40, "kind": "logistics", "slides": [], "gist": "exam date"},
    {"title": "BNF", "start": 40, "end": 100, "kind": "content", "slides": [], "gist": "grammar notation"},
    {"title": "Left recursion", "start": 100, "end": 160, "kind": "content", "slides": [], "gist": "LL loops"}]}
STUDY = {"summary": "Grammars and left recursion.\n\n- BNF writes grammars",
         "open_questions": ["What fixes left recursion?", " "],
         "cards": [{"kind": "cloze", "front": "{{c1::BNF}} is how we write grammars", "back": "notation", "t": 45},
                   {"kind": "qa", "front": "Why do LL parsers loop on left recursion?", "back": "they call it first",
                    "t": 120},
                   {"kind": "qa", "front": "Made-up card", "back": "nope", "t": 50}],
         "quiz": [{"kind": "mcq", "question": "Which side of the arrow?", "choices": ["left", "right"],
                   "answer": "left", "explanation": "the lecture said so", "t": 75},
                  {"kind": "mcq", "question": "Broken", "choices": ["a", "b"], "answer": "c", "explanation": "",
                   "t": 75},
                  {"kind": "short", "question": "What is left recursion?", "choices": [],
                   "answer": "A nonterminal deriving itself first.", "explanation": "", "t": 100}]}
CHECK = {"edits": [{"number": 3, "find": "**BNF**: how we write grammars",
                    "replace": "Left recursion: the nonterminal calls itself first.\n\n> [!check] the loop claim"},
                   {"number": 2, "find": "text that isn't in the notes", "replace": ""}],
         "summary": None, "drop_cards": [3], "drop_quiz": [], "issues": [
             {"where": "card 3", "verdict": "removed", "claim": "Made-up card", "reason": "not in the lecture"},
             {"where": "section 3", "verdict": "flagged", "claim": "loops forever", "reason": "audio unclear"}]}


def _section(prompt):
    if "► 2. Midterm" in prompt or "Midterm (" in prompt.split("►")[1][:40]:
        return "- [▶ 0:20](#t=20) Midterm next Thursday.\n\n> [!exam] midterm"
    # a heading it shouldn't add, a good link, and one far outside the section
    return "## BNF\n\n[▶ 0:40](#t=40) **BNF**: how we write grammars (see [▶ 9:59](#t=599))."


def test_claude_is_called_on_the_plan_without_an_api_key_or_your_setup(ll, monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-test-not-real")
    calls, envs = [], []
    _fake_claude(ll, monkeypatch, {"section": "hi"}, calls, envs)
    assert ll.claude_call("long prompt " * 5000, system=ll.SECTION_SYSTEM, lecture_id=1, kind="section") == "hi"
    kind, cmd, prompt = calls[0]
    assert "ANTHROPIC_API_KEY" not in envs[0] and "ANTHROPIC_AUTH_TOKEN" not in envs[0]
    assert "--bare" not in cmd  # it forces API-key auth
    assert {"-p", "--setting-sources=", "--strict-mcp-config", "--tools=", "--no-session-persistence"} <= set(cmd)
    assert cmd[cmd.index("--model") + 1] == "sonnet"
    assert prompt.startswith("long prompt") and "long prompt" not in " ".join(cmd)  # on stdin, not the command line


def test_the_same_claude_call_twice_is_answered_from_the_cache(ll, monkeypatch):
    calls = []
    _fake_claude(ll, monkeypatch, {"outline": OUTLINE}, calls)
    a = ll.claude_call("p", system=ll.OUTLINE_SYSTEM, schema=ll.OUTLINE_SCHEMA, kind="outline")
    b = ll.claude_call("p", system=ll.OUTLINE_SYSTEM, schema=ll.OUTLINE_SCHEMA, kind="outline")
    assert a == b == OUTLINE and len(calls) == 1
    ll.claude_call("other", system=ll.OUTLINE_SYSTEM, schema=ll.OUTLINE_SCHEMA, kind="outline")
    assert len(calls) == 2
    with ll.db() as con:
        assert [r[0] for r in con.execute("SELECT cached FROM claude_calls ORDER BY id")] == [0, 1, 0]


def test_the_usage_limit_parks_claude_calls_until_it_resets(ll, monkeypatch):
    calls = []
    reset = int(time.time()) + 7200
    _fake_claude(ll, monkeypatch, {"section": (True, f"Claude AI usage limit reached|{reset}")}, calls)
    with pytest.raises(ll.QuotaExceeded) as e:
        ll.claude_call("p", system=ll.SECTION_SYSTEM)
    assert e.value.reset_at == reset
    with pytest.raises(ll.QuotaExceeded):
        ll.claude_call("q", system=ll.SECTION_SYSTEM)
    assert len(calls) == 1  # parked: no second call while the limit lasts
    now = datetime(2026, 10, 28, 13, 0).timestamp()
    assert datetime.fromtimestamp(ll.quota_reset_at("5-hour limit reached · resets 3pm", now)).hour == 15
    assert ll.quota_reset_at("You've hit your limit · resets 9am", now) > now + 19 * 3600  # tomorrow
    assert ll.quota_reset_at("limit reached", now) == now + 3600


def test_a_claude_failure_says_what_happened(ll, monkeypatch):
    _fake_claude(ll, monkeypatch, {"section": (True, "Invalid model"),
                                   "study": subprocess.TimeoutExpired("claude", 5)})
    with pytest.raises(ll.ClaudeError, match="Invalid model"):
        ll.claude_call("p", system=ll.SECTION_SYSTEM)
    with pytest.raises(ll.ClaudeError, match="didn't answer"):
        ll.claude_call("p", system=ll.STUDY_SYSTEM, schema=ll.STUDY_SCHEMA)


def test_the_outline_is_cleaned_into_contiguous_sections():
    import lecturelab as ll
    out = ll.clean_outline([{"title": "B", "start": 300, "end": 1, "kind": "content", "slides": [3, 3, 0]},
                            {"title": "A", "start": 12, "end": 9, "kind": "content", "slides": [1]},
                            {"title": "sliver", "start": 305, "end": 306, "kind": "content", "slides": [4]},
                            {"title": "tail", "start": 595, "end": 600, "kind": "content", "slides": []}], 600)
    assert [(s["title"], s["start"], s["end"], s["slides"]) for s in out] == \
        [("A", 0.0, 300.0, [1]), ("B", 300.0, 600, [3, 4])]


def test_notes_are_written_checked_and_stored_with_cards_and_quiz(ll, monkeypatch):
    lid = _notes_lecture(ll)
    calls = []
    _fake_claude(ll, monkeypatch, {"outline": OUTLINE, "section": _section, "logistics": _section,
                                   "study": STUDY, "check": CHECK}, calls)
    res = ll.job_notes({"id": 0, "payload": {"lecture_id": lid}})
    assert [c[0] for c in calls] == ["outline", "logistics", "section", "section", "study", "check"]
    assert calls[-1][1][calls[-1][1].index("--model") + 1] == "opus"  # only the check pass
    assert res == {"sections": 3, "cards": 2, "quiz": 2, "removed": 1, "flagged": 2, "mode": "claude"}
    n = ll.lecture_notes(lid)
    assert n["state"] == "done" and not n["stale"] and n["open_questions"] == ["What fixes left recursion?"]
    assert [s["title"] for s in n["sections"]] == ["Midterm", "BNF", "Left recursion"]  # chatter left out
    bnf = n["sections"][1]["markdown"]
    assert not bnf.startswith("#") and "(#t=40)" in bnf and "(#t=599)" not in bnf  # a link outside it dropped
    assert "[!check]" in n["sections"][2]["markdown"]  # the check pass edited it
    assert "(#t=40)" in n["sections"][2]["markdown"]  # the rest of the section untouched
    assert any("wasn't applied" in i["reason"] for i in n["issues"])  # an edit that matched nothing
    with ll.db() as con:
        assert con.execute("SELECT status FROM lectures WHERE id=?", (lid,)).fetchone()[0] == "exported"
        assert [r[0] for r in con.execute("SELECT front FROM cards WHERE lecture_id=? ORDER BY id", (lid,))] == \
            ["{{c1::BNF}} is how we write grammars", "Why do LL parsers loop on left recursion?"]
        assert [r[0] for r in con.execute("SELECT question FROM quiz WHERE lecture_id=? ORDER BY idx", (lid,))] == \
            ["Which side of the arrow?", "What is left recursion?"]  # the mcq with no right choice is gone
        con.execute("UPDATE cards SET box=3, due='2026-11-01', reviews=4 WHERE kind='cloze'")
    # Regenerating changes nothing it doesn't have to: no new claude calls, your card progress kept
    calls.clear()
    ll.job_notes({"id": 0, "payload": {"lecture_id": lid}})
    assert calls == []
    with ll.db() as con:
        assert tuple(con.execute("SELECT box, due, reviews FROM cards WHERE kind='cloze'").fetchone()) == \
            (3, "2026-11-01", 4)
        assert con.execute("SELECT COUNT(*) FROM claude_calls WHERE cached=0").fetchone()[0] == 6
    with ll.db() as con:  # a new transcript makes the notes stale
        con.execute("UPDATE lectures SET transcript_source='colab-large-v3 (T4)' WHERE id=?", (lid,))
    assert ll.lecture_notes(lid)["stale"]


def test_notes_wait_out_the_usage_limit_and_resume_from_the_cache(ll, monkeypatch):
    lid = _notes_lecture(ll)
    calls = []
    reset = int(time.time()) + 3600
    answers = {"outline": OUTLINE, "section": _section, "logistics": _section,
               "study": (True, f"Claude AI usage limit reached|{reset}"), "check": CHECK}
    _fake_claude(ll, monkeypatch, answers, calls)
    jid = ll.enqueue_notes(lid)
    ll.run_job(ll.claim_job(["notes"]))
    job = ll.get_job(jid)
    assert job["status"] == "queued" and job["retry_at"] and "usage limit" in job["progress"]
    assert ll.lecture_notes(lid)["state"] == "waiting_quota"
    # the limit resets: only what didn't run before is asked again
    ll.meta_set("claude_quota_until", 0)
    answers["study"] = STUDY
    calls.clear()
    with ll.db() as con:
        con.execute("UPDATE jobs SET retry_at=NULL WHERE id=?", (jid,))
    ll.run_job(ll.claim_job(["notes"]))
    assert [c[0] for c in calls] == ["study", "check"]
    assert ll.get_job(jid)["status"] == "done" and ll.lecture_notes(lid)["state"] == "done"


def test_a_failed_notes_run_says_why_and_keeps_the_transcript(ll, monkeypatch):
    lid = _notes_lecture(ll)
    _fake_claude(ll, monkeypatch, {"outline": (True, "Invalid model name")})
    with pytest.raises(ll.ClaudeError):
        ll.job_notes({"id": 0, "payload": {"lecture_id": lid}})
    n = ll.lecture_notes(lid)
    assert n["state"] == "error" and "Invalid model name" in n["error"]
    assert ll.get_lecture(lid)["status"] == "transcribed"


def test_notes_follow_the_slides_once_they_are_settled(ll, monkeypatch):
    queued = []
    monkeypatch.setattr(ll, "enqueue_notes", lambda lid, mode="claude": queued.append(lid))
    cid = ll.find_course("COP4020")["id"]
    lid = _lec(ll, "Audio only", cid, segs=_said([1, 2]))  # no video, no deck, course not on Canvas
    ll.job_slides_then_notes({"id": 0, "payload": {"lecture_id": lid}})
    assert queued == [lid]
    with ll.db() as con:  # on Canvas and not scanned yet: its deck may still turn up
        con.execute("UPDATE courses SET canvas_id=11 WHERE id=?", (cid,))
    other = _lec(ll, "Audio only 2", cid, segs=_said([1, 2]))
    ll.job_slides_then_notes({"id": 0, "payload": {"lecture_id": other}})
    assert queued == [lid]
    ll.slides_scanned(cid)
    ll.job_slides_then_notes({"id": 0, "payload": {"lecture_id": other}})
    assert queued == [lid, other]


def test_notes_api_queues_and_reports(ll, monkeypatch):
    from fastapi.testclient import TestClient
    monkeypatch.setattr(ll, "wake_workers", lambda: None)
    lid = _notes_lecture(ll)
    empty = _lecture(ll, "queued_transcribe")
    c = TestClient(ll.create_app(), base_url=f"http://127.0.0.1:{ll.PORT}")  # no lifespan: no workers
    assert c.get(f"/api/lectures/{lid}/notes").json()["state"] == "none"
    assert c.post(f"/api/lectures/{lid}/notes", json={}).json()["state"] == "queued"
    assert c.post(f"/api/lectures/{empty}/notes", json={}).status_code == 409
    assert c.post(f"/api/lectures/{lid}/notes", json={"mode": "gpt"}).status_code == 422
    assert c.get("/api/lectures/99999/notes").status_code == 404
    with ll.db() as con:
        assert con.execute("SELECT COUNT(*) FROM jobs WHERE kind='notes'").fetchone()[0] == 1


# ---------------------------------------------------------------- Stage 4: study

def _notes_done(ll, monkeypatch):
    lid = _notes_lecture(ll)
    _fake_claude(ll, monkeypatch, {"outline": OUTLINE, "section": _section, "logistics": _section,
                                   "study": STUDY, "check": CHECK})
    ll.job_notes({"id": 0, "payload": {"lecture_id": lid}})
    return lid


def test_leitner_moves_cards_up_on_right_and_back_to_box_1_on_wrong(ll, monkeypatch):
    lid = _notes_done(ll, monkeypatch)
    due = ll.due_cards(lecture_id=lid)
    assert (due["total"], due["new"], due["due"], due["boxes"]) == (2, 2, 0, [0, 0, 0, 0, 0])  # new: no box yet
    a, b = (c["id"] for c in due["cards"])
    r = ll.review_card(a, True)
    assert (r["card"]["box"], r["card"]["due"]) == (2, ll._today(2)) and r["before"]["box"] == 1
    assert ll.due_cards(lecture_id=lid)["boxes"] == [0, 1, 0, 0, 0]
    r = ll.review_card(a, True)
    assert (r["card"]["box"], r["card"]["due"]) == (3, ll._today(4))
    for _ in range(5):
        r = ll.review_card(a, True)
    assert (r["card"]["box"], r["card"]["due"], r["card"]["reviews"]) == (5, ll._today(16), 7)
    r = ll.review_card(a, False)
    assert (r["card"]["box"], r["card"]["due"], r["card"]["lapses"]) == (1, ll._today(1), 1)
    back = ll.undo_review(a, r["before"])  # the UI's undo
    assert (back["box"], back["lapses"], back["reviews"]) == (5, 0, 7)
    assert [c["id"] for c in ll.due_cards(lecture_id=lid)["cards"]] == [b]  # a isn't due for 16 days


def test_quiz_marks_multiple_choice_and_takes_your_word_on_short_answers(ll, monkeypatch):
    lid = _notes_done(ll, monkeypatch)
    qs = ll.lecture_quiz(lid)["questions"]
    assert [q["kind"] for q in qs] == ["mcq", "short"] and "answer" not in qs[0]
    out = ll.grade_quiz(lid, [{"id": qs[0]["id"], "answer": "  LEFT "}, {"id": qs[1]["id"], "answer": "x",
                                                                         "correct": True}])
    assert (out["score"], out["total"]) == (2, 2) and out["results"][0]["explanation"]
    out = ll.grade_quiz(lid, [{"id": qs[0]["id"], "answer": "right"}])  # the short one unanswered
    assert (out["score"], out["total"]) == (0, 2)
    assert ll.lecture_quiz(lid)["last_attempt"]["score"] == 0


def test_quiz_check_shows_answers_without_saving_then_one_attempt_is_saved(ll, monkeypatch):
    from fastapi.testclient import TestClient
    lid = _notes_done(ll, monkeypatch)
    c = TestClient(ll.create_app(), base_url=f"http://127.0.0.1:{ll.PORT}")
    url = f"/api/lectures/{lid}/quiz"
    mcq, short = c.get(url).json()["questions"]
    out = c.post(url, json={"check": True, "answers": [{"id": mcq["id"], "answer": "left"},
                                                       {"id": short["id"], "answer": "my words"}]}).json()
    assert [r["correct"] for r in out["results"]] == [True, None] and out["results"][1]["answer"]
    assert (out["score"], out["unmarked"], out["saved"]) == (1, 1, False)
    assert c.get(url).json()["last_attempt"] is None  # checking saved nothing
    out = c.post(url, json={"answers": [{"id": mcq["id"], "answer": "left"},
                                        {"id": short["id"], "answer": "my words", "correct": True}]}).json()
    assert (out["score"], out["total"], out["saved"]) == (2, 2, True)
    last = c.get(url).json()["last_attempt"]
    assert (last["score"], last["total"]) == (2, 2) and last["at"]
    with ll.db() as con:
        assert con.execute("SELECT COUNT(*) FROM quiz_attempts WHERE lecture_id=?", (lid,)).fetchone()[0] == 1


def test_a_cloze_answer_with_double_colons_becomes_a_question(ll):
    c = ll.fix_card({"kind": "cloze", "front": "A growable array in C++ is {{c1::std::vector}}.",
                     "back": "From <vector>.", "t": 5})
    assert c["kind"] == "qa" and c["front"] == "Fill in the blank: A growable array in C++ is ____."
    assert c["back"] == "`std::vector`\n\nFrom <vector>." and c["t"] == 5
    plain = {"kind": "cloze", "front": "A {{c1::grammar}} has rules; {{c2::BNF}} writes them.", "back": "", "t": 0}
    assert ll.fix_card(plain) is plain
    multi = ll.fix_card({"kind": "cloze", "front": "Sort with {{c1::sort}} from {{c2::std::\nranges}}.", "back": "", "t": 0})
    assert multi["front"] == "Fill in the blank: Sort with ____ from ____." and multi["back"].startswith("sort; `std::")


def test_anki_export_is_a_clean_import_with_stable_ids_and_links_back(ll, monkeypatch):
    lid = _notes_done(ll, monkeypatch)
    with ll.db() as con:
        con.execute("UPDATE lectures SET date='2026-09-01' WHERE id=?", (lid,))
        con.execute("UPDATE cards SET back='a <b>\tc\nd' WHERE kind='qa'")
    tsv = ll.anki_tsv(lecture_id=lid)
    head, lines = tsv.splitlines()[:6], tsv.splitlines()[6:]
    assert head == ["#separator:tab", "#html:true", "#guid column:1", "#notetype column:2", "#deck column:3",
                    "#tags column:6"]
    assert len(lines) == 2 and all(len(x.split("\t")) == 6 for x in lines)
    cloze, qa = (x.split("\t") for x in lines)
    assert cloze[1:4] == ["Cloze", "LectureLab::COP4020::2026-09-01 Grammars", "{{c1::BNF}} is how we write grammars"]
    assert qa[1] == "Basic" and qa[4].startswith("a &lt;b&gt;    c<br>d<br><a href=\"http://127.0.0.1:")
    assert f'?t=120#/lecture/{lid}">▶ 2:00</a>' in qa[4]
    assert ll.anki_tsv(lecture_id=lid).splitlines()[6].split("\t")[0] == cloze[0]  # same GUID next time
    assert ll.anki_tsv(lecture_id=lid + 99).count("\n") == 6  # headers only


def test_search_finds_transcript_notes_and_slides_with_safe_snippets(ll, monkeypatch):
    lid = _notes_done(ll, monkeypatch)
    hits = ll.search("recursion")
    assert {h["kind"] for h in hits} >= {"transcript", "notes"} and all(h["lecture_id"] == lid for h in hits)
    tr = next(h for h in hits if h["kind"] == "transcript")
    assert tr["t"] == 100 and "<mark>recursion</mark>" in tr["snippet_html"]
    assert ll.search("parsers loop")  # stemming: "parser loops" matches
    assert ll.search("recurs")  # the last word is a prefix, for search-as-you-type
    assert ll.search("grammar") and ll.search("") == [] and ll.search('"; DROP') == []
    assert ll.search("recursion", course_id=ll.find_course("MML")["id"]) == []
    with ll.db() as con:
        con.execute("INSERT INTO segments(lecture_id, idx, start, end, text) VALUES (?, 99, 200, 210, ?)",
                    (lid, "the <script>alert(1)</script> grammar"))
    ll.index_lecture(lid)
    snip = next(h for h in ll.search("alert") if h["kind"] == "transcript")["snippet_html"]
    assert "<script>" not in snip and "&lt;script&gt;" in snip


def test_lectures_from_before_search_get_indexed(ll):
    cid = ll.find_course("COP4020")["id"]
    lid = _lec(ll, "Old one", cid, segs=[(0, 5, "monads are burritos according to nobody")])
    assert ll.search("burritos") == []
    assert ll.index_missing() == 1 and ll.search("burritos")[0]["lecture_id"] == lid
    assert ll.index_missing() == 0


def test_offline_notes_need_no_claude(ll, monkeypatch):
    monkeypatch.setattr(ll, "_claude_run", lambda *a: pytest.fail("the offline generator must not call claude"))
    cid = ll.find_course("COP4020")["id"]
    words = "grammar parser lexer token derivation production terminal nonterminal".split()
    segs = [(i * 20, i * 20 + 20, f"Now the {words[i % len(words)]} is the key idea we use in compilers today. "
             f"Remember that every {words[(i + 3) % len(words)]} matters here.") for i in range(40)]
    with ll.db() as con:
        con.execute("UPDATE courses SET vocabulary=? WHERE id=?", ("\n".join(words), cid))
    lid = _lec(ll, "Offline", cid, segs=segs)
    res = ll.job_notes({"id": 0, "payload": {"lecture_id": lid, "mode": "offline"}})
    assert res["mode"] == "offline" and res["sections"] >= 2 and res["cards"] > 0 and res["quiz"] > 0
    n = ll.lecture_notes(lid)
    assert n["generator"] == "offline" and "(#t=" in n["sections"][0]["markdown"]
    q = ll.lecture_quiz(lid)["questions"][0]
    assert len(q["choices"]) == 4 and "_____" in q["question"]


def test_study_api(ll, monkeypatch):
    from fastapi.testclient import TestClient
    lid = _notes_done(ll, monkeypatch)
    c = TestClient(ll.create_app(), base_url=f"http://127.0.0.1:{ll.PORT}")
    cards = c.get(f"/api/lectures/{lid}/cards").json()
    assert len(cards) == 2 and c.get("/api/lectures/999/cards").status_code == 404
    r = c.post(f"/api/cards/{cards[0]['id']}/review", json={"correct": True}).json()
    assert r["card"]["box"] == 2
    assert c.post(f"/api/cards/{cards[0]['id']}/undo", json={"before": r["before"]}).json()["box"] == 1
    assert c.post("/api/cards/999/review", json={"correct": True}).status_code == 404
    assert c.post(f"/api/cards/{cards[0]['id']}/undo", json={"before": {"box": 1}}).status_code == 422
    assert c.get("/api/review").json()["total"] == 2
    qs = c.get(f"/api/lectures/{lid}/quiz").json()["questions"]
    out = c.post(f"/api/lectures/{lid}/quiz", json={"answers": [{"id": qs[0]["id"], "answer": "left"}]}).json()
    assert out["score"] == 1
    a = c.get(f"/api/anki.tsv?lecture_id={lid}")
    assert a.status_code == 200 and "attachment" in a.headers["content-disposition"] and "#guid" in a.text
    assert c.get("/api/search?q=recursion").json()[0]["lecture_id"] == lid


# ---------------------------------------------------------------- Stage 5: export and sync

def test_notes_export_to_ai_notes_with_frontmatter_and_links_back(ll, monkeypatch):
    lid = _notes_done(ll, monkeypatch)  # saving notes exports them
    with ll.db() as con:
        con.execute("UPDATE lectures SET date='2026-09-01', title='Grammars: BNF \"and\" more' WHERE id=?", (lid,))
        con.execute("INSERT INTO sources(lecture_id, kind, url) VALUES (?, 'url', 'https://example.edu/rec/1')", (lid,))
    notes, tr = ll.export_lecture(lid)
    assert notes.parent == ll.NOTES_DIR / "COP4020"
    assert notes.name == "2026-09-01 - Grammars_ BNF _and_ more.md"
    assert tr.name == "2026-09-01 - Grammars_ BNF _and_ more - Transcript.md"
    md = notes.read_text(encoding="utf-8")
    assert md.startswith('---\ncourse: "COP4020"\n') and 'title: "Grammars: BNF \\"and\\" more"' in md
    assert 'source_url: "https://example.edu/rec/1"' in md and 'duration: "0:10:00"' in md
    assert "## Summary" in md and "## BNF" in md and "## Open questions" in md
    assert "(#t=" not in md and f"(http://127.0.0.1:{ll.PORT}/?t=40#/lecture/{lid})" in md
    assert "[Transcript](<2026-09-01 - Grammars_ BNF _and_ more - Transcript.md>)" in md
    t = tr.read_text(encoding="utf-8")
    assert 'kind: "transcript"' in t and "[0:01:40] left recursion means" in t
    # renamed: re-exported under the new name, the old files gone
    with ll.db() as con:
        con.execute("UPDATE lectures SET title='Context-free grammars' WHERE id=?", (lid,))
    new_notes, new_tr = ll.export_lecture(lid)
    assert new_notes.exists() and not notes.exists() and not tr.exists()
    assert ll.get_lecture(lid)["status"] == "exported"
    (ll.NOTES_DIR / "COP4020" / "mine.md").write_text("my own file", encoding="utf-8")
    ll.export_lecture(lid)
    assert (ll.NOTES_DIR / "COP4020" / "mine.md").exists()  # only files LectureLab wrote are ever removed


def test_export_pending_catches_up_after_a_failed_export(ll, monkeypatch):
    real = ll.export_lecture
    monkeypatch.setattr(ll, "export_lecture", lambda lid: (_ for _ in ()).throw(OSError("drive missing")))
    lid = _notes_done(ll, monkeypatch)
    assert ll.lecture_notes(lid)["state"] == "done"  # the notes are kept even though the export failed
    assert ll.export_pending() == 0
    monkeypatch.setattr(ll, "export_lecture", real)
    assert ll.export_pending() == 1 and ll.export_pending() == 0


def _quiet_sync(ll, monkeypatch, server=False):
    monkeypatch.setattr(ll, "canvas_state", lambda: {"state": "needs_signin"})
    monkeypatch.setattr(ll, "inbox_roots", lambda: [])
    monkeypatch.setattr(ll, "server_running", lambda: server)


def test_sync_does_everything_new_once_and_nothing_the_second_time(ll, monkeypatch):
    _quiet_sync(ll, monkeypatch)
    ran = []
    monkeypatch.setattr(ll, "job_transcribe", lambda job: ran.append("transcribe") or {"transcribed": []})
    ll.JOB_HANDLERS["transcribe"] = ll.job_transcribe
    _fake_claude(ll, monkeypatch, {"outline": OUTLINE, "section": _section, "logistics": _section,
                                   "study": STUDY, "check": CHECK})
    _lecture(ll, "queued_transcribe")
    lid = _notes_lecture(ll)  # transcribed before Stage 4, no notes yet
    out = ll.run_sync()
    assert out["transcribe"] == 1 and out["notes"] == 1 and out["canvas"] == "needs sign-in"
    assert ran == ["transcribe"] and ll.lecture_notes(lid)["state"] == "done"
    assert ll.export_paths(ll.get_lecture(lid))[0].exists()
    with ll.db() as con:  # the fake batch transcribed nothing; take it out of the picture
        con.execute("UPDATE lectures SET status='error' WHERE status='queued_transcribe'")
    assert ll.run_sync() == {"canvas": "needs sign-in"}  # nothing new
    assert ran == ["transcribe"]


def test_sync_leaves_the_queue_to_the_server_when_it_is_up(ll, monkeypatch):
    _quiet_sync(ll, monkeypatch, server=True)
    monkeypatch.setattr(ll, "run_job", lambda job: pytest.fail("the server runs jobs, not sync"))
    lid = _notes_lecture(ll)
    assert ll.run_sync() == {"canvas": "needs sign-in", "notes": 1}
    with ll.db() as con:
        assert con.execute("SELECT status FROM jobs WHERE kind='notes' AND lecture_id=?", (lid,)).fetchone()[0] == "queued"


def test_only_one_sync_runs_at_a_time(ll, monkeypatch):
    _quiet_sync(ll, monkeypatch)
    lock = ll.FileLock(ll.DATA / "sync.lock")
    assert lock.acquire()
    try:
        assert ll.run_sync() == {"skipped": "another sync is running"}
    finally:
        lock.release()


def test_sync_skips_failed_notes_and_lectures_still_waiting_for_slides(ll, monkeypatch):
    _quiet_sync(ll, monkeypatch, server=True)
    cid = ll.find_course("COP4020")["id"]
    failed = _notes_lecture(ll)
    with ll.db() as con:
        con.execute("UPDATE lectures SET notes_error='Claude Code failed: x' WHERE id=?", (failed,))
        con.execute("UPDATE courses SET canvas_id=11 WHERE id=?", (cid,))
    waiting = _lec(ll, "Has video, course not scanned", cid, segs=_said([1, 2]), video="media/z/original.mp4")
    assert ll.lecture_slides(waiting)["state"] == "waiting"
    assert ll.queue_missing_notes() == 0


def test_tidy_section_fixes_what_the_model_gets_wrong(ll):
    md = ("## Section title\n\n- [▶ 32:28](#t=1948) > [!unclear] The lecturer says it's confusing\n\n"
          "[▶ 1:00](#t=60): far outside\n\n> [!exam] kept as is")
    out = ll.tidy_section(md, 1800, 2000)
    assert out.splitlines()[0] == "> [!unclear] [▶ 32:28](#t=1948) The lecturer says it's confusing"
    assert "(1:00): far outside" in out and "> [!exam] kept as is" in out and "Section title" not in out
    assert ll.no_heading("## Summary\nThis lecture covers") == "This lecture covers"


# ---------------------------------------------------------------- review fixes

def test_sync_takes_a_new_inbox_file_in_one_pass(ll, monkeypatch, tmp_path):
    _quiet_sync(ll, monkeypatch, server=True)
    inbox = tmp_path / "inbox"
    (inbox / "COP4020").mkdir(parents=True)
    (inbox / "COP4020" / "lecture 9.mp4").write_bytes(b"\0" * 2048)
    monkeypatch.setattr(ll, "inbox_roots", lambda: [(inbox, "inbox")])
    assert ll.run_sync()["inbox"] == 1  # a fresh process every hour can't wait for a second look


def test_sync_recovers_a_job_left_running_by_a_dead_run(ll, monkeypatch):
    _quiet_sync(ll, monkeypatch)
    lid = _notes_lecture(ll)
    jid = ll.enqueue_notes(lid)
    with ll.db() as con:
        con.execute("UPDATE jobs SET status='running', owner_pid=99999999 WHERE id=?", (jid,))
    ran = []
    monkeypatch.setattr(ll, "run_job", lambda job: ran.append(job["id"]) or ll.job_update(job["id"], status="done"))
    out = ll.run_sync()
    assert out["recovered"] == 1 and jid in ran


def test_sync_gives_up_on_a_recording_that_failed_three_times(ll, monkeypatch):
    _quiet_sync(ll, monkeypatch, server=True)
    lid = _lecture(ll, "queued_transcribe")
    with ll.db() as con:
        con.execute("UPDATE lectures SET transcribe_fails=3 WHERE id=?", (lid,))
    out = ll.run_sync()
    assert "transcribe" not in out and out["transcribe_failing"] == 1
    new = _lecture(ll, "queued_transcribe")  # a new recording doesn't drag the stuck one along
    assert ll.run_sync()["transcribe"] == 1
    with ll.db() as con:
        payload = json.loads(con.execute("SELECT payload FROM jobs WHERE kind='transcribe'").fetchone()[0])
    assert payload["lecture_ids"] == [new]


def test_sync_writes_a_few_lectures_notes_at_a_time(ll, monkeypatch):
    _quiet_sync(ll, monkeypatch, server=True)
    for _ in range(5):
        _notes_lecture(ll)
    assert ll.run_sync()["notes"] == ll.SYNC_NOTES == 3
    assert "notes" not in ll.run_sync()  # three still waiting (the usage limit): no more on top


def test_exports_of_lectures_with_the_same_name_never_clobber_each_other(ll, monkeypatch):
    a = _notes_done(ll, monkeypatch)
    b = _notes_lecture(ll)  # same course, date and title
    ll.job_notes({"id": 0, "payload": {"lecture_id": b}})
    pa = ll.export_paths(ll.get_lecture(a))
    with ll.db() as con:
        pb = [ll.Path(p) for p in json.loads(con.execute("SELECT export_paths FROM notes WHERE lecture_id=?",
                                                      (b,)).fetchone()[0])]
    assert pa[0].exists() and pb[0].exists() and pa[0] != pb[0] and f"({b})" in pb[0].name
    with ll.db() as con:
        con.execute("UPDATE lectures SET title='Lecture 3.1 Parsing basics' WHERE id=?", (a,))
    assert ll.export_pending() == 1  # renamed: re-exported without being asked
    new = ll.export_paths(ll.get_lecture(a))[0]
    assert new.name.endswith("Lecture 3.1 Parsing basics.md") and new.exists() and not pa[0].exists()
    assert pb[0].exists() and pb[1].exists()  # b's files survive a's rename
    ll.delete_lecture(a)
    assert not new.exists() and pb[0].exists()  # deleting a lecture takes its exports with it


def test_notes_never_take_a_lecture_out_of_the_transcription_queue(ll, monkeypatch):
    lid = _notes_lecture(ll)
    with ll.db() as con:
        con.execute("UPDATE lectures SET status='queued_transcribe' WHERE id=?", (lid,))
    _fake_claude(ll, monkeypatch, {"outline": OUTLINE, "section": _section, "logistics": _section,
                                   "study": STUDY, "check": CHECK})
    ll.enqueue_notes(lid)
    assert ll.get_lecture(lid)["status"] == "queued_transcribe"
    ll.job_notes({"id": 0, "payload": {"lecture_id": lid}})
    assert ll.get_lecture(lid)["status"] == "queued_transcribe" and ll.lecture_notes(lid)["summary"]


def test_notes_stay_stored_when_the_export_breaks(ll, monkeypatch):
    monkeypatch.setattr(ll, "export_lecture", lambda lid: 1 / 0)
    lid = _notes_done(ll, monkeypatch)
    assert ll.lecture_notes(lid)["state"] == "done" and ll.get_lecture(lid)["notes_error"] is None


def test_regenerating_keeps_review_progress_and_offline_never_replaces_claude(ll, monkeypatch):
    from fastapi.testclient import TestClient
    lid = _notes_done(ll, monkeypatch)
    card = ll.due_cards(lecture_id=lid)["cards"][0]["id"]
    ll.review_card(card, True)
    ll.job_notes({"id": 0, "payload": {"lecture_id": lid}})  # every call cached: the same cards
    with ll.db() as con:
        assert con.execute("SELECT MAX(box) FROM cards WHERE lecture_id=?", (lid,)).fetchone()[0] == 2
    with pytest.raises(ValueError):
        ll.enqueue_notes(lid, "offline")
    c = TestClient(ll.create_app(), base_url=f"http://127.0.0.1:{ll.PORT}")
    assert c.post(f"/api/lectures/{lid}/notes", json={"mode": "offline"}).status_code == 409
    assert c.post(f"/api/lectures/{lid}/notes", json={"mode": "offline", "force": True}).status_code == 200


def test_claude_never_sees_a_paid_route(ll, monkeypatch):
    for k in ("anthropic_api_key", "ANTHROPIC_BASE_URL", "CLAUDE_CODE_USE_BEDROCK", "CLAUDE_CODE_USE_VERTEX",
              "AWS_BEARER_TOKEN_BEDROCK", "ANTHROPIC_CUSTOM_HEADERS"):
        monkeypatch.setenv(k, "x")
    env = ll.claude_env()
    assert not [k for k in env if k.upper().startswith(("ANTHROPIC_", "CLAUDE_CODE_USE_", "AWS_BEARER"))]
    assert env.get("PATH") or env.get("Path")


def test_an_empty_answer_is_never_cached(ll, monkeypatch):
    calls = []
    _fake_claude(ll, monkeypatch, {"section": ""}, calls)
    for _ in range(2):
        with pytest.raises(ll.ClaudeError):
            ll.claude_call("p", system=ll.SECTION_SYSTEM, kind="section")
    assert len(calls) == 2 and not list(ll.CLAUDE_CACHE.glob("*.json"))


def test_claude_signed_out_waits_instead_of_failing_every_lecture(ll, monkeypatch):
    calls = []
    _fake_claude(ll, monkeypatch, {"outline": (True, "Invalid API key · Please run /login")}, calls)
    lid = _notes_lecture(ll)
    with pytest.raises(ll.RetryLater):
        ll.job_notes({"id": ll.enqueue_notes(lid), "payload": {"lecture_id": lid}})
    lec = ll.get_lecture(lid)
    assert lec["notes_error"] is None and lec["status"] == "transcribed"
    with pytest.raises(ll.ClaudeUnavailable):  # parked: no second call for an hour
        ll.claude_call("p", system=ll.OUTLINE_SYSTEM, kind="outline")
    assert len(calls) == 1


def test_check_edits_that_match_nothing_leave_the_notes_alone(ll):
    secs = [{"markdown": "[▶ 0:10](#t=10) A is B. A is B.", "start": 0, "end": 60}]
    skipped = ll.apply_check_edits(secs, [{"number": 1, "find": "A is B.", "replace": ""},
                                          {"number": 1, "find": "nowhere", "replace": ""},
                                          {"number": 9, "find": "A", "replace": ""},
                                          {"number": 1, "find": secs[0]["markdown"], "replace": ""}])
    assert secs[0]["markdown"] == "[▶ 0:10](#t=10) A is B. A is B." and [n for n, _ in skipped] == [1, 1, 1]


def test_a_new_recording_makes_the_same_hours_batch(ll, monkeypatch):
    _quiet_sync(ll, monkeypatch)
    ran = []

    def fake_run(job):
        if job["kind"] == "ingest":
            _lecture(ll, "queued_transcribe")
        ran.append(job["kind"])
        ll.job_update(job["id"], status="done")
    monkeypatch.setattr(ll, "run_job", fake_run)
    ll.enqueue("ingest", {"path": "inbox:x.mp4"})  # what the inbox scan queues
    out = ll.run_sync()
    assert out["transcribe"] == 1 and ran == ["ingest", "transcribe"]


# ---------------------------------------------------------------- sweep regressions

def test_colab_off_refuses_an_explicit_colab_batch(ll, monkeypatch):
    monkeypatch.setattr(ll, "COLAB_ENABLED", False)
    monkeypatch.setattr(ll, "transcribe_batch_colab", lambda *a, **k: pytest.fail("Colab is off"))
    monkeypatch.setattr(ll, "stop_orphan_sessions", lambda *a, **k: pytest.fail("Colab is off"))
    lid = _lecture(ll, "queued_transcribe")
    with pytest.raises(RuntimeError, match="LECTURELAB_COLAB=off"):
        ll.job_transcribe({"id": 1, "payload": {"mode": "colab"}})
    lec = ll.get_lecture(lid)
    assert lec["status"] == "queued_transcribe" and lec["transcribe_fails"] == 0


def test_waiting_for_another_batch_does_not_use_up_attempts(ll):
    jid = ll.enqueue("transcribe", {"mode": "local"})
    held = ll.FileLock(ll.DATA / "transcribe.lock")
    assert held.acquire()
    try:
        for _ in range(4):
            with ll.db() as con:
                con.execute("UPDATE jobs SET retry_at=NULL WHERE id=?", (jid,))
            ll.run_job(ll.claim_job(["transcribe"]))
    finally:
        held.release()
    with ll.db() as con:  # the server dies while it finally runs: it's re-queued, not given up on
        con.execute("UPDATE jobs SET status='running', owner_pid=999999, attempts=attempts+1 WHERE id=?", (jid,))
    ll.requeue_orphaned_jobs()
    assert ll.get_job(jid)["status"] == "queued"


def test_a_zoom_passcode_in_the_link_never_shows_in_jobs(ll, monkeypatch):
    from fastapi.testclient import TestClient
    url = "https://ufl.zoom.us/rec/share/abc?pwd=s3cret&startTime=1"
    monkeypatch.setattr(ll, "JOB_HANDLERS", {**ll.JOB_HANDLERS, "ingest": lambda job: {"lectures": []}})
    queued = ll.enqueue("ingest", {"url": url})
    c = TestClient(ll.create_app(), base_url=f"http://127.0.0.1:{ll.PORT}")
    assert "s3cret" not in c.get("/api/jobs").text  # queued: the label is the link without its passcode
    ll.run_job(ll.claim_job(["ingest"]))
    assert "s3cret" not in json.dumps(ll.get_job(queued)["payload"])  # and the stored link loses it after the run
    assert "s3cret" not in c.get("/api/jobs?active=false").text


def test_lecture_details_can_clear_the_course_and_date_and_reject_impossible_dates(ll):
    from fastapi.testclient import TestClient
    c = TestClient(ll.create_app(), base_url=f"http://127.0.0.1:{ll.PORT}")
    lid = _lec(ll, "Grammars", ll.find_course("COP4020")["id"], date="2026-10-01")
    r = c.patch(f"/api/lectures/{lid}", json={"title": "  Parsing\n part 2 ", "date": None, "course_id": None})
    assert r.status_code == 200 and (r.json()["title"], r.json()["date"], r.json()["course_id"]) == \
        ("Parsing part 2", None, None)  # what the UI sends for "No course" and an empty date
    assert c.patch(f"/api/lectures/{lid}", json={"date": "2026-10-02"}).json()["title"] == "Parsing part 2"
    for bad in ("2024-02-30", "2024-13-01"):
        assert c.patch(f"/api/lectures/{lid}", json={"date": bad}).status_code == 400
        assert c.post("/api/recordings", json={"url": "https://x.example/a", "date": bad}).status_code == 400
    assert c.patch(f"/api/lectures/{lid}", json={"title": " \n "}).status_code == 400
    assert c.patch(f"/api/lectures/{lid}", json={"title": None}).json()["title"] == "Parsing part 2"
    assert ll.get_lecture(lid)["date"] == "2026-10-02"


def test_asking_to_transcribe_the_same_lecture_twice_makes_one_batch(ll, monkeypatch):
    from fastapi.testclient import TestClient
    monkeypatch.setattr(ll, "wake_workers", lambda: None)
    c = TestClient(ll.create_app(), base_url=f"http://127.0.0.1:{ll.PORT}")
    a, b = _lecture(ll, "error"), _lecture(ll, "error")
    first = c.post("/api/transcribe", json={"lecture_ids": [a], "mode": "colab"}).json()["job_id"]
    again = c.post("/api/transcribe", json={"lecture_ids": [a], "mode": "local"}).json()
    assert again["job_id"] == first and again["queued"] == "colab"  # says how the batch really runs
    assert c.post("/api/transcribe", json={"lecture_ids": [b], "mode": "auto"}).json()["job_id"] != first


def test_a_regenerate_that_stops_keeps_a_lecture_with_notes_ready(ll, monkeypatch):
    lid = _notes_done(ll, monkeypatch)
    assert ll.get_lecture(lid)["status"] == "exported"
    for err in (ll.ClaudeError("Claude Code failed: boom"), ll._claude_down("signed out")):
        def fail(*a, err=err, **k):
            raise err
        monkeypatch.setattr(ll, "generate_notes", fail)
        with pytest.raises((ll.ClaudeError, ll.RetryLater)):
            ll.job_notes({"id": ll.enqueue_notes(lid), "payload": {"lecture_id": lid}})
        assert ll.get_lecture(lid)["status"] == "exported"  # its notes are still there to study from
    monkeypatch.setattr(ll, "generate_notes", lambda *a, **k: (_ for _ in ()).throw(ll.ClaudeError("boom")))
    fresh = _notes_lecture(ll)  # a lecture without notes still goes back to transcribed
    with pytest.raises(ll.ClaudeError):
        ll.job_notes({"id": ll.enqueue_notes(fresh), "payload": {"lecture_id": fresh}})
    assert ll.get_lecture(fresh)["status"] == "transcribed"


def test_deleting_a_lecture_mid_batch_does_not_fail_the_rest(ll, monkeypatch):
    a, b = _lecture(ll, "queued_transcribe"), _lecture(ll, "queued_transcribe")

    def fake(lec, progress=None, model_name=None):
        if lec["id"] == a:
            ll.delete_lecture(a)  # you delete it while the CPU works on it
        return {"model": "base", "segments": [{"start": 0, "end": 1, "text": "hi", "words": []}]}
    monkeypatch.setattr(ll, "transcribe_local", fake)
    out = ll.job_transcribe({"id": 1, "payload": {"mode": "local"}})
    assert b in out["transcribed"]
    lec = ll.get_lecture(b)
    assert lec["status"] == "transcribed" and lec["transcribe_fails"] == 0
    assert not (ll.DATA / "transcripts" / f"{a}.json").exists()


def test_a_canvas_file_named_with_windows_reserved_characters_downloads(ll, monkeypatch, noise):
    class FakeSession:
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def httpx_cookies(self):
            return None
    monkeypatch.setattr(ll, "canvas_ready", lambda: None)
    monkeypatch.setattr(ll, "CanvasSession", FakeSession)
    c = ll.find_course("COP4020")
    with ll.db() as con:
        con.execute("INSERT INTO canvas_items(key, course_id, kind, title, url) VALUES "
                    "('file:7', ?, 'file', 'What is a PL?', 'canvas-file:7')", (c["id"],))
    for name in ('What is a PL?.wav', 'Lecture 3: TCP.wav', 'a|b "c".wav'):
        monkeypatch.setattr(ll, "canvas_file_source", lambda fid, name=name: ("https://x/f", name, {"title": "t"}))
        got = []

        def fake_stream(url, out, jar=None, max_bytes=None):
            out.parent.mkdir(parents=True, exist_ok=True)
            out.write_bytes(noise.read_bytes())  # Windows refuses ? and | here; ':' would make an ADS
            got.append(out)
            return out
        monkeypatch.setattr(ll, "stream_download", fake_stream)
        out = ll.job_canvas_pull({"id": 6, "payload": {"course_id": c["id"], "key": "file:7"}})
        assert out["lectures"][0]["lecture_id"] and ":" not in got[0].name


def test_a_read_only_file_in_the_inbox_is_still_taken_in(ll, noise, tmp_path, monkeypatch):
    import shutil
    import stat
    monkeypatch.setattr(ll, "LOCAL_INBOX", tmp_path / "inbox")
    monkeypatch.setenv("LECTURELAB_DRIVE_INBOX", "off")
    ll.scan_inboxes()
    f = tmp_path / "inbox" / "COP4020" / "view-only.wav"
    shutil.copy(noise, f)
    os.chmod(f, stat.S_IREAD)
    try:
        assert len(ll.scan_inboxes(require_stable=False)) == 1
    finally:
        os.chmod(f, stat.S_IREAD | stat.S_IWRITE)


def test_a_password_protected_deck_from_canvas_says_so(ll):
    import pymupdf as fitz
    rel = "Slides/test/locked.pdf"
    plain = _pdf(ll.DATA / "Slides/test/plain.pdf", ["Grammars: BNF"])
    doc = fitz.open(plain)
    doc.save(ll.DATA / rel, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="u", owner_pw="o")
    doc.close()
    did = _deck(ll, "locked", ll.find_course("COP4020")["id"], path=rel, source="canvas")
    with pytest.raises(ll.DeckError, match="password-protected"):
        ll.extract_deck(did)


def test_an_answer_of_the_wrong_shape_is_never_cached(ll, monkeypatch):
    calls = []
    _fake_claude(ll, monkeypatch, {"outline": {"chapters": []}}, calls)  # no "sections"
    for _ in range(2):
        with pytest.raises(ll.ClaudeError, match="wasn't the JSON"):
            ll.claude_call("p", system=ll.OUTLINE_SYSTEM, schema=ll.OUTLINE_SCHEMA, kind="outline")
    assert len(calls) == 2 and not list(ll.CLAUDE_CACHE.glob("*.json"))
    for stdout in ("null", "[1, 2]", "not json at all"):  # claude printed something that isn't its result object
        monkeypatch.setattr(ll, "_claude_run", lambda cmd, p, t, out=stdout: subprocess.CompletedProcess(cmd, 0, out, ""))
        with pytest.raises(ll.ClaudeError, match="Claude Code failed"):
            ll.claude_call(stdout, system=ll.SECTION_SYSTEM, kind="section")


def test_a_title_of_only_spaces_falls_back_to_the_file_name(ll, noise):
    lid = ll.ingest_file(noise, kind="file", title="   \n ")["lecture_id"]
    assert ll.get_lecture(lid)["title"] == "noise"


def test_the_cli_runs_a_job_again_once_it_stops_waiting(ll, monkeypatch):
    import threading
    tries = []

    def flaky(job):
        tries.append(1)
        if len(tries) == 1:  # e.g. another batch held the lock: it re-queues itself for a moment later
            ll.job_update(job["id"], status="queued", retry_at=ll._sqlnow(1), owner_pid=None)
            raise ll.RetryLater("busy")
        return {"ok": True}
    monkeypatch.setitem(ll.JOB_HANDLERS, "flaky", flaky)
    jid = ll.enqueue("flaky", {})
    out = {}
    t = threading.Thread(target=lambda: out.update(ll.run_job_now(jid)), daemon=True)
    t.start()
    t.join(15)
    assert out.get("status") == "done" and len(tries) == 2  # no server: the CLI must not wait forever


def test_stop_while_the_supervisor_waits_to_restart_starts_no_new_server(ll, monkeypatch, tmp_path):
    import cli
    flag = tmp_path / "supervisor.stop"
    monkeypatch.setattr(cli, "STOP_FLAG", flag)
    started = []

    class Crashed:
        pid = 1

        def wait(self):
            return 1  # the server died at once
    monkeypatch.setattr(cli.subprocess, "Popen", lambda *a, **k: started.append(1) or Crashed())
    monkeypatch.setattr(cli.time, "sleep", lambda s: flag.write_text("stop"))  # `lecturelab stop` meanwhile
    assert cli.cmd_supervise(None) == 0
    assert len(started) == 1 and not flag.exists()


def test_cli_output_redirected_to_a_file_survives_any_title():
    env = {k: v for k, v in os.environ.items() if k not in ("PYTHONIOENCODING", "PYTHONUTF8")}
    r = subprocess.run([sys.executable, "-c", "import cli; cli._safe_console(); print('Caf\u00e9 \u2615 lecture')"],
                       capture_output=True, env=env, cwd=os.path.dirname(os.path.dirname(__file__)))
    assert r.returncode == 0, r.stderr


def test_a_canvas_page_that_turns_forbidden_keeps_the_pages_before_it(ll):
    nxt = '<https://ufl.instructure.com/api/v1/courses/1/modules?page=2&per_page=100>; rel="next"'
    cs, _ = fake_session(ll, [FakeResponse(200, [{"id": 1}], link=nxt), FakeResponse(403, {"errors": []})])
    assert cs.get("courses/1/modules") == [{"id": 1}]


def test_releasing_a_missing_lecture_is_a_404(ll):
    from fastapi.testclient import TestClient
    c = TestClient(ll.create_app(), base_url=f"http://127.0.0.1:{ll.PORT}")
    assert c.post("/api/lectures/999/not-duplicate").status_code == 404


def test_a_canvas_scan_of_a_deleted_course_is_skipped(ll):
    assert ll.job_canvas_scan({"id": 1, "payload": {"course_id": 999}}) == {"skipped": "course deleted"}


def test_anki_rows_survive_carriage_returns(ll, monkeypatch):
    lid = _notes_done(ll, monkeypatch)
    with ll.db() as con:
        con.execute("UPDATE cards SET back='one\r\ntwo\rthree' WHERE kind='qa'")
    lines = ll.anki_tsv(lecture_id=lid).splitlines()[6:]  # splitlines breaks on \r, as Anki's reader does
    assert len(lines) == 2 and all(len(x.split("\t")) == 6 for x in lines)


def test_cli_notes_for_a_lecture_that_does_not_exist_says_so(ll, capsys):
    import cli
    assert cli.main(["notes", "999", "--offline"]) == 1
    assert "no such lecture" in capsys.readouterr().out


def test_cli_add_rejects_an_impossible_date_before_downloading(ll, monkeypatch, capsys):
    import cli
    jobs = []
    monkeypatch.setattr(ll, "enqueue", lambda *a, **k: jobs.append(a) or 1)
    assert cli.main(["add", "https://ufl.zoom.us/rec/share/abc", "--date", "2026-02-30"]) == 2
    assert "real day" in capsys.readouterr().out
    assert not jobs  # nothing was queued, so nothing was downloaded
